from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-49-SYSTEM06-ENTERPRISE-ERP-CONTRACT-REMEDIATION-V1"
BASE_HEAD="c168abd122d26681e7336a7f2ed4c0cb5750dd75"
S06="06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
DEV="ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx"
ERP="ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S06:"2f90b14fdcdc9eceeab0c648ed6f58faff8c727e",
DEV:"17760884cccefa1172aab85709be03cf31157d9e",
ERP:"3be4e9a0cc2c11e2b8180eee5f2bfd6db8cc1704",
LOGIC:"84336a8b6a7daa80c3201a730e3cca55c5a579fd",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":DEV,
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":ERP,
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

TARGETS={
"DEV-01-BTN-CAMPAIGN-APPROVE":{
 "page":"DEV-01","operation":"approveOutreachCampaign",
 "fields":{"payload_schema":"ApproveOutreachCampaignRequest","persistence_owner":"Existing Outreach Owner / CampaignApproval + Campaign Audit"},
 "basis":"Approval is distinct from dispatch. The exact Campaign/version/recipient refs and reviewer authority must resolve before approval is persisted."
},
"DEV-01-BTN-CAMPAIGN-CREATE":{
 "page":"DEV-01","operation":"createOutreachCampaign",
 "fields":{"payload_schema":"CreateOutreachCampaignRequest","persistence_owner":"Existing Outreach Owner / Campaign + Audit"},
 "basis":"Campaign persists an exact recipient_ref set, message/sender/schedule refs and governance metadata. Search/filter text is never itself the audience."
},
"DEV-01-BTN-CAMPAIGN-SEARCH":{
 "page":"DEV-01","operation":"searchProjection",
 "fields":{"payload_schema":"SearchEnterpriseAudienceProjectionRequest","persistence_owner":"Global Brain / SemanticIndexResult + Search Audit"},
 "basis":"Audience search persists derived search/index result pointers and scope only; it must dereference exact source/evidence and never create a second Company/Audience truth store."
},
"DEV-01-BTN-CANDIDATE-CREATE":{
 "page":"DEV-01","operation":"createCandidate",
 "fields":{"payload_schema":"CreateEnterpriseCompanyCandidateRequest","persistence_owner":"Existing Enterprise Development Owner / CompanyCandidate + Evidence/Audit"},
 "basis":"Candidate creation records unresolved company evidence under the existing enterprise owner. Candidate is not CompanyMaster and cannot bypass identity resolution."
},
"DEV-01-BTN-CANDIDATE-DECIDE":{
 "page":"DEV-01","operation":"decideCandidate",
 "fields":{"payload_schema":"DecideEnterpriseCandidateRequest"},
 "basis":"Decision must bind exact candidate/version and review evidence. Existing persistence public.context_candidates remains unchanged."
},
"DEV-01-BTN-CR-CREATE":{
 "page":"DEV-01","operation":"createChangeRequest",
 "fields":{"payload_schema":"CreateEnterpriseChangeRequestRequest","persistence_owner":"Existing Governed Change Request Owner / ChangeRequest + Audit"},
 "basis":"Change Request records requested correction/change scope and audit lineage; it does not itself mutate canonical enterprise state."
},
"DEV-01-BTN-DIRECTORY-SEARCH":{
 "page":"DEV-01","operation":"searchProjection",
 "fields":{"payload_schema":"SearchEnterpriseDirectoryProjectionRequest","persistence_owner":"Global Brain / SemanticIndexResult + Search Audit"},
 "basis":"Directory search is a derived projection/search contract and must not create a duplicate CompanyMaster or independent directory truth store."
},
"DEV-01-BTN-DISCOVERY-PAUSE":{
 "page":"DEV-01","operation":"pauseCompanyDiscovery",
 "fields":{"payload_schema":"PauseCompanyDiscoveryRequest","persistence_owner":"Shared Acquisition / Knowledge Owner / DiscoveryJob State + Audit"},
 "basis":"Pause changes only the exact DiscoveryJob state with actor/reason/version audit; it does not replace the shared acquisition engine."
},
"DEV-01-BTN-DISCOVERY-RESUME":{
 "page":"DEV-01","operation":"resumeCompanyDiscovery",
 "fields":{"payload_schema":"ResumeCompanyDiscoveryRequest","persistence_owner":"Shared Acquisition / Knowledge Owner / DiscoveryJob State + Audit"},
 "basis":"Resume requires exact paused job/version and preserves prior attempts/evidence."
},
"DEV-01-BTN-DISCOVERY-START":{
 "page":"DEV-01","operation":"startCompanyDiscovery",
 "fields":{"payload_schema":"StartCompanyDiscoveryRequest","persistence_owner":"Shared Acquisition / Knowledge Owner / DiscoveryJob + Audit"},
 "basis":"Start creates/binds a DiscoveryJob from an authorized active profile and allowed source scope; connector absence remains BLOCK/RETRY, not simulated success."
},
"DEV-01-BTN-DISCOVERY-STOP":{
 "page":"DEV-01","operation":"stopCompanyDiscovery",
 "fields":{"payload_schema":"StopCompanyDiscoveryRequest","persistence_owner":"Shared Acquisition / Knowledge Owner / DiscoveryJob State + Audit"},
 "basis":"Stop closes the legal target job scope while preserving historical evidence and replay/idempotency lineage."
},
"DEV-01-BTN-EMAIL-DISPATCH":{
 "page":"DEV-01","operation":"dispatchOutreachCampaign",
 "fields":{"payload_schema":"DispatchOutreachCampaignRequest","persistence_owner":"Existing Outreach Owner / Campaign Dispatch + IndividualDelivery + Audit"},
 "basis":"Dispatch fans out approved Campaign recipients into exactly-one-To IndividualDelivery units. Provider acceptance and Production connector E2E remain separate runtime evidence."
},
"DEV-01-BTN-EXPORT":{
 "page":"DEV-01","operation":"exportProjection",
 "fields":{"payload_schema":"ExportEnterpriseProjectionRequest","persistence_owner":"Existing Projection/Export Owner / Enterprise Export Artifact + Export Audit"},
 "basis":"Export persists an authorized derived artifact and audit lineage only; permission/classification/privacy/suppression gates remain mandatory."
},
"DEV-01-BTN-KILL-SWITCH":{
 "page":"DEV-01","operation":"setKillSwitch",
 "fields":{"payload_schema":"SetEnterpriseOutreachKillSwitchRequest","persistence_owner":"Existing Operations Governance / Kill Switch State + Audit"},
 "basis":"Kill switch stops eligible pending/in-flight legal scope and records actor/reason/time. Already-completed sends remain immutable historical facts."
},
"DEV-01-BTN-MERGE":{
 "page":"DEV-01","operation":"mergeCompanyCandidate",
 "fields":{"payload_schema":"MergeCompanyCandidateRequest","persistence_owner":"Existing Enterprise Development Owner / CompanyMaster + CompanyMergeDecision + Lineage Audit"},
 "basis":"Effectful merge requires prior preview, exact survivor/merged refs and no blocking conflict; aliases/evidence/contacts/history/suppression lineage must be preserved."
},
"DEV-01-BTN-MERGE-PREVIEW":{
 "page":"DEV-01","operation":"previewCompanyMerge",
 "fields":{"payload_schema":"PreviewCompanyMergeRequest","persistence_owner":"Existing Enterprise Development Owner / CompanyMergeCandidate + Merge Preview Evidence"},
 "basis":"Preview records candidate identity conflicts/survivor/lineage evidence only and cannot itself merge or delete CompanyMaster history."
},
"DEV-01-BTN-REFRESH":{
 "page":"DEV-01","operation":"refreshProjection",
 "fields":{"payload_schema":"RefreshEnterpriseProjectionRequest","persistence_owner":"Global Brain / Enterprise Projection + Projection Audit"},
 "basis":"Refresh rematerializes derived enterprise projection state from canonical owners; projection cannot become a second Company/Campaign/Delivery truth."
},
"ERP-01-BTN-CONNECTOR-CREATE":{
 "page":"ERP-01","operation":"createERPConnector",
 "fields":{"payload_schema":"CreateERPConnectorRequest","persistence_owner":"ERPConnectorService / ERP Connector Configuration + Audit"},
 "basis":"Connector configuration persists provider/adapter/entity scope/mapping/secret-reference metadata only. Raw secret plaintext must remain outside this persistence boundary."
},
"ERP-01-BTN-CONNECTOR-UPDATE":{
 "page":"ERP-01","operation":"updateERPConnector",
 "fields":{"payload_schema":"UpdateERPConnectorRequest","persistence_owner":"ERPConnectorService / ERP Connector Configuration + Audit"},
 "basis":"Update changes exact connector configuration/version with audit. Secret material remains referenced, not copied into Word/log/audit payload."
},
"ERP-01-BTN-CONNECTOR-VALIDATE":{
 "page":"ERP-01","operation":"validateERPConnector",
 "fields":{"payload_schema":"ValidateERPConnectorRequest","persistence_owner":"ERPConnectorService / ERP Connector Validation Evidence"},
 "basis":"Validation persists bounded connectivity/schema/capability evidence for an exact connector/version; a pass is not equivalent to Production business-data correctness."
},
"ERP-01-BTN-EXPORT":{
 "page":"ERP-01","operation":"exportProjection",
 "fields":{"payload_schema":"ExportERPProjectionRequest","persistence_owner":"Existing Projection/Export Owner / ERP Export Artifact + Export Audit"},
 "basis":"ERP export persists derived projection/export artifact and audit only; ERP/Finance remains financial truth and export cannot mutate accounting facts."
},
"ERP-01-BTN-MAPPING-VALIDATE":{
 "page":"ERP-01","operation":"validateERPMapping",
 "fields":{"payload_schema":"ValidateERPMappingRequest","persistence_owner":"ERPConnectorService / ERP Mapping Validation Evidence"},
 "basis":"Mapping validation records exact mapping-version/entity-scope result/evidence without silently changing canonical financial data."
},
"ERP-01-BTN-SNAPSHOT-REFRESH":{
 "page":"ERP-01","operation":"refreshERPSnapshot",
 "fields":{"payload_schema":"RefreshERPSnapshotRequest","persistence_owner":"ERPConnectorService / ERP Snapshot + Freshness/Audit"},
 "basis":"Snapshot refresh persists immutable snapshot identity/freshness/completeness lineage sourced from the connector; it does not fabricate finance values."
},
"ERP-01-BTN-SYNC-CREATE":{
 "page":"ERP-01","operation":"createERPSyncJob",
 "fields":{"payload_schema":"CreateERPSyncJobRequest","persistence_owner":"ERPConnectorService / ERP Sync Job + Audit"},
 "basis":"Sync job persists exact connector/mapping/scope/snapshot target and idempotency/audit metadata. Job creation is not sync completion."
},
"ERP-01-BTN-SYNC-RETRY":{
 "page":"ERP-01","operation":"retryERPSync",
 "fields":{"payload_schema":"RetryERPSyncRequest","persistence_owner":"ERPConnectorService / ERP Sync Attempt + Failure/Audit"},
 "basis":"Retry binds the failed sync job/failure/version and creates auditable retry-attempt lineage; prior failure evidence must not be erased."
},
}

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out
def blob(path):
    b=Path(path).read_bytes();return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():
    got=blob(f);assert got==s,(f,got,s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""
assert len(TARGETS)==25
assert sum(len(x["fields"]) for x in TARGETS.values())==49
assert sum("payload_schema" in x["fields"] for x in TARGETS.values())==25
assert sum("persistence_owner" in x["fields"] for x in TARGETS.values())==24

page_files={"DEV-01":DEV,"ERP-01":ERP}
pre_pages={p:compose(parse_controls(fn)) for p,fn in page_files.items()}
pre_owner=compose(parse_controls(S06))
for uid,s in TARGETS.items():
    for m,label in [(pre_pages[s["page"]],"PAGE"),(pre_owner,"OWNER")]:
        assert uid in m,(label,uid,"MISSING")
        r=m[uid]
        assert r["operation"]==s["operation"],(label,uid,r["operation"],s["operation"])
        assert r["runtime_status"]=="EFFECTFUL_EXACT",(label,uid,r["runtime_status"])
        assert not missing(r["method_path"]),(label,uid,"METHOD_MISSING")
        for field,new in s["fields"].items():
            assert missing(r[field]),(label,uid,field,r[field],new)
assert pre_pages["DEV-01"]["DEV-01-BTN-CANDIDATE-DECIDE"]["persistence_owner"]=="public.context_candidates"

def patch(path,target_uids):
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact_index(h,"Control UID")
        if ci is None:continue
        idx={"payload_schema":hfind(h,"payload / schema","payload","schema"),"persistence_owner":hfind(h,"persistence owner")}
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals):continue
            uid=vals[ci]
            if uid not in target_uids:continue
            for field,new in TARGETS[uid]["fields"].items():
                fi=idx[field]
                if fi is None or fi>=len(row.cells):continue
                old=norm(row.cells[fi].text)
                if old==new:hits[(uid,field)]+=1;continue
                assert missing(old),(path,uid,field,old,new)
                row.cells[fi].text=new;hits[(uid,field)]+=1
    expected={(u,f) for u in target_uids for f in TARGETS[u]["fields"]}
    assert set(hits)==expected,(path,"PATCH_MISMATCH",sorted(expected-set(hits)),dict(hits))
    d.save(path);Document(path)
    return {f"{u}:{f}":n for (u,f),n in hits.items()}

patches={
"DEV_PAGE":patch(DEV,[u for u,s in TARGETS.items() if s["page"]=="DEV-01"]),
"ERP_PAGE":patch(ERP,[u for u,s in TARGETS.items() if s["page"]=="ERP-01"]),
"SYSTEM06_OWNER":patch(S06,list(TARGETS)),
}

def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.33);sec.bottom_margin=Inches(.33);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=3.55):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for rr in rows:
        cells=t.add_row().cells
        for i,v in enumerate(rr):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
    return t

for path,title,prefix,uids in [
(S06,"Batch 49 · System 06 Enterprise / ERP Canonical Contracts","SYSTEM06",list(TARGETS)),
(DEV,"Batch 49 · DEV-01 Enterprise Contract Binding Ledger","PAGE::DEV-01",[u for u,s in TARGETS.items() if s["page"]=="DEV-01"]),
(ERP,"Batch 49 · ERP-01 Connector / Sync Contract Binding Ledger","PAGE::ERP-01",[u for u,s in TARGETS.items() if s["page"]=="ERP-01"]),
]:
    d=Document(path);add_landscape(d);d.add_heading(title,level=1)
    p=d.add_paragraph();p.add_run(f"[{MARK}::{prefix}] ").bold=True
    p.add_run("These bindings close only Payload/Schema and Persistence Owner gaps. Existing Method/Path, Action, Gate, Permission, Runtime Owner and Runtime Status remain unchanged. Projection/search/export state stays derived; enterprise canonical entities reuse their existing owners; ERP connector secrets remain references only; no Production/provider execution is claimed.")
    rr=[]
    for uid in uids:
        s=TARGETS[uid]
        rr.append([uid,s["operation"],s["fields"].get("payload_schema","UNCHANGED"),s["fields"].get("persistence_owner","UNCHANGED"),s["basis"]])
    add_table(d,["Control UID","Operation","Payload / Schema","Persistence Owner","Required Semantics"],rr)
    d.save(path);Document(path)

post_pages={p:compose(parse_controls(fn)) for p,fn in page_files.items()}
post_owner=compose(parse_controls(S06))
for uid,s in TARGETS.items():
    for m,label in [(post_pages[s["page"]],"PAGE"),(post_owner,"OWNER")]:
        for f,v in s["fields"].items():
            assert m[uid][f]==v,(label,uid,f,m[uid][f],v)
        assert m[uid]["method_path"]==pre_pages[s["page"]][uid]["method_path"],(label,uid,"METHOD_CHANGED")

all_maps={p:compose(parse_controls(fn)) for p,fn in PAGES.items()}
remaining=[]
for page,m in all_maps.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if missing(op) or st not in {"EFFECTFUL_EXACT","READ_EXACT"}:continue
        if missing(r.get("method_path","")):remaining.append(("method_path",page,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("payload_schema","")):remaining.append(("payload_schema",page,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("persistence_owner","")):remaining.append(("persistence_owner",page,uid))
rc=collections.Counter(f for f,_,_ in remaining)
assert len(remaining)==190,(len(remaining),rc)
assert rc=={"method_path":62,"payload_schema":73,"persistence_owner":55},rc

logic=Document(LOGIC);add_landscape(logic);logic.add_heading("Batch 49 · System 06 Enterprise / ERP Contract Closure",level=1)
p=logic.add_paragraph();p.add_run(f"[{MARK}] ").bold=True
p.add_run("System 06 closes all 49 remaining DEV-01/ERP-01 Payload/Persistence cells. Company/Campaign/Discovery reuse existing canonical owners, projection/search/export remain derived, ERP connector/sync state remains under ERPConnectorService, and raw secret plaintext is excluded from the persistence contract. Word completion is not Production connector/provider proof.")
add_table(logic,["Metric","Value","Result"],[
["Pre contract denominator",239,"Method 62 + Payload 98 + Persistence 79"],
["System 06 cells closed",49,"Payload 25 + Persistence 24"],
["Post contract denominator",190,"Method 62 + Payload 73 + Persistence 55"],
["Method/Path changed",0,"All existing exact routes preserved"],
["New duplicate enterprise truth stores",0,"Existing owners reused"],
["Raw ERP/provider secret plaintext persisted",0,"References only"],
["Runtime / Production execution claimed","False","Design-contract closure only"],
],4.0)
machine={
"marker":MARK,"base_head":BASE_HEAD,"pre_contract_cells":239,"system06_cells_closed":49,
"payload_schema_closed":25,"persistence_owner_closed":24,"method_path_changed":0,
"post_contract_cells":190,"remaining_method_path":62,"remaining_payload_schema":73,"remaining_persistence_owner":55,
"duplicate_enterprise_truth_stores_created":0,"raw_secret_plaintext_persisted":0,"runtime_execution_claimed":False,
}
logic.add_paragraph("BATCH49_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

changed=[S06,DEV,ERP,LOGIC]
report={"machine":machine,"targets":TARGETS,"patches":patches,"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed}}
Path("__batch49_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH49="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
