from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib

MARK="ACPOS-20260921-BATCH-15-OWNER-SPECIFIC-READ-PROJECTION-CLOSURE-V1"
S01="01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx"
S06="06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
DB="ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
ERP="ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
 S01:"89447bc4f2a3b87f7bd9157cb2a2d8a4081dbea1",
 S06:"7396aa6f5fab6661e6a4a6813ed6d840d836bc11",
 DB:"741bd95a20b44252725732c0158d2439364cb83e",
 ERP:"a0c75d345291cdc5a6c8131a3a2120473957ea48",
 LOGIC:"0867830ad75ace7a7570529f35e035694ff82ffe",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":DB,
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":ERP,
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()
for f,s in EXPECTED_SHA.items():
    assert blob(f)==s,(f,blob(f),s)

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def display_only_type(t):
    u=(t or "").upper()
    return u in {"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"} or any(x in u for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(h,"control uid")
        if ci is None:continue
        idx={
          "control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),
          "method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),
          "persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status"),
        }
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"table":ti+1,"row":ri,"headers":h})
            out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out

# Batch 13/14 existing rules.
def is_wb_na(r):
    return missing(r.get("action","")) and r.get("type")=="SECTION_OPEN" and r.get("gate")=="PAGE_READ" and r.get("permission")=="workspace.dashboard.view" and r.get("payload_schema")=="N/A_READ_PROJECTION" and r.get("operation")=="getDashboardReadModel" and r.get("method_path")=="GET /v1/dashboard/read-model" and r.get("runtime_owner")=="DASHBOARD_READ_MODEL" and r.get("runtime_status")=="READ_EXACT"
def is_aiapi_na(r):
    return missing(r.get("action","")) and r.get("type")=="BUTTON_OR_ROW_ACTION" and r.get("payload_schema")=="AIAPI Page Operation-specific Form / Provider Profile Field Contract" and not missing(r.get("permission","")) and not missing(r.get("operation","")) and not missing(r.get("method_path","")) and not missing(r.get("runtime_owner","")) and r.get("runtime_status") in {"EFFECTFUL_EXACT","READ_EXACT"}
def is_edit_na(r):
    return missing(r.get("operation","")) and r.get("runtime_status")=="LOCAL_WORKING_DRAFT_EXACT" and r.get("method_path")=="NO_PUBLIC_API_BY_AUTHORITY" and not missing(r.get("action","")) and not missing(r.get("gate","")) and not missing(r.get("permission","")) and not missing(r.get("runtime_owner",""))
def is_iam_op_na(r):
    return missing(r.get("operation","")) and r.get("runtime_status")=="ORCHESTRATES_EXISTING_EXACT_OPERATIONS" and not missing(r.get("action","")) and not missing(r.get("gate","")) and not missing(r.get("permission",""))
def is_iam_owner_na(r):
    return missing(r.get("runtime_owner","")) and r.get("runtime_status")=="ORCHESTRATES_EXISTING_EXACT_OPERATIONS" and not missing(r.get("action","")) and not missing(r.get("gate","")) and not missing(r.get("permission",""))

# Batch 15 structural rules.
DB_OWNER="Database Read Model / DB-01"
ERP_CONNECTOR_OWNER="ERPConnectorService"
ERP_FINANCE_OWNER="FinanceGuardrailService"
ERP_GROUPS={
 "ERP-01-ACT-READ-CONNECTOR":("ERP-01-GATE-CONNECTOR-READ","ERP-01-PERM-CONNECTOR",ERP_CONNECTOR_OWNER),
 "ERP-01-ACT-READ-FINANCE":("ERP-01-GATE-FINANCE","ERP-01-PERM-FINANCE",ERP_FINANCE_OWNER),
 "ERP-01-ACT-READ-SYNC":("ERP-01-GATE-SYNC-READ","ERP-01-PERM-SYNC",ERP_CONNECTOR_OWNER),
}
def is_db_read_model_operation_na(r):
    return (
      missing(r.get("operation","")) and missing(r.get("method_path",""))
      and r.get("runtime_status")=="UI_LOCAL_OR_READ_SOURCE"
      and r.get("runtime_owner")==DB_OWNER
      and not missing(r.get("action","")) and not missing(r.get("gate","")) and not missing(r.get("permission",""))
    )
def is_erp_read_field_operation_na(r):
    action=r.get("action","")
    if action not in ERP_GROUPS:return False
    gate,perm,owner=ERP_GROUPS[action]
    return (
      r.get("type")=="READONLY"
      and missing(r.get("operation","")) and missing(r.get("method_path",""))
      and r.get("runtime_status")=="UI_LOCAL_OR_READ_SOURCE"
      and r.get("gate")==gate and r.get("permission")==perm and r.get("runtime_owner")==owner
    )

def classify_pre(page,field,r):
    v=r.get(field,"");st=r.get("runtime_status","")
    if not missing(v):return "BOUND"
    if "UI_LOCAL_EXACT" in st:
        if field in {"operation","runtime_owner"}:return "LEGITIMATE_NA_UI_LOCAL"
        if field=="action" and display_only_type(r.get("type","")):return "LEGITIMATE_NA_UI_LOCAL"
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in st and field=="operation":return "CLOSED_BY_OWNER_LOCAL_COMMAND"
    if "READ_EXACT" in st and display_only_type(r.get("type","")):
        if field=="action":return "LEGITIMATE_NA_READ_PRESENTATION"
        if field=="operation" and not missing(r.get("runtime_owner","")):return "READ_OWNER_BOUND_OPERATION_UNSPECIFIED"
    if field=="runtime_owner" and any(x in st for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return "KNOWN_RUNTIME_BLOCKER"
    if v=="SOURCE_NOT_DEFINED":return "SOURCE_RESOLUTION_REQUIRED"
    if page=="WB-01" and field=="action" and is_wb_na(r):return "LEGITIMATE_NA_READ_PROJECTION_ACTION"
    if page=="AIAPI-01" and field=="action" and is_aiapi_na(r):return "LEGITIMATE_NA_DIRECT_OPERATION_ACTION"
    if page=="EDIT-01" and field=="operation" and is_edit_na(r):return "LEGITIMATE_NA_LOCAL_WORKING_DRAFT_OPERATION"
    if page=="IAM-01" and field=="operation" and is_iam_op_na(r):return "LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_OPERATION"
    if page=="IAM-01" and field=="runtime_owner" and is_iam_owner_na(r):return "LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_RUNTIME_OWNER"
    return "DEFINITION_BINDING_GAP"
def classify_post(page,field,r):
    pre=classify_pre(page,field,r)
    if pre!="DEFINITION_BINDING_GAP":return pre
    if page=="DB-01" and field=="operation" and is_db_read_model_operation_na(r):
        return "LEGITIMATE_NA_DB_READ_MODEL_OPERATION"
    if page=="ERP-01" and field=="operation" and is_erp_read_field_operation_na(r):
        return "LEGITIMATE_NA_ERP_READ_PROJECTION_FIELD_OPERATION"
    return pre
def inventory(classifier):
    rows=[];counts=collections.Counter()
    for page,fn in PAGES.items():
        bm=compose(parse_controls(fn))
        for uid,r in bm.items():
            for field in ["action","gate","permission","operation","runtime_owner"]:
                cls=classifier(page,field,r)
                counts[cls]+=1
                if cls=="DEFINITION_BINDING_GAP":rows.append((page,uid,field))
    return rows,counts

pre_rows,pre_counts=inventory(classify_pre)
assert len(pre_rows)==164,len(pre_rows)

# Baseline target sets.
db_map=compose(parse_controls(DB));erp_map=compose(parse_controls(ERP))
db_ops=sorted(uid for uid,r in db_map.items() if is_db_read_model_operation_na(r))
assert len(db_ops)==8,(len(db_ops),db_ops)

connector_uids=sorted(uid for uid,r in erp_map.items() if r.get("type")=="READONLY" and r.get("action")=="ERP-01-ACT-READ-CONNECTOR" and r.get("gate")=="ERP-01-GATE-CONNECTOR-READ" and r.get("permission")=="ERP-01-PERM-CONNECTOR" and r.get("runtime_status")=="UI_LOCAL_OR_READ_SOURCE")
finance_uids=sorted(uid for uid,r in erp_map.items() if r.get("type")=="READONLY" and r.get("action")=="ERP-01-ACT-READ-FINANCE" and r.get("gate")=="ERP-01-GATE-FINANCE" and r.get("permission")=="ERP-01-PERM-FINANCE" and r.get("runtime_status")=="UI_LOCAL_OR_READ_SOURCE")
sync_uids=sorted(uid for uid,r in erp_map.items() if r.get("type")=="READONLY" and r.get("action")=="ERP-01-ACT-READ-SYNC" and r.get("gate")=="ERP-01-GATE-SYNC-READ" and r.get("permission")=="ERP-01-PERM-SYNC" and r.get("runtime_status")=="UI_LOCAL_OR_READ_SOURCE")
assert (len(connector_uids),len(finance_uids),len(sync_uids))==(7,7,6),(len(connector_uids),len(finance_uids),len(sync_uids))
for uid in connector_uids+finance_uids:
    assert missing(erp_map[uid]["runtime_owner"]),(uid,erp_map[uid]["runtime_owner"])
for uid in sync_uids:
    assert erp_map[uid]["runtime_owner"]==ERP_CONNECTOR_OWNER,(uid,erp_map[uid]["runtime_owner"])
for uid in connector_uids+finance_uids+sync_uids:
    assert missing(erp_map[uid]["operation"]),(uid,erp_map[uid]["operation"])

# Validate owner-source contract in System 06 before defining read projection ownership.
s06_before=compose(parse_controls(S06))
# Exact connector service ownership: all currently bound connector permission operations use ERPConnectorService.
connector_bound=[r for r in s06_before.values() if r.get("permission")=="ERP-01-PERM-CONNECTOR" and not missing(r.get("operation",""))]
assert connector_bound, "NO_BOUND_CONNECTOR_OPERATIONS"
assert set(r.get("runtime_owner") for r in connector_bound)=={ERP_CONNECTOR_OWNER},[(r["control"],r["operation"],r["runtime_owner"]) for r in connector_bound]
assert {"createERPConnector","updateERPConnector","validateERPConnector","validateERPMapping"}.issubset(set(r.get("operation") for r in connector_bound))

# Exact finance READ ownership: all READ_EXACT finance operations use FinanceGuardrailService.
finance_read=[r for r in s06_before.values() if r.get("permission")=="ERP-01-PERM-FINANCE" and r.get("runtime_status")=="READ_EXACT" and not missing(r.get("operation",""))]
assert set(r.get("runtime_owner") for r in finance_read)=={ERP_FINANCE_OWNER},[(r["control"],r["operation"],r["runtime_owner"]) for r in finance_read]
assert {"getERPFinanceFactPack","getERPForecast","getERPCapacityGuardrails"}.issubset(set(r.get("operation") for r in finance_read))
# Effectful export is explicitly not used as READONLY field owner evidence.
assert any(r.get("operation")=="exportProjection" and r.get("runtime_status")=="EFFECTFUL_EXACT" for r in s06_before.values())

# Exact sync READ operations use ERPConnectorService, but per-field Operation remains N/A because fields are presentation nodes.
sync_read=[r for r in s06_before.values() if r.get("permission")=="ERP-01-PERM-SYNC" and r.get("runtime_status")=="READ_EXACT" and not missing(r.get("operation",""))]
assert {"getERPFailure","getERPSyncStatus"}.issubset(set(r.get("operation") for r in sync_read))
assert set(r.get("runtime_owner") for r in sync_read)=={ERP_CONNECTOR_OWNER}

# Patch exact Runtime Owner values into ERP page and System06 canonical product rows.
owner_bindings={**{uid:ERP_CONNECTOR_OWNER for uid in connector_uids},**{uid:ERP_FINANCE_OWNER for uid in finance_uids}}
def patch_product_runtime_owner(path,bindings):
    d=Document(path);seen=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        ci=exact_index(headers,"Control UID");ri=exact_index(headers,"Runtime Owner");usi=exact_index(headers,"UID Status")
        if ci is None or ri is None or usi is None:continue
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals) or ri>=len(vals) or usi>=len(vals):continue
            uid=vals[ci]
            if uid not in bindings:continue
            assert vals[usi]=="EXISTING_SOURCE_UID",(path,uid,vals[usi])
            old=vals[ri]
            if missing(old):
                row.cells[ri].text=bindings[uid];seen[uid]+=1
            else:
                assert old==bindings[uid],(path,uid,old,bindings[uid])
    assert set(seen)==set(bindings),(path,"PATCH_MISMATCH",sorted(set(bindings)-set(seen)),dict(seen))
    d.save(path);Document(path)
    return dict(seen)

page_patch=patch_product_runtime_owner(ERP,owner_bindings)
owner_patch=patch_product_runtime_owner(S06,owner_bindings)

# Recompute targets after owner binding.
erp_after_binding=compose(parse_controls(ERP))
erp_ops=sorted(uid for uid,r in erp_after_binding.items() if is_erp_read_field_operation_na(r))
assert len(erp_ops)==20,(len(erp_ops),erp_ops)
expected_removed={("DB-01",uid,"operation") for uid in db_ops}
expected_removed|={("ERP-01",uid,"operation") for uid in erp_ops}
expected_removed|={("ERP-01",uid,"runtime_owner") for uid in connector_uids+finance_uids}
assert len(expected_removed)==42,len(expected_removed)

post_rows,post_counts=inventory(classify_post)
removed=set(pre_rows)-set(post_rows);added=set(post_rows)-set(pre_rows)
assert removed==expected_removed,{"missing":sorted(expected_removed-removed),"unexpected":sorted(removed-expected_removed)}
assert not added,sorted(added)
assert len(post_rows)==122,len(post_rows)

# Helpers.
def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.33);sec.bottom_margin=Inches(.33);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=4.5):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for rr in p.runs:rr.font.size=Pt(fs)
    return t

# System01 owner-specific DB rule.
d=Document(S01)
txt="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
assert MARK not in txt,"S01_BATCH15_ALREADY_APPLIED"
add_landscape(d)
d.add_heading("Batch 15 · DB Read-Model Operation Applicability Rule",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::SYSTEM01] ").bold=True
p.add_run("For the eight exact DB-01 controls listed below, Runtime Status=UI_LOCAL_OR_READ_SOURCE, Runtime Owner=Database Read Model / DB-01, Method/Path is absent, and Action/Gate/Permission are already explicit. These controls operate inside the DB read-model/UI context and do not own an independent registered Operation. Operation is non-applicable and MUST remain '—'. This is a Word-definition closure only and does not assert a deployed runtime route.")
add_table(d,["Target UID","Field","Runtime Owner","Runtime Status","Applicability"],[
 [uid,"Operation",DB_OWNER,"UI_LOCAL_OR_READ_SOURCE","LEGITIMATE_NA_DB_READ_MODEL_OPERATION"] for uid in db_ops
],4.5)
d.save(S01);Document(S01)

# DB page applicability ledger.
d=Document(DB)
txt="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
assert MARK not in txt,"DB_BATCH15_ALREADY_APPLIED"
add_landscape(d)
d.add_heading("Batch 15 · DB Read-Model Operation Applicability Ledger",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::PAGE::DB-01] ").bold=True
p.add_run("Applicability-only ledger. Product Control Registry rows remain unchanged; Operation stays '—'. Target UID is used so this evidence table is not a second product-control denominator.")
add_table(d,["Target UID","Field","Current Cell","Owner","Handling"],[
 [uid,"Operation","—",DB_OWNER,"DO_NOT_INVENT_REGISTERED_OPERATION"] for uid in db_ops
],4.8)
d.save(DB);Document(DB)

# System06 ERP owner-specific read projection rule.
d=Document(S06)
txt="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
assert MARK not in txt,"S06_BATCH15_ALREADY_APPLIED"
add_landscape(d)
d.add_heading("Batch 15 · ERP Owner-Specific Read Projection Closure",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::SYSTEM06] ").bold=True
p.add_run("ERP READONLY field controls are projection/presentation nodes, not independent execution entrypoints. Connector read fields are owned by ERPConnectorService because the exact Connector canonical boundary is uniformly owned by ERPConnectorService. Finance read fields are owned by FinanceGuardrailService because all exact READ_EXACT Finance operations are owned by FinanceGuardrailService; the separate EFFECTFUL exportProjection owner is excluded from this read-owner rule. Sync read fields already carry ERPConnectorService. For all twenty exact READONLY rows, field-level Operation is non-applicable and MUST remain '—'; source acquisition occurs through the owner service/read operations, not through a per-field Operation.")

add_table(d,["Read Group","Exact Action","Gate","Permission","Resolved Runtime Owner","Operation Handling"],[
 ["Connector","ERP-01-ACT-READ-CONNECTOR","ERP-01-GATE-CONNECTOR-READ","ERP-01-PERM-CONNECTOR",ERP_CONNECTOR_OWNER,"PER_FIELD_OPERATION_NA"],
 ["Finance","ERP-01-ACT-READ-FINANCE","ERP-01-GATE-FINANCE","ERP-01-PERM-FINANCE",ERP_FINANCE_OWNER,"PER_FIELD_OPERATION_NA"],
 ["Sync","ERP-01-ACT-READ-SYNC","ERP-01-GATE-SYNC-READ","ERP-01-PERM-SYNC",ERP_CONNECTOR_OWNER,"PER_FIELD_OPERATION_NA"],
],4.5)

add_table(d,["Target UID","Group","Field","Resolved Value / Applicability","Status"],[
 *[[uid,"Connector","Runtime Owner",ERP_CONNECTOR_OWNER,"BOUND_BY_BATCH_15"] for uid in connector_uids],
 *[[uid,"Finance","Runtime Owner",ERP_FINANCE_OWNER,"BOUND_BY_BATCH_15"] for uid in finance_uids],
 *[[uid,"Connector","Operation","N/A","LEGITIMATE_NA_ERP_READ_PROJECTION_FIELD_OPERATION"] for uid in connector_uids],
 *[[uid,"Finance","Operation","N/A","LEGITIMATE_NA_ERP_READ_PROJECTION_FIELD_OPERATION"] for uid in finance_uids],
 *[[uid,"Sync","Operation","N/A","LEGITIMATE_NA_ERP_READ_PROJECTION_FIELD_OPERATION"] for uid in sync_uids],
],4.0)
d.save(S06);Document(S06)

# ERP page ledger after actual owner writes.
d=Document(ERP)
txt="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
assert MARK not in txt,"ERP_BATCH15_ALREADY_APPLIED"
add_landscape(d)
d.add_heading("Batch 15 · ERP Read Projection Binding / Applicability Ledger",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::PAGE::ERP-01] ").bold=True
p.add_run("Fourteen Runtime Owner cells are bound in the existing product registry; twenty READONLY field Operation cells remain '—' by explicit non-applicability. No API route or per-field Operation is invented.")
add_table(d,["Target UID","Group","Runtime Owner","Operation","Handling"],[
 *[[uid,"Connector",ERP_CONNECTOR_OWNER,"—","OWNER_BOUND / OPERATION_NA"] for uid in connector_uids],
 *[[uid,"Finance",ERP_FINANCE_OWNER,"—","OWNER_BOUND / OPERATION_NA"] for uid in finance_uids],
 *[[uid,"Sync",ERP_CONNECTOR_OWNER,"—","EXISTING_OWNER / OPERATION_NA"] for uid in sync_uids],
],4.3)
d.save(ERP);Document(ERP)

# Verify product and owner rows after writes.
erp_final=compose(parse_controls(ERP));s06_final=compose(parse_controls(S06))
for uid in connector_uids:
    assert erp_final[uid]["runtime_owner"]==ERP_CONNECTOR_OWNER and s06_final[uid]["runtime_owner"]==ERP_CONNECTOR_OWNER
for uid in finance_uids:
    assert erp_final[uid]["runtime_owner"]==ERP_FINANCE_OWNER and s06_final[uid]["runtime_owner"]==ERP_FINANCE_OWNER
for uid in erp_ops:
    assert missing(erp_final[uid]["operation"]) and missing(s06_final[uid]["operation"]),(uid,erp_final[uid]["operation"],s06_final[uid]["operation"])
for uid in db_ops:
    assert missing(compose(parse_controls(DB))[uid]["operation"]),(uid,"DB_OPERATION_MUTATED")

# Central logic.
logic=Document(LOGIC)
txt="\n".join([p.text for p in logic.paragraphs]+[c.text for t in logic.tables for row in t.rows for c in row.cells])
assert MARK not in txt,"LOGIC_BATCH15_ALREADY_APPLIED"
add_landscape(logic)
logic.add_heading("Batch 15 · Owner-Specific UI-Local / Read-Source Resolution",level=1)
p=logic.add_paragraph();p.add_run("["+MARK+"] ").bold=True
p.add_run("Batch 15 resolves all 42 remaining UI_LOCAL_OR_READ_SOURCE gaps by owner-specific rules instead of a global assumption. DB read-model controls close Operation as N/A within their explicit read-model owner boundary. ERP READONLY projection fields close per-field Operation as N/A; Connector and Finance rows also receive exact Runtime Owner bindings derived from their canonical owner boundaries. This is design-contract closure only; it is not Runtime/Production execution evidence.")

add_table(logic,["Resolution class","Count","Result"],[
 ["Pre-batch Definition Binding Gap",164,"Current Batch-14 denominator."],
 ["DB read-model Operation N/A",8,"Owner-bound local/read-model controls; no registered per-control Operation."],
 ["ERP READONLY field Operation N/A",20,"Field-level presentation/projection nodes; no per-field Operation."],
 ["ERP Connector Runtime Owner binding",7,ERP_CONNECTOR_OWNER],
 ["ERP Finance Runtime Owner binding",7,ERP_FINANCE_OWNER],
 ["Total gaps closed",42,"28 explicit N/A + 14 exact Runtime Owner bindings."],
 ["Post-batch Definition Binding Gap",122,"121 unique-owner canonical-spec gaps + 1 preserved SYS Gate conflict."],
 ["Operation values invented",0,"All 28 target Operation cells remain '—'."],
 ["Runtime Owner values written",14,"Only existing canonical owner names are used."],
],4.6)

logic.add_heading("Owner-Specific Evidence Boundary",level=2)
add_table(logic,["Owner scope","Evidence","Accepted consequence"],[
 ["DB-01","Eight exact rows already bind Database Read Model / DB-01 and have no Method/Path.","Per-control Operation is N/A; no API/operation invented."],
 ["ERP Connector","All currently bound ERP-01-PERM-CONNECTOR operations are owned by ERPConnectorService.","Seven connector READONLY fields bind ERPConnectorService; field-level Operation N/A."],
 ["ERP Finance READ","All READ_EXACT ERP-01-PERM-FINANCE operations are owned by FinanceGuardrailService.","Seven finance READONLY fields bind FinanceGuardrailService; field-level Operation N/A."],
 ["ERP Finance Export","exportProjection is EFFECTFUL_EXACT and uses a different export owner.","Explicitly excluded from READONLY field owner inference."],
 ["ERP Sync READ","getERPFailure/getERPSyncStatus are READ_EXACT and owned by ERPConnectorService.","Existing six field owners retained; field-level Operation N/A."],
],4.3)

logic.add_heading("Batch 15 Fail-Closed Boundary",level=2)
add_table(logic,["Rule","Requirement"],[
 ["B15-01","UI_LOCAL_OR_READ_SOURCE is not globally N/A. Only the exact DB and ERP owner-specific structures defined in this batch are closed."],
 ["B15-02","ERP field-level Operation N/A does not mean the source has no read operation; it means a READONLY field is not itself an execution endpoint."],
 ["B15-03","Runtime Owner bindings may use only canonical owner names already established by the matching domain's exact operations."],
 ["B15-04","No Method/Path, API ID, or Operation is synthesized for DB or ERP field controls."],
 ["B15-05","The preserved SYS-01-BTN-NAV-OPEN Gate conflict remains fail-closed."],
],4.7)

machine={
 "marker":MARK,
 "pre_definition_gaps":164,
 "db_operation_na":8,
 "erp_operation_na":20,
 "erp_connector_runtime_owner_bound":7,
 "erp_finance_runtime_owner_bound":7,
 "gaps_closed":42,
 "post_definition_gaps":122,
 "remaining_unique_owner_remediation":121,
 "preserved_sys_gate_conflict":1,
 "operation_values_invented":0,
 "runtime_owner_values_written":14,
 "runtime_execution_claimed":False,
}
logic.add_paragraph("BATCH15_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

report={
 "machine":machine,
 "db_operation_targets":db_ops,
 "erp_connector_uids":connector_uids,
 "erp_finance_uids":finance_uids,
 "erp_sync_uids":sync_uids,
 "erp_operation_targets":erp_ops,
 "runtime_owner_bindings":owner_bindings,
 "page_patch_counts":page_patch,
 "canonical_owner_patch_counts":owner_patch,
 "removed_from_definition_gap":[{"page":p,"uid":u,"field":f} for p,u,f in sorted(removed)],
 "pre_classification_counts":dict(pre_counts),
 "post_classification_counts":dict(post_counts),
 "output_hashes":{f:blob(f) for f in [S01,S06,DB,ERP,LOGIC]},
}
Path("__batch15_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH15_REMEDIATION="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
