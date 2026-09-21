from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import hashlib,json,re

MARK="ACPOS-20260921-BATCH-03-EXHAUSTIVE-DENOMINATOR-LEDGER-V1"
TARGET="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
TARGET_SHA="61afda116af678fd293e8f5917fe130ebf72fdf8"

PAGES={
"AIAPI-01":("ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx","ef610d7f6f2bcc72890f1ba563c063f737cae60e"),
"ASSET-01":("ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","71c6c4b8b734a51db5f83b73fc76d455973b8775"),
"CORE-01":("ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","a140edcba1d76d5d38b088ec856f6c2bd4692978"),
"DB-01":("ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","741bd95a20b44252725732c0158d2439364cb83e"),
"DEV-01":("ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx","a98a7eedabcf27af1ecb7015591b2a7c85ecb8de"),
"EDIT-01":("ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","f478bbe5e80b4b26b73e131536db740724364003"),
"ERP-01":("ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx","87e3086adc0f705a8e13ac9c73dabb2b30cafa33"),
"IAM-01":("ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx","e1845f6997cd87abf4f38a24a247a30d41cd065b"),
"INFO-01":("ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","9c8af306b835c49d7216824c82f420c0c4d8e5b1"),
"KB-01":("ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx","e46ef542a5be7249f3b5f60ff308bdfd8e26bacb"),
"QA-01":("ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","dd5533f7ab0ffec491f2d6ab16c86c84cead8a7c"),
"SG-02":("ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx","dc21e0b85a0058953d127b0f9b42c9b500fd19ba"),
"SOC-01":("ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx","82112c2a3759b4d5b3ebce4e1ed184198ac0d0b2"),
"STR-01":("ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","2273aa79858e3a0b4bb7c9ff99cd192d54220bc4"),
"SYS-01":("ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx","e2be8ed3fec0eed66fa53135d5bad5b4a963ff7a"),
"VIDEO-01":("ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","5c524d869b8e255b1706fbc3faa05918d0ab9524"),
"WB-01":("ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","d262022d97d3138dae98b0c45c59bd859f3600fd"),
"ADMIN-STR-01":("ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx","b1102528b53aa1a9230f610ac386fd208df0a916"),
}

def git_blob_sha(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

assert git_blob_sha(TARGET)==TARGET_SHA,(git_blob_sha(TARGET),TARGET_SHA)
for uid,(fn,sha) in PAGES.items():
    a=git_blob_sha(fn)
    assert a==sha,(uid,a,sha)

def norm(x):
    return re.sub(r"\s+"," ",x.replace("\n"," | ").strip())

def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None

def get(vals,i):
    if i is None or i>=len(vals):return ""
    return norm(vals[i])

def extract_controls(doc):
    best={}
    refs={}
    for ti,t in enumerate(doc.tables):
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(headers,"control uid")
        if ci is None:continue
        idx={
          "uid":ci,
          "label":hfind(headers,"label","顯示名稱"),
          "action":hfind(headers,"action uid"),
          "gate":hfind(headers,"gate"),
          "permission":hfind(headers,"permission","auth resource"),
          "payload":hfind(headers,"payload","schema"),
          "operation":hfind(headers,"operation"),
          "method":hfind(headers,"method","path"),
          "runtime_owner":hfind(headers,"runtime owner"),
          "persistence_owner":hfind(headers,"persistence owner"),
          "runtime_status":hfind(headers,"runtime status"),
        }
        count=0
        for r in t.rows[1:]:
            vals=[c.text for c in r.cells]
            uid=get(vals,idx["uid"])
            if not uid or uid in {"—","-"}:continue
            if not re.search(r"[A-Za-z0-9]",uid):continue
            row={k:get(vals,v) for k,v in idx.items()}
            score=sum(1 for k,v in row.items() if k!="uid" and v not in {"","—","-"})
            if uid not in best or score>best[uid]["_score"]:
                row["_score"]=score;row["_table"]=ti+1
                best[uid]=row
            count+=1
        if count:
            refs[ti+1]={"header":" || ".join(headers),"rows":count}
    for r in best.values():r.pop("_score",None)
    return best,refs

def all_text(doc):
    return "\n".join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])

def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}

summaries=[]
page_rows={}
for page_uid,(fn,blob) in PAGES.items():
    d=Document(fn);txt=all_text(d)
    controls,refs=extract_controls(d)
    ordered=[controls[k] for k in sorted(controls)]
    serialized=json.dumps(ordered,ensure_ascii=False,sort_keys=True,separators=(",",":")).encode()
    reg_hash=hashlib.sha256(serialized).hexdigest()
    effectful=blocked=no_route=ui_local=read_exact=0
    ma=mg=mp=mo=mr=0
    source_nd=0
    for r in ordered:
        status=r["runtime_status"]
        if "UI_LOCAL_EXACT" in status:ui_local+=1
        if "READ_EXACT" in status:read_exact+=1
        if any(x in status for x in ["EFFECTFUL","OWNER_ORCHESTRATED","SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED"]):effectful+=1
        if any(x in status for x in ["BLOCKED","NOT_EXECUTED"]):blocked+=1
        if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in status:no_route+=1
        if missing(r["action"]):ma+=1
        if missing(r["gate"]):mg+=1
        if missing(r["permission"]):mp+=1
        if missing(r["operation"]) and "UI_LOCAL_EXACT" not in status and "READ_EXACT" not in status and "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" not in status:mo+=1
        if missing(r["runtime_owner"]) and any(x in status for x in ["EFFECTFUL","OWNER_ORCHESTRATED","BLOCKED","NOT_EXECUTED"]):mr+=1
        source_nd += sum(1 for v in r.values() if v=="SOURCE_NOT_DEFINED")
    summaries.append([
      page_uid,fn,blob,len(ordered),",".join("T"+str(x) for x in sorted(refs)) or "NO_CONTROL_UID_TABLE",
      reg_hash,effectful,blocked,no_route,ui_local,read_exact,ma,mg,mp,mo,mr,source_nd,
      "YES" if ("L01" in txt and "L24" in txt) else "NO",
      "YES" if ("S01" in txt and "S30" in txt) else "NO",
      "YES" if ("H01" in txt and "H12" in txt) else "NO",
      "YES" if ("TY01" in txt and "TY15" in txt) else "NO",
      "HUMAN_REVIEW" if ("Human Visual Review" in txt or "PENDING HUMAN REVIEW" in txt or "PENDING_HUMAN_REVIEW" in txt) else "NO_FLAG",
      "FALSE" if ("DESIGN_FROZEN=FALSE" in txt or "DESIGN_FROZEN = FALSE" in txt) else "NO_FALSE_FLAG",
    ])
    page_rows[page_uid]=ordered

doc=Document(TARGET)
if MARK in all_text(doc):
    raise SystemExit("ALREADY_MATERIALIZED")

s=doc.add_section(WD_SECTION.NEW_PAGE)
s.orientation=WD_ORIENT.LANDSCAPE
s.page_width,s.page_height=s.page_height,s.page_width
s.top_margin=Inches(.35);s.bottom_margin=Inches(.35);s.left_margin=Inches(.3);s.right_margin=Inches(.3)
doc.add_heading("18-Page Exhaustive Denominator / Evidence Ledger",level=1)
p=doc.add_paragraph()
p.add_run(f"[{MARK}] ").bold=True
p.add_run("This ledger is generated from the exact current 18 page Word control registries. It does not turn missing Runtime, Human Review, or source-defined fields into PASS. It establishes a reproducible denominator/evidence snapshot so subsequent remediation works only on verified gaps.")

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr()
    e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def table(headers,rows,fs=4.9):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row):c[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for pp in cell.paragraphs:
                for rr in pp.runs:rr.font.size=Pt(fs)
    return t

doc.add_heading("Ledger Contract",level=2)
table(["Rule","Requirement"],[
["Source snapshot","Each page row binds exact Word blob SHA + deterministic normalized control-registry SHA-256."],
["Denominator","Unique Control UID rows found in current tables whose header explicitly contains Control UID. Duplicate UIDs across repeated construction tables are de-duplicated by the richest source row."],
["Direct binding diagnostics","Missing Action/Gate/Permission/Operation/Runtime Owner counts are raw source diagnostics. They are not automatically defects when the runtime status makes the field legitimately N/A; remediation must classify them against current authority."],
["Runtime truth","SPEC_EXACT_RUNTIME_BLOCKED / NOT_EXECUTED and OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE remain truthful states; this ledger does not materialize Runtime."],
["Page-level L/S/H/TY","YES only when the page itself contains both boundary IDs for that audit family. NO means formal page-level materialization remains to be cross-walked; it does not by itself erase existing domain content."],
["Human/Freeze","Human review and DESIGN_FROZEN flags are evidence boundaries and cannot be auto-approved by this ledger."],
],5.5)

doc.add_heading("18-Page Denominator Summary",level=2)
table([
"Page","File","Blob SHA","Unique Controls","Source Tables","Registry SHA256","Effectful-like","Runtime blocked","No-public-route","UI local","Read exact",
"Missing Action","Missing Gate","Missing Permission","Missing Operation*","Missing Runtime Owner*","SOURCE_NOT_DEFINED","L01-L24","S01-S30","H01-H12","TY01-TY15","Visual Review","Design Frozen"
],summaries,3.7)

totals={
"pages":len(summaries),
"unique_controls":sum(int(r[3]) for r in summaries),
"effectful":sum(int(r[6]) for r in summaries),
"blocked":sum(int(r[7]) for r in summaries),
"no_route":sum(int(r[8]) for r in summaries),
"missing_action":sum(int(r[11]) for r in summaries),
"missing_gate":sum(int(r[12]) for r in summaries),
"missing_permission":sum(int(r[13]) for r in summaries),
"missing_operation":sum(int(r[14]) for r in summaries),
"missing_runtime_owner":sum(int(r[15]) for r in summaries),
"source_not_defined":sum(int(r[16]) for r in summaries),
}
doc.add_heading("Fresh Denominator Totals / Next Remediation Boundary",level=2)
table(["Metric","Value","Interpretation"],[
["Pages",totals["pages"],"Fixed by current physical page package for this snapshot; future runs must re-resolve."],
["Unique Control UIDs",totals["unique_controls"],"Current source-grounded control denominator from explicit Control UID tables."],
["Effectful-like rows",totals["effectful"],"Rows requiring runtime/owner truth or an explicit owner-orchestrated boundary."],
["Runtime blocked / not executed rows",totals["blocked"],"Must remain blocked until fresh implementation evidence closes them."],
["Owner-orchestrated no-public-route rows",totals["no_route"],"Valid only with exact owner/local-command contract; no invented public API."],
["Raw missing Action",totals["missing_action"],"Classify against control type; do not auto-fill."],
["Raw missing Gate",totals["missing_gate"],"Classify against control type; do not auto-fill."],
["Raw missing Permission",totals["missing_permission"],"Classify against control type; do not auto-fill."],
["Raw missing Operation*",totals["missing_operation"],"*Excludes UI_LOCAL_EXACT, READ_EXACT and owner-orchestrated no-public-route rows from missing-operation count."],
["Raw missing Runtime Owner*",totals["missing_runtime_owner"],"*Only effectful/owner-orchestrated/blocked-like rows counted."],
["SOURCE_NOT_DEFINED cells",totals["source_not_defined"],"Must be resolved from canonical source or remain decision-required; never guessed."],
["Formal page-level L/S/H/TY materialization","See 18-page summary","Next batch must cross-walk verified source evidence; a generic global rule alone is not page-level exhaustive closure."],
],5.2)

# Persist a compact machine-readable payload in plain text for deterministic audit parsing.
payload={"marker":MARK,"totals":totals,"pages":[]}
for r in summaries:
    payload["pages"].append({
      "page_uid":r[0],"file":r[1],"blob_sha":r[2],"unique_controls":r[3],"source_tables":r[4],"control_registry_sha256":r[5],
      "effectful_like":r[6],"runtime_blocked":r[7],"no_public_route":r[8],"missing_action":r[11],"missing_gate":r[12],
      "missing_permission":r[13],"missing_operation":r[14],"missing_runtime_owner":r[15],"source_not_defined":r[16],
      "L":r[17],"S":r[18],"H":r[19],"TY":r[20],"visual_review":r[21],"design_frozen":r[22]
    })
doc.add_paragraph("MACHINE_LEDGER_JSON="+json.dumps(payload,ensure_ascii=False,sort_keys=True,separators=(",",":")))

doc.save(TARGET)
Document(TARGET)
print(json.dumps(payload,ensure_ascii=False,indent=2))
