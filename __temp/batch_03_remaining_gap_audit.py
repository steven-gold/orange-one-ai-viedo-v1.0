from pathlib import Path
from docx import Document
import json,re,hashlib

PAGES=[
"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
]
assert len(PAGES)==18

def all_text(d):
    return "\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells])

def sha1blob(p):
    b=Path(p).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for fn in PAGES:
    d=Document(fn); t=all_text(d)
    # Detect likely control registry rows by first cell UID-like patterns.
    controls=set()
    effectful=0
    blocked=0
    no_route=0
    source_not_defined=t.count("SOURCE_NOT_DEFINED")
    for tbl in d.tables:
        for row in tbl.rows:
            vals=[c.text.strip() for c in row.cells]
            if not vals: continue
            first=vals[0]
            if re.match(r"^(?:CTRL-|[A-Z0-9]+-01-(?:BTN|INP|TGL|ACT|CTL|CONTROL)-)",first):
                controls.add(first)
                joined=" || ".join(vals)
                if any(x in joined for x in ["EFFECTFUL_EXACT","OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE","SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED"]):
                    effectful+=1
                if any(x in joined for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):
                    blocked+=1
                if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in joined:
                    no_route+=1
    result={
      "file":fn,
      "blob_sha":sha1blob(fn),
      "control_rows_detected":len(controls),
      "effectful_like_rows":effectful,
      "runtime_blocked_rows":blocked,
      "owner_orchestrated_no_public_route_rows":no_route,
      "source_not_defined_mentions":source_not_defined,
      "has_L01_L24":("L01" in t and "L24" in t),
      "has_S01_S30":("S01" in t and "S30" in t),
      "has_H01_H12":("H01" in t and "H12" in t),
      "has_TY01_TY15":("TY01" in t and "TY15" in t),
      "has_function_chain_terms":all(x in t for x in ["Trigger","Validation","Operation","Owner","Failure","Recovery"]),
      "has_runtime_claim_guard":("NOT_EXECUTED" in t or "IMPLEMENTATION_REQUIRED" in t or "SPEC_EXACT_RUNTIME_BLOCKED" in t),
      "has_human_visual_review":("Human Visual Review" in t or "PENDING HUMAN REVIEW" in t or "PENDING_HUMAN_REVIEW" in t),
      "design_frozen_false":("DESIGN_FROZEN=FALSE" in t or "DESIGN_FROZEN = FALSE" in t),
      "batch01_marker":"ACPOS-20260921-BATCH-01-FOUNDATION-CLOSURE-V1" in t,
      "batch02_marker":"ACPOS-20260921-BATCH-02-CONSUMER-CONTRACT-CLOSURE-V1" in t,
    }
    print("AUDIT_PAGE "+json.dumps(result,ensure_ascii=False))
