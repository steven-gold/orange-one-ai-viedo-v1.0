from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-61-EDIT-ORCHESTRATION-LEGACY-CLEANUP-V1"
BASE_HEAD="bfa065ca445959d9d872cbefdf6c35c00379a106"
D02="02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
EDIT="ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
D02:"a90b38cca7ab723edce372b19aa7ead8d5ef4aba",
S05:"6743dc331393ad7c09da21f8e89981cb1e5b599d",
EDIT:"b7ceae3a930fec4225f0d75cd746679919c192e9",
LOGIC:"95cd3e10a170729ce8c4a37fde2cdecbaa402e4f",
}
TARGET="EDIT-01-BTN-FLOW-START"
NEW={
"operation":"dispatchEditingFlowStartContinue",
"method_path":"NO_PUBLIC_API_BY_AUTHORITY",
"payload_schema":"EditingFlowStartContinueDispatchContext",
"persistence_owner":"DELEGATED_TO_SELECTED_EDIT_PORT_PERSISTENCE",
}
ACTION="EDIT-01-ACT-FLOW-START-CONTINUE"
GATE="EDIT-01-GATE-CONTEXT-INTEGRITY"
PORT_CREATE="EDIT-01-PORT-EDIT-RUN-CREATE"
PORT_ASSEMBLY="EDIT-01-PORT-ASSEMBLY-COMPLETE"
OLD_COMPOSITES=[
"createEditingRuntimeRun | completeAssembly",
"POST /v1/editing-runtime-runs | POST /v1/editing-runtime-runs/{runId}/assembly",
"CreateEditingRuntimeRunRequest | CompleteEditingAssemblyRequest",
]
LEGACY_RE=re.compile(r"(?<![A-Za-z0-9])R(?:9(?:\.0\.1)?|5)(?![A-Za-z0-9])",re.I)
AMBIG_LOCATOR_RE=re.compile(r"\bt\d+/r(?:5|9)\b",re.I)

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def blob(path):
    b=Path(path).read_bytes();return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()
def all_text(path):
    d=Document(path);xs=[norm(p.text) for p in d.paragraphs if norm(p.text)]
    for t in d.tables:
        for row in t.rows:
            xs.extend(norm(c.text) for c in row.cells if norm(c.text))
    return "\n".join(xs)
def replace_in_cell(cell,old,new):
    if old not in cell.text:return 0
    # Preserve cell paragraph structure conservatively, replacing text in runs when possible.
    hits=0
    for p in cell.paragraphs:
        for run in p.runs:
            if old in run.text:
                hits+=run.text.count(old);run.text=run.text.replace(old,new)
    if hits==0 and old in cell.text:
        txt=cell.text.replace(old,new);cell.text=txt;hits=1
    return hits

for f,s in EXPECTED_SHA.items():
    got=blob(f);assert got==s,(f,got,s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","-c","core.quotePath=false","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

# 1) Remove R9/R5 legacy naming and ambiguous locator spellings.
cleanup={}
d=Document(D02);hits=0
for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            hits+=replace_in_cell(cell,"舊 R9/歷史 master","歷史/舊版 master")
assert hits==1,("D02_LEGACY_HITS",hits)
d.save(D02);Document(D02);cleanup[D02]={"legacy_name_replacements":hits}

d=Document(LOGIC);rep={}
locator_map={
"t26/r4":"table 26 / row 4","t73/r5":"table 73 / row 5",
"t26/r5":"table 26 / row 5","t73/r6":"table 73 / row 6",
"t26/r9":"table 26 / row 9","t73/r7":"table 73 / row 7",
"t26/r3":"table 26 / row 3","t73/r9":"table 73 / row 9",
}
for old,new in locator_map.items():
    n=0
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:n+=replace_in_cell(cell,old,new)
    if n:rep[old]={"new":new,"count":n}
assert sum(x["count"] for x in rep.values())==8,rep
d.save(LOGIC);Document(LOGIC);cleanup[LOGIC]={"locator_replacements":rep}

# 2) Patch exact target control row in EDIT page + System05 owner.
def patch_target_row(path):
    d=Document(path);patched=0;details=[]
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact(h,"Control UID")
        if ci is None:continue
        oi=hfind(h,"operation");mi=hfind(h,"method / path","method","path");pi=hfind(h,"payload / schema","payload","schema")
        pei=hfind(h,"persistence owner");ai=hfind(h,"action uid");gi=hfind(h,"gate uid","gate");permi=hfind(h,"permission","auth resource")
        rti=hfind(h,"runtime status");roi=hfind(h,"runtime owner")
        if None in (oi,mi,pi,pei,ai,gi,permi,rti,roi):continue
        for rno,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals) or vals[ci]!=TARGET:continue
            assert vals[ai]==ACTION,(path,"ACTION",vals[ai])
            assert vals[gi]==GATE,(path,"GATE",vals[gi])
            assert vals[permi]=="EDITING_USE",(path,"UI_PERMISSION",vals[permi])
            assert vals[roi]=="EDITING/ORCHESTRATION",(path,"RUNTIME_OWNER",vals[roi])
            assert vals[rti]=="EFFECTFUL_EXACT",(path,"STATUS",vals[rti])
            assert vals[oi]==OLD_COMPOSITES[0],(path,"OLD_OP",vals[oi])
            assert vals[mi]==OLD_COMPOSITES[1],(path,"OLD_METHOD",vals[mi])
            assert vals[pi]==OLD_COMPOSITES[2],(path,"OLD_PAYLOAD",vals[pi])
            row.cells[oi].text=NEW["operation"]
            row.cells[mi].text=NEW["method_path"]
            row.cells[pi].text=NEW["payload_schema"]
            row.cells[pei].text=NEW["persistence_owner"]
            patched+=1;details.append({"table":ti,"row":rno})
    assert patched==1,(path,"TARGET_PATCH_COUNT",patched,details)
    d.save(path);Document(path)
    return details

target_patch={S05:patch_target_row(S05),EDIT:patch_target_row(EDIT)}

# 3) Replace old Batch56 composite ledger row with one orchestration identity.
def patch_ledger(path):
    d=Document(path);hits=0;where=[]
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact(h,"Control UID")
        if ci is None:continue
        oi=hfind(h,"operation");mi=hfind(h,"method / path","method","path");pi=hfind(h,"payload / schema","payload","schema");pei=hfind(h,"persistence owner");si=hfind(h,"semantics")
        for rno,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals) or vals[ci]!=TARGET:continue
            # Skip main exact row; it has Runtime Status headers and is already patched.
            if hfind(h,"runtime status") is not None:continue
            if oi is not None and oi<len(vals) and vals[oi]==OLD_COMPOSITES[0]:
                row.cells[oi].text=NEW["operation"]
                if mi is not None:row.cells[mi].text=NEW["method_path"]
                if pi is not None:row.cells[pi].text=NEW["payload_schema"]
                if pei is not None:row.cells[pei].text=NEW["persistence_owner"]
                if si is not None:
                    row.cells[si].text=("State-dependent orchestration only. Exactly one atomic port is selected per invocation: "
                    "no valid editing run → EDIT-01-PORT-EDIT-RUN-CREATE; valid bound run in Stage 01 Assembly with no assembly completion → "
                    "EDIT-01-PORT-ASSEMBLY-COMPLETE; all other/stale/blocked states → no atomic mutation. "
                    "EDITING_USE controls UI availability, while the selected effectful port additionally requires editing.task.execute. "
                    "The orchestration itself has no public API and delegates persistence to the selected port.")
                hits+=1;where.append({"table":ti,"row":rno})
    assert hits==1,(path,"LEDGER_PATCH_COUNT",hits,where)
    d.save(path);Document(path)
    return where
ledger_patch={S05:patch_ledger(S05),EDIT:patch_ledger(EDIT)}

# 4) Append explicit bidirectional orchestration contract to both owner and page.
def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.32);sec.bottom_margin=Inches(.32);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=3.6):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for rr in rows:
        cs=t.add_row().cells
        for i,v in enumerate(rr):cs[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
    return t

branches=[
["A · Bootstrap / no valid run",
 "No valid editing runtime run is bound to the current Project/Topic/Task + Locked Blueprint + Production Package + Input Manifest fingerprint.",
 PORT_CREATE,"createEditingRuntimeRun","POST /v1/editing-runtime-runs","CreateEditingRuntimeRunRequest","editing.task.execute",
 "acpos_runtime.editing_runtime_runs_runtime",
 "Bind returned runId to the exact input fingerprint/source versions; remain in Stage 01 Assembly. MUST NOT also call completeAssembly in the same dispatch."],
["B · Continue Stage 01 Assembly",
 "A valid runId is already bound to the same exact input fingerprint/source versions; current stage is EDIT-01-STAGE-01-ASSEMBLY; assembly completion for the bound stage/version is absent.",
 PORT_ASSEMBLY,"completeAssembly","POST /v1/editing-runtime-runs/{runId}/assembly","CompleteEditingAssemblyRequest","editing.task.execute",
 "acpos_runtime.editing_runtime_steps_runtime; public.editing_timelines; editing run state",
 "Persist the assembly step/version/evidence. Transition beyond Assembly still requires Evaluation PASS + explicit Stage Confirm; this dispatch MUST NOT auto-confirm or jump to AUDIO."],
["C · Not eligible / stale / later state",
 "Run exists but source fingerprint/version is stale or mismatched, assembly is already completed, current stage is not Stage 01 Assembly, or a blocking gate/state applies.",
 "NONE","NONE","NONE","NONE","N/A","NONE",
 "Fail closed or render the correct current-stage control. No create/complete mutation is allowed from this action."]
]
bridges=[
["Control","EDIT-01-BTN-FLOW-START","UI control; label 開始 / 繼續","Must resolve one legal branch only."],
["Action",ACTION,"Single orchestration action identity","One action → one orchestration operation."],
["Operation",NEW["operation"],"State-dependent dispatcher; not an atomic persistence API","Exactly-one port selection; no public route."],
["Method / Path",NEW["method_path"],"No third public endpoint is created","Atomic ports retain their existing routes."],
["Payload",NEW["payload_schema"],"Contains dispatch context only: exact project/topic/task/run/input-fingerprint/source-version/stage refs as applicable","Does not replace either atomic request schema."],
["UI Permission","EDITING_USE","Page/button availability prerequisite","Not sufficient for effectful port execution."],
["Runtime Permission","editing.task.execute","Required by both atomic effectful ports","Selected port MUST authorize this permission server-side."],
["Gate",GATE,"Requires Project/Topic/Task + Locked Blueprint + Production Package + Input Manifest + exact fingerprint","Failure disables stage execution; no port call."],
["Persistence",NEW["persistence_owner"],"Orchestrator does not own business persistence","Selected port owns the actual state mutation/evidence."],
]
reverse=[
[PORT_CREATE,"createEditingRuntimeRun","POST /v1/editing-runtime-runs","Consumer for this UI flow: "+NEW["operation"],"Branch A only"],
[PORT_ASSEMBLY,"completeAssembly","POST /v1/editing-runtime-runs/{runId}/assembly","Consumer for this UI flow: "+NEW["operation"],"Branch B only"],
]
for path,title,prefix in [
(S05,"Batch 61 · EDIT Start/Continue State-Dependent Orchestration Contract","SYSTEM05_CANONICAL_OWNER"),
(EDIT,"Batch 61 · EDIT-01 Start/Continue Bidirectional Binding Contract","PAGE_EDIT-01"),
]:
    d=Document(path);add_landscape(d);d.add_heading(title,level=1)
    p=d.add_paragraph();p.add_run(f"[{MARK}::{prefix}] ").bold=True
    p.add_run("The UI action is an orchestration dispatcher, not two atomic operations concatenated into one field. "
              "Each invocation selects exactly one existing atomic Port or performs no mutation. No third public API is introduced. "
              "UI authorization and runtime execution authorization are both required at their own layers.")
    add_table(d,["Layer","Identity / Value","Role","Invariant"],bridges,3.6)
    d.add_heading("Dispatch branches",level=2)
    add_table(d,["Branch","Exact eligibility condition","Selected Port","Atomic Operation","Method / Path","Atomic Payload","Runtime Permission","Atomic Persistence / Effect","Required result"],branches,3.05)
    d.add_heading("Reverse consumer binding",level=2)
    add_table(d,["Port UID","Operation","Method / Path","UI-flow consumer","Allowed branch"],reverse,3.8)
    p=d.add_paragraph()
    p.add_run("Hard invariants: ").bold=True
    p.add_run("never call both atomic ports in one dispatch; stale/mismatched fingerprint or source version fails closed; "
              "create-run does not imply assembly completion; assembly completion does not imply Stage Confirm; "
              "Stage 01 → AUDIO still requires Evaluation PASS + explicit human Stage Confirm; runtime/Production execution is not claimed by this Word contract.")
    d.save(path);Document(path)

# 5) Persist package-level resolution ledger in System Logic.
d=Document(LOGIC);add_landscape(d);d.add_heading("Batch 61 · EDIT Orchestration + Legacy Token Cleanup",level=1)
p=d.add_paragraph();p.add_run(f"[{MARK}] ").bold=True
p.add_run("Resolved the only compound EFFECTFUL_EXACT control found by the full-package audit. "
          "EDIT-01-BTN-FLOW-START is now a single state-dependent orchestration operation that delegates to exactly one existing atomic Port. "
          "The control retains EDITING_USE as UI availability, while selected effectful Ports require editing.task.execute. "
          "Legacy revision-token wording and ambiguous table/row locator strings identified by Batch 59 were rewritten without changing their referents.")
add_table(d,["Check","Before","After","Resolution"],[
["EDIT target Operation","createEditingRuntimeRun | completeAssembly",NEW["operation"],"Single orchestration identity"],
["EDIT target Method / Path",OLD_COMPOSITES[1],NEW["method_path"],"No third public API; atomic Port routes remain"],
["EDIT target Payload",OLD_COMPOSITES[2],NEW["payload_schema"],"Dispatch context only; atomic schemas remain at Ports"],
["EDIT target Persistence","3-table union",NEW["persistence_owner"],"Persistence delegated to selected atomic Port"],
["Port selection","Implicit / ambiguous","Exactly-one state-dependent branch","No double invocation"],
["Permission bridge","EDITING_USE in control; editing.task.execute only on Ports","Both layers explicitly required","UI entitlement ≠ runtime execute authorization"],
["Legacy revision token","Explicit old revision label present","Generic historical/legacy wording","No obsolete revision identifier remains"],
["Ambiguous locator shorthand","tNN/rNN strings containing r5/r9","table NN / row NN","No revision-token false positive"],
],3.8)
machine={
"marker":MARK,"base_head":BASE_HEAD,"changed_docs":4,
"legacy_revision_name_replacements":1,"ambiguous_locator_replacements":8,
"target_control_rows_patched":2,"historical_ledger_rows_patched":2,
"orchestration_operation":NEW["operation"],"public_api_created":False,
"atomic_ports_reused":2,"dispatch_exactly_one_port":True,
"ui_permission":"EDITING_USE","runtime_port_permission":"editing.task.execute",
"runtime_execution_claimed":False,
}
d.add_paragraph("BATCH61_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
d.save(LOGIC);Document(LOGIC)

# 6) Exact post-mutation assertions across all 31 docs.
texts={p.name:all_text(p) for p in Path(".").glob("*.docx")}
all_join="\n".join(texts.values())
legacy_hits=LEGACY_RE.findall(all_join)
ambiguous_hits=AMBIG_LOCATOR_RE.findall(all_join)
assert not legacy_hits,legacy_hits
assert not ambiguous_hits,ambiguous_hits

for path in [S05,EDIT]:
    txt=texts[path]
    for old in OLD_COMPOSITES:assert old not in txt,(path,"OLD_COMPOSITE_REMAINS",old)
    assert NEW["operation"] in txt and NEW["method_path"] in txt and NEW["payload_schema"] in txt
    assert PORT_CREATE in txt and PORT_ASSEMBLY in txt
    assert "editing.task.execute" in txt and "EDITING_USE" in txt

# Verify main exact row atomized.
def exact_target(path):
    d=Document(path);hits=[]
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact(h,"Control UID");rsi=hfind(h,"runtime status")
        if ci is None or rsi is None:continue
        oi=hfind(h,"operation");mi=hfind(h,"method / path","method","path");pi=hfind(h,"payload / schema","payload","schema");pei=hfind(h,"persistence owner")
        ai=hfind(h,"action uid");gi=hfind(h,"gate uid","gate");permi=hfind(h,"permission","auth resource")
        for rno,row in enumerate(t.rows[1:],2):
            v=[norm(c.text) for c in row.cells]
            if ci<len(v) and v[ci]==TARGET and v[rsi]=="EFFECTFUL_EXACT":
                hits.append({"table":ti,"row":rno,"operation":v[oi],"method":v[mi],"payload":v[pi],"persistence":v[pei],"action":v[ai],"gate":v[gi],"permission":v[permi]})
    assert len(hits)==1,(path,hits)
    x=hits[0]
    assert x["operation"]==NEW["operation"] and x["method"]==NEW["method_path"] and x["payload"]==NEW["payload_schema"] and x["persistence"]==NEW["persistence_owner"],(path,x)
    assert " | " not in x["operation"] and " | " not in x["method"] and " | " not in x["payload"],(path,x)
    return x

post_target={S05:exact_target(S05),EDIT:exact_target(EDIT)}

changed=[D02,S05,EDIT,LOGIC]
report={"machine":machine,"cleanup":cleanup,"target_patch":target_patch,"ledger_patch":ledger_patch,
        "post_target":post_target,"changed_docs":changed,"output_blob_sha":{f:blob(Path(f)) for f in changed}}
Path("__batch61_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH61="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
