from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-38-SYSTEM03-PERSISTENCE-REMEDIATION-V1"
BASE_HEAD="ab8f013ee54404217953d87a3f7f51f54e88b83d"
S03="03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
AIAPI="ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S03:"491b86d7c6c6feda0ee77f94dd8cd0487ba65817",
AIAPI:"8b1a176ec4def4fcdff15be24385357c7b41599b",
LOGIC:"acd89c38c57751bcb22d02b849c9c8c1e35556d4",
}
PAGES={
"AIAPI-01":AIAPI,
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]
BINDINGS={
"AIAPI-01-BTN-CONFIGURE-GOVERNED-RESOURCE":{
 "operation":"configureGovernedResource",
 "persistence_owner":"Resolved Governed Resource Canonical Owner / Governance Audit",
 "basis":"The command configures an exact governed resource. AIAPI/IAM runtime must not steal domain persistence ownership; the exact resource owner persists state while governance audit preserves actor/correlation/version evidence."
},
"AIAPI-01-BTN-APPROVE-GOVERNED-RESOURCE":{
 "operation":"approveGovernedResource",
 "persistence_owner":"Resolved Governed Resource Canonical Owner / Approval Audit",
 "basis":"Approval writes the governed resource's approved/version state under its canonical owner and preserves separate approver/rationale/audit evidence."
},
"AIAPI-01-BTN-RUN-SANDBOX-TEST":{
 "operation":"runSandboxTest",
 "persistence_owner":"ACPOS-AI-03 Sandbox Evidence / Audit Lineage",
 "basis":"Sandbox may persist isolated test evidence and lineage only; Production secret/output writes and Production-pass claims remain forbidden."
},
"CTRL-ADMIN-AIAPI-09-ACT-02-ACT-KILL-SWITCH":{
 "operation":"setKillSwitch",
 "persistence_owner":"ACPOS-AI-03 Operational Control State / Audit Lineage",
 "basis":"Kill switch is a server-side operational gate that dominates execution; state transition and authorized reason/correlation/audit evidence must remain reconstructable."
},
}

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out
def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():
    got=blob(f);assert got==s,(f,got,s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

pre_page=compose(parse_controls(AIAPI));pre_owner=compose(parse_controls(S03))
for uid,s in BINDINGS.items():
    for m,label in [(pre_page,"PAGE"),(pre_owner,"OWNER")]:
        assert uid in m,(label,uid,"MISSING")
        r=m[uid]
        assert r["operation"]==s["operation"],(label,uid,r["operation"],s["operation"])
        assert r["runtime_status"]=="EFFECTFUL_EXACT",(label,uid,r["runtime_status"])
        assert missing(r["persistence_owner"]),(label,uid,r["persistence_owner"])

def patch(path):
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact_index(h,"Control UID");pi=hfind(h,"persistence owner")
        if ci is None or pi is None:continue
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals):continue
            uid=vals[ci]
            if uid not in BINDINGS:continue
            old=norm(row.cells[pi].text);new=BINDINGS[uid]["persistence_owner"]
            if old==new:hits[uid]+=1;continue
            assert missing(old),(path,uid,old,new)
            row.cells[pi].text=new;hits[uid]+=1
    assert set(hits)==set(BINDINGS),(path,dict(hits),set(BINDINGS)-set(hits))
    d.save(path);Document(path)
    return dict(hits)

patches={"SYSTEM03_OWNER":patch(S03),"AIAPI_PAGE":patch(AIAPI)}

def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.33);sec.bottom_margin=Inches(.33);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=4.1):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for rr in rows:
        cells=t.add_row().cells
        for i,v in enumerate(rr):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
    return t

for path,title,prefix in [
(S03,"Batch 38 · System 03 Persistence Ownership Contract","SYSTEM03"),
(AIAPI,"Batch 38 · AIAPI Persistence Binding Ledger","PAGE::AIAPI-01"),
]:
    d=Document(path);add_landscape(d);d.add_heading(title,level=1)
    p=d.add_paragraph();p.add_run(f"[{MARK}::{prefix}] ").bold=True
    p.add_run("These four Persistence Owner bindings preserve existing Method/Path, Payload/Schema, Runtime Owner and Runtime Status. Governed-resource persistence remains with the exact resolved resource owner; sandbox and kill-switch evidence remain AI-03 operational/audit concerns. No Production execution claim is made.")
    rr=[[uid,s["operation"],s["persistence_owner"],s["basis"]] for uid,s in BINDINGS.items()]
    add_table(d,["Control UID","Operation","Persistence Owner","Required Semantics"],rr,3.8)
    d.save(path);Document(path)

post_page=compose(parse_controls(AIAPI));post_owner=compose(parse_controls(S03))
for uid,s in BINDINGS.items():
    assert post_page[uid]["persistence_owner"]==s["persistence_owner"],(uid,post_page[uid]["persistence_owner"])
    assert post_owner[uid]["persistence_owner"]==s["persistence_owner"],(uid,post_owner[uid]["persistence_owner"])

all_maps={p:compose(parse_controls(fn)) for p,fn in PAGES.items()}
remaining=[]
for page,m in all_maps.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if st in {"EFFECTFUL_EXACT","READ_EXACT"} and not missing(op) and missing(r.get("method_path","")):
            remaining.append(("method_path",page,uid))
        if st=="EFFECTFUL_EXACT" and not missing(op) and missing(r.get("payload_schema","")):
            remaining.append(("payload_schema",page,uid))
        if st=="EFFECTFUL_EXACT" and not missing(op) and missing(r.get("persistence_owner","")):
            remaining.append(("persistence_owner",page,uid))
rc=collections.Counter(f for f,_,_ in remaining)
assert len(remaining)==286,(len(remaining),rc)
assert rc=={"method_path":63,"payload_schema":119,"persistence_owner":104},rc

logic=Document(LOGIC);add_landscape(logic);logic.add_heading("Batch 38 · System 03 Persistence Closure",level=1)
p=logic.add_paragraph();p.add_run(f"[{MARK}] ").bold=True
p.add_run("System 03 closes four persistence-owner cells without inventing physical DB table names. Governed-resource operations delegate durable state to the exact resolved resource owner with governance audit; sandbox and kill-switch operations persist only their operational/evidence lineage under AI-03. Runtime completion remains unclaimed.")
add_table(logic,["Metric","Value","Result"],[
["Pre contract denominator",290,"Method 63 + Payload 119 + Persistence 108"],
["System 03 Persistence Owner cells closed",4,"4 exact semantic owner bindings"],
["Post contract denominator",286,"Method 63 + Payload 119 + Persistence 104"],
["Physical DB table names invented",0,"None"],
["Runtime execution claimed","False","Design-contract closure only"],
],4.2)
machine={"marker":MARK,"base_head":BASE_HEAD,"pre_contract_cells":290,"system03_persistence_closed":4,
"post_contract_cells":286,"remaining_method_path":63,"remaining_payload_schema":119,"remaining_persistence_owner":104,
"physical_db_tables_invented":0,"runtime_execution_claimed":False}
logic.add_paragraph("BATCH38_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

changed=[S03,AIAPI,LOGIC]
report={"machine":machine,"bindings":BINDINGS,"patches":patches,"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed}}
Path("__batch38_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH38="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
