from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260921-BATCH-17-REMAINING-DEFINITION-CLOSURE-V1"
BASE_HEAD="98171c9989147690db2acd75228923d91bb2d92b"
S01="01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx"
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S06="06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S09="09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
DEV="ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx"
EDIT="ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
KB="ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx"
SYS="ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"

PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":DEV,
"EDIT-01":EDIT,
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":KB,
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":SYS,
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]
EXPECTED_SHA={
S01:"e016ed9ce67c5bcf8210c64751322ba6af16f57e",
S05:"26008b9e2244316a6da2341e59eca25c79783ce7",
S06:"ef6e424423c3ea282a0d625c953a0f608245637a",
S09:"0acb96e6a71f9de5368a65c1a68480d3e08a79fd",
DEV:"3f6fdc5b0c71257850e6b504b6ccafee1eeaa41e",
EDIT:"7fe89816d0678d9b7f91d147b9571ef89d77ae6d",
KB:"e46ef542a5be7249f3b5f60ff308bdfd8e26bacb",
SYS:"77b5ba6c34e02b96d957587b1852a231612ffa2c",
LOGIC:"13e896b5d5280df9e4d7414529dfd6d8cb793702",
}

DEV_UIDS=[f"DEV-01-BTN-STAGE-{i}" for i in range(1,6)]
EDIT_UIDS=["EDIT-01-PNL-API-CANDIDATE","EDIT-01-PNL-FINAL-PREVIEW"]
KB_UIDS=[
"KB-01-CTL-VIEW-EXPERIENCE","KB-01-CTL-VIEW-OVERVIEW","KB-01-CTL-VIEW-REVIEW",
"KB-01-CTL-VIEW-SEARCH","KB-01-CTL-VIEW-SOURCE",
]
SYS_GATE_BINDINGS={
"SYS-01-BTN-CR-CREATE":"SYS-01-GATE-CR-CREATE",
"SYS-01-BTN-NAV-OPEN":"SOURCE_PAGE_GATE",
"SYS-01-BTN-SANDBOX-TEST":"SYS-01-GATE-SANDBOX-TEST",
}

def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()
for f,s in EXPECTED_SHA.items():
    got=blob(f);assert got==s,(f,got,s)
assert subprocess.check_output(["git","rev-parse","HEAD"],text=True).strip()!=BASE_HEAD
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""
assert len(list(Path(".").glob("*.docx")))==31

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def display_only_type(t):
    u=(t or "").upper()
    return u in {"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"} or any(x in u for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={
          "control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status"),
        }
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"table":ti+1,"row":ri,"headers":h});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out

# Existing closure rules from Batch 13-16.
def is_wb_na(r): return missing(r.get("action","")) and r.get("type")=="SECTION_OPEN" and r.get("gate")=="PAGE_READ" and r.get("permission")=="workspace.dashboard.view" and r.get("payload_schema")=="N/A_READ_PROJECTION" and r.get("operation")=="getDashboardReadModel" and r.get("method_path")=="GET /v1/dashboard/read-model" and r.get("runtime_owner")=="DASHBOARD_READ_MODEL" and r.get("runtime_status")=="READ_EXACT"
def is_aiapi_na(r): return missing(r.get("action","")) and r.get("type")=="BUTTON_OR_ROW_ACTION" and r.get("payload_schema")=="AIAPI Page Operation-specific Form / Provider Profile Field Contract" and not missing(r.get("permission","")) and not missing(r.get("operation","")) and not missing(r.get("method_path","")) and not missing(r.get("runtime_owner","")) and r.get("runtime_status") in {"EFFECTFUL_EXACT","READ_EXACT"}
def is_edit_local_op_na(r): return missing(r.get("operation","")) and r.get("runtime_status")=="LOCAL_WORKING_DRAFT_EXACT" and r.get("method_path")=="NO_PUBLIC_API_BY_AUTHORITY" and not missing(r.get("action","")) and not missing(r.get("gate","")) and not missing(r.get("permission","")) and not missing(r.get("runtime_owner",""))
def is_iam_op_na(r): return missing(r.get("operation","")) and r.get("runtime_status")=="ORCHESTRATES_EXISTING_EXACT_OPERATIONS" and not missing(r.get("action","")) and not missing(r.get("gate","")) and not missing(r.get("permission",""))
def is_iam_owner_na(r): return missing(r.get("runtime_owner","")) and r.get("runtime_status")=="ORCHESTRATES_EXISTING_EXACT_OPERATIONS" and not missing(r.get("action","")) and not missing(r.get("gate","")) and not missing(r.get("permission",""))
DB_OWNER="Database Read Model / DB-01";ERP_CONNECTOR_OWNER="ERPConnectorService";ERP_FINANCE_OWNER="FinanceGuardrailService"
ERP_GROUPS={
 "ERP-01-ACT-READ-CONNECTOR":("ERP-01-GATE-CONNECTOR-READ","ERP-01-PERM-CONNECTOR",ERP_CONNECTOR_OWNER),
 "ERP-01-ACT-READ-FINANCE":("ERP-01-GATE-FINANCE","ERP-01-PERM-FINANCE",ERP_FINANCE_OWNER),
 "ERP-01-ACT-READ-SYNC":("ERP-01-GATE-SYNC-READ","ERP-01-PERM-SYNC",ERP_CONNECTOR_OWNER),
}
def is_db_read_model_operation_na(r): return missing(r.get("operation","")) and missing(r.get("method_path","")) and r.get("runtime_status")=="UI_LOCAL_OR_READ_SOURCE" and r.get("runtime_owner")==DB_OWNER and not missing(r.get("action","")) and not missing(r.get("gate","")) and not missing(r.get("permission",""))
def is_erp_read_field_operation_na(r):
    action=r.get("action","")
    if action not in ERP_GROUPS:return False
    gate,perm,owner=ERP_GROUPS[action]
    return r.get("type")=="READONLY" and missing(r.get("operation","")) and missing(r.get("method_path","")) and r.get("runtime_status")=="UI_LOCAL_OR_READ_SOURCE" and r.get("gate")==gate and r.get("permission")==perm and r.get("runtime_owner")==owner

def classify_pre(page,field,r):
    v=r.get(field,"");st=r.get("runtime_status","")
    if not missing(v):return "BOUND"
    if "UI_LOCAL_EXACT" in st:
        if field in {"operation","runtime_owner"}:return "LEGITIMATE_NA_UI_LOCAL"
        if field=="action" and display_only_type(r.get("type","")):return "LEGITIMATE_NA_UI_LOCAL"
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in st and field=="operation":return "CLOSED_BY_OWNER_LOCAL_COMMAND"
    if "READ_EXACT" in st and display_only_type(r.get("type","")):
        if field=="action":return "LEGITIMATE_NA_READ_PRESENTATION"
        if field=="operation" and not missing(r.get("runtime_owner","")):return "READ_OWNER_BOUND_OPERATION_UNSPECIFIED"
    if field=="runtime_owner" and any(x in st for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return "KNOWN_RUNTIME_BLOCKER"
    if v=="SOURCE_NOT_DEFINED":return "SOURCE_RESOLUTION_REQUIRED"
    if page=="WB-01" and field=="action" and is_wb_na(r):return "LEGITIMATE_NA_READ_PROJECTION_ACTION"
    if page=="AIAPI-01" and field=="action" and is_aiapi_na(r):return "LEGITIMATE_NA_DIRECT_OPERATION_ACTION"
    if page=="EDIT-01" and field=="operation" and is_edit_local_op_na(r):return "LEGITIMATE_NA_LOCAL_WORKING_DRAFT_OPERATION"
    if page=="IAM-01" and field=="operation" and is_iam_op_na(r):return "LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_OPERATION"
    if page=="IAM-01" and field=="runtime_owner" and is_iam_owner_na(r):return "LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_RUNTIME_OWNER"
    if page=="DB-01" and field=="operation" and is_db_read_model_operation_na(r):return "LEGITIMATE_NA_DB_READ_MODEL_OPERATION"
    if page=="ERP-01" and field=="operation" and is_erp_read_field_operation_na(r):return "LEGITIMATE_NA_ERP_READ_PROJECTION_FIELD_OPERATION"
    return "DEFINITION_BINDING_GAP"

def inventory(classifier):
    rows=[];counts=collections.Counter();maps={}
    for page,fn in PAGES.items():
        bm=compose(parse_controls(fn));maps[page]=bm
        for uid,r in bm.items():
            for field in ["action","gate","permission","operation","runtime_owner"]:
                cls=classifier(page,field,r);counts[cls]+=1
                if cls=="DEFINITION_BINDING_GAP":rows.append((page,uid,field))
    return rows,counts,maps

pre_rows,pre_counts,pre_maps=inventory(classify_pre)
assert len(pre_rows)==20,pre_rows
assert collections.Counter(f for _,_,f in pre_rows)=={"action":7,"gate":8,"permission":5}

# Validate the exact 17 N/A structures before defining applicability.
for uid in DEV_UIDS:
    r=pre_maps["DEV-01"][uid]
    assert r["type"]=="SEGMENT_BUTTON" and r["runtime_status"]=="UI_LOCAL_EXACT"
    assert r["gate"]=="DEV-01-GATE-PAGE" and not missing(r["action"])
    assert missing(r["permission"]) and missing(r["operation"]) and missing(r["method_path"]) and missing(r["runtime_owner"])
for uid in EDIT_UIDS:
    r=pre_maps["EDIT-01"][uid]
    assert r["type"].lower()=="preview" and r["runtime_status"]=="READ_EXACT"
    assert r["permission"]=="EDITING_VIEW" and not missing(r["gate"]) and not missing(r["operation"]) and r["runtime_owner"]=="CLIENT/ORCHESTRATION"
    assert missing(r["action"])
for uid in KB_UIDS:
    r=pre_maps["KB-01"][uid]
    assert r["type"]=="TAB" and r["runtime_status"]=="UI_LOCAL_EXACT"
    assert r["permission"]=="knowledge.read" and r["runtime_owner"]=="PAGE_UI_STATE"
    assert missing(r["action"]) and missing(r["gate"]) and missing(r["operation"]) and missing(r["method_path"])

# Validate SYS conflict-resolution context and new-definition eligibility.
sysmap=pre_maps["SYS-01"]
cand=sysmap["SYS-01-BTN-CANDIDATE-CREATE"]
assert cand["gate"]=="CURRENT_VIEW_AND_SOURCE_ACTION_GATE" and cand["operation"]=="createCandidate"
nav=sysmap["SYS-01-BTN-NAV-OPEN"]
assert nav["action"]=="ACT-NAV-OPEN" and nav["permission"]=="ops.read" and nav["operation"]=="getUiProjection"
assert nav["runtime_owner"]=="Projection / source page read owner" and nav["runtime_status"]=="UI_LOCAL_OR_READ_SOURCE" and missing(nav["gate"])
cr=sysmap["SYS-01-BTN-CR-CREATE"]
assert cr["action"]=="ACT-CR-CREATE" and cr["permission"]=="core.change.create" and cr["operation"]=="createChangeRequest"
assert cr["method_path"]=="POST /v1/system/changes/{SYSTEM_CHANGE_ID}/requests" and cr["runtime_owner"]=="SystemEngineer" and cr["runtime_status"]=="EFFECTFUL_EXACT" and missing(cr["gate"])
sb=sysmap["SYS-01-BTN-SANDBOX-TEST"]
assert sb["action"]=="SYS-01-ACT-SANDBOX-TEST" and sb["permission"]=="system.test.execute" and sb["operation"]=="runSandboxTest"
assert sb["method_path"]=="POST /v1/system/changes/{SYSTEM_CHANGE_ID}/sandbox-tests" and sb["runtime_owner"]=="AIAPI_PROVIDER_COMMAND_RUNTIME" and sb["runtime_status"]=="EFFECTFUL_EXACT" and missing(sb["gate"])

# SOURCE_PAGE_GATE is an existing registered gate for getUiProjection source-page reads.
source_page_examples=[]
for page,fn in PAGES.items():
    for uid,r in compose(parse_controls(fn)).items():
        if r.get("gate")=="SOURCE_PAGE_GATE" and r.get("operation")=="getUiProjection":
            source_page_examples.append((page,uid,r.get("permission"),r.get("runtime_status")))
assert source_page_examples, "SOURCE_PAGE_GATE_GET_UI_PROJECTION_EVIDENCE_MISSING"

def patch_field(path,bindings,field):
    d=Document(path);seen=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=exact_index(h,"Control UID");usi=exact_index(h,"UID Status");fi=exact_index(h,{"gate":"Gate","action":"Action UID","permission":"Permission / Auth Resource"}.get(field,field))
        if ci is None or usi is None or fi is None:continue
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals) or usi>=len(vals):continue
            uid=vals[ci]
            if uid not in bindings:continue
            assert vals[usi]=="EXISTING_SOURCE_UID",(path,uid,vals[usi])
            old=norm(row.cells[fi].text);new=bindings[uid]
            if missing(old):row.cells[fi].text=new
            else:assert old==new,(path,uid,field,old,new)
            seen[uid]+=1
    assert set(seen)==set(bindings),(path,field,"PATCH_MISMATCH",sorted(set(bindings)-set(seen)),dict(seen))
    d.save(path);Document(path)
    return dict(seen)

sys_page_patch=patch_field(SYS,SYS_GATE_BINDINGS,"gate")
sys_owner_patch=patch_field(S09,SYS_GATE_BINDINGS,"gate")

def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.33);sec.bottom_margin=Inches(.33);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=4.4):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for rr in rows:
        cells=t.add_row().cells
        for i,v in enumerate(rr):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
    return t

# Canonical owner N/A declarations.
d=Document(S06);add_landscape(d);d.add_heading("Batch 17 · DEV Local Stage Selector Permission Applicability",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::SYSTEM06::DEV] ").bold=True
p.add_run("The five DEV-01 stage selector controls are UI_LOCAL_EXACT segment selectors. They change local/page stage focus only, have no Operation, Method/Path, Runtime Owner, Persistence Owner, or external effect. DEV-01-GATE-PAGE is already the page interaction boundary. Therefore a separate control-level Permission/Auth Resource is non-applicable and MUST remain '—'; no permission identifier may be invented.")
add_table(d,["Target UID","Field","Action","Gate","Runtime Status","Applicability"],[[u,"Permission/Auth Resource",pre_maps["DEV-01"][u]["action"],"DEV-01-GATE-PAGE","UI_LOCAL_EXACT","LEGITIMATE_NA_DEV_LOCAL_STAGE_PERMISSION"] for u in DEV_UIDS],4.2)
d.save(S06);Document(S06)

d=Document(DEV);add_landscape(d);d.add_heading("Batch 17 · DEV Local Stage Permission Applicability Ledger",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::PAGE::DEV-01] ").bold=True
p.add_run("The existing five Permission cells remain '—'. This ledger makes the non-applicability explicit; it does not create a new authorization resource.")
add_table(d,["Target UID","Permission Cell","Handling"],[[u,"—","LEGITIMATE_NA_DEV_LOCAL_STAGE_PERMISSION"] for u in DEV_UIDS],4.7)
d.save(DEV);Document(DEV)

d=Document(S05);add_landscape(d);d.add_heading("Batch 17 · EDIT Read Preview Action Applicability",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::SYSTEM05::EDIT] ").bold=True
p.add_run("Candidate Preview and Final Preview are READ_EXACT projection panels. Their registered read Operation, Gate, Permission and Runtime Owner already define acquisition/visibility. The panel itself does not own a separate Action UID. Action is non-applicable and MUST remain '—'.")
add_table(d,["Target UID","Operation","Gate","Permission","Runtime Owner","Applicability"],[[u,pre_maps["EDIT-01"][u]["operation"],pre_maps["EDIT-01"][u]["gate"],"EDITING_VIEW","CLIENT/ORCHESTRATION","LEGITIMATE_NA_EDIT_READ_PREVIEW_ACTION"] for u in EDIT_UIDS],4.1)
d.save(S05);Document(S05)

d=Document(EDIT);add_landscape(d);d.add_heading("Batch 17 · EDIT Preview Action Applicability Ledger",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::PAGE::EDIT-01] ").bold=True
p.add_run("The two preview panel Action cells remain '—' because the read Operation is the execution/read identity; no parallel Action UID is created.")
add_table(d,["Target UID","Action Cell","Operation","Result"],[[u,"—",pre_maps["EDIT-01"][u]["operation"],"LEGITIMATE_NA_EDIT_READ_PREVIEW_ACTION"] for u in EDIT_UIDS],4.4)
d.save(EDIT);Document(EDIT)

d=Document(S01);add_landscape(d);d.add_heading("Batch 17 · KB Local Tab Action / Gate Applicability",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::SYSTEM01::KB] ").bold=True
p.add_run("The five KB-01 VIEW tabs are UI_LOCAL_EXACT TAB controls owned by PAGE_UI_STATE. They only switch the local Knowledge workbench view; knowledge.read is the existing authorization boundary. They own no Operation, Method/Path, Runtime mutation, or persistence effect. Separate control-level Action UID and Gate UID are non-applicable and MUST remain '—'.")
add_table(d,["Target UID","Fields","Permission","Runtime Owner","Runtime Status","Applicability"],[[u,"Action + Gate","knowledge.read","PAGE_UI_STATE","UI_LOCAL_EXACT","LEGITIMATE_NA_KB_LOCAL_TAB_ACTION_GATE"] for u in KB_UIDS],4.2)
d.save(S01);Document(S01)

d=Document(KB);add_landscape(d);d.add_heading("Batch 17 · KB Local Tab Applicability Ledger",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::PAGE::KB-01] ").bold=True
p.add_run("Action and Gate remain '—' for the exact five local tabs. knowledge.read remains the read access boundary; no execution endpoint or transition gate is synthesized.")
add_table(d,["Target UID","Action","Gate","Permission","Handling"],[[u,"—","—","knowledge.read","LEGITIMATE_NA_KB_LOCAL_TAB_ACTION_GATE"] for u in KB_UIDS],4.4)
d.save(KB);Document(KB)

# System 09 canonical Gate contracts, including the historical conflict resolution.
d=Document(S09);add_landscape(d);d.add_heading("Batch 17 · SYS Canonical Gate Contracts / Conflict Resolution",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::SYSTEM09::SYS] ").bold=True
p.add_run("System 09 is the Canonical Owner for these three SYS-01 controls. Two effectful controls had no Gate defined and now receive explicit owner-defined Gate UIDs. SYS-01-BTN-NAV-OPEN had a preserved conflict between CURRENT_VIEW_AND_SOURCE_ACTION_GATE and SOURCE_PAGE_GATE. The later exact binding of operation=getUiProjection and Runtime Owner='Projection / source page read owner', together with ops.read and UI_LOCAL_OR_READ_SOURCE, disambiguates the control to the source-page read domain; SOURCE_PAGE_GATE is therefore the canonical Gate. The mutation-oriented CURRENT_VIEW_AND_SOURCE_ACTION_GATE is not transferred.")
add_table(d,["Control UID","Canonical Gate","Authority Basis","Precondition / Failure Semantics"],[
["SYS-01-BTN-CR-CREATE","SYS-01-GATE-CR-CREATE","Canonical owner new definition for createChangeRequest","SYSTEM_CHANGE_ID resolves; core.change.create authorized; request validates against existing createChangeRequest contract. Missing/invalid precondition => fail closed."],
["SYS-01-BTN-NAV-OPEN","SOURCE_PAGE_GATE","Resolved historical conflict by getUiProjection + source-page read owner + ops.read","Referenced source-page/projection target must resolve under read authority. No mutation; unresolved source => fail closed."],
["SYS-01-BTN-SANDBOX-TEST","SYS-01-GATE-SANDBOX-TEST","Canonical owner new definition for runSandboxTest","SYSTEM_CHANGE_ID resolves; system.test.execute authorized; sandbox remains non-production and provider command runtime is eligible. Failure => no production promotion and fail closed."],
],4.0)
d.save(S09);Document(S09)

d=Document(SYS);add_landscape(d);d.add_heading("Batch 17 · SYS Gate Binding Ledger",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::PAGE::SYS-01] ").bold=True
p.add_run("Three existing Gate cells are now bound from System 09 Canonical Owner. The NAV-OPEN historical conflict is resolved to SOURCE_PAGE_GATE; the two effectful Gate identities are newly defined by the canonical owner without changing existing Action, Permission, Operation, Method/Path or Runtime Owner.")
add_table(d,["Control UID","Gate","Operation","Permission","Runtime Owner","Result"],[[u,g,compose(parse_controls(SYS))[u]["operation"],compose(parse_controls(SYS))[u]["permission"],compose(parse_controls(SYS))[u]["runtime_owner"],"BOUND_BY_BATCH_17"] for u,g in SYS_GATE_BINDINGS.items()],4.0)
d.save(SYS);Document(SYS)

# Batch17 post-classifier.
def classify_post(page,field,r):
    pre=classify_pre(page,field,r)
    if pre!="DEFINITION_BINDING_GAP":return pre
    if page=="DEV-01" and r.get("control") in DEV_UIDS and field=="permission":
        return "LEGITIMATE_NA_DEV_LOCAL_STAGE_PERMISSION"
    if page=="EDIT-01" and r.get("control") in EDIT_UIDS and field=="action":
        return "LEGITIMATE_NA_EDIT_READ_PREVIEW_ACTION"
    if page=="KB-01" and r.get("control") in KB_UIDS and field in {"action","gate"}:
        return "LEGITIMATE_NA_KB_LOCAL_TAB_ACTION_GATE"
    return pre

post_rows,post_counts,post_maps=inventory(classify_post)
# inventory() reparses current page files; exact SYS Gate values are now bound.
assert len(post_rows)==0,post_rows
for uid,g in SYS_GATE_BINDINGS.items():
    assert post_maps["SYS-01"][uid]["gate"]==g,(uid,post_maps["SYS-01"][uid]["gate"],g)
s09map=compose(parse_controls(S09))
for uid,g in SYS_GATE_BINDINGS.items():
    assert s09map[uid]["gate"]==g,(uid,s09map[uid]["gate"],g)
for uid in DEV_UIDS: assert missing(post_maps["DEV-01"][uid]["permission"])
for uid in EDIT_UIDS: assert missing(post_maps["EDIT-01"][uid]["action"])
for uid in KB_UIDS:
    assert missing(post_maps["KB-01"][uid]["action"]) and missing(post_maps["KB-01"][uid]["gate"])

# Central logic closure.
logic=Document(LOGIC);add_landscape(logic);logic.add_heading("Batch 17 · Final Definition Binding Closure",level=1)
p=logic.add_paragraph();p.add_run("["+MARK+"] ").bold=True
p.add_run("Batch 17 resolves the final 20 Definition Binding gaps after Batch 16. Seventeen fields are closed by exact Canonical Owner non-applicability semantics rather than invented identifiers. Three SYS Gate fields are bound by System 09: two new canonical Gate definitions and one historical conflict resolution. No Operation, Runtime Owner, Method/Path, API route or permission resource is fabricated.")
add_table(logic,["Closure Class","Count","Result"],[
["Pre-batch Definition Binding Gap",20,"Action 7 + Gate 8 + Permission 5"],
["DEV local stage Permission N/A",5,"Exact five UI_LOCAL_EXACT segment selectors; cells remain '—'."],
["EDIT read preview Action N/A",2,"Exact two READ_EXACT preview panels; cells remain '—'."],
["KB local tab Action/Gate N/A",10,"Exact five UI_LOCAL_EXACT TAB controls × Action/Gate; cells remain '—'."],
["SYS canonical Gate bindings",3,"CR create + NAV open + Sandbox test."],
["Historical SYS NAV Gate conflict",1,"Resolved: SOURCE_PAGE_GATE; CURRENT_VIEW_AND_SOURCE_ACTION_GATE rejected for this source-page read control."],
["Post-batch Definition Binding Gap",0,"Fresh full 18-page classifier denominator."],
["New API routes / Operation / Runtime Owner / Permission invented",0,"None."],
],4.2)
logic.add_heading("SYS NAV Conflict Resolution Evidence",level=2)
add_table(logic,["Item","Evidence"],[
["Historical conflicting values","CURRENT_VIEW_AND_SOURCE_ACTION_GATE | SOURCE_PAGE_GATE"],
["Current exact Operation","getUiProjection"],
["Current exact Runtime Owner","Projection / source page read owner"],
["Current Permission","ops.read"],
["Current Runtime Status","UI_LOCAL_OR_READ_SOURCE"],
["Existing registered relation","SOURCE_PAGE_GATE is used by source-page getUiProjection reads."],
["Canonical decision","SOURCE_PAGE_GATE"],
],4.4)
machine={
"marker":MARK,"base_head":BASE_HEAD,"pre_definition_gaps":20,"dev_permission_na":5,"edit_action_na":2,
"kb_action_gate_na":10,"sys_gate_bindings":3,"new_canonical_gate_definitions":2,
"resolved_sys_nav_conflict":1,"sys_nav_gate":"SOURCE_PAGE_GATE","na_reclassified":17,
"gaps_closed":20,"post_definition_gaps":0,"invented_api_routes":0,"invented_operations":0,
"invented_runtime_owners":0,"invented_permissions":0,"runtime_execution_claimed":False,
}
logic.add_paragraph("BATCH17_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

changed=[S01,S05,S06,S09,DEV,EDIT,KB,SYS,LOGIC]
report={
"machine":machine,
"source_page_gate_examples":[{"page":p,"uid":u,"permission":perm,"runtime_status":st} for p,u,perm,st in source_page_examples],
"sys_gate_bindings":SYS_GATE_BINDINGS,
"sys_page_patch":sys_page_patch,"sys_owner_patch":sys_owner_patch,
"changed_docs":changed,
"output_blob_sha":{f:blob(f) for f in changed},
}
Path("__batch17_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH17_REMEDIATION="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
