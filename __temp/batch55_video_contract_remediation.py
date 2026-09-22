from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-55-VIDEO-CONTRACT-REMEDIATION-V1"
BASE_HEAD="661151bef10a0d7f401530e8f8fff9809994e99c"
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
VIDEO="ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S05:"7c95ef12473051d7b5b5d0a9eb19dfe8645383d9",
VIDEO:"379af52edba5700b4cd8d6566dcd25f49cbb43a0",
LOGIC:"79e0c130b7cbdd6afeee714017e46b22aa8c8125",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx","ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx","EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx","IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx","STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx","VIDEO-01":VIDEO,
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx"}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

TARGETS={
"VIDEO-01-BTN-EXECUTE":{
 "operation":"startVideoGenerationFlow","route":"POST /v1/video-runtime-runs","payload":"StartVideoGenerationFlowRequest",
 "persistence":"System 05 / ProductionWorkItem + VideoOutputVersion Generation Lineage + Audit",
 "basis":"Starts generation only from exact admitted production input/bindings. Job/run creation is not output confirmation; provider/runtime success requires separate evidence."
},
"VIDEO-01-BTN-CONFIRM":{
 "operation":"confirmVideoCandidate","route":"POST /v1/video-runtime-runs/{runId}/candidates/{videoOutputVersionId}/confirm","payload":"ConfirmVideoCandidateRequest",
 "persistence":"System 05 / VideoOutputVersion Confirmed State + Decision Evidence/Audit",
 "basis":"Confirmation binds exact candidate/output version and authorized decision evidence, producing immutable confirmed lineage without overwriting prior VideoOutputVersion."
},
"VIDEO-01-BTN-LOCK":{
 "operation":"lockVideoVersion","route":"POST /v1/state-commands/videoversion/lock","payload":"LockVideoVersionRequest",
 "persistence":"public.production_output_version_locks",
 "basis":"Reuses exact Current lock route/persistence. Lock binds exact video output version/hash and governed decision evidence; request/command is not self-approval."
},
"VIDEO-01-BTN-HANDOFF":{
 "operation":"handoffVideoToEdit","route":"POST /v1/video-runtime-runs/{runId}/handoff-to-edit","payload":"HandoffVideoToEditRequest",
 "persistence":"System 05 / HandoffManifest + VideoOutputVersion Provenance/Audit",
 "basis":"VIDEO→EDIT handoff must include exact VideoOutputVersion + Scene/Shot lineage + Asset/Script/DNA refs + continuity + rights/policy; successor admission remains separate."
},
"VIDEO-01-BTN-GEN-CORRECTION":{
 "operation":"generateCorrectionScriptCandidate","route":"POST /v1/state-commands/correctionscript/generate","payload":"GenerateVideoCorrectionScriptCandidateRequest",
 "persistence":"public.correction_script_versions",
 "basis":"Reuses exact Current correction-script generation route. Persists a candidate/version with exact source/output/findings/context refs; candidate is not approved correction authority."
},
"VIDEO-01-BTN-APPROVE-CORRECTION":{
 "operation":"approveCorrectionScriptCandidate","route":"POST /v1/state-commands/correctionscript/approve","payload":"ApproveVideoCorrectionScriptCandidateRequest",
 "persistence":"public.correction_script_versions",
 "basis":"Reuses exact Current approval route. Human/authorized approval binds exact correction script candidate/version; prior versions remain immutable."
},
"VIDEO-01-BTN-EXEC-CORRECTION":{
 "operation":"executeVideoCorrection","route":"POST /v1/video-runtime-runs/{runId}/corrections/execute","payload":"ExecuteVideoCorrectionRequest",
 "persistence":"System 05 / VideoOutputVersion + Correction/Delta Lineage + Audit",
 "basis":"Executes only an approved correction script/delta against exact selected source version. Correction produces a new VideoOutputVersion and preserves unchanged shot lineage."
},
"VIDEO-01-BTN-RETRY":{
 "operation":"retryVideoTask","route":"POST /v1/video-runtime-runs/{runId}/tasks/{taskId}/retry","payload":"RetryVideoTaskRequest",
 "persistence":"System 05 / ProductionWorkItem Retry Attempt + Video Generation Audit",
 "basis":"Retry binds exact failed task/attempt/version and preserves failure evidence; it cannot erase prior attempts or change scope silently."
},
"VIDEO-01-BTN-EVALUATE":{
 "operation":"evaluateVideoCandidate","route":"POST /v1/video-runtime-runs/{runId}/candidates/{videoOutputVersionId}/evaluations","payload":"EvaluateVideoCandidateRequest",
 "persistence":"System 05 / EvaluationResult + Video Evidence/Audit",
 "basis":"Evaluation records score/evidence/findings against exact criteria and VideoOutputVersion. Score alone cannot auto-confirm, lock or handoff."
},
"VIDEO-01-BTN-FINDING":{
 "operation":"createVideoFinding","route":"POST /v1/video-runtime-runs/{runId}/candidates/{videoOutputVersionId}/findings","payload":"CreateVideoFindingRequest",
 "persistence":"System 05 / Finding + Video Evidence/Audit",
 "basis":"Finding binds exact source VideoOutputVersion/shot or scene lineage, category, severity, evidence and expected correction; no silent acceptance."
},
}

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        for i,h in enumerate(hs):
            if needle.lower() in h:return i
    return None
def exact_index(headers,name):
    for i,h in enumerate(headers):
        if norm(h).lower()==name.lower():return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
        "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),"permission":hfind(h,"permission","auth resource"),
        "payload_schema":hfind(h,"payload / schema","payload","schema"),"operation":hfind(h,"operation"),
        "method_path":hfind(h,"method / path","method","path"),"runtime_owner":hfind(h,"runtime owner"),
        "persistence_owner":hfind(h,"persistence owner"),"runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()};rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out
def collect_uid_operations(path):
    d=Document(path);out=collections.defaultdict(set)
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ui=hfind(h,"control uid","target uid");oi=hfind(h,"operation")
        if ui is None or oi is None:continue
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ui>=len(vals) or oi>=len(vals):continue
            uid,op=vals[ui],vals[oi]
            if uid and uid not in {"—","-"} and op and op not in {"—","-"}:out[uid].add(op)
    return out
def blob(path):
    b=Path(path).read_bytes();return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():assert blob(f)==s,(f,blob(f),s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""
assert len(TARGETS)==10

pre_page=compose(parse_controls(VIDEO));pre_owner=compose(parse_controls(S05));owner_ops=collect_uid_operations(S05)
existing_complete={"VIDEO-01-BTN-LOCK","VIDEO-01-BTN-APPROVE-CORRECTION"}
existing_method={"VIDEO-01-BTN-LOCK","VIDEO-01-BTN-GEN-CORRECTION","VIDEO-01-BTN-APPROVE-CORRECTION"}
for uid,s in TARGETS.items():
    assert pre_page[uid]["operation"]==s["operation"],("PAGE",uid,pre_page[uid]["operation"],s["operation"])
    assert s["operation"] in owner_ops.get(uid,set()),("OWNER_OPERATION_MISSING",uid,s["operation"])
    for m,label in [(pre_page,"PAGE"),(pre_owner,"OWNER")]:
        r=m[uid];assert r["runtime_status"]=="EFFECTFUL_EXACT",(label,uid,r["runtime_status"])
        assert missing(r["payload_schema"]),(label,uid,"PAYLOAD_ALREADY",r["payload_schema"])
        if uid in existing_method:
            assert r["method_path"]==s["route"],(label,uid,"METHOD",r["method_path"],s["route"])
        else:
            assert missing(r["method_path"]),(label,uid,"METHOD_ALREADY",r["method_path"])
        if uid in existing_complete:
            assert r["persistence_owner"]==s["persistence"],(label,uid,"PERSIST",r["persistence_owner"],s["persistence"])
        else:
            assert missing(r["persistence_owner"]),(label,uid,"PERSIST_ALREADY",r["persistence_owner"])

def patch(path):
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact_index(h,"Control UID")
        if ci is None:continue
        idx={"method_path":hfind(h,"method / path","method","path"),"payload_schema":hfind(h,"payload / schema","payload","schema"),"persistence_owner":hfind(h,"persistence owner")}
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if uid not in TARGETS:continue
            s=TARGETS[uid];assign={"payload_schema":s["payload"]}
            if uid not in existing_method:assign["method_path"]=s["route"]
            if uid not in existing_complete:assign["persistence_owner"]=s["persistence"]
            for f,new in assign.items():
                i=idx[f]
                if i is None or i>=len(row.cells):continue
                old=norm(row.cells[i].text)
                if old==new:hits[(uid,f)]+=1;continue
                assert missing(old),(path,uid,f,old,new)
                row.cells[i].text=new;hits[(uid,f)]+=1
    exp={(u,"payload_schema") for u in TARGETS}|{(u,"method_path") for u in TARGETS if u not in existing_method}|{(u,"persistence_owner") for u in TARGETS if u not in existing_complete}
    assert set(hits)==exp,(path,sorted(exp-set(hits)),dict(hits));d.save(path);Document(path);return {f"{u}:{f}":n for (u,f),n in hits.items()}

patches={"SYSTEM05_OWNER":patch(S05),"VIDEO_PAGE":patch(VIDEO)}

def landscape(d):
    s=d.add_section(WD_SECTION.NEW_PAGE);s.orientation=WD_ORIENT.LANDSCAPE;s.page_width,s.page_height=s.page_height,s.page_width
    s.top_margin=Inches(.33);s.bottom_margin=Inches(.33);s.left_margin=Inches(.28);s.right_margin=Inches(.28)
def table(d,h,rows,fs=3.45):
    t=d.add_table(rows=1,cols=len(h));t.style="Table Grid"
    tr=t.rows[0]._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");tr.append(e)
    for i,x in enumerate(h):
        t.rows[0].cells[i].text=str(x);sh=OxmlElement("w:shd");sh.set(qn("w:fill"),"EDE9FE");t.rows[0].cells[i]._tc.get_or_add_tcPr().append(sh)
    for rr in rows:
        c=t.add_row().cells
        for i,x in enumerate(rr):c[i].text=str(x)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
for path,title,prefix in [(S05,"Batch 55 · VIDEO Canonical Contract Registry","SYSTEM05"),(VIDEO,"Batch 55 · VIDEO-01 Contract Binding Ledger","PAGE::VIDEO-01")]:
    d=Document(path);landscape(d);d.add_heading(title,1)
    d.add_paragraph(f"[{MARK}::{prefix}] Existing Video lock and correction-script state-command routes are preserved. Seven previously undefined VIDEO routes are formal System 05 design contracts. VideoOutputVersion remains immutable/versioned; correction creates new lineage; evaluation cannot auto-confirm; handoff requires exact successor-consumable provenance. No provider/runtime deployment is claimed.")
    rows=[[u,s["operation"],s["route"],s["payload"],s["persistence"],s["basis"]] for u,s in TARGETS.items()]
    table(d,["Control UID","Operation","Method / Path","Payload / Schema","Persistence Owner","Semantics"],rows)
    d.save(path);Document(path)

post_page=compose(parse_controls(VIDEO));post_owner=compose(parse_controls(S05))
for uid,s in TARGETS.items():
    for m,label in [(post_page,"PAGE"),(post_owner,"OWNER")]:
        assert m[uid]["method_path"]==s["route"],(label,uid,m[uid]["method_path"],s["route"])
        assert m[uid]["payload_schema"]==s["payload"],(label,uid,m[uid]["payload_schema"],s["payload"])
        assert m[uid]["persistence_owner"]==s["persistence"],(label,uid,m[uid]["persistence_owner"],s["persistence"])

allm={p:compose(parse_controls(fn)) for p,fn in PAGES.items()};rem=[]
for p,m in allm.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if missing(op) or st not in {"EFFECTFUL_EXACT","READ_EXACT"}:continue
        if missing(r.get("method_path","")):rem.append(("method_path",p,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("payload_schema","")):rem.append(("payload_schema",p,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("persistence_owner","")):rem.append(("persistence_owner",p,uid))
rc=collections.Counter(x[0] for x in rem);assert len(rem)==61,(len(rem),rc);assert rc=={"method_path":18,"payload_schema":23,"persistence_owner":20},rc

d=Document(LOGIC);landscape(d);d.add_heading("Batch 55 · VIDEO Contract Closure",1)
machine={"marker":MARK,"base_head":BASE_HEAD,"pre_contract_cells":86,"video_cells_closed":25,"method_path_closed":7,"payload_schema_closed":10,"persistence_owner_closed":8,
"post_contract_cells":61,"remaining_method_path":18,"remaining_payload_schema":23,"remaining_persistence_owner":20,
"existing_state_command_routes_reused":3,"new_video_design_routes":7,"runtime_execution_claimed":False}
table(d,["Metric","Value"],[["Pre denominator","86"],["VIDEO cells closed","25 = Method 7 + Payload 10 + Persistence 8"],["Post denominator","61 = Method 18 + Payload 23 + Persistence 20"],["Existing exact state-command routes reused","3"],["New VIDEO design routes","7"],["Runtime execution claimed","False"]],4.0)
d.add_paragraph("BATCH55_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")));d.save(LOGIC);Document(LOGIC)
changed=[S05,VIDEO,LOGIC]
Path("__batch55_remediation_report.json").write_text(json.dumps({"machine":machine,"targets":TARGETS,"patches":patches,"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed}},ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH55="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
