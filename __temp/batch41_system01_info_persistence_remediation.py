from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-41-SYSTEM01-INFO-PERSISTENCE-REMEDIATION-V1"
BASE_HEAD="8cde9f690e825356763a36ef3dd0146d841fea1f"
S01="01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx"
INFO="ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S01:"1d26673c7fdee7b99e93b7963d5d21d8e8aba0c5",
INFO:"9c8af306b835c49d7216824c82f420c0c4d8e5b1",
LOGIC:"87705597acbe11577a7fd09da01ba454791408ce",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":INFO,
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]
BINDINGS={
"INFO-01-BTN-REFRESH":{
 "operation":"refreshProjection",
 "persistence_owner":"Global Brain / ACPOSStateProjection + Projection Audit",
 "basis":"Refresh rematerializes a derived projection for an exact domain version. It must not create a second canonical domain row; projection state/version and audit lineage remain reconstructable."
},
"INFO-01-BTN-SEARCH":{
 "operation":"searchProjection",
 "persistence_owner":"Global Brain / SemanticIndexResult + Search Audit",
 "basis":"Search persists only search/index result pointers, scope and audit/evidence needed to dereference exact source/version/evidence. Search score/result is not canonical truth."
},
"INFO-01-BTN-EXPORT":{
 "operation":"exportProjection",
 "persistence_owner":"Global Brain / Export Artifact + Export Audit",
 "basis":"Export persists an authorized derived export artifact and audit lineage for the exact projection/scope. Export output does not become a second canonical owner for domain data."
},
}

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out
def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():
    got=blob(f);assert got==s,(f,got,s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

pre_page=compose(parse_controls(INFO));pre_owner=compose(parse_controls(S01))
for uid,s in BINDINGS.items():
    for m,label in [(pre_page,"PAGE"),(pre_owner,"OWNER")]:
        assert uid in m,(label,uid,"MISSING")
        r=m[uid]
        assert r["operation"]==s["operation"],(label,uid,r["operation"],s["operation"])
        assert r["runtime_status"]=="EFFECTFUL_EXACT",(label,uid,r["runtime_status"])
        assert missing(r["persistence_owner"]),(label,uid,r["persistence_owner"])
        assert not missing(r["method_path"]),(label,uid,"METHOD_MISSING")
        assert not missing(r["payload_schema"]),(label,uid,"PAYLOAD_MISSING")

def patch(path):
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact_index(h,"Control UID");pi=hfind(h,"persistence owner")
        if ci is None or pi is None:continue
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals):continue
            uid=vals[ci]
            if uid not in BINDINGS:continue
            old=norm(row.cells[pi].text);new=BINDINGS[uid]["persistence_owner"]
            if old==new:hits[uid]+=1;continue
            assert missing(old),(path,uid,old,new)
            row.cells[pi].text=new;hits[uid]+=1
    assert set(hits)==set(BINDINGS),(path,dict(hits),set(BINDINGS)-set(hits))
    d.save(path);Document(path)
    return dict(hits)

patches={"SYSTEM01_OWNER":patch(S01),"INFO_PAGE":patch(INFO)}

def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.33);sec.bottom_margin=Inches(.33);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=4.0):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for rr in rows:
        cells=t.add_row().cells
        for i,v in enumerate(rr):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
    return t

for path,title,prefix in [
(S01,"Batch 41 · Global Brain Projection Persistence Ownership","SYSTEM01"),
(INFO,"Batch 41 · INFO-01 Projection Persistence Binding Ledger","PAGE::INFO-01"),
]:
    d=Document(path);add_landscape(d);d.add_heading(title,level=1)
    p=d.add_paragraph();p.add_run(f"[{MARK}::{prefix}] ").bold=True
    p.add_run("Projection/search/export persistence is derived/audit state only. Domain canonical business rows remain under their original owners. No physical table names are invented and no exported/search result is promoted to canonical truth.")
    rr=[[uid,s["operation"],s["persistence_owner"],s["basis"]] for uid,s in BINDINGS.items()]
    add_table(d,["Control UID","Operation","Persistence Owner","Required Semantics"],rr,3.75)
    d.save(path);Document(path)

post_page=compose(parse_controls(INFO));post_owner=compose(parse_controls(S01))
for uid,s in BINDINGS.items():
    assert post_page[uid]["persistence_owner"]==s["persistence_owner"],(uid,post_page[uid]["persistence_owner"])
    assert post_owner[uid]["persistence_owner"]==s["persistence_owner"],(uid,post_owner[uid]["persistence_owner"])

all_maps={p:compose(parse_controls(fn)) for p,fn in PAGES.items()}
remaining=[]
for page,m in all_maps.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if missing(op) or st not in {"EFFECTFUL_EXACT","READ_EXACT"}:continue
        if missing(r.get("method_path","")):remaining.append(("method_path",page,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("payload_schema","")):remaining.append(("payload_schema",page,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("persistence_owner","")):remaining.append(("persistence_owner",page,uid))
rc=collections.Counter(f for f,_,_ in remaining)
assert len(remaining)==283,(len(remaining),rc)
assert rc=={"method_path":63,"payload_schema":119,"persistence_owner":101},rc

logic=Document(LOGIC);add_landscape(logic);logic.add_heading("Batch 41 · System 01 INFO Persistence Closure",level=1)
p=logic.add_paragraph();p.add_run(f"[{MARK}] ").bold=True
p.add_run("System 01 closes the three remaining INFO-01 persistence-owner cells. Refresh persists derived projection state/audit, Search persists semantic index result pointers/audit, and Export persists derived export artifact/audit. None becomes a second domain canonical truth owner.")
add_table(logic,["Metric","Value","Result"],[
["Pre contract denominator",286,"Method 63 + Payload 119 + Persistence 104"],
["System 01 Persistence Owner cells closed",3,"Refresh/Search/Export"],
["Post contract denominator",283,"Method 63 + Payload 119 + Persistence 101"],
["Physical DB table names invented",0,"None"],
["Second canonical domain owner created",0,"None"],
["Runtime execution claimed","False","Design-contract closure only"],
],4.1)
machine={"marker":MARK,"base_head":BASE_HEAD,"pre_contract_cells":286,"system01_persistence_closed":3,
"post_contract_cells":283,"remaining_method_path":63,"remaining_payload_schema":119,"remaining_persistence_owner":101,
"physical_db_tables_invented":0,"second_canonical_domain_owner_created":0,"runtime_execution_claimed":False}
logic.add_paragraph("BATCH41_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

changed=[S01,INFO,LOGIC]
report={"machine":machine,"bindings":BINDINGS,"patches":patches,"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed}}
Path("__batch41_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH41="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
