from pathlib import Path
from docx import Document
import json,re,collections,subprocess,zipfile,xml.etree.ElementTree as ET,hashlib,unicodedata,math,shutil,os
from PIL import Image

MARK="ACPOS-20260922-BATCH-58-FULL-31-DOCX-CONTENT-ENCODING-AUDIT-V1"
BASE_HEAD="85fd2a890eb78bf050cc4b7616439e484d940ca6"
ROOT=Path(".")
OUT=Path("__batch58_full_audit")
PDFDIR=OUT/"pdf"; PNGDIR=OUT/"png"; CONTACT=OUT/"contact"
for p in [OUT,PDFDIR,PNGDIR,CONTACT]:p.mkdir(parents=True,exist_ok=True)

DOCS=sorted(ROOT.glob("*.docx"))
assert len(DOCS)==31,len(DOCS)
assert subprocess.check_output(["git","-c","core.quotePath=false","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

PAGE_DOCS={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}

MOJI_PATTERNS=[
"\ufffd","ï¿½","â€™","â€œ","â€","â€“","â€”","â€¦","Ã¤","Ã¥","Ã©","Ã¨","Ã","Â©","Â®","Â ","ðŸ"
]
HIGH_PLACEHOLDER_PATTERNS=[
r"\bTODO\b",r"\bTBD\b",r"\bFIXME\b",r"\bLOREM\b",r"\bPLACEHOLDER\b",
r"SOURCE_NOT_DEFINED",r"\bXXX\b",r"\?\?\?",r"待補",r"待定",r"尚未定義",r"未定義"
]
LEGACY_PATTERNS=[r"\bR9(?:\.0\.1)?\b",r"\bR5\b",r"038-\d",r"Batch 0(?:3[0-9]|4[0-9])"]
CJK_RE=re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")
METHOD_RE=re.compile(r"^(GET|POST|PUT|PATCH|DELETE|HEAD|OPTIONS)\s+/\S+$")
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

def norm(x):
    return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):
    return v in {"","—","-","SOURCE_NOT_DEFINED","UNCHANGED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":str(path),"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        vals_by={}
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            vals_by[f]=vals
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        base["_values"]=vals_by;out[uid]=base
    return out
def blob_sha(path):
    b=path.read_bytes();return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()
def source_text_from_docx(path):
    chunks=[]
    with zipfile.ZipFile(path) as z:
        names=[n for n in z.namelist() if re.match(r"word/(document|header\d+|footer\d+|footnotes|endnotes)\.xml$",n)]
        ns="{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        for n in names:
            root=ET.fromstring(z.read(n))
            chunks.extend((e.text or "") for e in root.iter(ns+"t"))
    return "\n".join(chunks)
def all_xml_integrity(path):
    errs=[];count=0
    try:
        with zipfile.ZipFile(path) as z:
            bad=z.testzip()
            if bad:errs.append(f"ZIP_CRC:{bad}")
            for n in z.namelist():
                if n.endswith(".xml") or n.endswith(".rels"):
                    count+=1
                    try:ET.fromstring(z.read(n))
                    except Exception as e:errs.append(f"{n}:{type(e).__name__}:{e}")
    except Exception as e:errs.append(f"ZIP_OPEN:{type(e).__name__}:{e}")
    return count,errs
def suspicious_unicode(txt):
    out=[]
    for i,ch in enumerate(txt):
        cp=ord(ch);cat=unicodedata.category(ch)
        if ch=="\ufeff":out.append({"kind":"BOM_IN_TEXT","index":i,"codepoint":f"U+{cp:04X}"})
        elif 0xE000<=cp<=0xF8FF:out.append({"kind":"PRIVATE_USE","index":i,"codepoint":f"U+{cp:04X}"})
        elif cat=="Cc" and ch not in "\n\r\t":out.append({"kind":"CONTROL_CHAR","index":i,"codepoint":f"U+{cp:04X}"})
        elif ch in {"\u200b","\u200c","\u200d","\u2060"}:out.append({"kind":"ZERO_WIDTH","index":i,"codepoint":f"U+{cp:04X}"})
        if len(out)>=100:break
    return out

report={
 "marker":MARK,"base_head":BASE_HEAD,"documents":{},"package":{},"content_findings":[],
 "contract_findings":{},"cross_document_findings":{}
}
critical=[];warnings=[]
total_pages=0;total_cjk_src=0;total_cjk_pdf=0

# Source + render audit every document.
for idx,path in enumerate(DOCS,1):
    fn=path.name
    xml_count,xml_errors=all_xml_integrity(path)
    src=source_text_from_docx(path)
    src_moji=[p for p in MOJI_PATTERNS if p in src]
    uni=suspicious_unicode(src)
    placeholders=[]
    for pat in HIGH_PLACEHOLDER_PATTERNS:
        hits=list(re.finditer(pat,src,re.I))
        if hits:placeholders.append({"pattern":pat,"count":len(hits),"samples":[src[max(0,m.start()-80):m.end()+120] for m in hits[:5]]})
    legacy=[]
    for pat in LEGACY_PATTERNS:
        hits=list(re.finditer(pat,src,re.I))
        if hits:legacy.append({"pattern":pat,"count":len(hits),"samples":[src[max(0,m.start()-80):m.end()+120] for m in hits[:5]]})
    if xml_errors:critical.append({"file":fn,"kind":"OOXML_INTEGRITY","detail":xml_errors})
    if src_moji:critical.append({"file":fn,"kind":"SOURCE_MOJIBAKE","detail":src_moji})
    severe_uni=[x for x in uni if x["kind"] in {"BOM_IN_TEXT","PRIVATE_USE","CONTROL_CHAR"}]
    if severe_uni:critical.append({"file":fn,"kind":"SOURCE_UNICODE_ANOMALY","detail":severe_uni})
    if placeholders:warnings.append({"file":fn,"kind":"PLACEHOLDER_CANDIDATE","detail":placeholders})
    if legacy:warnings.append({"file":fn,"kind":"LEGACY_REFERENCE_CANDIDATE","detail":legacy})

    # LibreOffice -> PDF.
    proc=subprocess.run(["libreoffice","--headless","--convert-to","pdf","--outdir",str(PDFDIR),str(path)],
                        text=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,timeout=180)
    pdf=PDFDIR/(path.stem+".pdf")
    convert_ok=proc.returncode==0 and pdf.exists() and pdf.stat().st_size>0
    if not convert_ok:
        critical.append({"file":fn,"kind":"PDF_CONVERSION_FAILED","detail":proc.stdout[-4000:]})
        pdf_text="";pages=0;blank=[];edge=[];pdf_moji=[];ratio=None;fonts=[]
    else:
        pdf_text=subprocess.check_output(["pdftotext","-layout",str(pdf),"-"],text=True,errors="strict")
        pdf_moji=[p for p in MOJI_PATTERNS if p in pdf_text]
        if pdf_moji:critical.append({"file":fn,"kind":"PDF_TEXT_MOJIBAKE","detail":pdf_moji})
        src_cjk=len(CJK_RE.findall(src));pdf_cjk=len(CJK_RE.findall(pdf_text))
        total_cjk_src+=src_cjk;total_cjk_pdf+=pdf_cjk
        ratio=(pdf_cjk/src_cjk) if src_cjk else None
        if src_cjk and ratio<0.97:
            critical.append({"file":fn,"kind":"CJK_LOSS_AFTER_CONVERSION","detail":{"source_cjk":src_cjk,"pdf_cjk":pdf_cjk,"ratio":ratio}})
        # fonts
        try:
            pf=subprocess.check_output(["pdffonts",str(pdf)],text=True,errors="replace").splitlines()
            fonts=pf[:80]
        except Exception as e:fonts=[f"ERROR:{e}"]
        # render pages.
        out=PNGDIR/path.stem;out.mkdir(parents=True,exist_ok=True)
        subprocess.check_call(["pdftoppm","-png","-r","72",str(pdf),str(out/"page")],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        imgs=sorted(out.glob("page-*.png"));pages=len(imgs);total_pages+=pages
        blank=[];edge=[];thumbs=[]
        for pi,imf in enumerate(imgs,1):
            im=Image.open(imf).convert("L");w,h=im.size
            bbox=im.point(lambda x:255 if x<245 else 0).getbbox()
            if bbox is None:blank.append(pi)
            elif bbox[0]<=1 or bbox[1]<=1 or bbox[2]>=w-1 or bbox[3]>=h-1:edge.append({"page":pi,"bbox":bbox,"size":(w,h)})
            th=Image.open(imf).convert("RGB");th.thumbnail((165,235));thumbs.append(th.copy())
        if blank:critical.append({"file":fn,"kind":"BLANK_RENDERED_PAGES","detail":blank})
        if edge:critical.append({"file":fn,"kind":"EDGE_CLIP_CANDIDATE","detail":edge[:30]})
        # compact contact sheet.
        cols=5;cw=175;ch=245;rows=math.ceil(len(thumbs)/cols)
        sheet=Image.new("RGB",(cols*cw,max(ch,rows*ch)),"white")
        for i,im in enumerate(thumbs):sheet.paste(im,((i%cols)*cw+5,(i//cols)*ch+5))
        sheet.save(CONTACT/(path.stem+"_contact.jpg"),quality=80)
        shutil.rmtree(out)

    report["documents"][fn]={
      "git_blob_sha":blob_sha(path),"bytes":path.stat().st_size,"xml_parts_checked":xml_count,
      "xml_errors":xml_errors,"source_chars":len(src),"source_cjk":len(CJK_RE.findall(src)),
      "source_mojibake_patterns":src_moji,"unicode_anomalies":uni,
      "placeholder_candidates":placeholders,"legacy_reference_candidates":legacy,
      "pdf_conversion_ok":convert_ok,"pdf_bytes":pdf.stat().st_size if convert_ok else 0,
      "pdf_pages":pages,"pdf_cjk":len(CJK_RE.findall(pdf_text)) if convert_ok else 0,
      "cjk_preservation_ratio":ratio,"pdf_mojibake_patterns":pdf_moji if convert_ok else [],
      "blank_pages":blank,"edge_clip_candidates":edge,"pdffonts_head":fonts,
    }
    print(f"AUDITED {idx:02d}/31 {fn} pages={pages} cjk_ratio={ratio}")

# Contract semantics across the 18 page docs.
operation_missing=[];operation_placeholders=[];action_unbound=[];gate_gaps=[];permission_gaps=[];runtime_owner_gaps=[];contract_gaps=[];conflicts=[];route_conflicts=[];passive_reads=[];direct_ops=[]
for page,fn in PAGE_DOCS.items():
    m=compose(parse_controls(fn));routes=collections.defaultdict(set)
    for uid,r in m.items():
        st=r.get("runtime_status","")
        if st not in {"EFFECTFUL_EXACT","READ_EXACT"}:continue
        op=r.get("operation","");typ=norm(r.get("type","")).upper();act=r.get("action","")
        if missing(op):
            if st=="READ_EXACT" and (typ=="READONLY" or missing(act) or norm(act).upper().endswith("ACT-NOOP")):
                passive_reads.append((page,uid,typ,act));continue
            operation_missing.append((page,uid,st,typ,act));continue
        if re.search(r"(placeholder|pending|todo|tbd|source_not_defined)",op,re.I):operation_placeholders.append((page,uid,op))
        if missing(act):direct_ops.append((page,uid,op,typ,st))
        if missing(r.get("gate","")):gate_gaps.append((page,uid,op))
        if missing(r.get("permission","")):permission_gaps.append((page,uid,op))
        if missing(r.get("runtime_owner","")):runtime_owner_gaps.append((page,uid,op))
        if missing(r.get("method_path","")):contract_gaps.append(("method_path",page,uid,op))
        else:
            mp=r["method_path"]
            if mp not in {"LOCAL LOCAL","NO_PUBLIC_API_BY_AUTHORITY"} and not METHOD_RE.match(mp):
                critical.append({"file":fn,"kind":"MALFORMED_METHOD_PATH","detail":{"page":page,"uid":uid,"operation":op,"method_path":mp}})
        if st=="EFFECTFUL_EXACT":
            if missing(r.get("payload_schema","")):contract_gaps.append(("payload_schema",page,uid,op))
            if missing(r.get("persistence_owner","")):contract_gaps.append(("persistence_owner",page,uid,op))
        for field in ["operation","runtime_status","method_path","payload_schema","persistence_owner","runtime_owner"]:
            vals=r["_values"].get(field,[])
            if len(vals)>1:conflicts.append((page,uid,field,vals))
        mp=r.get("method_path","")
        if not missing(mp) and mp not in {"LOCAL LOCAL","NO_PUBLIC_API_BY_AUTHORITY"}:routes[mp].add(op)
    for mp,ops in routes.items():
        if len(ops)>1:route_conflicts.append((page,mp,sorted(ops)))

for kind,arr in [
("OPERATION_MISSING",operation_missing),("OPERATION_PLACEHOLDER",operation_placeholders),
("GATE_GAP",gate_gaps),("PERMISSION_GAP",permission_gaps),("RUNTIME_OWNER_GAP",runtime_owner_gaps),
("CONTRACT_GAP",contract_gaps),("MULTI_VALUE_CONFLICT",conflicts),("ROUTE_OPERATION_CONFLICT",route_conflicts)]:
    if arr:critical.append({"kind":kind,"detail":arr[:200],"count":len(arr)})

report["contract_findings"]={
 "operation_missing":operation_missing,"operation_placeholders":operation_placeholders,
 "gate_gaps":gate_gaps,"permission_gaps":permission_gaps,"runtime_owner_gaps":runtime_owner_gaps,
 "contract_gaps":contract_gaps,"multi_value_conflicts":conflicts,"route_operation_conflicts":route_conflicts,
 "passive_read_bindings_without_operation":passive_reads,"direct_operation_bindings_without_action_uid":direct_ops,
}

# High-confidence table residue audit: exact runtime rows must not contain high placeholder values in binding fields.
binding_residue=[]
for page,fn in PAGE_DOCS.items():
    for r in parse_controls(fn):
        if r.get("runtime_status") not in {"EFFECTFUL_EXACT","READ_EXACT"}:continue
        for field in ["operation","method_path","payload_schema","persistence_owner","runtime_owner","gate","permission"]:
            v=r.get(field,"")
            if v and any(re.search(p,v,re.I) for p in HIGH_PLACEHOLDER_PATTERNS):
                binding_residue.append((page,r["control"],field,v))
if binding_residue:critical.append({"kind":"CANONICAL_BINDING_PLACEHOLDER_RESIDUE","detail":binding_residue[:200],"count":len(binding_residue)})

# Cross-document named UID conflicts across the 18 pages are tracked by page namespace; same UID across same page occurrences handled above.
# Package/root checks.
root_names=subprocess.check_output(["git","-c","core.quotePath=false","ls-tree","--name-only",BASE_HEAD],text=True).splitlines()
non_docx=[x for x in root_names if not x.lower().endswith(".docx")]
if len(root_names)!=31 or non_docx:critical.append({"kind":"PACKAGE_ROOT_SCOPE","detail":{"root_items":len(root_names),"non_docx":non_docx}})

report["package"]={
 "root_items":len(root_names),"docx_files":len(DOCS),"non_docx_root_items":non_docx,
 "total_rendered_pages":total_pages,"total_source_cjk":total_cjk_src,"total_pdf_cjk":total_cjk_pdf,
 "critical_findings":len(critical),"warning_findings":len(warnings),
 "audit_status":"PASS" if not critical else "ISSUES_FOUND",
}
report["content_findings"]=critical
report["warning_candidates"]=warnings
report["binding_placeholder_residue"]=binding_residue

(OUT/"batch58_full_audit_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
# TSV-like concise findings.
lines=["severity\tfile\tkind\tdetail"]
for x in critical:
    lines.append("CRITICAL\t"+str(x.get("file","PACKAGE"))+"\t"+x["kind"]+"\t"+json.dumps(x.get("detail"),ensure_ascii=False,separators=(",",":")))
for x in warnings:
    lines.append("WARNING\t"+str(x.get("file","PACKAGE"))+"\t"+x["kind"]+"\t"+json.dumps(x.get("detail"),ensure_ascii=False,separators=(",",":")))
(OUT/"batch58_findings.tsv").write_text("\n".join(lines)+"\n",encoding="utf-8")
summary=report["package"]|{
 "contract_operation_missing":len(operation_missing),"contract_operation_placeholders":len(operation_placeholders),
 "contract_gate_gaps":len(gate_gaps),"contract_permission_gaps":len(permission_gaps),
 "contract_runtime_owner_gaps":len(runtime_owner_gaps),"contract_gaps":len(contract_gaps),
 "contract_multi_value_conflicts":len(conflicts),"route_operation_conflicts":len(route_conflicts),
 "binding_placeholder_residue":len(binding_residue),"passive_reads":len(passive_reads),"direct_operation_bindings":len(direct_ops),
}
(OUT/"batch58_summary.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH58_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
