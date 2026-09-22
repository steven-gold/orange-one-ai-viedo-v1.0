from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-56-ASSET-FINAL-CONTRACT-REMEDIATION-V1"
BASE_HEAD="fff310848afaec420a7052f83baf035815bbb2aa"
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
ASSET="ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S05:"71ebc53154cccbd39e1eac345496dbe2c6ae0ed0",
ASSET:"b6d3aa5bfa9d1c6f43cf602ba427c442daeca830",
LOGIC:"ef81d0649e0dca073ae92e4c80d323003b400ae1",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":ASSET,
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

TARGETS={
"ASSET-01-BTN-EXECUTE":{
 "operation":"startAssetProductionFlow","route":"POST /v1/asset-runtime-runs","payload":"StartAssetProductionFlowRequest",
 "persistence":"System 05 / ProductionWorkItem + AssetVersion Generation Lineage + Audit",
 "basis":"Starts ASSET production only from exact admitted CORE/Topic/Asset inputs and bindings. Run creation is not output confirmation or provider success."
},
"ASSET-01-TXT-CORRECTION-REQUEST":{
 "operation":"generateCorrectionScriptCandidate","route":"POST /v1/state-commands/correctionscript/generate","payload":"GenerateAssetCorrectionScriptCandidateRequest",
 "persistence":"public.correction_script_versions",
 "basis":"The correction request text participates in the governed correction-script candidate command; candidate/version preserves exact source, finding and context refs and is not approved correction authority."
},
"ASSET-01-BTN-CORRECTION-GENERATE":{
 "operation":"generateCorrectionScriptCandidate","route":"POST /v1/state-commands/correctionscript/generate","payload":"GenerateAssetCorrectionScriptCandidateRequest",
 "persistence":"public.correction_script_versions",
 "basis":"Reuses exact Current state-command route. Generated correction script is a candidate/version only; approval remains separate."
},
"ASSET-01-BTN-CORRECTION-APPROVE":{
 "operation":"approveCorrectionScriptCandidate","route":"POST /v1/state-commands/correctionscript/approve","payload":"ApproveAssetCorrectionScriptCandidateRequest",
 "persistence":"public.correction_script_versions",
 "basis":"Reuses exact Current approval route/persistence. Approval binds exact candidate/version and authorized human decision evidence."
},
"ASSET-01-BTN-CORRECTION-EXECUTE":{
 "operation":"executeAssetCorrection","route":"POST /v1/asset-runtime-runs/{runId}/corrections/execute","payload":"ExecuteAssetCorrectionRequest",
 "persistence":"System 05 / AssetVersion + PatchVersion/Delta Lineage + Audit",
 "basis":"Executes only an approved correction against exact source/version/region/layer refs. Correction creates a new AssetVersion or bounded PatchVersion and never overwrites confirmed source."
},
"ASSET-01-BTN-LAYER-DOC-CREATE":{
 "operation":"createAssetLayerDocument","route":"POST /v1/asset-runtime-runs/{runId}/layer-documents","payload":"CreateAssetLayerDocumentRequest",
 "persistence":"System 05 / CompositionDocument + CompositionVersion",
 "basis":"Creates the logical composition tree and first immutable CompositionVersion from exact source bindings."
},
"ASSET-01-BTN-LAYER-DOC-UPDATE":{
 "operation":"updateAssetLayerDocument","route":"PATCH /v1/asset-runtime-runs/{runId}/layer-documents/{compositionDocumentId}","payload":"UpdateAssetLayerDocumentRequest",
 "persistence":"System 05 / CompositionDocument + New CompositionVersion Lineage",
 "basis":"Updates create a new composition version from expected document/version; historical composition snapshots remain immutable."
},
"ASSET-01-BTN-LAYER-ADD":{
 "operation":"addAssetLayer","route":"POST /v1/asset-runtime-runs/{runId}/layer-documents/{compositionDocumentId}/layers","payload":"AddAssetLayerRequest",
 "persistence":"System 05 / LayerVersion + CompositionVersion",
 "basis":"Adds a new versioned layer and composition snapshot with exact source/ref provenance."
},
"ASSET-01-BTN-LAYER-DELETE":{
 "operation":"deleteAssetLayer","route":"DELETE /v1/asset-runtime-runs/{runId}/layer-documents/{compositionDocumentId}/layers/{layerId}","payload":"DeleteAssetLayerRequest",
 "persistence":"System 05 / LayerVersion Tombstone + CompositionVersion",
 "basis":"Deletion is versioned/tombstoned through a new composition state; prior LayerVersion remains recoverable/auditable."
},
"ASSET-01-BTN-LAYER-DUPLICATE":{
 "operation":"duplicateAssetLayer","route":"POST /v1/asset-runtime-runs/{runId}/layer-documents/{compositionDocumentId}/layers/{layerId}/duplicate","payload":"DuplicateAssetLayerRequest",
 "persistence":"System 05 / LayerVersion + CompositionVersion",
 "basis":"Duplicate creates a new layer identity/version from exact source layer lineage; it does not alias the same mutable row."
},
"ASSET-01-BTN-LAYER-REORDER":{
 "operation":"reorderAssetLayers","route":"POST /v1/asset-runtime-runs/{runId}/layer-documents/{compositionDocumentId}/layers/reorder","payload":"ReorderAssetLayersRequest",
 "persistence":"System 05 / CompositionVersion + Layer Ordering Lineage",
 "basis":"Reorder creates a new composition version from exact expected layer order/version and preserves layer identities."
},
"ASSET-01-CTL-LAYER-PROPERTIES":{
 "operation":"updateAssetLayerProperties","route":"PATCH /v1/asset-runtime-runs/{runId}/layer-documents/{compositionDocumentId}/layers/{layerId}/properties","payload":"UpdateAssetLayerPropertiesRequest",
 "persistence":"System 05 / LayerVersion + CompositionVersion",
 "basis":"Opacity/blend/transform changes create a new LayerVersion/composition snapshot; unaffected layers/content remain preserved."
},
"ASSET-01-CTL-LAYER-MASK":{
 "operation":"updateAssetLayerMask","route":"PUT /v1/asset-runtime-runs/{runId}/layer-documents/{compositionDocumentId}/layers/{layerId}/mask","payload":"UpdateAssetLayerMaskRequest",
 "persistence":"System 05 / LayerVersion + SpatialRegion + CompositionVersion",
 "basis":"Mask/semantic-region update binds an exact SpatialRegion and creates versioned layer/composition state; content outside authorized region must remain preserved."
},
"ASSET-01-BTN-PATCH-CREATE":{
 "operation":"createAssetPatch","route":"POST /v1/asset-runtime-runs/{runId}/layer-documents/{compositionDocumentId}/patches","payload":"CreateAssetPatchRequest",
 "persistence":"System 05 / PatchVersion + SourceBinding",
 "basis":"Patch is a bounded delta against exact source hash/version/region. Source mismatch blocks and the same PatchVersion cannot be applied twice to the same target state."
},
"ASSET-01-BTN-PATCH-PREVIEW":{
 "operation":"previewAssetPatch","route":"POST /v1/asset-runtime-runs/{runId}/layer-documents/{compositionDocumentId}/patches/{patchVersionId}/preview","payload":"PreviewAssetPatchRequest",
 "persistence":"System 05 / PatchVersion Candidate + Preview Evidence (non-applied)",
 "basis":"Preview materializes bounded preview evidence only; it must not alter the canonical composition or count as patch acceptance."
},
"ASSET-01-BTN-PATCH-ACCEPT":{
 "operation":"acceptAssetPatch","route":"POST /v1/asset-runtime-runs/{runId}/layer-documents/{compositionDocumentId}/patches/{patchVersionId}/accept","payload":"AcceptAssetPatchRequest",
 "persistence":"System 05 / PatchVersion + LayerVersion + CompositionVersion",
 "basis":"Acceptance applies exact non-stale PatchVersion once, producing new version lineage while preserving unaffected content."
},
"ASSET-01-BTN-PATCH-REJECT":{
 "operation":"rejectAssetPatch","route":"POST /v1/asset-runtime-runs/{runId}/layer-documents/{compositionDocumentId}/patches/{patchVersionId}/reject","payload":"RejectAssetPatchRequest",
 "persistence":"System 05 / PatchVersion Rejected State + Decision/Audit",
 "basis":"Reject records an explicit decision against exact patch version/evidence and leaves source composition unchanged."
},
"ASSET-01-BTN-RETRY":{
 "operation":"retryAssetTask","route":"POST /v1/asset-runtime-runs/{runId}/tasks/{taskId}/retry","payload":"RetryAssetTaskRequest",
 "persistence":"System 05 / ProductionWorkItem Retry Attempt + Asset Runtime Audit",
 "basis":"Retry binds exact failed task/attempt/version and preserves prior failure evidence; it cannot silently change scope/source binding."
},
"ASSET-01-BTN-EVALUATE":{
 "operation":"evaluateAssetCandidate","route":"POST /v1/asset-runtime-runs/{runId}/candidates/{assetVersionId}/evaluations","payload":"EvaluateAssetCandidateRequest",
 "persistence":"System 05 / EvaluationResult + Asset Evidence/Audit",
 "basis":"Evaluation records exact criteria results/evidence/findings for the candidate AssetVersion. Score alone cannot auto-confirm, lock or handoff."
},
"ASSET-01-BTN-CONFIRM":{
 "operation":"confirmAssetCandidate","route":"POST /v1/asset-runtime-runs/{runId}/candidates/{assetVersionId}/confirm","payload":"ConfirmAssetCandidateRequest",
 "persistence":"System 05 / AssetVersion Confirmed State + Decision/Audit",
 "basis":"Confirmation selects exact candidate version under authorized decision evidence and creates immutable confirmed AssetVersion lineage."
},
"ASSET-01-BTN-RESTORE-AS-NEW":{
 "operation":"restoreAssetVersionAsNewDraft","route":"POST /v1/state-commands/assetversion/restoreasnewdraft","payload":"RestoreAssetVersionAsNewDraftRequest",
 "persistence":"public.asset_version_restore_drafts",
 "basis":"Reuses exact Current route/persistence. Restore always creates a new draft/version lineage from exact source; historical AssetVersion is never overwritten."
},
"ASSET-01-BTN-LOCK":{
 "operation":"lockAssetVersion","route":"POST /v1/state-commands/assetversion/lock","payload":"LockAssetVersionRequest",
 "persistence":"public.production_output_version_locks",
 "basis":"Reuses exact Current route/persistence. Lock binds exact AssetVersion/hash and authorized decision evidence; it does not imply downstream handoff completion."
},
"ASSET-01-BTN-HANDOFF":{
 "operation":"handoffAssetToVideo","route":"POST /v1/asset-runtime-runs/{runId}/handoff-to-video","payload":"HandoffAssetToVideoRequest",
 "persistence":"System 05 / HandoffManifest + AssetVersion Provenance/Audit",
 "basis":"ASSET→VIDEO handoff contains exact AssetVersion UIDs/versions/checksums + layer/patch refs + score/evaluation + rights + continuity + output contract. Successor admission remains separate."
},
}

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        for i,h in enumerate(hs):
            if needle.lower() in h:return i
    return None
def exact_index(headers,name):
    for i,h in enumerate(headers):
        if norm(h).lower()==name.lower():return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
        "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),"permission":hfind(h,"permission","auth resource"),
        "payload_schema":hfind(h,"payload / schema","payload","schema"),"operation":hfind(h,"operation"),
        "method_path":hfind(h,"method / path","method","path"),"runtime_owner":hfind(h,"runtime owner"),
        "persistence_owner":hfind(h,"persistence owner"),"runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()};rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out
def collect_uid_operations(path):
    d=Document(path);out=collections.defaultdict(set)
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ui=hfind(h,"control uid","target uid");oi=hfind(h,"operation")
        if ui is None or oi is None:continue
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ui>=len(vals) or oi>=len(vals):continue
            uid,op=vals[ui],vals[oi]
            if uid and uid not in {"—","-"} and op and op not in {"—","-"}:out[uid].add(op)
    return out
def blob(path):
    b=Path(path).read_bytes();return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():assert blob(f)==s,(f,blob(f),s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""
assert len(TARGETS)==23

pre_page=compose(parse_controls(ASSET));pre_owner=compose(parse_controls(S05));owner_ops=collect_uid_operations(S05)
existing_method={
"ASSET-01-TXT-CORRECTION-REQUEST","ASSET-01-BTN-CORRECTION-GENERATE","ASSET-01-BTN-CORRECTION-APPROVE",
"ASSET-01-BTN-RESTORE-AS-NEW","ASSET-01-BTN-LOCK"}
existing_persistence={"ASSET-01-BTN-CORRECTION-APPROVE","ASSET-01-BTN-RESTORE-AS-NEW","ASSET-01-BTN-LOCK"}
for uid,s in TARGETS.items():
    assert pre_page[uid]["operation"]==s["operation"],("PAGE",uid,pre_page[uid]["operation"],s["operation"])
    assert s["operation"] in owner_ops.get(uid,set()),("OWNER_OPERATION_MISSING",uid,s["operation"],sorted(owner_ops.get(uid,set())))
    for m,label in [(pre_page,"PAGE"),(pre_owner,"OWNER")]:
        r=m[uid];assert r["runtime_status"]=="EFFECTFUL_EXACT",(label,uid,r["runtime_status"])
        assert missing(r["payload_schema"]),(label,uid,"PAYLOAD_ALREADY",r["payload_schema"])
        if uid in existing_method:
            assert r["method_path"]==s["route"],(label,uid,"METHOD",r["method_path"],s["route"])
        else:
            assert missing(r["method_path"]),(label,uid,"METHOD_ALREADY",r["method_path"])
        if uid in existing_persistence:
            assert r["persistence_owner"]==s["persistence"],(label,uid,"PERSIST",r["persistence_owner"],s["persistence"])
        else:
            assert missing(r["persistence_owner"]),(label,uid,"PERSIST_ALREADY",r["persistence_owner"])

def patch(path):
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact_index(h,"Control UID")
        if ci is None:continue
        idx={"method_path":hfind(h,"method / path","method","path"),"payload_schema":hfind(h,"payload / schema","payload","schema"),"persistence_owner":hfind(h,"persistence owner")}
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if uid not in TARGETS:continue
            s=TARGETS[uid];assign={"payload_schema":s["payload"]}
            if uid not in existing_method:assign["method_path"]=s["route"]
            if uid not in existing_persistence:assign["persistence_owner"]=s["persistence"]
            for f,new in assign.items():
                i=idx[f]
                if i is None or i>=len(row.cells):continue
                old=norm(row.cells[i].text)
                if old==new:hits[(uid,f)]+=1;continue
                assert missing(old),(path,uid,f,old,new)
                row.cells[i].text=new;hits[(uid,f)]+=1
    exp={(u,"payload_schema") for u in TARGETS}|{(u,"method_path") for u in TARGETS if u not in existing_method}|{(u,"persistence_owner") for u in TARGETS if u not in existing_persistence}
    assert len(exp)==61,len(exp)
    assert set(hits)==exp,(path,sorted(exp-set(hits)),dict(hits));d.save(path);Document(path);return {f"{u}:{f}":n for (u,f),n in hits.items()}

patches={"SYSTEM05_OWNER":patch(S05),"ASSET_PAGE":patch(ASSET)}

def landscape(d):
    s=d.add_section(WD_SECTION.NEW_PAGE);s.orientation=WD_ORIENT.LANDSCAPE;s.page_width,s.page_height=s.page_height,s.page_width
    s.top_margin=Inches(.33);s.bottom_margin=Inches(.33);s.left_margin=Inches(.28);s.right_margin=Inches(.28)
def table(d,h,rows,fs=3.2):
    t=d.add_table(rows=1,cols=len(h));t.style="Table Grid"
    tr=t.rows[0]._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");tr.append(e)
    for i,x in enumerate(h):
        t.rows[0].cells[i].text=str(x);sh=OxmlElement("w:shd");sh.set(qn("w:fill"),"EDE9FE");t.rows[0].cells[i]._tc.get_or_add_tcPr().append(sh)
    for rr in rows:
        c=t.add_row().cells
        for i,x in enumerate(rr):c[i].text=str(x)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
for path,title,prefix in [(S05,"Batch 56 · ASSET Final Canonical Contract Registry","SYSTEM05"),(ASSET,"Batch 56 · ASSET-01 Final Contract Binding Ledger","PAGE::ASSET-01")]:
    d=Document(path);landscape(d);d.add_heading(title,1)
    d.add_paragraph(f"[{MARK}::{prefix}] Existing correction-script, restore-as-new and version-lock state-command routes are preserved. Eighteen previously undefined ASSET routes are formal System 05 design contracts. Layer/Patch/Composition/Asset versions remain immutable/versioned; preview is non-applied; source hash/version mismatch fails closed; handoff carries exact provenance. No provider/runtime deployment is claimed.")
    rows=[[u,s["operation"],s["route"],s["payload"],s["persistence"],s["basis"]] for u,s in TARGETS.items()]
    table(d,["Control UID","Operation","Method / Path","Payload / Schema","Persistence Owner","Semantics"],rows)
    d.save(path);Document(path)

post_page=compose(parse_controls(ASSET));post_owner=compose(parse_controls(S05))
for uid,s in TARGETS.items():
    for m,label in [(post_page,"PAGE"),(post_owner,"OWNER")]:
        assert m[uid]["method_path"]==s["route"],(label,uid,"METHOD",m[uid]["method_path"],s["route"])
        assert m[uid]["payload_schema"]==s["payload"],(label,uid,"PAYLOAD",m[uid]["payload_schema"],s["payload"])
        assert m[uid]["persistence_owner"]==s["persistence"],(label,uid,"PERSIST",m[uid]["persistence_owner"],s["persistence"])

# Fresh global denominator over all 18 pages. This is the closure assertion, not arithmetic-only.
allm={p:compose(parse_controls(fn)) for p,fn in PAGES.items()}
remaining=[]
for p,m in allm.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if missing(op) or st not in {"EFFECTFUL_EXACT","READ_EXACT"}:continue
        if missing(r.get("method_path","")):remaining.append(("method_path",p,uid,op))
        if st=="EFFECTFUL_EXACT" and missing(r.get("payload_schema","")):remaining.append(("payload_schema",p,uid,op))
        if st=="EFFECTFUL_EXACT" and missing(r.get("persistence_owner","")):remaining.append(("persistence_owner",p,uid,op))
assert remaining==[],remaining[:50]

# ASSET same-route conflict check permits same operation on multiple controls (e.g. correction request text + button).
route_ops=collections.defaultdict(set)
for uid,r in post_page.items():
    mp=r.get("method_path","");op=r.get("operation","")
    if missing(mp) or mp in {"LOCAL LOCAL","NO_PUBLIC_API_BY_AUTHORITY"}:continue
    route_ops[mp].add(op)
collisions={k:sorted(v) for k,v in route_ops.items() if len(v)>1}
assert not collisions,collisions

d=Document(LOGIC);landscape(d);d.add_heading("Batch 56 · ASSET Final Contract Closure",1)
machine={
"marker":MARK,"base_head":BASE_HEAD,"pre_contract_cells":61,"asset_cells_closed":61,
"method_path_closed":18,"payload_schema_closed":23,"persistence_owner_closed":20,
"post_contract_cells":0,"remaining_method_path":0,"remaining_payload_schema":0,"remaining_persistence_owner":0,
"fresh_global_remaining_rows":0,"existing_state_command_routes_reused":5,"new_asset_design_routes":18,
"same_page_route_operation_conflicts":0,"runtime_execution_claimed":False,
}
table(d,["Metric","Value"],[
["Pre denominator","61"],["ASSET cells closed","61 = Method 18 + Payload 23 + Persistence 20"],
["Fresh global post denominator","0"],["Existing exact state-command control bindings reused","5 controls"],
["New ASSET design routes","18"],["Same-route different-operation conflicts","0"],["Runtime execution claimed","False"]],4.0)
d.add_paragraph("BATCH56_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")));d.save(LOGIC);Document(LOGIC)

changed=[S05,ASSET,LOGIC]
Path("__batch56_remediation_report.json").write_text(json.dumps({"machine":machine,"targets":TARGETS,"patches":patches,"changed_docs":changed,"fresh_remaining_rows":remaining,"output_blob_sha":{f:blob(f) for f in changed}},ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH56="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
