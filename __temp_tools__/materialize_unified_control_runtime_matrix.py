from pathlib import Path
import hashlib, json, re
from collections import defaultdict
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('.')
MARK='ACPOS-20260921-UNIFIED-CONTROL-RUNTIME-MATRIX-V2'\nOLD_HEADING='Unified Page / Control / Action / API / DB / Runtime Construction Matrix'
TARGETS={
'01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx':[
 'ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx','ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx'],
'02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx':[
 'ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx','ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx','ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx','ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx'],
'03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx':['ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx'],
'04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx':['ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx'],
'05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx':['ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx'],
'06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx':['ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx','ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx'],
'07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx':['ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx'],
'08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx':['ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx'],
'09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx':['ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx','ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx','ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx','ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx','ACPOS_LOGIN_IDENTITY_ENTRY_Mother_Basic_Design_OPTIMIZED.docx','ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx','ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx'],
}

def n(s): return re.sub(r'\s+',' ',str(s or '').strip())
def k(s): return re.sub(r'[^a-z0-9]+',' ',n(s).lower()).strip()
def slug(s):
    s=re.sub(r'([a-z0-9])([A-Z])',r'\1-\2',n(s))
    return re.sub(r'[^A-Za-z0-9]+','-',s).strip('-').upper()[:64] or 'ACTION'
def rows(t): return [[n(c.text) for c in r.cells] for r in t.rows]
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def page_uid(d,name):
    for t in d.tables:
        for r in rows(t)[:10]:
            if len(r)>1 and k(r[0])=='page uid': return r[1]
    m=re.search(r'ACPOS_([^_]+)',name); return m.group(1) if m else Path(name).stem
def pfx(uid): return re.sub(r'^(workspace:|admin:|identity:)','',uid).replace(':','-').upper()
def header_map(r): return {k(x):i for i,x in enumerate(r)}
def cell(r,hm,*keys):
    for q in keys:
        i=hm.get(k(q))
        if i is not None and i<len(r): return r[i]
    return ''
def repeat(row):
    trPr=row._tr.get_or_add_trPr(); x=OxmlElement('w:tblHeader'); x.set(qn('w:val'),'true'); trPr.append(x)
def shade(c):
    tcPr=c._tc.get_or_add_tcPr(); x=OxmlElement('w:shd'); x.set(qn('w:fill'),'EDE9FE'); tcPr.append(x)
def add_table(doc,hdr,data,fs=6.3):
    t=doc.add_table(rows=1,cols=len(hdr)); t.style='Table Grid'; repeat(t.rows[0])
    for i,v in enumerate(hdr): t.rows[0].cells[i].text=str(v); shade(t.rows[0].cells[i])
    for r in data:
        c=t.add_row().cells
        for i,v in enumerate(r): c[i].text=str(v)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs: run.font.size=Pt(fs)
    return t
def heading(doc,s,l=2): doc.add_heading(s,level=l)

# Build a global operation -> exact method/path registry from all uploaded page docs.
all_sources=sorted(ROOT.glob('ACPOS_*.docx'))
op_paths=defaultdict(list)
source_context=defaultdict(list)
for sp in all_sources:
    d=Document(sp)
    pu=page_uid(d,sp.name)
    for ti,t in enumerate(d.tables,1):
        rr=rows(t)
        if not rr: continue
        hm=header_map(rr[0])
        # direct table header
        oi=next((hm[x] for x in hm if x in {'operation','api operation','operation uid'}),None)
        mi=next((hm[x] for x in hm if x in {'method path','method','api runtime','api route','endpoint'}),None)
        for r in rr[1:]:
            joined=' | '.join(r)
            for i,v in enumerate(r):
                if re.match(r'^(GET|POST|PATCH|PUT|DELETE)\s+/',v,re.I):
                    op=r[i-1] if i>0 else ''
                    if op: op_paths[op].append((v,sp.name,ti))
            if oi is not None and mi is not None and oi<len(r) and mi<len(r):
                op=r[oi]; val=r[mi]
                if op and (re.match(r'^(GET|POST|PATCH|PUT|DELETE)\s+/',val,re.I) or '/v1/' in val):
                    op_paths[op].append((val,sp.name,ti))
            # index exact source context by every camelCase op token and visible action phrase
            for v in r:
                if re.match(r'^[a-z][A-Za-z0-9]{3,}$',v):
                    source_context[v].append((pu,sp.name,ti,joined))

# Canonical shared exact paths. These are already materialized in first closure.
op_paths.update({
 'sendConversationMessage':[('POST /v1/conversations/{conversationId}/messages','2026-09-21 shared closure',0)],
 'stopConversationGeneration':[('POST /v1/conversations/{conversationId}/generation/stop','2026-09-21 shared closure',0)],
 'createSystemChange':[('POST /v1/system-changes','2026-09-21 shared closure',0)],
 'createChangeCandidate':[('POST /v1/system-changes/{changeId}/candidates','2026-09-21 shared closure',0)],
 'runSystemValidation':[('POST /v1/validation-runs','2026-09-21 shared closure',0)],
 'getUiProjection':[('GET /v1/ui-projections/{pageUid}','2026-09-21 shared closure',0)],
})

TYPE_WORDS={'BUTTON','PRIMARY_BUTTON','SECONDARY_BUTTON','TOGGLE','SEARCH','TEXTAREA','SELECT','CHECKBOX','CHECKBOX_GROUP','ROW_ACTION','MENU','LINK','TAB','SEGMENT','ATTACHMENT','TABLE','LIST','READONLY','READ_ONLY','READONLY_CARD','INPUT','CONFIRMATION'}
UI_ONLY_HINT=re.compile(r'noop|view|open|filter|mode|select|detail|tab|drawer|read|search',re.I)

def canonical_path(op, action_uid, effect=''):
    vals=op_paths.get(op,[])
    if vals:
        uniq=[]
        for v,src,ti in vals:
            if v not in [x[0] for x in uniq]: uniq.append((v,src,ti))
        return uniq[0][0],f'SOURCE:{uniq[0][1]}#T{uniq[0][2]}'
    eff=(effect or '').upper()
    opv=(op or '')
    if 'UI_ONLY' in eff or 'CONTEXT_STATE' in eff or 'UI_NAVIGATION' in eff:
        return 'N/A — UI LOCAL ONLY','NO_EFFECTFUL_API'
    readish=('READ' in eff or 'READONLY' in eff or bool(re.match(r'^(get|list|search|compare|preview|fetch|trace|inspect|view)',opv,re.I)))
    if readish:
        return f'GET /v1/queries/{action_uid}', 'NEW_QUERY_AUTHORITY_AFTER_GLOBAL_SOURCE_EXHAUSTION'
    return f'POST /v1/commands/{action_uid}', 'NEW_COMMAND_AUTHORITY_AFTER_GLOBAL_SOURCE_EXHAUSTION'

def db_owner(pu,op,effect=''):
    s=(op+' '+effect).lower()
    mapping=[
      ('conversation','acpos_conversation.messages'),('provider','acpos_runtime.provider_model_profiles / execution lineage'),
      ('route','acpos_runtime.provider_route_decisions'),('credential','secret-manager reference + acpos_runtime.provider_credential_refs'),
      ('blueprint','acpos_core.blueprint_versions'),('script','acpos_core.canonical_script_versions'),('dna','acpos_core.dna_versions'),
      ('handoff','acpos_core.department_handoffs'),('asset','acpos_media.asset_versions'),('layer','acpos_media.layer_documents'),
      ('video','acpos_media.video_versions'),('edit','acpos_media.edit_versions'),('qa','acpos_qa.qa_runs / qa_findings'),
      ('campaign','acpos_enterprise.outreach_campaigns'),('delivery','acpos_enterprise.delivery_attempts'),('opportunity','acpos_enterprise.opportunities'),
      ('signal','acpos_enterprise.market_signals'),('company','acpos_enterprise.companies'),('erp','acpos_erp.connectors / sync_jobs'),
      ('social','acpos_social.account_bindings / market_targets / publish_requests'),('publish','acpos_social.publish_requests / publish_attempts'),
      ('strategy','acpos_strategy.strategy_candidates / reviews / decisions'),('candidate','acpos_system.change_candidates'),
      ('change','acpos_system.system_changes / change_candidates'),('validation','acpos_system.validation_runs'),
      ('permission','acpos_iam.permission_assignments / resource_grants'),('knowledge','acpos_knowledge.knowledge_items / experience_records'),
      ('audit','acpos_audit.audit_events')
    ]
    for x,y in mapping:
        if x in s:return y
    defaults={'CORE-01':'acpos_core.projects','ASSET-01':'acpos_media.asset_versions','VIDEO-01':'acpos_media.video_versions','EDIT-01':'acpos_media.edit_versions','QA-01':'acpos_qa.qa_runs','workspace:STR-01':'acpos_strategy.strategy_candidates','admin:STR-01':'acpos_strategy.strategy_candidates','admin:SYS-01':'acpos_system.system_changes','admin:IAM-01':'acpos_iam.resource_grants','admin:DEV-01':'acpos_enterprise.opportunities','admin:ERP-01':'acpos_erp.connectors','admin:SOC-01':'acpos_social.publish_requests','admin:AIAPI-01':'acpos_runtime.provider_model_profiles','admin:KB-01':'acpos_knowledge.knowledge_items','admin:DB-01':'acpos_system.trace_edges','workspace:INFO-01':'acpos_knowledge.context_candidates','workspace:WB-01':'READ_MODEL / NO_WRITE','identity:login':'acpos_iam.accounts'}
    return defaults.get(pu,'DOMAIN_CANONICAL_OWNER')

def runtime_class(path,op,effect):
    x=(path+' '+op+' '+effect).lower()
    if 'n/a' in x:return 'UI_LOCAL / READ_ONLY'
    if 'conversation' in x:return 'SHARED_CONVERSATION_CORE'
    if 'provider' in x or 'execution-job' in x:return 'SHARED_AI_EXECUTION_RUNTIME'
    if 'layer' in x or 'video' in x:return 'MEDIA_RUNTIME'
    if path.startswith('GET '):return 'READ_PROJECTION_RUNTIME'
    if '/commands/' in path:return 'SHARED_TYPED_COMMAND_RUNTIME'
    return 'DOMAIN_COMMAND_RUNTIME'

def control_records(sp):
    d=Document(sp); pu=page_uid(d,sp.name); pref=pfx(pu)
    recs={}
    action_defs={}; ports={}; function_rows=[]; human_ops={}
    # First pass action/port/function registries.
    for ti,t in enumerate(d.tables,1):
        rr=rows(t)
        if not rr:continue
        hm=header_map(rr[0]); h=set(hm)
        if 'action uid' in h:
            for r in rr[1:]:
                au=cell(r,hm,'Action UID')
                if au:
                    action_defs[au]={
                      'label':cell(r,hm,'Label','Name'),'permission':cell(r,hm,'Permission','Permission UID'),
                      'gate':cell(r,hm,'Gate','Gate UID'),'effect':cell(r,hm,'Effect','Effect Class'),
                      'portop':cell(r,hm,'Port / Owner','Port/Owner','Port','Operation','Owner'),'context':' | '.join(r),'table':ti}
        if 'port uid' in h:
            for r in rr[1:]:
                po=cell(r,hm,'Port UID')
                if po:
                    ports[po]={'operation':cell(r,hm,'Operation'),'path':cell(r,hm,'Method / Path','Method','Path','Endpoint'),'permission':cell(r,hm,'Permission'),'context':' | '.join(r),'table':ti}
        if 'function' in h and any(x in h for x in {'api runtime','api route','operation'}):
            for r in rr[1:]:
                function_rows.append(r)
        if ('control' in h or 'control uid' in h) and 'operation' in h:
            for r in rr[1:]:
                label=cell(r,hm,'Control','Label'); op=cell(r,hm,'Operation')
                if label and op:human_ops[k(label)]=op
    # Second pass every row with explicit control UID.
    uid_rx=re.compile(r'(?<![A-Z0-9])([A-Za-z0-9:-]+-(?:BTN|CTL|INP|TGL|FLD|VIEW|LST|CARD|TBL|SEL|CHK|TAB|CMP)-[A-Za-z0-9-]+)')
    for ti,t in enumerate(d.tables,1):
        rr=rows(t)
        if not rr:continue
        hm=header_map(rr[0]); h=set(hm)
        for r in rr[1:]:
            joined=' | '.join(r)
            uids=uid_rx.findall(joined)
            if not uids:continue
            for uid in uids:
                rec=recs.setdefault(uid,{'page_uid':pu,'source':sp.name,'control_uid':uid,'section':'','type':'','label':'','visible_when':'','enabled_when':'','action_uid':'','operation':'','gate':'','permission':'','effect':'','source_contexts':[]})
                rec['source_contexts'].append(f'T{ti}: {joined}')
                if 'control uid' in h:
                    rec['section']=rec['section'] or cell(r,hm,'Sec','Section','Section UID')
                    rec['type']=rec['type'] or cell(r,hm,'Type','Type / Label','Control Type')
                    rec['label']=rec['label'] or cell(r,hm,'Label','Name','Type / Label')
                    rec['visible_when']=rec['visible_when'] or cell(r,hm,'Visible When','Visibility Condition')
                    rec['permission']=rec['permission'] or cell(r,hm,'Permission','Permission UID')
                    rec['action_uid']=rec['action_uid'] or cell(r,hm,'Action','Action UID')
                    rec['gate']=rec['gate'] or cell(r,hm,'Gate','Gate UID')
                    ao=cell(r,hm,'Action / Operation','Operation','Action / Effect')
                    if ao and not rec['action_uid']:
                        if re.match(r'^[A-Z0-9:-]+-ACT-',ao):rec['action_uid']=ao
                        elif re.match(r'^[a-z][A-Za-z0-9]+$',ao):rec['operation']=ao
                        else: rec['operation']=rec['operation'] or human_ops.get(k(ao),'')
                    rec['enabled_when']=rec['enabled_when'] or cell(r,hm,'State / Permission','Gate / UX Rule','Invariant / Recovery')
                else:
                    # alternate tables keyed by UID
                    rec['visible_when']=rec['visible_when'] or cell(r,hm,'Visible When')
                    rec['permission']=rec['permission'] or cell(r,hm,'Permission')
                    ae=cell(r,hm,'Action / Effect','Action / Operation','Operation')
                    if ae:
                        if re.match(r'^[A-Z0-9:-]+-ACT-',ae): rec['action_uid']=rec['action_uid'] or ae
                        elif re.match(r'^[a-z][A-Za-z0-9]+$',ae): rec['operation']=rec['operation'] or ae
                        else:
                            rec['operation']=rec['operation'] or human_ops.get(k(ae),'')
                            if not rec['operation']:
                                # common SYS phrase -> shared operation mapping
                                phrase=k(ae)
                                phrase_map={'send message':'sendConversationMessage','stop current generation':'stopConversationGeneration','create candidate':'createChangeCandidate','create change request':'createSystemChange','run sandbox test':'runSystemValidation'}
                                rec['operation']=phrase_map.get(phrase,rec['operation'])
                    rec['enabled_when']=rec['enabled_when'] or cell(r,hm,'Invariant / Recovery','Visual State Rule')
    # If source uses a Core Control Registry without UIDs, materialize button/control UID per existing operation.
    for ti,t in enumerate(d.tables,1):
        rr=rows(t)
        if not rr:continue
        hm=header_map(rr[0]); h=set(hm)
        if 'control' in h and 'operation' in h and 'control uid' not in h:
            for r in rr[1:]:
                label=cell(r,hm,'Control'); op=cell(r,hm,'Operation')
                if not label or not op:continue
                uid=f'{pref}-BTN-{slug(op)}'
                rec=recs.setdefault(uid,{'page_uid':pu,'source':sp.name,'control_uid':uid,'section':'SOURCE_REGISTERED_ACTION_DOCK_OR_ROW','type':'BUTTON_OR_ROW_ACTION','label':label,'visible_when':'SOURCE_GATE','enabled_when':cell(r,hm,'Gate / UX Rule'),'action_uid':'','operation':op,'gate':'','permission':'','effect':'','source_contexts':[]})
                rec['source_contexts'].append(f'T{ti}: '+' | '.join(r))
    # Join actions/ports, and finalize exact row.
    out=[]
    for uid,rec in sorted(recs.items()):
        au=rec['action_uid']
        ad=action_defs.get(au,{}) if au else {}
        rec['gate']=rec['gate'] or ad.get('gate','')
        rec['permission']=rec['permission'] or ad.get('permission','')
        rec['effect']=ad.get('effect','')
        po=ad.get('portop','')
        if po in ports:
            pr=ports[po]; rec['operation']=rec['operation'] or pr['operation']; exact_path=pr['path']; path_basis=f'SOURCE:{sp.name}#T{pr["table"]}'
            rec['permission']=rec['permission'] or pr['permission']
        else:
            if not rec['operation'] and po and re.match(r'^[a-z][A-Za-z0-9]+$',po):rec['operation']=po
            exact_path=''; path_basis=''
        if not rec['operation']:
            # Read-only controls get explicit projection binding rather than guessed business command.
            if any(x in (rec['type']+' '+rec['label']).upper() for x in ['READONLY','READ_ONLY','FIELD','LIST','TABLE','CARD','VIEW']):
                rec['operation']='getUiProjection'
                rec['effect']=rec['effect'] or 'READ_ONLY'
            else:
                rec['operation']='UI_LOCAL_STATE'
                rec['effect']=rec['effect'] or 'UI_ONLY'
        if not au:
            if rec['operation']=='UI_LOCAL_STATE':
                au=f'{pref}-ACT-UI-LOCAL-STATE'
            else:
                au=f'{pref}-ACT-{slug(rec["operation"])}'
            rec['action_uid']=au
            action_status='NEW_ACTION_UID_FROM_EXISTING_CONTROL_OR_OPERATION'
        else:
            action_status='EXISTING_ACTION_UID'
        effu=(rec['effect'] or '').upper()
        if rec['operation']=='UI_LOCAL_STATE':
            if 'READ' in effu or 'READONLY' in effu:
                rec['operation']='getUiProjection'
            elif not any(x in effu for x in ['UI_ONLY','CONTEXT_STATE','UI_NAVIGATION']):
                rec['operation']=au
        if exact_path:
            path=exact_path
        else:
            path,path_basis=canonical_path(rec['operation'],au,rec['effect'])
        if not rec['gate']:
            rec['gate']=f'{pref}-GATE-PAGE'
            gate_status='NEW_SHARED_PAGE_GATE_FROM_SOURCE_PAGE_AUTH'
        else:gate_status='EXISTING_GATE'
        if not rec['permission']:
            rec['permission']='SOURCE_PAGE_PERMISSION / SERVER_RECHECK'
        trigger='CLICK'
        typ=(rec['type'] or '').upper()
        if 'INPUT' in typ or 'TEXTAREA' in typ or 'SEARCH' in typ:trigger='INPUT_OR_SUBMIT'
        elif 'TOGGLE' in typ or 'SELECT' in typ or 'CHECK' in typ:trigger='CHANGE'
        payload=f'{rec["action_uid"]} typed schema + source contract fields + actor/scope/target + correlation_id'
        if path.startswith(('POST ','PATCH ','PUT ','DELETE ')):
            payload+=' + expected_version(if versioned) + idempotency_key'
        success='SOURCE_PAGE_STATE_REGISTRY / exact server result'
        failure='SOURCE_PAGE_ERROR_REGISTRY / BLOCKED|ERROR|VERSION_CONFLICT as applicable'
        feedback=rec['enabled_when'] or 'Registered success/error/disabled reason; no fabricated success'
        next_step='SAME_PAGE_CONTEXT'
        xx=(rec['label']+' '+rec['operation']).lower()
        if 'handoff' in xx:
            next_step='SOURCE_HANDOFF_TARGET_WITH_EXACT_CONTEXT'
        elif 'open' in xx or 'nav' in xx:
            next_step='SOURCE_REGISTERED_NAVIGATION_TARGET'
        evidence='acpos_audit.audit_events + correlation_id + operation/result identity'
        ctx=' || '.join(rec['source_contexts'])[:700]
        out.append({
          **rec,'action_uid':au,'action_status':action_status,'gate_status':gate_status,'api_path':path,'api_basis':path_basis,
          'trigger':trigger,'payload':payload,'runtime':runtime_class(path,rec['operation'],rec['effect']),'db_owner':db_owner(pu,rec['operation'],rec['effect']),
          'success_state':success,'failure_state':failure,'ui_feedback':feedback,'next_step':next_step,'evidence':evidence,'source_context':ctx
        })
    return pu,out

report={'marker':MARK,'pages':{},'targets':{}}
def strip_old_matrix(doc):
    body=doc._element.body
    start=False
    for child in list(body):
        if child.tag==qn('w:p'):
            txt=''.join(t.text or '' for t in child.iter() if t.tag==qn('w:t'))
            if OLD_HEADING in txt:
                start=True
        if start and child.tag != qn('w:sectPr'):
            body.remove(child)

for target,sources in TARGETS.items():
    tp=ROOT/target; doc=Document(tp)
    alltxt='\n'.join(p.text for p in doc.paragraphs)
    if MARK in alltxt:
        report['targets'][target]={'status':'ALREADY_PRESENT'}
        continue
    strip_old_matrix(doc)
    doc.add_page_break()
    heading(doc,OLD_HEADING,1)
    doc.add_paragraph(f'[{MARK}] This matrix is the direct construction lookup. Existing source identifiers are preserved. New Action/Gate/Control identities are created only when exhaustive uploaded-source scan found the semantic control/operation but no exact UID/path. Existing exact Method/Path always wins; otherwise effectful operations use POST /v1/commands/{{actionUid}} through the Shared Typed Command Runtime. UI/read-only controls explicitly bind to N/A or projection runtime and MUST NOT create writes.')
    target_rows=0
    for s in sources:
        sp=ROOT/s
        if not sp.exists():
            raise RuntimeError(f'MISSING_SOURCE:{s}')
        pu,recs=control_records(sp)
        report['pages'][pu]={'source':s,'source_sha256':sha(sp),'control_rows':len(recs)}
        heading(doc,f'{pu} — Unified Construction Bindings',2)
        if not recs:
            # Domain/service documents may have operation-only registries. Build from exact API operations.
            sd=Document(sp)
            ops=[]
            seen=set()
            for t in sd.tables:
                rr=rows(t)
                for r in rr:
                    for i,v in enumerate(r):
                        if re.match(r'^(GET|POST|PATCH|PUT|DELETE)\s+/',v,re.I):
                            op=r[i-1] if i>0 else ''
                            if op and (op,v) not in seen:
                                seen.add((op,v)); au=f'{pfx(pu)}-ACT-{slug(op)}'
                                ops.append([pu,'SYSTEM_OR_CONTEXT_ACTION',f'{pfx(pu)}-BTN-{slug(op)}','BUTTON_OR_SYSTEM_TRIGGER','SOURCE_GATE','SOURCE_GATE',op,au,'NEW_UID_FROM_EXISTING_OPERATION',v,'SOURCE_EXACT',runtime_class(v,op,''),db_owner(pu,op,''),'SOURCE_SCHEMA + common envelope','SOURCE_STATE_REGISTRY','SOURCE_ERROR_REGISTRY','Registered result/error','SOURCE_NEXT_OR_SAME_PAGE','Audit/correlation'])
            if ops:
                add_table(doc,['Page','Region','Control UID','Type','Visible','Enabled/Gate','Operation','Action UID','Action UID Status','API','API Basis','Runtime','DB/Provider','Payload','Success','Failure','UI Feedback','Next/Handoff','Evidence'],ops,5.5)
                target_rows+=len(ops)
            else:
                doc.add_paragraph('No user-visible/effectful control registry exists in this source page. Construction MUST treat it as system/read-only/support projection; do not invent a UI control. Existing source tables in the preceding closure remain authoritative.')
            continue
        data=[]
        for r in recs:
            data.append([
              r['page_uid'],r['section'] or 'SOURCE_REGION',r['control_uid'],r['type'] or 'SOURCE_CONTROL',
              r['visible_when'] or 'AUTHORIZED_PAGE_CONTEXT',r['gate'],r['trigger'],r['operation'],r['action_uid'],r['action_status'],
              r['permission'],r['payload'],r['api_path'],r['api_basis'],r['runtime'],r['db_owner'],r['success_state'],r['failure_state'],
              r['ui_feedback'],r['next_step'],r['evidence']
            ])
        add_table(doc,['Page UID','Region','Control UID','Control Type','Visible When','Enabled/Gate','Trigger','Operation','Action UID','Action UID Status','Permission','Required Payload','API Method/Path','API Basis','Runtime','DB/Provider Owner','Success State','Failure State','UI Feedback/Recovery','Next Page/Handoff','Evidence'],data,5.1)
        target_rows+=len(data)
    heading(doc,'Construction Lookup Invariant',2)
    doc.add_paragraph('施工 AI MUST resolve a UI action from exactly one row of this matrix. If a source control is UI_ONLY/READ_ONLY, API/DB write is explicitly N/A and creating a write endpoint is forbidden. If API Basis is NEW_QUERY_AUTHORITY_AFTER_GLOBAL_SOURCE_EXHAUSTION, use only the Shared Typed Query Runtime GET endpoint shown. If API Basis is NEW_COMMAND_AUTHORITY_AFTER_GLOBAL_SOURCE_EXHAUSTION, use only the Shared Typed Command Runtime POST endpoint shown. Page-local alternative endpoints are forbidden.')
    doc.save(tp)
    Document(tp)
    report['targets'][target]={'status':'UPDATED','sha256':sha(tp),'rows':target_rows}

Path('__authority_extract__/UNIFIED_CONSTRUCTION_BINDING_REPORT.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
