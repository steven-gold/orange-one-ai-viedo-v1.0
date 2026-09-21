from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib

MARK="ACPOS-20260921-BATCH-07-CANONICAL-OWNER-PRECEDENCE-V1"
TARGET="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
TARGET_SHA="e23a64df918342da298d9fc192b076473ed3d074"

SYSTEMS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}

def blob_sha(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()
assert blob_sha(TARGET)==TARGET_SHA,(blob_sha(TARGET),TARGET_SHA)

def norm(x): return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def get(vals,i): return "" if i is None or i>=len(vals) else norm(vals[i])
def missing(v): return v in {"","—","-","SOURCE_NOT_DEFINED"}

DISPLAY_ONLY_TYPES={"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"}
def display_only(r):
    typ=r.get("type","").upper()
    return typ in DISPLAY_ONLY_TYPES or any(x in typ for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]
ALIASES={
 "action":[r"^action uid$",r"^action id$",r"^action_uid$",r"^action$"],
 "gate":[r"^gate uid$",r"^gate id$",r"^gate_uid$",r"^gate$",r"^gate / policy$",r"^gate/policy$"],
 "permission":[r"^permission$",r"^permission resource$",r"^auth resource$",r"^authorization resource$",r"^required permission$"],
 "operation":[r"^operation$",r"^operation id$",r"^operation uid$",r"^operation_id$",r"^api operation$",r"^service operation$"],
 "runtime_owner":[r"^runtime owner$",r"^runtime_owner$",r"^registered runtime owner$",r"^runtime service owner$",r"^runtime / owner$",r"^runtime/owner$"],
}
def alias_indices(headers,field):
    out=[]
    for i,h in enumerate(headers):
        hl=norm(h).lower()
        if any(re.match(p,hl) for p in ALIASES[field]):out.append(i)
    return out

def parse_page_rows(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(h,"control uid")
        if ci is None:continue
        idx={
          "control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),"operation":hfind(h,"operation"),
          "runtime_owner":hfind(h,"runtime owner"),"runtime_status":hfind(h,"runtime status")
        }
        for ri,r in enumerate(t.rows[1:],2):
            vals=[c.text for c in r.cells]
            uid=get(vals,ci)
            if not uid or uid in {"—","-"} or not re.search(r"[A-Za-z0-9]",uid):continue
            rec={k:get(vals,idx[k]) for k in FIELDS}
            rec.update({"source":str(path),"table":ti+1,"row":ri,"headers":h,"indices":idx})
            out.append(rec)
    return out

def compose(rows):
    groups=collections.defaultdict(list)
    for r in rows:groups[r["control"]].append(r)
    out={}
    for uid,rs in groups.items():
        ranked=sorted(rs,key=lambda r:sum(1 for k in FIELDS[1:] if not missing(r.get(k,""))),reverse=True)
        base=dict(ranked[0])
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        base["same_uid_rows"]=rs
        out[uid]=base
    return out

def classify(field,row):
    val=row.get(field,"");status=row.get("runtime_status","")
    if not missing(val):return "BOUND"
    if "UI_LOCAL_EXACT" in status:
        if field in {"operation","runtime_owner"}:return "LEGITIMATE_NA_UI_LOCAL"
        if field=="action" and display_only(row):return "LEGITIMATE_NA_UI_LOCAL"
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in status and field=="operation":return "CLOSED_BY_OWNER_LOCAL_COMMAND"
    if "READ_EXACT" in status and display_only(row):
        if field=="action":return "LEGITIMATE_NA_READ_PRESENTATION"
        if field=="operation" and not missing(row.get("runtime_owner","")):return "READ_OWNER_BOUND_OPERATION_UNSPECIFIED"
    if field=="runtime_owner" and any(x in status for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return "KNOWN_RUNTIME_BLOCKER"
    if val=="SOURCE_NOT_DEFINED":return "SOURCE_RESOLUTION_REQUIRED"
    return "DEFINITION_BINDING_GAP"

# Current unresolved denominator: 310 + 1 separately preserved SYS Gate conflict.
queue=[]
for page,fn in PAGES.items():
    rows=parse_page_rows(fn);bm=compose(rows)
    for uid,t in sorted(bm.items()):
        for field in ["action","gate","permission","operation","runtime_owner"]:
            if classify(field,t)!="DEFINITION_BINDING_GAP":continue
            if page=="SYS-01" and uid=="SYS-01-BTN-NAV-OPEN" and field=="gate":continue
            queue.append({
              "page":page,"file":fn,"uid":uid,"field":field,"type":t.get("type",""),
              "label":t.get("label",""),"action":t.get("action",""),"gate":t.get("gate",""),
              "permission":t.get("permission",""),"operation":t.get("operation",""),
              "runtime_owner":t.get("runtime_owner",""),"runtime_status":t.get("runtime_status",""),
              "field_column_present":any(r["indices"].get(field) is not None for r in t["same_uid_rows"])
            })
assert len(queue)==310,len(queue)

# Read all System Normative rows generically.
system_rows=[]
for sf in SYSTEMS:
    d=Document(sf)
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        if not any(headers):continue
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if not any(vals):continue
            system_rows.append({"source":sf,"table":ti+1,"row":ri,"headers":headers,"values":vals})

ledger=[]
counts=collections.Counter()
field_counts=collections.defaultdict(collections.Counter)
page_counts=collections.defaultdict(collections.Counter)
alias_counts=collections.Counter()
alias_unique_rows=[]

for q in queue:
    # exact Control UID owner rows require exact cell equality, never substring.
    owner_rows=[]
    for sr in system_rows:
        if q["uid"] in sr["values"]:
            owner_rows.append(sr)

    owner_vals=collections.defaultdict(list)
    for sr in owner_rows:
        for i in alias_indices(sr["headers"],q["field"]):
            if i>=len(sr["values"]):continue
            v=sr["values"][i]
            if missing(v):continue
            owner_vals[v].append({"source":sr["source"],"table":sr["table"],"row":sr["row"],"header":sr["headers"][i]})

    if len(owner_vals)>1:
        owner_class="CANONICAL_OWNER_ROW_CONFLICT"
    elif len(owner_vals)==1:
        owner_class="CANONICAL_OWNER_ROW_BOUND"
    elif owner_rows:
        owner_class="CANONICAL_OWNER_ROW_FIELD_UNDEFINED"
    else:
        owner_class="CANONICAL_OWNER_ROW_ABSENT"

    # Re-evaluate cross-domain alias candidates only to prove they are non-authoritative.
    keys=[]
    for key in ["action","gate","operation"]:
        v=q.get(key,"")
        if not missing(v):keys.append((key,v))
    alias_vals=collections.defaultdict(list)
    for sr in system_rows:
        matched=[{"key":k,"value":v} for k,v in keys if v in sr["values"]]
        if not matched:continue
        # rows that own this exact Control UID are not cross-domain.
        exact_owner=(q["uid"] in sr["values"])
        for i in alias_indices(sr["headers"],q["field"]):
            if i>=len(sr["values"]):continue
            v=sr["values"][i]
            if missing(v):continue
            alias_vals[v].append({
              "source":sr["source"],"table":sr["table"],"row":sr["row"],
              "header":sr["headers"][i],"matched":matched,"exact_owner":exact_owner
            })
    if len(alias_vals)==1:
        av=next(iter(alias_vals))
        alias_status="UNIQUE_ALIAS_VALUE"
        alias_unique_rows.append((q,av,alias_vals[av]))
    elif len(alias_vals)>1:
        alias_status="ALIAS_VALUE_CONFLICT"
    else:
        alias_status="NO_ALIAS_VALUE"
    alias_counts[alias_status]+=1

    # Owner-row precedence: only an exact owner-row value can authorize remediation.
    if owner_class=="CANONICAL_OWNER_ROW_BOUND":
        disposition="ADMISSIBLE_DIRECT_BINDING"
    elif owner_class=="CANONICAL_OWNER_ROW_CONFLICT":
        disposition="BLOCKED_OWNER_CONFLICT"
    elif owner_class=="CANONICAL_OWNER_ROW_FIELD_UNDEFINED":
        disposition="BLOCKED_OWNER_FIELD_UNDEFINED"
    else:
        disposition="BLOCKED_NO_CANONICAL_OWNER_ROW"

    rec=q|{
      "owner_class":owner_class,
      "owner_row_count":len(owner_rows),
      "owner_sources":sorted(set(r["source"] for r in owner_rows)),
      "owner_values":{k:v for k,v in owner_vals.items()},
      "alias_status":alias_status,
      "alias_values":sorted(alias_vals.keys()),
      "disposition":disposition
    }
    ledger.append(rec)
    counts[owner_class]+=1
    field_counts[q["field"]][owner_class]+=1
    page_counts[q["page"]][owner_class]+=1

assert counts==collections.Counter({
 "CANONICAL_OWNER_ROW_FIELD_UNDEFINED":282,
 "CANONICAL_OWNER_ROW_ABSENT":28
}),counts
assert alias_counts==collections.Counter({
 "UNIQUE_ALIAS_VALUE":46,
 "ALIAS_VALUE_CONFLICT":54,
 "NO_ALIAS_VALUE":210
}),alias_counts
assert sum(1 for r in ledger if r["disposition"]=="ADMISSIBLE_DIRECT_BINDING")==0
assert all(r["page"]=="WB-01" for r in ledger if r["owner_class"]=="CANONICAL_OWNER_ROW_ABSENT")
assert collections.Counter(r["field"] for r in ledger if r["owner_class"]=="CANONICAL_OWNER_ROW_ABSENT")==collections.Counter({"action":14,"gate":14})

# Explicit regression example: shared Action/Operation must not override SYS owner row blank Gate.
ex=[r for r in ledger if r["page"]=="SYS-01" and r["uid"]=="SYS-01-BTN-CR-CREATE" and r["field"]=="gate"]
assert len(ex)==1 and ex[0]["owner_class"]=="CANONICAL_OWNER_ROW_FIELD_UNDEFINED",ex
assert "DEV-01-GATE-MESSAGE-REVIEW" in ex[0]["alias_values"],ex

doc=Document(TARGET)
alltxt="\n".join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])
assert MARK not in alltxt,"BATCH07_ALREADY_PRESENT"

sec=doc.add_section(WD_SECTION.NEW_PAGE)
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width,sec.page_height=sec.page_height,sec.page_width
sec.top_margin=Inches(.35);sec.bottom_margin=Inches(.35);sec.left_margin=Inches(.3);sec.right_margin=Inches(.3)

doc.add_heading("Batch 07 · Canonical Owner-Row Precedence / Cross-Domain Non-Transfer",level=1)
p=doc.add_paragraph()
p.add_run("["+MARK+"] ").bold=True
p.add_run("Canonical binding authority is owned by the exact Control UID owner row. A value found only through a shared Action, Gate, Operation, label, route, or neighboring control MUST NOT be transferred into another control when the exact owner row leaves that field undefined or when no exact owner row exists. This is a fail-closed rule.")

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def table(headers,rows,fs=4.5):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row):c[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for pp in cell.paragraphs:
                for rr in pp.runs:rr.font.size=Pt(fs)
    return t

doc.add_heading("Owner Precedence Rules",level=2)
table(["Rule","Requirement"],[
["OP-01 Exact owner wins","Only the exact Control UID owner row may directly authorize a missing Action/Gate/Permission/Operation/Runtime Owner value."],
["OP-02 Owner blank is authoritative absence","If the exact owner row exists but the target field is blank/—, shared keys from another control or domain cannot fill it."],
["OP-03 No owner row = no inheritance","If no exact Control UID owner row exists, shared Operation/Gate/Action values are evidence for investigation only, not a binding source."],
["OP-04 Cross-domain non-transfer","Page/domain-specific Gate, Permission, Runtime Owner, Operation, and Action values must not transfer across controls merely because one shared key matches."],
["OP-05 Conflict stays blocked","If multiple candidate values exist, preserve conflict; no majority vote or first-match resolution."],
["OP-06 Schema absence is not permission to add","A page table missing the field column does not authorize adding a value unless exact canonical owner authority exists."],
],5.1)

doc.add_heading("Batch 07 Reclassification Result",level=2)
table(["Classification","Count","Disposition"],[
["CANONICAL_OWNER_ROW_FIELD_UNDEFINED",282,"BLOCKED_OWNER_FIELD_UNDEFINED"],
["CANONICAL_OWNER_ROW_ABSENT",28,"BLOCKED_NO_CANONICAL_OWNER_ROW"],
["Unique shared-key alias value observed",46,"NON-AUTHORITATIVE under OP-01/02/03"],
["Shared-key alias conflict",54,"NON-AUTHORITATIVE + conflict"],
["No alias value",210,"No candidate value"],
["Admissible direct binding",0,"No page binding mutation in Batch 07"],
],5.2)

doc.add_heading("Field Breakdown",level=2)
table(["Missing Field","Owner Field Undefined","Owner Row Absent"],[
 [f,field_counts[f].get("CANONICAL_OWNER_ROW_FIELD_UNDEFINED",0),field_counts[f].get("CANONICAL_OWNER_ROW_ABSENT",0)]
 for f in ["action","gate","permission","operation","runtime_owner"]
],5.2)

doc.add_heading("Regression Example / Why Shared-Key Propagation Is Forbidden",level=2)
table(["Control","Exact Owner State","Cross-domain Alias","Correct Result"],[
["SYS-01-BTN-CR-CREATE / Gate","System 09 exact Control UID row exists; Gate is undefined","DEV-01 row sharing createChangeRequest + related action exposes DEV-01-GATE-MESSAGE-REVIEW","KEEP SYS Gate UNRESOLVED. DEV Gate must not be copied into SYS."]
],5.0)

doc.add_heading("310-row Canonical Owner Resolution Ledger",level=2)
table(["Page","Control UID","Missing Field","Page Field Schema","Owner Classification","Owner Source","Alias Status","Alias Values","Disposition"],[
[
 r["page"],r["uid"],r["field"],
 "COLUMN_PRESENT" if r["field_column_present"] else "COLUMN_ABSENT",
 r["owner_class"],
 "; ".join(r["owner_sources"]) if r["owner_sources"] else "NONE",
 r["alias_status"],
 " | ".join(r["alias_values"]) if r["alias_values"] else "NONE",
 r["disposition"]
] for r in ledger
],3.6)

doc.add_heading("WB-01 Canonical Owner Gap",level=2)
table(["Finding","Count","Required next authority action"],[
["WB compact/synthetic Control UID has no exact System owner row; Action missing",14,"Resolve canonical Control UID/owner mapping before any Action binding."],
["WB compact/synthetic Control UID has no exact System owner row; Gate missing",14,"Resolve canonical Control UID/owner mapping before any Gate binding."],
],5.0)

machine={
 "marker":MARK,
 "denominator":310,
 "owner_classification":dict(counts),
 "alias_observation":dict(alias_counts),
 "admissible_direct_binding":0,
 "wb_owner_row_absent":{"action":14,"gate":14},
 "preserved_sys_nav_gate_conflict":True,
 "regression_control":"SYS-01-BTN-CR-CREATE",
 "cross_domain_gate_not_transferred":"DEV-01-GATE-MESSAGE-REVIEW"
}
doc.add_paragraph("BATCH07_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
doc.save(TARGET)
Document(TARGET)

Path("__batch07_owner_precedence_report.json").write_text(json.dumps({
 "machine":machine,"ledger":ledger
},ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH07_OWNER_PRECEDENCE="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
