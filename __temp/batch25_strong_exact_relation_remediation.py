from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-25-STRONG-EXACT-RELATION-REMEDIATION-V1"
BASE_HEAD="035770f8e4d3cea1e8d0df9595cb47b25df53902"
S02="02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S09="09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
CORE="ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
STR="ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
SYS="ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx"
IAM="ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S02:"79f7e465d6f03352337de93eb62cd0f62ba47375",
S09:"5753584e25067569e3e6978e41241774e3de543d",
CORE:"4b967452f4c092276e0b65c95da8e0ec55c80ea9",
STR:"5e1caba99dda7c8d69109ae4b60e9eb1fa782b09",
SYS:"c986ad173890dcf936496933395593b0455ad5c4",
LOGIC:"96e4f93ab86bf77a4b9fe45171ae5fa9cfa6d967",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":CORE,
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":IAM,
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":STR,
"SYS-01":SYS,
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]
BINDINGS={
("CORE-01","CORE-01-FLD-MESSAGE","method_path"):"POST /v1/conversations/{CONVERSATION_ID}/messages",
("CORE-01","CORE-01-BTN-SEND","method_path"):"POST /v1/conversations/{CONVERSATION_ID}/messages",
("STR-01","STR-01-INP-MESSAGE","method_path"):"POST /v1/conversations/{CONVERSATION_ID}/messages",
("STR-01","STR-01-BTN-SEND","method_path"):"POST /v1/conversations/{CONVERSATION_ID}/messages",
("STR-01","STR-01-BTN-STOP","method_path"):"POST /v1/conversations/{CONVERSATION_ID}/generation/stop",
("SYS-01","SYS-01-BTN-SANDBOX-TEST","payload_schema"):"AIAPI Page Operation-specific Form / Provider Profile Field Contract",
}

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        fv={}
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            fv[f]=vals
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        base["_field_values"]=fv;out[uid]=base
    return out
def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()
for f,s in EXPECTED_SHA.items():
    got=blob(f);assert got==s,(f,got,s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

# Source authority checks: unique same-operation field value and same Runtime Status.
s02=parse_controls(S02);s09=parse_controls(S09)
def source_values(rows,op,field,status):
    return sorted(set(r[field] for r in rows if r.get("operation")==op and r.get("runtime_status")==status and not missing(r.get(field,""))))
assert source_values(s02,"sendConversationMessage","method_path","EFFECTFUL_EXACT")==["POST /v1/conversations/{CONVERSATION_ID}/messages"]
assert source_values(s02,"stopConversationGeneration","method_path","EFFECTFUL_EXACT")==["POST /v1/conversations/{CONVERSATION_ID}/generation/stop"]
assert source_values(s09,"runSandboxTest","payload_schema","EFFECTFUL_EXACT")==["AIAPI Page Operation-specific Form / Provider Profile Field Contract"]

pre_maps={p:compose(parse_controls(fn)) for p,fn in PAGES.items()}
for (page,uid,field),value in BINDINGS.items():
    r=pre_maps[page][uid]
    assert r["runtime_status"]=="EFFECTFUL_EXACT",(page,uid,field,r["runtime_status"])
    assert missing(r[field]),(page,uid,field,r[field])
    if "SEND" in uid or uid=="CORE-01-FLD-MESSAGE" or uid=="STR-01-INP-MESSAGE":
        assert r["operation"]=="sendConversationMessage",(page,uid,r["operation"])
    if uid=="STR-01-BTN-STOP": assert r["operation"]=="stopConversationGeneration"
    if uid=="SYS-01-BTN-SANDBOX-TEST": assert r["operation"]=="runSandboxTest"

# Isolate the nine Batch24 unique candidates that are not safe due runtime-status / authority-status mismatch.
mismatches=[
{"page":"IAM-01","uid":"IAM-01-BTN-AUDIT","field":"method_path","operation":"getUiProjection","target_status":"EFFECTFUL_EXACT","source_status":"READ_EXACT","reason":"READ_SOURCE_TARGET_EFFECTFUL_STATUS_MISMATCH"},
{"page":"STR-01","uid":"STR-01-TBL-COMPARE","field":"method_path","operation":"compareCandidates","target_status":"EFFECTFUL_EXACT","source_status":"SPEC_EXACT_RUNTIME_BLOCKED / unspecified","reason":"BLOCKED_SOURCE_TARGET_EFFECTFUL_STATUS_MISMATCH"},
{"page":"STR-01","uid":"STR-01-TBL-COMPARE","field":"payload_schema","operation":"compareCandidates","target_status":"EFFECTFUL_EXACT","source_status":"unspecified","reason":"SOURCE_STATUS_NOT_EQUAL_TARGET"},
{"page":"STR-01","uid":"STR-01-TBL-COMPARE","field":"persistence_owner","operation":"compareCandidates","target_status":"EFFECTFUL_EXACT","source_status":"unspecified","reason":"SOURCE_STATUS_NOT_EQUAL_TARGET_AND_SOURCE_NOT_DEFINED_EMBEDDED"},
{"page":"STR-01","uid":"STR-01-BTN-COMPARE","field":"method_path","operation":"compareCandidates","target_status":"EFFECTFUL_EXACT","source_status":"SPEC_EXACT_RUNTIME_BLOCKED / unspecified","reason":"BLOCKED_SOURCE_TARGET_EFFECTFUL_STATUS_MISMATCH"},
{"page":"STR-01","uid":"STR-01-BTN-COMPARE","field":"payload_schema","operation":"compareCandidates","target_status":"EFFECTFUL_EXACT","source_status":"unspecified","reason":"SOURCE_STATUS_NOT_EQUAL_TARGET"},
{"page":"STR-01","uid":"STR-01-BTN-COMPARE","field":"persistence_owner","operation":"compareCandidates","target_status":"EFFECTFUL_EXACT","source_status":"unspecified","reason":"SOURCE_STATUS_NOT_EQUAL_TARGET_AND_SOURCE_NOT_DEFINED_EMBEDDED"},
{"page":"STR-01","uid":"STR-01-BTN-ADOPT","field":"method_path","operation":"adoptAsContextCandidate","target_status":"EFFECTFUL_EXACT","source_status":"SPEC_EXACT_RUNTIME_BLOCKED / unspecified","reason":"BLOCKED_SOURCE_TARGET_EFFECTFUL_STATUS_MISMATCH"},
{"page":"STR-01","uid":"STR-01-BTN-ADOPT","field":"payload_schema","operation":"adoptAsContextCandidate","target_status":"EFFECTFUL_EXACT","source_status":"SPEC_EXACT_RUNTIME_BLOCKED / unspecified","reason":"BLOCKED_SOURCE_TARGET_EFFECTFUL_STATUS_MISMATCH"},
]
assert len(mismatches)==9

def patch_page(path,page):
    wanted={(uid,field):v for (p,uid,field),v in BINDINGS.items() if p==page}
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=exact_index(h,"Control UID");usi=exact_index(h,"UID Status")
        if ci is None:continue
        idx={"method_path":hfind(h,"method / path","method","path"),"payload_schema":hfind(h,"payload / schema","payload","schema")}
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals):continue
            uid=vals[ci]
            for field in ["method_path","payload_schema"]:
                key=(uid,field)
                if key not in wanted:continue
                fi=idx[field]
                if fi is None: continue
                if usi is not None and usi<len(vals):assert vals[usi]=="EXISTING_SOURCE_UID",(path,uid,vals[usi])
                old=norm(row.cells[fi].text);assert missing(old),(path,uid,field,old)
                row.cells[fi].text=wanted[key];hits[key]+=1
    assert set(hits)==set(wanted),(path,sorted(set(wanted)-set(hits)),dict(hits))
    d.save(path);Document(path)
    return {f"{u}:{fld}": n for (u,fld),n in hits.items()}

patches={
"CORE-01":patch_page(CORE,"CORE-01"),
"STR-01":patch_page(STR,"STR-01"),
"SYS-01":patch_page(SYS,"SYS-01"),
}

def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.33);sec.bottom_margin=Inches(.33);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=4.4):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for rr in rows:
        cells=t.add_row().cells
        for i,v in enumerate(rr):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
    return t

for page,path in [("CORE-01",CORE),("STR-01",STR),("SYS-01",SYS)]:
    d=Document(path);add_landscape(d);d.add_heading("Batch 25 · Strong Exact Relation Binding Ledger",level=1)
    p=d.add_paragraph();p.add_run(f"[{MARK}::PAGE::{page}] ").bold=True
    p.add_run("Only same Canonical Owner + same Operation + identical Runtime Status relations are materialized. No shared-Gate inference, label matching, blocked-source promotion, or cross-owner transfer is permitted.")
    rr=[]
    for (pge,uid,field),value in sorted(BINDINGS.items()):
        if pge==page:
            r=compose(parse_controls(path))[uid]
            rr.append([uid,field,r["operation"],r["runtime_status"],value,"STRONG_EXACT_RELATION_BOUND"])
    add_table(d,["Control UID","Field","Operation","Runtime Status","Bound Value","Result"],rr,4.1)
    d.save(path);Document(path)

# Recompute denominator after six safe bindings.
post_maps={p:compose(parse_controls(fn)) for p,fn in PAGES.items()}
remaining=[]
for page,m in post_maps.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if st in {"EFFECTFUL_EXACT","READ_EXACT"} and not missing(op) and missing(r.get("method_path","")):
            remaining.append(("method_path",page,uid))
        if st=="EFFECTFUL_EXACT" and not missing(op) and missing(r.get("payload_schema","")):
            remaining.append(("payload_schema",page,uid))
        if st=="EFFECTFUL_EXACT" and not missing(op) and missing(r.get("persistence_owner","")):
            remaining.append(("persistence_owner",page,uid))
rc=collections.Counter(f for f,_,_ in remaining)
assert len(remaining)==335,len(remaining)
assert rc=={"method_path":75,"payload_schema":136,"persistence_owner":124},rc
for (page,uid,field),value in BINDINGS.items():
    assert post_maps[page][uid][field]==value,(page,uid,field,post_maps[page][uid][field],value)

logic=Document(LOGIC);add_landscape(logic);logic.add_heading("Batch 25 · Strong Exact Relation Remediation / Runtime-Status Guard",level=1)
p=logic.add_paragraph();p.add_run("["+MARK+"] ").bold=True
p.add_run("Batch 24 found 15 same-owner same-Operation unique candidate values. Batch 25 applies only the six candidates whose source relation has the same Runtime Status as the target and contains no unresolved-source sentinel. Nine candidates are deliberately rejected from automatic materialization because source and target runtime semantics differ or the candidate embeds unresolved authority.")
add_table(logic,["Class","Count","Handling"],[
["Batch 24 unique candidates",15,"Revalidated against Runtime Status and unresolved-source sentinels."],
["Strong exact relations applied",6,"5 Method/Path + 1 Payload/Schema."],
["Runtime/authority mismatch isolated",9,"No Word field value copied; next batch must resolve status/authority first."],
["Pre contract-cell denominator",341,"Method 80 + Payload 137 + Persistence 124."],
["Post contract-cell denominator",335,"Method 75 + Payload 136 + Persistence 124."],
],4.4)
add_table(logic,["Page","Control UID","Field","Operation","Target Status","Source Status / Issue","Decision"],[[x["page"],x["uid"],x["field"],x["operation"],x["target_status"],x["source_status"],x["reason"]] for x in mismatches],3.8)
machine={
"marker":MARK,"base_head":BASE_HEAD,"pre_contract_cells":341,"strong_exact_candidates_applied":6,
"method_path_bound":5,"payload_schema_bound":1,"persistence_owner_bound":0,
"runtime_authority_mismatch_isolated":9,"post_contract_cells":335,
"remaining_method_path":75,"remaining_payload_schema":136,"remaining_persistence_owner":124,
"runtime_execution_claimed":False,
}
logic.add_paragraph("BATCH25_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

changed=[CORE,STR,SYS,LOGIC]
report={"machine":machine,"bindings":[{"page":p,"uid":u,"field":f,"value":v} for (p,u,f),v in BINDINGS.items()],
"mismatches":mismatches,"patches":patches,"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed}}
Path("__batch25_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH25="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
