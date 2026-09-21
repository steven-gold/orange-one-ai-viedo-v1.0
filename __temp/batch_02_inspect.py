from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import hashlib, json

MARK="ACPOS-20260921-BATCH-02-CONSUMER-CONTRACT-CLOSURE-V1"
FILES={
"SOC":("ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx","154ca20b3592728f11f475ee8dc8c22466a02524"),
"EDIT":("ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","334a37b3ca39f3356f6c9e83ef8572be409ba35b"),
"ADMIN_STR":("ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx","c3993113999e839380498839a72797d28d671c2a"),
"S05":("05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx","8e3c4e7c4ba36a094293bcbe0d56ee05c1d2e6b9"),
"S08":("08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx","95102b7a6c4496355cdb6d2ce069d0125c0cb5ef"),
}

def blob_sha(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for key,(fn,expected) in FILES.items():
    actual=blob_sha(fn)
    if actual!=expected:
        raise SystemExit(f"INPUT_CHANGED:{fn}:{actual}!={expected}")

def all_text(doc):
    vals=[p.text for p in doc.paragraphs]
    vals += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join(vals)

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr()
    e=OxmlElement("w:tblHeader"); e.set(qn("w:val"),"true"); trPr.append(e)

def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr()
    e=OxmlElement("w:shd"); e.set(qn("w:fill"),fill); tcPr.append(e)

def table(doc,headers,rows,fs=5.8):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; repeat_header(t.rows[0])
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=str(h); shade(t.rows[0].cells[i])
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs: r.font.size=Pt(fs)
    return t

def landscape(doc):
    s=doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation=WD_ORIENT.LANDSCAPE
    s.page_width,s.page_height=s.page_height,s.page_width
    s.top_margin=Inches(.4); s.bottom_margin=Inches(.4)
    s.left_margin=Inches(.35); s.right_margin=Inches(.35)

def head(doc,title):
    landscape(doc)
    doc.add_heading(title,level=1)
    p=doc.add_paragraph()
    p.add_run(f"[{MARK}] ").bold=True
    p.add_run("本節只補齊既有 Current source 已可證明的設計契約。既有 Control / Action / Gate / Permission / Owner / Operation / Route 優先；SOURCE_NOT_DEFINED、NO_PUBLIC_API_ID_IN_CURRENT_AUTHORITY、SPEC_EXACT_RUNTIME_BLOCKED 不得由 Word 自行發明或升格為 Runtime 完成。")

QA_FIELDS=[
("exact_content_ref","Exact reviewed content/release candidate ref; alias/latest-only forbidden."),
("output_version_ref","Exact reviewed content/output version."),
("artifact_checksum","Must match the QA-reviewed artifact/content manifest."),
("criteria_profile_uid","Applicable governed QA profile."),
("criteria_version_id","Exact APPROVED criteria version."),
("quality_gate_policy_version","Exact gate policy version paired with the scorecard."),
("scorecard_id","Fresh scorecard for the exact version."),
("qa_gate_status","Must equal PASS for governed publish eligibility."),
("required_checks","All applicable hard checks PASS, including Rights/Policy/Provenance/Checksum where required."),
("open_blocking_findings","Must equal 0."),
("evidence_refs","Required-check / review / recheck / manual-case evidence lineage."),
]

def patch_soc(fn):
    d=Document(fn)
    if MARK in all_text(d): return
    head(d,"SOC-01 Publish QA Eligibility Consumer Contract")
    d.add_paragraph("Existing publish authority remains unchanged: SOC-01-BTN-PUBLISH → SOC-01-ACT-PUBLISH-REQUEST → SOC-01-GATE-PUBLISH → SOC-01-PERM-PUBLISH → requestSocialTargetPublish → POST /v1/social/targets/{targetId}/publish → PublishingService. This batch adds the missing governed QA eligibility bindings to that SAME publish gate; it does not create a second publish API.")
    table(d,["QA eligibility binding","Required contract"],QA_FIELDS,5.8)
    table(d,["Gate condition","Required behavior"],[
        ["Exact target/content/account/policy valid","Continue to QA eligibility validation."],
        ["Missing exact version/checksum/criteria/scorecard/evidence","BLOCK before requestSocialTargetPublish."],
        ["Artifact/content changed after QA","Previous scorecard becomes stale; BLOCK until fresh recheck."],
        ["qa_gate_status != PASS","BLOCK."],
        ["required hard check FAIL or open_blocking_findings > 0","BLOCK regardless of aggregate score or business approval."],
        ["Eligibility PASS","May invoke existing requestSocialTargetPublish only; local request acceptance becomes PENDING_EXTERNAL, not POSTED."],
        ["External result","Only real adapter callback/result can transition external publish truth to POSTED/failed/other recorded state."],
    ],5.7)
    d.save(fn)

EDIT_ROWS=[
("EDIT-01-BTN-ADD-MEDIA","EDIT-01-ACT-MEDIA-IMPORT","EDIT-01-GATE-STANDALONE","EDITING_UPLOAD","ASSET","media_ref/file_ref, standalone_session_ref, media_type, source_metadata","Import accepted into owner-scoped session media; projection refreshes.","Validation/import failure; preserve source and surface error."),
("EDIT-01-BTN-API-CANCEL","EDIT-01-ACT-API-CANCEL","EDIT-01-GATE-CORRECTION","EDITING_API_EXECUTE","EDITING","job_ref, reason?, correlation_id","Eligible job enters CANCEL_REQUESTED/CANCELLED according to owner truth.","Non-cancellable/terminal/unknown job → no fake cancellation."),
("EDIT-01-BTN-CORR-EXECUTE","EDIT-01-ACT-CORRECTION-EXECUTE","EDIT-01-GATE-CORRECTION-TRANSLATED","EDITING_API_EXECUTE","EDITING/API","correction_candidate_ref, exact_working_version, affected_range, evidence_refs","Owner starts governed correction execution and returns job/version ref.","Missing translated candidate/version/evidence → BLOCK; owner error retained."),
("EDIT-01-BTN-CORR-GENERATE","EDIT-01-ACT-CORRECTION-CANDIDATE-GENERATE","EDIT-01-GATE-CORRECTION","EDITING_API_EXECUTE","EDITING/AI","finding_refs, exact_working_version, affected_scope, requirement_context","New correction-script candidate ref produced; no direct mutation of final output.","No finding/context/owner response → BLOCK/FAIL without fabricated candidate."),
("EDIT-01-BTN-EVAL-RECHECK-FULL","EDIT-01-ACT-EVAL-RECHECK-FULL","EDIT-01-GATE-EVALUATION-RUN","EDITING_STAGE_EVALUATION","EDITING/EVALUATION","exact_timeline_version, criteria_version_id, policy_version, evidence_context","Fresh full-Timeline scorecard/recheck ref.","Stale version/missing criteria/evidence → BLOCK."),
("EDIT-01-BTN-EVAL-RECHECK-SELECTED","EDIT-01-ACT-EVAL-RECHECK-SELECTED","EDIT-01-GATE-EVALUATION-RUN","EDITING_STAGE_EVALUATION","EDITING/EVALUATION","exact_timeline_version, selected_finding_refs, selected_range_refs, criteria_version_id","Fresh selected-scope recheck; only eligible findings may close.","Selection/version mismatch → no finding closure."),
("EDIT-01-BTN-LIPSYNC-RETRY","EDIT-01-ACT-LIPSYNC-RETRY-RANGE","EDIT-01-GATE-DIALOGUE-SYNC","EDITING_VOICE_EXECUTE","VOICE/AUDIO","dialogue_binding_ref, selected_range, voice_version_ref, timing_fingerprint","New lipsync attempt/job ref bound to exact dialogue timing input.","Stale voice/timing fingerprint/range mismatch → BLOCK."),
("EDIT-01-BTN-RETURN-BLUEPRINT","EDIT-01-ACT-RETURN-BLUEPRINT","EDIT-01-GATE-BLUEPRINT-RETURN","EDITING_USE","EDITING/ORCHESTRATION","exact_blueprint_ref, issue_refs, return_reason, evidence_refs","Return request/audit event routes to canonical upstream owner; EDIT remains on current safe version.","Missing exact blueprint/issue/evidence → BLOCK; no direct upstream overwrite."),
("EDIT-01-BTN-UPLOAD-SPEC","EDIT-01-ACT-STANDALONE-SPEC-IMPORT","EDIT-01-GATE-STANDALONE","EDITING_UPLOAD","EDITING","standalone_session_ref, spec_file_ref/text_ref, content_type, checksum?","Spec/script becomes owner-scoped standalone input ref; no Project/Blueprint claim.","Invalid/unsupported input → no import and no fabricated project lineage."),
("EDIT-01-BTN-VOICE-GENERATE","EDIT-01-ACT-VOICE-GENERATE","EDIT-01-GATE-STANDALONE","EDITING_VOICE_EXECUTE","VOICE/AUDIO","standalone_session_ref, script/spec_ref, voice_config_ref, segment_scope","Voice owner returns generation job/output ref for standalone session.","Missing script/config/owner response → BLOCK/FAIL; no fake voice output."),
]

def edit_invocation_section(d,title):
    head(d,title)
    d.add_paragraph("Current registry truth is 22 Ports / 20 unique operations. The ten rows below are not assigned a new public API. Their existing source status remains OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE with NO_PUBLIC_API_ID_IN_CURRENT_AUTHORITY. Invocation is therefore defined as an owner-orchestrated local command envelope keyed by the EXISTING Action UID.")
    table(d,["Control UID","Existing Action UID","Gate","Permission","Owner","Local command minimum payload","Success contract","Failure / fail-closed"],EDIT_ROWS,4.9)
    table(d,["Invocation invariant","Rule"],[
        ["Command identity","command_id = existing Action UID; control_uid + gate + permission + owner must be preserved."],
        ["Envelope","action_uid, control_uid, exact_context_ref, correlation_id, idempotency_key when effectful/retryable, operation-specific payload, evidence_refs when governed."],
        ["No invented public API","NO_PUBLIC_API_ID_IN_CURRENT_AUTHORITY remains true. UI/implementation MUST NOT synthesize REST/GraphQL endpoints from this Word."],
        ["Authorization","Existing permission must be checked before owner dispatch; UI visibility is not authorization."],
        ["State truth","Only owner/runtime response may advance state or produce output/job/version refs."],
        ["Audit","Persist action/control/context/correlation + success/failure + exact output/job/version refs where applicable."],
        ["Implementation status","Design invocation contract is CLOSED; callable Runtime/Production acceptance remains NOT_EXECUTED / IMPLEMENTATION_REQUIRED until fresh implementation evidence exists."],
        ["Historical count","Any 15 ports / 13 operations derived_validation value is stale historical metadata; Current Word construction denominator is 22 ports / 20 unique operations."],
    ],5.5)

def patch_edit(fn):
    d=Document(fn)
    if MARK in all_text(d): return
    edit_invocation_section(d,"EDIT-01 Owner-Orchestrated Local Command Invocation Closure")
    d.save(fn)

STR_ROWS=[
("CTRL-ADMIN-STR-03-ACT-01-ACT-DRAFT-SAVE","saveDraft","POST /v1/drafts","SaveDraftRequest","admin:STR-03::ACT-DRAFT-SAVE","IAM_RUNTIME","SOURCE_NOT_DEFINED","Persistence owner must be canonically bound + runtime authorization/persistence tested.","SPEC_EXACT_RUNTIME_BLOCKED"),
("CTRL-ADMIN-STR-04-ACT-03-ACT-EXPORT","exportProjection","POST /v1/exports","ExportProjectionRequest","admin:STR-04::ACT-EXPORT","INFO_COMMAND_RUNTIME","SOURCE_NOT_DEFINED","BLOCKED_NO_CANONICAL_EXPORT_PERSISTENCE_OWNER must be resolved + tested.","SPEC_EXACT_RUNTIME_BLOCKED"),
("CTRL-ADMIN-STR-05-ACT-01-ACT-CANDIDATE-CREATE","createCandidate","POST /v1/candidates","CreateCandidateRequest","admin:STR-05::ACT-CANDIDATE-CREATE","SOURCE_NOT_DEFINED","SOURCE_NOT_DEFINED","Canonical runtime + persistence owner must be resolved; permission must be enforced there.","SPEC_EXACT_RUNTIME_BLOCKED"),
("CTRL-ADMIN-STR-06-ACT-01-ACT-CANDIDATE-COMPARE","compareCandidates","GET /v1/candidates/compare","No-form read contract","admin:STR-06::ACT-CANDIDATE-COMPARE","SOURCE_NOT_DEFINED","No mutation expected; read/runtime owner still SOURCE_NOT_DEFINED","Canonical read/runtime owner + authorization enforcement must be resolved/tested.","SPEC_EXACT_RUNTIME_BLOCKED"),
("CTRL-ADMIN-STR-06-ACT-02-ACT-CANDIDATE-DECIDE","rejectStrategyCandidate","POST /v1/ai/strategy/candidates/{strategyCandidateId}/reject","RejectStrategyCandidateRequest","admin:STR-06::ACT-CANDIDATE-DECIDE","SOURCE_NOT_DEFINED","SOURCE_NOT_DEFINED","Canonical runtime + persistence owner and state-transition/audit handling must be resolved/tested.","SPEC_EXACT_RUNTIME_BLOCKED"),
("CTRL-ADMIN-STR-06-ACT-03-ACT-ADOPT-CONTEXT","adoptAsContextCandidate","POST /v1/state-commands/strategycandidate/adoptascontextcandidate","AdoptAsContextCandidateRequest","admin:STR-06::ACT-ADOPT-CONTEXT","SOURCE_NOT_DEFINED","SOURCE_NOT_DEFINED","Canonical runtime + persistence/context-state owner and authorization enforcement must be resolved/tested.","SPEC_EXACT_RUNTIME_BLOCKED"),
]

def str_boundary_section(d,title):
    head(d,title)
    d.add_paragraph("These six operations already have exact source-preserving control/operation contracts. This closure makes the missing runtime/persistence boundary explicit instead of inventing owners. Existing permission resources are authoritative design inputs, but enablement remains forbidden until the runtime enforcement path is materially implemented and tested.")
    table(d,["Control UID","Operation","Method / Path","Payload","Permission","Current Runtime Owner","Current Persistence Owner","Required closure before enablement","Current status"],STR_ROWS,4.7)
    table(d,["Rule","Required behavior"],[
        ["Owner resolution","Search and bind existing canonical runtime/persistence owner first. Do not create a parallel Strategy runtime merely to clear the table."],
        ["Authorization","Existing permission resource must be enforced at the materialized runtime boundary; a Word permission cell alone is not runtime evidence."],
        ["Persistence","Mutation operations must identify canonical persistence owner/transaction boundary before enablement."],
        ["Read-only compare","GET comparison must still have canonical runtime/read owner and authorization evidence; absence of mutation does not waive owner binding."],
        ["UI state","Keep blocked controls disabled/fail-closed. Do not treat exact route/payload as implementation completion."],
        ["Fresh evidence","Only current-revision implementation + authorization + persistence/read + audit validation may change SPEC_EXACT_RUNTIME_BLOCKED."],
    ],5.5)

def patch_admin(fn):
    d=Document(fn)
    if MARK in all_text(d): return
    str_boundary_section(d,"ADMIN STR Runtime / Persistence / Authorization Boundary Closure")
    d.save(fn)

def patch_s05(fn):
    d=Document(fn)
    if MARK in all_text(d): return
    edit_invocation_section(d,"System 05 Synchronization · EDIT Owner-Orchestrated Invocation Contract")
    d.add_paragraph("System 05 must consume the same 10-row EDIT invocation semantics. It MUST NOT replace these local owner-orchestrated commands with invented public API IDs. Existing exact API-backed Creative Production operations elsewhere in System 05 remain unchanged.")
    d.save(fn)

def patch_s08(fn):
    d=Document(fn)
    if MARK in all_text(d): return
    str_boundary_section(d,"System 08 Synchronization · Strategy Admin Runtime Boundary")
    d.add_paragraph("Front-office STR-01 operations that are already EFFECTFUL_EXACT remain distinct from the six ADMIN STR rows above. A front-office shared-service binding is not evidence that the admin route has a materialized runtime/persistence owner. Do not transfer completion credit across those boundaries.")
    d.save(fn)

patch_soc(FILES["SOC"][0])
patch_edit(FILES["EDIT"][0])
patch_admin(FILES["ADMIN_STR"][0])
patch_s05(FILES["S05"][0])
patch_s08(FILES["S08"][0])

checks={
FILES["SOC"][0]:["SOC-01 Publish QA Eligibility Consumer Contract","requestSocialTargetPublish","artifact_checksum","criteria_version_id","scorecard_id","open_blocking_findings","PENDING_EXTERNAL"],
FILES["EDIT"][0]:["EDIT-01 Owner-Orchestrated Local Command Invocation Closure","22 Ports / 20 unique operations","OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE","EDIT-01-BTN-ADD-MEDIA","EDIT-01-BTN-VOICE-GENERATE","NO_PUBLIC_API_ID_IN_CURRENT_AUTHORITY","NOT_EXECUTED / IMPLEMENTATION_REQUIRED"],
FILES["ADMIN_STR"][0]:["ADMIN STR Runtime / Persistence / Authorization Boundary Closure","saveDraft","exportProjection","createCandidate","compareCandidates","rejectStrategyCandidate","adoptAsContextCandidate","SPEC_EXACT_RUNTIME_BLOCKED","SOURCE_NOT_DEFINED"],
FILES["S05"][0]:["System 05 Synchronization","EDIT-01-BTN-ADD-MEDIA","EDIT-01-BTN-VOICE-GENERATE","NO_PUBLIC_API_ID_IN_CURRENT_AUTHORITY"],
FILES["S08"][0]:["System 08 Synchronization","saveDraft","exportProjection","createCandidate","compareCandidates","rejectStrategyCandidate","adoptAsContextCandidate","SPEC_EXACT_RUNTIME_BLOCKED"],
}
for fn,needles in checks.items():
    t=all_text(Document(fn))
    if MARK not in t: raise SystemExit(f"MARKER_MISSING:{fn}")
    for n in needles:
        if n not in t: raise SystemExit(f"SEMANTIC_MISSING:{fn}:{n}")

print(json.dumps({"marker":MARK,"files":[v[0] for v in FILES.values()],"status":"SEMANTIC_PRECHECK_PASS"},ensure_ascii=False,indent=2))
