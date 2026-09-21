from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import hashlib,re,json

WB="ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED={
 WB:"64df8aae2c94f26607e47189f24a7416b13fd964",
 LOGIC:"39933113a0bf800991cecbe1aa4b04dd0cd93066",
}
MARK="ACPOS-20260921-BATCH-11-WB-SYNTHETIC-DENOMINATOR-CLEANUP-V1"

def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

def text_snapshot(doc):
    return "\n".join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])

def set_cell_margins(cell,top=20,bottom=20,left=30,right=30):
    tc=cell._tc
    tcPr=tc.get_or_add_tcPr()
    tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar=OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m,v in [("top",top),("bottom",bottom),("left",left),("right",right)]:
        node=tcMar.find(qn("w:"+m))
        if node is None:
            node=OxmlElement("w:"+m)
            tcMar.append(node)
        node.set(qn("w:w"),str(v))
        node.set(qn("w:type"),"dxa")

for f,s in EXPECTED.items():
    assert blob(f)==s,(f,blob(f),s)

# WB layout-only compaction.
wb=Document(WB)
wb_before=text_snapshot(wb)
target=None
for t in wb.tables:
    if not t.rows: continue
    headers=[c.text.strip() for c in t.rows[0].cells]
    if headers[:2]==["Compact Alias","Canonical UID"] and "Disposition" in headers:
        target=t
        break
assert target is not None,"BATCH11_MAPPING_TABLE_NOT_FOUND"

for ri,row in enumerate(target.rows):
    for cell in row.cells:
        set_cell_margins(cell,top=12,bottom=12,left=24,right=24)
        for p in cell.paragraphs:
            p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.0
            for r in p.runs:
                r.font.size=Pt(4.6 if ri else 4.8)

# Compact the trailing explanatory paragraph without changing text.
for p in wb.paragraphs:
    if p.text.startswith("Action UID remains"):
        p.paragraph_format.space_before=Pt(2)
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1.0
        for r in p.runs:r.font.size=Pt(8)

# Tighten only the final Batch-11 section.
sec=wb.sections[-1]
sec.top_margin=Inches(.25);sec.bottom_margin=Inches(.25);sec.left_margin=Inches(.3);sec.right_margin=Inches(.3)
wb.save(WB)
wb2=Document(WB)
assert text_snapshot(wb2)==wb_before,"WB_TEXT_CHANGED_DURING_LAYOUT_ONLY_FIX"

# System Logic layout-only compaction.
logic=Document(LOGIC)
logic_before=text_snapshot(logic)
found=False
for p in logic.paragraphs:
    if p.text.startswith("BATCH11_MACHINE_JSON="):
        found=True
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1.0
        for r in p.runs:
            r.font.size=Pt(4)
        break
assert found,"BATCH11_MACHINE_JSON_NOT_FOUND"
sec=logic.sections[-1]
sec.top_margin=Inches(.28);sec.bottom_margin=Inches(.22);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
logic.save(LOGIC)
logic2=Document(LOGIC)
assert text_snapshot(logic2)==logic_before,"LOGIC_TEXT_CHANGED_DURING_LAYOUT_ONLY_FIX"

report={
 "marker":MARK,
 "layout_only":True,
 "text_unchanged":True,
 "wb_blob":blob(WB),
 "logic_blob":blob(LOGIC),
}
Path("__batch11_layout_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH11_LAYOUT="+json.dumps(report,ensure_ascii=False,sort_keys=True))
