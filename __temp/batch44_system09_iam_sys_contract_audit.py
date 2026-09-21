from pathlib import Path
from docx import Document
import json,re,subprocess

BASE_HEAD="05ccc1e130d0369039bb249a213d29bf6ffbdbfd"
S09="09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
IAM="ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx"
SYS="ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx"
OPS=["saveDraft","validateDraft","previewAuthorizationImpact","createCandidate","createChangeRequest","runSandboxTest"]
KW=["draft","validate","authorization","permission","role","policy","impact","candidate","change request","system change","sandbox","persistence","persist","audit","evidence","approval","preview","change planner","system engineer","runtime"]

assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
rx=re.compile("|".join(re.escape(x) for x in OPS+KW),re.I)

def collect(path):
    d=Document(path);out={"paragraphs":[],"rows":[]}
    for i,p in enumerate(d.paragraphs):
        t=norm(p.text)
        if t and rx.search(t):out["paragraphs"].append({"index":i,"text":t[:3500]})
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];joined=" | ".join(vals)
            if joined and rx.search(joined):
                out["rows"].append({"table":ti+1,"row":ri,"headers":h,"values":vals,"joined":joined[:6000]})
    return out

owner=collect(S09);iam=collect(IAM);sys=collect(SYS)
summary={
"operations":6,
"owner_paragraph_hits":len(owner["paragraphs"]),"owner_table_hits":len(owner["rows"]),
"iam_paragraph_hits":len(iam["paragraphs"]),"iam_table_hits":len(iam["rows"]),
"sys_paragraph_hits":len(sys["paragraphs"]),"sys_table_hits":len(sys["rows"]),
}
report={"marker":"ACPOS-20260922-BATCH-44-SYSTEM09-IAM-SYS-CONTRACT-DESIGN-AUDIT-V1","base_head":BASE_HEAD,
"summary":summary,"operations":OPS,"owner_evidence":owner,"iam_evidence":iam,"sys_evidence":sys}
Path("__batch44_system09_iam_sys_contract_audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
Path("__batch44_summary.txt").write_text(json.dumps(summary,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print("BATCH44_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
