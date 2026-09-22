from pathlib import Path
from docx import Document
import json,re,hashlib,subprocess,collections

MARK="ACPOS-20260922-BATCH-63-STALE-COMPOSITE-LEXICAL-CLEANUP-V1"
BASE_HEAD="d3eeb687015fe1f0220fce0a99629e129e5e5d39"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
EDIT="ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
EXPECTED_LOGIC_SHA="bdd7ea86fb6d26f1764769e172bf67e9bc4308e2"
TARGET="EDIT-01-BTN-FLOW-START"
ACTION="EDIT-01-ACT-FLOW-START-CONTINUE"
OP="dispatchEditingFlowStartContinue"
METHOD="NO_PUBLIC_API_BY_AUTHORITY"
PAYLOAD="EditingFlowStartContinueDispatchContext"
PERSIST="DELEGATED_TO_SELECTED_EDIT_PORT_PERSISTENCE"
GATE="EDIT-01-GATE-CONTEXT-INTEGRITY"
UI_PERMISSION="EDITING_USE"
RUNTIME_OWNER="EDITING/ORCHESTRATION"
PORTS={
"EDIT-01-PORT-EDIT-RUN-CREATE":("createEditingRuntimeRun","POST /v1/editing-runtime-runs","editing.task.execute"),
"EDIT-01-PORT-ASSEMBLY-COMPLETE":("completeAssembly","POST /v1/editing-runtime-runs/{runId}/assembly","editing.task.execute"),
}
REPLACEMENTS={
"createEditingRuntimeRun | completeAssembly":"REMOVED_COMPOSITE_OPERATION_BINDING",
"POST /v1/editing-runtime-runs | POST /v1/editing-runtime-runs/{runId}/assembly":"REMOVED_COMPOSITE_ROUTE_BINDING",
"CreateEditingRuntimeRunRequest | CompleteEditingAssemblyRequest":"REMOVED_COMPOSITE_PAYLOAD_BINDING",
}
LEGACY_RE=re.compile(r"(?<![A-Za-z0-9])R(?:9(?:\.0\.1)?|5)(?![A-Za-z0-9])",re.I)
AMBIG_LOCATOR_RE=re.compile(r"\bt\d+/r(?:5|9)\b",re.I)

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED","UNCHANGED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def blob(path):
    b=Path(path).read_bytes();return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()
def full_text(path):
    d=Document(path);parts=[]
    parts.extend(norm(p.text) for p in d.paragraphs if norm(p.text))
    for t in d.tables:
        for row in t.rows:parts.append(" | ".join(norm(c.text) for c in row.cells))
    return "\n".join(parts)
def control_exact(path,uid):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
             "permission":hfind(h,"permission","auth resource"),"payload":hfind(h,"payload / schema","payload","schema"),
             "operation":hfind(h,"operation"),"method":hfind(h,"method / path","method","path"),
             "runtime_owner":hfind(h,"runtime owner"),"persistence":hfind(h,"persistence owner"),
             "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals) or vals[ci]!=uid:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":str(path),"table":ti,"row":ri});out.append(rec)
    exact=[r for r in out if r["runtime_status"]=="EFFECTFUL_EXACT"]
    return exact
def port_rows(path,uid):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];pi=hfind(h,"port uid")
        if pi is None:continue
        oi=hfind(h,"operation");mi=hfind(h,"method / path","method","path");peri=hfind(h,"permission")
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if pi<len(vals) and vals[pi]==uid:
                out.append({"file":str(path),"table":ti,"row":ri,
                            "operation":vals[oi] if oi is not None and oi<len(vals) else "",
                            "method":vals[mi] if mi is not None and mi<len(vals) else "",
                            "permission":vals[peri] if peri is not None and peri<len(vals) else ""})
    return out
def reverse_rows(path,uid):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];pi=hfind(h,"port uid");ci=hfind(h,"ui-flow consumer","consumer")
        if pi is None or ci is None:continue
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if pi<len(vals) and vals[pi]==uid:
                out.append({"file":str(path),"table":ti,"row":ri,"consumer":vals[ci] if ci<len(vals) else ""})
    return out

assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","-c","core.quotePath=false","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""
assert blob(LOGIC)==EXPECTED_LOGIC_SHA,(blob(LOGIC),EXPECTED_LOGIC_SHA)

# Precondition: legacy revision aliases and ambiguous locator shorthand are already absent package-wide.
pre_legacy=[];pre_locators=[]
for p in sorted(Path(".").glob("*.docx")):
    txt=full_text(p)
    pre_legacy.extend((p.name,m.group(0)) for m in LEGACY_RE.finditer(txt))
    pre_locators.extend((p.name,m.group(0)) for m in AMBIG_LOCATOR_RE.finditer(txt))
assert not pre_legacy,pre_legacy
assert not pre_locators,pre_locators

# Remove only the stale lexical "before" examples from System Logic. Do not alter the atomic Port contracts.
d=Document(LOGIC)
counts=collections.Counter()
for node in d._element.iter():
    if not str(node.tag).endswith("}t") or node.text is None:continue
    for old,new in REPLACEMENTS.items():
        if old in node.text:
            n=node.text.count(old)
            node.text=node.text.replace(old,new)
            counts[old]+=n
assert counts==collections.Counter({k:1 for k in REPLACEMENTS}),dict(counts)
d.save(LOGIC);Document(LOGIC)

# Package-wide postconditions.
legacy=[];locators=[];stale=[]
for p in sorted(Path(".").glob("*.docx")):
    txt=full_text(p)
    for m in LEGACY_RE.finditer(txt):legacy.append((p.name,m.group(0)))
    for m in AMBIG_LOCATOR_RE.finditer(txt):locators.append((p.name,m.group(0)))
    for old in REPLACEMENTS:
        if old in txt:stale.append((p.name,old,txt.count(old)))
assert not legacy,legacy
assert not locators,locators
assert not stale,stale

# Re-verify the exact two-way EDIT orchestration contract was not touched.
target={}
expected={"action":ACTION,"gate":GATE,"permission":UI_PERMISSION,"payload":PAYLOAD,"operation":OP,
          "method":METHOD,"runtime_owner":RUNTIME_OWNER,"persistence":PERSIST,"runtime_status":"EFFECTFUL_EXACT"}
for p in [S05,EDIT]:
    rows=control_exact(p,TARGET);assert len(rows)==1,(p,rows)
    r=rows[0]
    for k,v in expected.items():assert r[k]==v,(p,k,r[k],v)
    target[p]=r

port_report={};reverse_report={}
for uid,(op,method,perm) in PORTS.items():
    found=[]
    for p in sorted(Path(".").glob("*.docx")):found.extend(port_rows(p,uid))
    exact=[r for r in found if r["operation"]==op and r["method"]==method and r["permission"]==perm]
    assert exact,(uid,found)
    port_report[uid]=exact
    rev=[]
    for p in [S05,EDIT]:rev.extend(reverse_rows(p,uid))
    assert len(rev)>=2,(uid,rev)
    assert all(OP in x["consumer"] for x in rev),(uid,rev)
    reverse_report[uid]=rev

# Verify branch/no-mutation and permission layering remain explicit.
orchestration_text=full_text(S05)+"\n"+full_text(EDIT)
for token in [PORTS.keys().__iter__().__next__(),list(PORTS)[1],OP,"editing.task.execute","EDITING_USE"]:
    assert token in orchestration_text,token
assert re.search(r"no atomic mutation|NO\s+MUTATION",orchestration_text,re.I),"NO_MUTATION_BRANCH_MISSING"
assert "No third public API" in orchestration_text or "no third public API" in orchestration_text

machine={
"marker":MARK,"base_head":BASE_HEAD,
"legacy_r9_r5_remaining":0,"ambiguous_locator_remaining":0,
"stale_composite_lexical_remaining":0,
"stale_composite_replacements":sum(counts.values()),
"target_exact_rows":2,"atomic_ports_verified":2,"reverse_binding_rows":sum(len(v) for v in reverse_report.values()),
"orchestration_operation":OP,"public_api_created":False,
"dispatch_exactly_one_port":True,"runtime_execution_claimed":False,
}
report={"machine":machine,"replacements":dict(counts),"target":target,"ports":port_report,"reverse":reverse_report,
        "changed_docs":[LOGIC],"output_blob_sha":{LOGIC:blob(LOGIC)}}
Path("__batch63_cleanup_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH63="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
