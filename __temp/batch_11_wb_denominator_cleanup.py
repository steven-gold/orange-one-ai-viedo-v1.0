from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib

MARK="ACPOS-20260921-BATCH-11-WB-SYNTHETIC-DENOMINATOR-CLEANUP-V1"
WB="ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
S01="01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx"
EXPECTED_SHA={
    WB:"23a0df921875a8c3a027c531585b2b79c635caaa",
    LOGIC:"a8d3fff5ec93fe176fe70154df191880410e0e25",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":WB,
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]

def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():
    assert blob(f)==s,(f,blob(f),s)

def norm(x):
    return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())

def missing(v):
    return v in {"","—","-","SOURCE_NOT_DEFINED"}

def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:
                return i
    return None

def display_only_type(t):
    u=(t or "").upper()
    return u in {"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"} or any(x in u for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:
            continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(headers,"control uid")
        if ci is None:
            continue
        idx={
            "control":ci,
            "type":hfind(headers,"type"),
            "label":hfind(headers,"label","顯示名稱"),
            "action":hfind(headers,"action uid"),
            "gate":hfind(headers,"gate uid","gate"),
            "permission":hfind(headers,"permission","auth resource"),
            "operation":hfind(headers,"operation"),
            "runtime_owner":hfind(headers,"runtime owner"),
            "runtime_status":hfind(headers,"runtime status"),
        }
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:
                continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"table":ti+1,"row":ri,"headers":headers})
            out.append(rec)
    return out

def compose(rows):
    groups=collections.defaultdict(list)
    for r in rows:
        groups[r["control"]].append(r)
    out={}
    for uid,rs in groups.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:
                base[f]=vals[0]
        base["same_uid_rows"]=rs
        out[uid]=base
    return out

def is_definition_gap(field,r):
    v=r.get(field,"");st=r.get("runtime_status","")
    if not missing(v):
        return False
    if "UI_LOCAL_EXACT" in st:
        if field in {"operation","runtime_owner"}:
            return False
        if field=="action" and display_only_type(r.get("type","")):
            return False
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in st and field=="operation":
        return False
    if "READ_EXACT" in st and display_only_type(r.get("type","")):
        if field=="action":
            return False
        if field=="operation" and not missing(r.get("runtime_owner","")):
            return False
    if field=="runtime_owner" and any(x in st for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):
        return False
    if v=="SOURCE_NOT_DEFINED":
        return False
    return True

def gap_inventory():
    rows=[]
    for page,fn in PAGES.items():
        bm=compose(parse_controls(fn))
        for uid,r in bm.items():
            for field in ["action","gate","permission","operation","runtime_owner"]:
                if is_definition_gap(field,r):
                    rows.append((page,uid,field))
    return rows

pre=gap_inventory()
assert len(pre)==277,len(pre)

# Locate the Batch-08 synthetic mapping table.
wb=Document(WB)
synthetic_table=None
for t in wb.tables:
    if not t.rows:
        continue
    headers=[norm(c.text) for c in t.rows[0].cells]
    if "WB Control UID" not in headers:
        continue
    rows=[]
    for row in t.rows[1:]:
        vals=[norm(c.text) for c in row.cells]
        if vals and vals[0].startswith("CTRL-"):
            rows.append(vals)
    if len(rows)==14:
        synthetic_table=t
        break
assert synthetic_table is not None,"BATCH08_SYNTHETIC_TABLE_NOT_FOUND"

# Read exact canonical product rows from WB and System 01.
wb_product={}
for r in parse_controls(WB):
    if r["headers"] and any(h.lower()=="control uid" for h in r["headers"]):
        if r["control"].startswith("CTRL-WORKSPACE-WB-01-"):
            wb_product[r["control"]]=r
s01_rows=parse_controls(S01)
s01_candidates=[
    r for r in s01_rows
    if r["permission"]=="workspace.dashboard.view"
    and r["operation"]=="getDashboardReadModel"
    and r["runtime_owner"]=="DASHBOARD_READ_MODEL"
    and r["control"].startswith("CTRL-WORKSPACE-WB-01-")
]
s01_by_uid={r["control"]:r for r in s01_candidates}

maps=[]
for row in synthetic_table.rows[1:]:
    vals=[norm(c.text) for c in row.cells]
    compact=vals[0] if vals else ""
    if not compact.startswith("CTRL-"):
        continue
    if "…" in compact:
        suffix=compact.split("…",1)[1]
    elif "..." in compact:
        suffix=compact.split("...",1)[1]
    else:
        suffix=compact
    matches=[uid for uid in s01_by_uid if uid.endswith(suffix)]
    assert len(matches)==1,(compact,matches)
    full=matches[0]
    assert full in wb_product,(compact,full,"FULL_UID_NOT_IN_WB_PRODUCT_REGISTRY")
    wr=wb_product[full];sr=s01_by_uid[full]
    assert wr["gate"]=="PAGE_READ",(full,wr["gate"])
    assert sr["gate"]=="PAGE_READ",(full,sr["gate"])
    assert missing(wr["action"]) and missing(sr["action"]),(full,wr["action"],sr["action"])
    assert wr["type"]=="SECTION_OPEN" and sr["type"]=="SECTION_OPEN",(full,wr["type"],sr["type"])
    assert wr["runtime_status"]=="READ_EXACT" and sr["runtime_status"]=="READ_EXACT",(full,wr["runtime_status"],sr["runtime_status"])
    maps.append({
        "compact_alias":compact,
        "canonical_uid":full,
        "canonical_owner":S01,
        "gate":"PAGE_READ",
        "action_uid":"—",
        "type":"SECTION_OPEN",
        "runtime_status":"READ_EXACT",
        "wb_product_table":wr["table"],
        "wb_product_row":wr["row"],
        "system01_table":sr["table"],
        "system01_row":sr["row"],
    })
assert len(maps)==14,len(maps)
assert len(set(x["canonical_uid"] for x in maps))==14

# Critical correction: this governance table is not a product Control Registry.
headers=[norm(c.text) for c in synthetic_table.rows[0].cells]
ci=headers.index("WB Control UID")
synthetic_table.rows[0].cells[ci].text="WB Compact Alias (Non-Control)"

# Rename the historical heading without destroying audit history.
for p in wb.paragraphs:
    if "Batch 08 · WB Canonical Control UID / Owner Mapping Register" in norm(p.text):
        p.text="Batch 08 · WB Compact Alias / Owner Mapping Register (superseded by Batch 11 denominator correction)"
        break

# Persist a clean alias mapping that intentionally avoids any "Control UID" header.
sec=wb.add_section(WD_SECTION.NEW_PAGE)
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width,sec.page_height=sec.page_height,sec.page_width
sec.top_margin=Inches(.4);sec.bottom_margin=Inches(.4);sec.left_margin=Inches(.4);sec.right_margin=Inches(.4)
wb.add_heading("Batch 11 · WB Compact Alias → Canonical UID Resolution",level=1)
p=wb.add_paragraph()
p.add_run("["+MARK+"] ").bold=True
p.add_run("The 14 compact aliases below are governance aliases only. They are not product controls and MUST NOT enter the Control denominator. Each resolves one-to-one to an already-existing WB product registry row and the same System 01 canonical owner row. No Action/Gate/Permission/Operation/Runtime Owner business value is created or changed by this batch.")
t=wb.add_table(rows=1,cols=7);t.style="Table Grid"
for i,h in enumerate(["Compact Alias","Canonical UID","Canonical Owner","WB Registry","System 01 Registry","Existing Gate","Disposition"]):
    t.rows[0].cells[i].text=h
for x in maps:
    r=t.add_row().cells
    r[0].text=x["compact_alias"]
    r[1].text=x["canonical_uid"]
    r[2].text=Path(x["canonical_owner"]).name
    r[3].text=f'table {x["wb_product_table"]} / row {x["wb_product_row"]}'
    r[4].text=f'table {x["system01_table"]} / row {x["system01_row"]}'
    r[5].text=x["gate"]
    r[6].text="ALIAS_ONLY / EXCLUDE_FROM_CONTROL_DENOMINATOR"
for row in t.rows:
    for cell in row.cells:
        for pp in cell.paragraphs:
            for rr in pp.runs:
                rr.font.size=Pt(6)

wb.add_paragraph("Action UID remains '—' on the actual canonical product rows and is not remediated in Batch 11. Those canonical rows remain part of the separate unique-owner specification-remediation denominator.")
wb.save(WB)
Document(WB)

post=gap_inventory()
assert len(post)==249,(len(post),post[-30:])
removed=set(pre)-set(post)
added=set(post)-set(pre)
assert len(removed)==28,(len(removed),sorted(removed))
assert not added,sorted(added)
assert all(page=="WB-01" and uid.startswith("CTRL-…") and field in {"action","gate"} for page,uid,field in removed),sorted(removed)

# Verify no compact alias is parsed as a control after correction.
post_wb_controls=compose(parse_controls(WB))
assert not any(uid.startswith("CTRL-…") for uid in post_wb_controls),[uid for uid in post_wb_controls if uid.startswith("CTRL-…")]
assert all(x["canonical_uid"] in post_wb_controls for x in maps)

# Central current-state supersession.
logic=Document(LOGIC)
ltxt="\n".join([p.text for p in logic.paragraphs]+[c.text for tb in logic.tables for row in tb.rows for c in row.cells])
assert MARK not in ltxt,"BATCH11_ALREADY_PRESENT"
sec=logic.add_section(WD_SECTION.NEW_PAGE)
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width,sec.page_height=sec.page_height,sec.page_width
sec.top_margin=Inches(.35);sec.bottom_margin=Inches(.35);sec.left_margin=Inches(.3);sec.right_margin=Inches(.3)
logic.add_heading("Batch 11 · WB Synthetic Denominator Cleanup / Canonical Alias Resolution",level=1)
p=logic.add_paragraph()
p.add_run("["+MARK+"] ").bold=True
p.add_run("Batch 11 proves that the 28 prior WB CANONICAL_OWNER_ROW_ABSENT findings were not 28 product binding defects. They were produced because the Batch-08 governance mapping table used a header containing 'Control UID', causing the scanner to ingest 14 compact aliases as a second product registry. All 14 canonical WB controls already existed in WB table 26 and System 01 table 58.")

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def table(headers,rows,fs=4.6):
    t=logic.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row):
            c[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for pp in cell.paragraphs:
                for rr in pp.runs:
                    rr.font.size=Pt(fs)
    return t

table(["Item","Pre","Post","Decision"],[
    ["Definition Binding Gap",277,249,"28 synthetic WB alias rows removed from denominator; no product binding value changed."],
    ["WB owner-absent findings",28,0,"All were scanner ingestion of a governance alias table, not missing product owners."],
    ["Canonical WB product controls",14,14,"Unchanged; already present with full UID in WB registry and System 01."],
    ["Unique-owner unresolved remediation",248,248,"Unchanged; includes the real canonical Action/spec gaps."],
    ["SYS Gate conflict",1,1,"Unchanged / fail-closed."],
    ["Business binding mutation",0,0,"Action/Gate/Permission/Operation/Runtime Owner values untouched."],
],5.0)

logic.add_heading("14 Alias Resolutions",level=2)
table(["Compact Alias","Canonical UID","Canonical Owner","WB Registry","System 01 Registry","Gate","Status"],[
    [x["compact_alias"],x["canonical_uid"],Path(x["canonical_owner"]).name,f't{x["wb_product_table"]}/r{x["wb_product_row"]}',f't{x["system01_table"]}/r{x["system01_row"]}',x["gate"],"ALIAS_ONLY"]
    for x in maps
],4.0)

logic.add_heading("Scanner Invariant",level=2)
table(["Rule","Requirement"],[
    ["WB-ALIAS-01","Governance/audit/mapping tables MUST NOT use a header that can be parsed as product 'Control UID' unless the rows are actual product controls."],
    ["WB-ALIAS-02","Compact/display aliases are non-authoritative identifiers and MUST NOT enter product denominators."],
    ["WB-ALIAS-03","Canonical UID resolution requires one-to-one exact suffix mapping plus identical Permission + Operation + Runtime Owner anchors and an existing canonical product row."],
    ["WB-ALIAS-04","Alias cleanup does not close the real canonical Action UID remediation; that remains separate work."],
],5.0)

machine={
    "marker":MARK,
    "pre_definition_gaps":277,
    "synthetic_wb_findings_removed":28,
    "post_definition_gaps":249,
    "wb_owner_absent":0,
    "unique_owner_unresolved_remediation":248,
    "preserved_sys_gate_conflict":1,
    "canonical_alias_mappings":14,
    "business_binding_mutations":0,
}
logic.add_paragraph("BATCH11_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC)
Document(LOGIC)

report={
    "machine":machine,
    "mappings":maps,
    "removed_findings":[{"page":a,"uid":b,"field":c} for a,b,c in sorted(removed)],
    "output_hashes":{WB:blob(WB),LOGIC:blob(LOGIC)},
}
Path("__batch11_cleanup_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH11_REMEDIATION="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
