from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib

MARK="ACPOS-20260921-BATCH-13-EXPLICIT-NA-ACTION-SEMANTICS-V1"
S01="01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx"
S03="03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
WB="ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
AIAPI="ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"

EXPECTED_SHA={
 S01:"b023c835f78565ab378eed11f01ec685571919e8",
 S03:"d5867d11d6b47b66c7966e6911774e142eb6ef69",
 WB:"7c97669fdb5d45dbe7956f1116aae70d3a0a8d5c",
 AIAPI:"e6a1a42e87459edd82ef109099508eb011adb29b",
 LOGIC:"54edc6c24e7140f628c4d50402a28a07a7415ac7",
}

PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":WB,
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():
    assert blob(f)==s,(f,blob(f),s)

def norm(x):
    return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())

def missing(v):
    return v in {"","—","-","SOURCE_NOT_DEFINED"}

def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:
                return i
    return None

def display_only_type(t):
    u=(t or "").upper()
    return u in {"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"} or any(x in u for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:
            continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(headers,"control uid")
        if ci is None:
            continue
        idx={
          "control":ci,
          "type":hfind(headers,"type"),
          "label":hfind(headers,"label","顯示名稱"),
          "action":hfind(headers,"action uid"),
          "gate":hfind(headers,"gate uid","gate"),
          "permission":hfind(headers,"permission","auth resource"),
          "payload_schema":hfind(headers,"payload / schema","payload","schema"),
          "operation":hfind(headers,"operation"),
          "method_path":hfind(headers,"method / path","method","path"),
          "runtime_owner":hfind(headers,"runtime owner"),
          "persistence_owner":hfind(headers,"persistence owner"),
          "runtime_status":hfind(headers,"runtime status"),
        }
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:
                continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"table":ti+1,"row":ri,"headers":headers})
            out.append(rec)
    return out

def compose(rows):
    groups=collections.defaultdict(list)
    for r in rows:
        groups[r["control"]].append(r)
    out={}
    for uid,rs in groups.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:
                base[f]=vals[0]
        out[uid]=base
    return out

# Baseline Batch-12 classifier (without Batch-13 semantics).
def classify_pre(field,r):
    v=r.get(field,"");st=r.get("runtime_status","")
    if not missing(v):
        return "BOUND"
    if "UI_LOCAL_EXACT" in st:
        if field in {"operation","runtime_owner"}:
            return "LEGITIMATE_NA_UI_LOCAL"
        if field=="action" and display_only_type(r.get("type","")):
            return "LEGITIMATE_NA_UI_LOCAL"
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in st and field=="operation":
        return "CLOSED_BY_OWNER_LOCAL_COMMAND"
    if "READ_EXACT" in st and display_only_type(r.get("type","")):
        if field=="action":
            return "LEGITIMATE_NA_READ_PRESENTATION"
        if field=="operation" and not missing(r.get("runtime_owner","")):
            return "READ_OWNER_BOUND_OPERATION_UNSPECIFIED"
    if field=="runtime_owner" and any(x in st for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):
        return "KNOWN_RUNTIME_BLOCKER"
    if v=="SOURCE_NOT_DEFINED":
        return "SOURCE_RESOLUTION_REQUIRED"
    return "DEFINITION_BINDING_GAP"

def is_wb_read_projection_action_na(r):
    return (
        missing(r.get("action",""))
        and r.get("type","")=="SECTION_OPEN"
        and r.get("gate","")=="PAGE_READ"
        and r.get("permission","")=="workspace.dashboard.view"
        and r.get("payload_schema","")=="N/A_READ_PROJECTION"
        and r.get("operation","")=="getDashboardReadModel"
        and r.get("method_path","")=="GET /v1/dashboard/read-model"
        and r.get("runtime_owner","")=="DASHBOARD_READ_MODEL"
        and r.get("runtime_status","")=="READ_EXACT"
    )

def is_aiapi_direct_operation_action_na(r):
    return (
        missing(r.get("action",""))
        and r.get("type","")=="BUTTON_OR_ROW_ACTION"
        and r.get("payload_schema","")=="AIAPI Page Operation-specific Form / Provider Profile Field Contract"
        and not missing(r.get("permission",""))
        and not missing(r.get("operation",""))
        and not missing(r.get("method_path",""))
        and not missing(r.get("runtime_owner",""))
        and r.get("runtime_status","") in {"EFFECTFUL_EXACT","READ_EXACT"}
    )

def classify_post(page,field,r):
    pre=classify_pre(field,r)
    if pre!="DEFINITION_BINDING_GAP":
        return pre
    if field=="action" and page=="WB-01" and is_wb_read_projection_action_na(r):
        return "LEGITIMATE_NA_READ_PROJECTION_ACTION"
    if field=="action" and page=="AIAPI-01" and is_aiapi_direct_operation_action_na(r):
        return "LEGITIMATE_NA_DIRECT_OPERATION_ACTION"
    return pre

def inventory(classifier):
    rows=[]
    counts=collections.Counter()
    for page,fn in PAGES.items():
        bm=compose(parse_controls(fn))
        for uid,r in bm.items():
            for field in ["action","gate","permission","operation","runtime_owner"]:
                cls=classifier(page,field,r) if classifier==classify_post else classifier(field,r)
                counts[cls]+=1
                if cls=="DEFINITION_BINDING_GAP":
                    rows.append((page,uid,field))
    return rows,counts

pre_rows,pre_counts=inventory(classify_pre)
assert len(pre_rows)==239,len(pre_rows)

# Determine exactly the 30 Batch-13 N/A rows.
wb_map=compose(parse_controls(WB))
aiapi_map=compose(parse_controls(AIAPI))
wb_targets=[uid for uid,r in wb_map.items() if is_wb_read_projection_action_na(r)]
aiapi_targets=[uid for uid,r in aiapi_map.items() if is_aiapi_direct_operation_action_na(r)]
assert len(wb_targets)==14,(len(wb_targets),wb_targets)
assert len(aiapi_targets)==16,(len(aiapi_targets),aiapi_targets)
expected_na={("WB-01",uid,"action") for uid in wb_targets}|{("AIAPI-01",uid,"action") for uid in aiapi_targets}
assert expected_na.issubset(set(pre_rows)),sorted(expected_na-set(pre_rows))

post_rows,post_counts=inventory(classify_post)
removed=set(pre_rows)-set(post_rows)
added=set(post_rows)-set(pre_rows)
assert removed==expected_na,{"missing":sorted(expected_na-removed),"unexpected":sorted(removed-expected_na)}
assert not added,sorted(added)
assert len(post_rows)==209,len(post_rows)

# Canonical owner evidence checks.
s01_map=compose(parse_controls(S01))
for uid in wb_targets:
    assert uid in s01_map,uid
    assert is_wb_read_projection_action_na(s01_map[uid]),(uid,s01_map[uid])

s03_map=compose(parse_controls(S03))
for uid in aiapi_targets:
    assert uid in s03_map,uid
    assert is_aiapi_direct_operation_action_na(s03_map[uid]),(uid,s03_map[uid])

# Confirm the WB page glossary explicitly describes Section Open as the same registered READ_ONLY operation.
wb_doc_for_glossary=Document(WB)
wb_all="\n".join([p.text for p in wb_doc_for_glossary.paragraphs]+[c.text for t in wb_doc_for_glossary.tables for row in t.rows for c in row.cells])
assert "Section Open = 同一 registered READ_ONLY operation" in wb_all,"WB_SECTION_OPEN_GLOSSARY_MISSING"

# Helpers.
def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.34);sec.bottom_margin=Inches(.34);sec.left_margin=Inches(.30);sec.right_margin=Inches(.30)

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr()
    e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)

def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr()
    e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)

def add_table(doc,headers,rows,fs=4.8):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for r in p.runs:r.font.size=Pt(fs)
    return t

# Add owner rules and exact ledgers.
for fn,title,targets,rule_text,app_class in [
    (
      S01,
      "Batch 13 · READ Projection Action Applicability Rule",
      wb_targets,
      "For SECTION_OPEN controls with READ_EXACT, Payload/Schema=N/A_READ_PROJECTION, exact GET getDashboardReadModel, PAGE_READ, workspace.dashboard.view and DASHBOARD_READ_MODEL, the registered READ_ONLY Operation is the interaction boundary. Action UID is non-applicable and MUST remain '—'; no separate effectful Action UID may be invented.",
      "LEGITIMATE_NA_READ_PROJECTION_ACTION"
    ),
    (
      S03,
      "Batch 13 · AIAPI Direct-Operation Action Applicability Rule",
      aiapi_targets,
      "For AIAPI BUTTON_OR_ROW_ACTION controls governed by the exact 'AIAPI Page Operation-specific Form / Provider Profile Field Contract', with exact Permission, Operation, Method/Path and Runtime Owner already bound, the Operation contract is the execution identity. If no existing Action UID is explicitly bound, Action UID is non-applicable and MUST remain '—'; a parallel Action UID MUST NOT be invented.",
      "LEGITIMATE_NA_DIRECT_OPERATION_ACTION"
    ),
]:
    d=Document(fn)
    text_all="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
    local=MARK+"::"+Path(fn).name
    assert local not in text_all,(fn,"ALREADY_APPLIED")
    add_landscape(d)
    d.add_heading(title,level=1)
    p=d.add_paragraph();p.add_run("["+local+"] ").bold=True;p.add_run(rule_text)
    add_table(d,["Target UID","Field","Applicability","Canonical Rule","Status"],[
      [uid,"Action UID","N/A",app_class,"CLOSED_AS_EXPLICIT_NA_BY_BATCH_13"] for uid in sorted(targets)
    ],4.6)
    d.save(fn);Document(fn)

# Page-level applicability ledgers; deliberately use Target UID, not Control UID, to avoid denominator ingestion.
for fn,page,targets,app_class in [
    (WB,"WB-01",wb_targets,"LEGITIMATE_NA_READ_PROJECTION_ACTION"),
    (AIAPI,"AIAPI-01",aiapi_targets,"LEGITIMATE_NA_DIRECT_OPERATION_ACTION"),
]:
    d=Document(fn)
    text_all="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
    local=MARK+"::PAGE::"+page
    assert local not in text_all,(fn,"PAGE_LEDGER_ALREADY_APPLIED")
    add_landscape(d)
    d.add_heading("Batch 13 · Action UID Applicability Ledger",level=1)
    p=d.add_paragraph();p.add_run("["+local+"] ").bold=True
    p.add_run("This ledger records applicability only. Product Control Registry rows remain unchanged; Action UID stays '—'. These rows MUST NOT be parsed as additional product controls.")
    add_table(d,["Target UID","Field","Current Cell","Applicability Class","Required Handling"],[
      [uid,"Action UID","—",app_class,"DO_NOT_INVENT_ACTION_UID"] for uid in sorted(targets)
    ],4.8)
    d.save(fn);Document(fn)

# Central System Logic current-state rule.
logic=Document(LOGIC)
logic_text="\n".join([p.text for p in logic.paragraphs]+[c.text for t in logic.tables for row in t.rows for c in row.cells])
assert MARK not in logic_text,"LOGIC_ALREADY_APPLIED"
add_landscape(logic)
logic.add_heading("Batch 13 · Explicit N/A Action Semantics",level=1)
p=logic.add_paragraph();p.add_run("["+MARK+"] ").bold=True
p.add_run("Batch 13 does not fill missing Action UID cells. It closes only two exact structural classes where current Canonical Authority proves that the already-bound read/direct Operation is the interaction or execution identity and a second Action UID would be invented duplication.")

add_table(logic,["Class","Exact eligibility","Meaning"],[
["LEGITIMATE_NA_READ_PROJECTION_ACTION","WB SECTION_OPEN + READ_EXACT + PAGE_READ + workspace.dashboard.view + Payload/Schema=N/A_READ_PROJECTION + getDashboardReadModel + GET /v1/dashboard/read-model + DASHBOARD_READ_MODEL","Registered READ_ONLY operation is the read-projection boundary; Action UID remains —."],
["LEGITIMATE_NA_DIRECT_OPERATION_ACTION","AIAPI BUTTON_OR_ROW_ACTION + exact Operation-specific Form contract + non-empty Permission/Operation/Method-Path/Runtime Owner + EFFECTFUL_EXACT or READ_EXACT","Direct Operation contract is execution identity; no parallel Action UID unless explicitly pre-existing."],
],4.5)

add_table(logic,["Item","Count","Result"],[
["Pre-batch Definition Binding Gap",239,"Fresh Batch-12 denominator."],
["WB explicit N/A Action",14,"Reclassified; no Action UID written."],
["AIAPI direct-operation N/A Action",16,"Reclassified; no Action UID written."],
["Total Batch-13 N/A reclassification",30,"No business binding mutation."],
["Post-batch Definition Binding Gap",209,"208 unique-owner canonical-spec gaps + 1 preserved SYS Gate conflict."],
["Business binding values mutated",0,"All product registry Action UID cells remain unchanged."],
],5.0)

logic.add_heading("30 Exact N/A Rows",level=2)
add_table(logic,["Page","Target UID","Field","Class"],[
  *[["WB-01",uid,"Action UID","LEGITIMATE_NA_READ_PROJECTION_ACTION"] for uid in sorted(wb_targets)],
  *[["AIAPI-01",uid,"Action UID","LEGITIMATE_NA_DIRECT_OPERATION_ACTION"] for uid in sorted(aiapi_targets)],
],4.0)

logic.add_heading("Batch 13 Fail-Closed Boundary",level=2)
add_table(logic,["Rule","Requirement"],[
["B13-01","A blank Action UID is not N/A by itself; the full exact structural eligibility rule must match."],
["B13-02","UI similarity, labels, first-match behavior, or nearby rows cannot create N/A eligibility."],
["B13-03","The remaining 208 unique-owner gaps retain DEFINITION_BINDING_GAP until their Canonical Owner explicitly defines a value or an equally explicit non-applicability rule."],
["B13-04","The preserved SYS-01-BTN-NAV-OPEN Gate conflict remains fail-closed and is outside this N/A batch."],
],5.0)

machine={
 "marker":MARK,
 "pre_definition_gaps":239,
 "wb_na_action":14,
 "aiapi_na_action":16,
 "na_reclassified":30,
 "post_definition_gaps":209,
 "remaining_unique_owner_remediation":208,
 "preserved_sys_gate_conflict":1,
 "business_binding_mutations":0,
}
logic.add_paragraph("BATCH13_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

# Final evidence: product registry rows still contain Action UID '—'.
for fn,targets in [(WB,wb_targets),(AIAPI,aiapi_targets)]:
    bm=compose(parse_controls(fn))
    for uid in targets:
        assert uid in bm,uid
        assert missing(bm[uid]["action"]),(fn,uid,bm[uid]["action"])

report={
 "machine":machine,
 "wb_targets":sorted(wb_targets),
 "aiapi_targets":sorted(aiapi_targets),
 "removed_from_definition_gap":[{"page":p,"uid":u,"field":f} for p,u,f in sorted(removed)],
 "pre_classification_counts":dict(pre_counts),
 "post_classification_counts":dict(post_counts),
 "output_hashes":{f:blob(f) for f in [S01,S03,WB,AIAPI,LOGIC]},
}
Path("__batch13_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH13_REMEDIATION="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
