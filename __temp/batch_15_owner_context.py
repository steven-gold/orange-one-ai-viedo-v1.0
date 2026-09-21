from docx import Document
import json,re

FILES={
"S01":"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"S06":"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"DB":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"LOGIC":"ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx",
}
TERMS={
"S01":[
"Database Read Model / DB-01","DB_READ","DB_TRACE","DB_MIGRATION_READ","DB_AUDIT_READ","DB_INTEGRITY_READ",
"UI_LOCAL_OR_READ_SOURCE","read source","projection","read model"
],
"S06":[
"getERPSyncStatus","getERPFailure","getERPFinanceFactPack","getERPForecast","getERPCapacityGuardrails","exportProjection",
"ERPConnectorService","FinanceGuardrailService","Existing Projection/Export owner",
"ERP-01-ACT-READ-CONNECTOR","ERP-01-ACT-READ-FINANCE","ERP-01-ACT-READ-SYNC",
"UI_LOCAL_OR_READ_SOURCE","Connector ID","Connection Status","Entity Scope","Mapping Version","Provider Key","Secret Reference ID","Adapter Key",
"Capacity","Cashflow","Cost","Forecast","Guardrails","Recommendation Boundary","Revenue",
"Completeness","Currency","Timezone","Failure ID","Freshness At","Last Sync Status","Snapshot ID"
],
}
def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
out={}
for key in ["S01","S06"]:
    fn=FILES[key];terms=TERMS[key];d=Document(fn);hits=[]
    for i,p in enumerate(d.paragraphs,1):
        t=norm(p.text)
        if t and any(term.lower() in t.lower() for term in terms):
            hits.append({"kind":"paragraph","index":i,"text":t})
    for ti,tb in enumerate(d.tables,1):
        if not tb.rows:continue
        headers=[norm(c.text) for c in tb.rows[0].cells]
        for ri,row in enumerate(tb.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];joined=" | ".join(vals)
            if joined and any(term.lower() in joined.lower() for term in terms):
                hits.append({"kind":"table","table":ti,"row":ri,"headers":headers,"values":vals})
    out[key]=hits
print("BATCH15_OWNER_CONTEXT="+json.dumps(out,ensure_ascii=False,sort_keys=True))
