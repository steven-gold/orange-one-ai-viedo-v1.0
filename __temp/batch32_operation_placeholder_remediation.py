from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-32-OPERATION-PLACEHOLDER-REMEDIATION-V1"
BASE_HEAD="d9030ee81b5bc44be9f418bf1f4382bcdb3691da"
S02="02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
QA="ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
SYS="ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S02:"79f7e465d6f03352337de93eb62cd0f62ba47375",
S05:"1481e73c1532600f62bcaf8e42100ff22db940db",
QA:"bd92d55b7df0d635990c27eef24145c75cbbb4a1",
SYS:"0bd0f7718ab9643f020bbbdf369a1d73dd11c321",
LOGIC:"5ba0854d366da186150fb8c469c50e73a1869135",
}
BINDINGS={
"QA-01-BTN-MANUAL-MODIFY":{
 "operation":"modifyManualReview",
 "owner":S05,"page":QA,
 "action":"QA-01-ACT-MANUAL-MODIFY","gate":"QA-01-GATE-MANUAL-DECISION","permission":"QA_USE",
 "runtime_owner":"ManualReviewService","runtime_status":"EFFECTFUL_EXACT",
 "semantics":"Modify the decision/content of the exact manual review case under an authorized reviewer. The operation MUST NOT bypass required manual review and MUST fail closed when the case, reviewer authority, or current review state cannot be resolved.",
},
"QA-01-BTN-MANUAL-PASS":{
 "operation":"passManualReview",
 "owner":S05,"page":QA,
 "action":"QA-01-ACT-MANUAL-PASS","gate":"QA-01-GATE-MANUAL-DECISION","permission":"QA_USE",
 "runtime_owner":"ManualReviewService","runtime_status":"EFFECTFUL_EXACT",
 "semantics":"Record an authorized manual-review PASS for the exact review case. Machine score alone MUST NOT satisfy this operation; the reviewer decision and case identity must be explicit, auditable, and fail closed when unresolved.",
},
"SYS-01-BTN-ATTACH":{
 "operation":"attachConversationContext",
 "owner":S02,"page":SYS,
 "action":"SYS-01-ACT-CONVERSATION-ATTACH","gate":"SYS-01-GATE-CONVERSATION-COMPOSER","permission":"system.ai.use",
 "runtime_owner":"Shared Conversation Core","runtime_status":"EFFECTFUL_EXACT",
 "semantics":"Attach an already-resolved typed attachment/context reference to the current conversation composer/context binding. The Conversation Core does not become owner of the raw attachment object; permission/classification and exact attachment/context identity must resolve before inclusion. Failure is fail-closed.",
},
}
PLACEHOLDER_RE=re.compile(r"^(RESOLVE_FROM_|SOURCE_NOT_DEFINED$|OPERATION_NOT_DEFINED$|TBD$|TODO$)",re.I)

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def is_placeholder_op(v):
    v=norm(v)
    return bool(PLACEHOLDER_RE.search(v)) or "DURING_LOGIC_PHASE" in v or "EXISTING_GOVERNED_" in v
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in ["action","gate","permission","operation","method_path","runtime_owner","persistence_owner","runtime_status"]))
        out[uid]=base
    return out
def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():
    got=blob(f);assert got==s,(f,got,s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

# Preconditions.
for uid,spec in BINDINGS.items():
    for path in [spec["owner"],spec["page"]]:
        rows=[r for r in parse_controls(path) if r["control"]==uid]
        assert rows,(path,uid,"MISSING")
        ops=sorted(set(r["operation"] for r in rows if not missing(r["operation"])))
        assert ops and all(is_placeholder_op(x) for x in ops),(path,uid,ops)
        for r in rows:
            if not missing(r.get("action","")): assert r["action"]==spec["action"],(path,uid,r["action"])
            if not missing(r.get("gate","")): assert r["gate"]==spec["gate"],(path,uid,r["gate"])
            if not missing(r.get("permission","")): assert r["permission"]==spec["permission"],(path,uid,r["permission"])
            if not missing(r.get("runtime_status","")): assert r["runtime_status"]==spec["runtime_status"],(path,uid,r["runtime_status"])

def patch_operation(path,uids):
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=exact_index(h,"Control UID");oi=hfind(h,"operation")
        if ci is None or oi is None:continue
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals):continue
            uid=vals[ci]
            if uid not in uids:continue
            old=norm(row.cells[oi].text)
            if old==BINDINGS[uid]["operation"]:
                hits[uid]+=1;continue
            assert is_placeholder_op(old),(path,uid,old)
            row.cells[oi].text=BINDINGS[uid]["operation"];hits[uid]+=1
    assert set(hits)==set(uids),(path,dict(hits),set(uids)-set(hits))
    d.save(path);Document(path)
    return dict(hits)

patches={
"S05_OWNER":patch_operation(S05,["QA-01-BTN-MANUAL-MODIFY","QA-01-BTN-MANUAL-PASS"]),
"QA_PAGE":patch_operation(QA,["QA-01-BTN-MANUAL-MODIFY","QA-01-BTN-MANUAL-PASS"]),
"S02_OWNER":patch_operation(S02,["SYS-01-BTN-ATTACH"]),
"SYS_PAGE":patch_operation(SYS,["SYS-01-BTN-ATTACH"]),
}

def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.33);sec.bottom_margin=Inches(.33);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=4.4):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for rr in rows:
        cells=t.add_row().cells
        for i,v in enumerate(rr):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
    return t

# Canonical Owner contracts.
for owner,uids,title in [
(S05,["QA-01-BTN-MANUAL-MODIFY","QA-01-BTN-MANUAL-PASS"],"Batch 32 · QA Manual Review Canonical Operation Contracts"),
(S02,["SYS-01-BTN-ATTACH"],"Batch 32 · Conversation Attachment Canonical Operation Contract"),
]:
    d=Document(owner);add_landscape(d);d.add_heading(title,level=1)
    p=d.add_paragraph();p.add_run(f"[{MARK}::{Path(owner).name}] ").bold=True
    p.add_run("These Operation identities replace temporary RESOLVE_FROM_* placeholders. This batch defines only the canonical Operation identity and operation semantics. Method/Path, Payload/Schema, and Persistence Owner remain unresolved unless separately established by Current authority; runtime execution is not claimed.")
    rr=[]
    for uid in uids:
        s=BINDINGS[uid]
        rr.append([uid,s["operation"],s["action"],s["gate"],s["permission"],s["runtime_owner"],s["runtime_status"],s["semantics"]])
    add_table(d,["Control UID","Canonical Operation","Action","Gate","Permission","Runtime Owner","Runtime Status","Required Semantics"],rr,3.6)
    d.save(owner);Document(owner)

# Page ledgers.
for pagepath,uids,pagename in [
(QA,["QA-01-BTN-MANUAL-MODIFY","QA-01-BTN-MANUAL-PASS"],"QA-01"),
(SYS,["SYS-01-BTN-ATTACH"],"SYS-01"),
]:
    d=Document(pagepath);add_landscape(d);d.add_heading("Batch 32 · Operation Placeholder Closure Ledger",level=1)
    p=d.add_paragraph();p.add_run(f"[{MARK}::PAGE::{pagename}] ").bold=True
    p.add_run("The page Operation field is synchronized to the Canonical Owner operation. No transport/payload/persistence value is synthesized by this batch.")
    add_table(d,["Control UID","Operation","Runtime Owner","Runtime Status","Downstream Contract State"],[
        [uid,BINDINGS[uid]["operation"],BINDINGS[uid]["runtime_owner"],BINDINGS[uid]["runtime_status"],"Method/Path + Payload/Schema + Persistence Owner remain for later remediation"]
        for uid in uids
    ],4.1)
    d.save(pagepath);Document(pagepath)

# Post-check.
for uid,spec in BINDINGS.items():
    for path in [spec["owner"],spec["page"]]:
        vals=sorted(set(r["operation"] for r in parse_controls(path) if r["control"]==uid and not missing(r["operation"])))
        assert vals==[spec["operation"]],(path,uid,vals)

# Global placeholder scan across the two affected pages and owners.
for path in [S02,S05,QA,SYS]:
    for r in parse_controls(path):
        if r["control"] in BINDINGS:
            assert not is_placeholder_op(r["operation"]),(path,r["control"],r["operation"])

logic=Document(LOGIC);add_landscape(logic);logic.add_heading("Batch 32 · Operation Semantic Placeholder Closure",level=1)
p=logic.add_paragraph();p.add_run(f"[{MARK}] ").bold=True
p.add_run("Batch 30 proved three Current controls still used RESOLVE_FROM_* text in the Operation field. Batch 31 found no reusable exact operation identity but confirmed each Canonical Owner's naming and semantic vocabulary. Batch 32 replaces those placeholders with owner-defined operations while deliberately leaving transport, payload, and persistence unresolved.")
add_table(logic,["Page","Control UID","Old Placeholder","Canonical Operation","Owner","Result"],[
["QA-01","QA-01-BTN-MANUAL-MODIFY","RESOLVE_FROM_EXISTING_GOVERNED_MANUAL_REVIEW_AUTHORITY","modifyManualReview","System 05 / ManualReviewService","OPERATION_PLACEHOLDER_CLOSED"],
["QA-01","QA-01-BTN-MANUAL-PASS","RESOLVE_FROM_EXISTING_GOVERNED_MANUAL_REVIEW_AUTHORITY","passManualReview","System 05 / ManualReviewService","OPERATION_PLACEHOLDER_CLOSED"],
["SYS-01","SYS-01-BTN-ATTACH","RESOLVE_FROM_CURRENT_CONVERSATION_CORE_DURING_LOGIC_PHASE","attachConversationContext","System 02 / Shared Conversation Core","OPERATION_PLACEHOLDER_CLOSED"],
],3.9)
machine={
"marker":MARK,"base_head":BASE_HEAD,"operation_placeholder_controls_pre":3,"operation_placeholder_controls_post":0,
"canonical_operations_defined":3,"transport_defined":0,"payload_defined":0,"persistence_owner_defined":0,
"remaining_contract_cells":324,"runtime_execution_claimed":False,
}
logic.add_paragraph("BATCH32_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

changed=[S02,S05,QA,SYS,LOGIC]
report={"machine":machine,"patches":patches,"bindings":{uid:{k:v for k,v in s.items() if k not in {"owner","page"}} for uid,s in BINDINGS.items()},
"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed}}
Path("__batch32_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH32="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
