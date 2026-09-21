from pathlib import Path
from docx import Document
import json,re,subprocess

BASE_HEAD="a0eca38e60aebc788e94529306cb609466f31a81"
S08="08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
STR="ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
OPS=["submitStrategyReview","adoptAsContextCandidate"]
KW=["strategy","review","candidate","adopt","context","decision","persistence","persist","audit","approval","human","proposal","context candidate","decision record","review state"]

assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
rx=re.compile("|".join(re.escape(x) for x in OPS+KW),re.I)

def collect(path):
    d=Document(path);out={"paragraphs":[],"rows":[]}
    for i,p in enumerate(d.paragraphs):
        t=norm(p.text)
        if t and rx.search(t):out["paragraphs"].append({"index":i,"text":t[:3500]})
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];joined=" | ".join(vals)
            if joined and rx.search(joined):
                out["rows"].append({"table":ti+1,"row":ri,"headers":h,"values":vals,"joined":joined[:6000]})
    return out

owner=collect(S08);page=collect(STR)
summary={"operations":2,"owner_paragraph_hits":len(owner["paragraphs"]),"owner_table_hits":len(owner["rows"]),
"page_paragraph_hits":len(page["paragraphs"]),"page_table_hits":len(page["rows"])}
report={"marker":"ACPOS-20260922-BATCH-42-SYSTEM08-STRATEGY-CONTRACT-DESIGN-AUDIT-V1","base_head":BASE_HEAD,
"summary":summary,"operations":OPS,"owner_evidence":owner,"page_evidence":page}
Path("__batch42_system08_strategy_contract_audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
Path("__batch42_summary.txt").write_text(json.dumps(summary,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print("BATCH42_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
