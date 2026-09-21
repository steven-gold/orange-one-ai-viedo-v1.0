from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-47-SYSTEM07-SOCIAL-CONTRACT-REMEDIATION-V1"
BASE_HEAD="c195fca67d80c7303123f131b699a6739056cb06"
S07="07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
SOC="ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={S07:"bbdfb4e75af7b5b98ebcf777365e3ba174844507",SOC:"82112c2a3759b4d5b3ebce4e1ed184198ac0d0b2",LOGIC:"48da8cd34db0223e3349af6cabee266a50762a5a"}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx","ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx","EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx","IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":SOC,"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx","VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx"}
FIELDS=["control","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]
T={
"SOC-01-BTN-ACCOUNT-BIND":("bindSocialAccount",{"payload_schema":"BindSocialAccountRequest","persistence_owner":"System 07 / SocialAccountBinding + Audit"}),
"SOC-01-BTN-ACCOUNT-UNBIND":("unbindSocialAccount",{"payload_schema":"UnbindSocialAccountRequest","persistence_owner":"System 07 / SocialAccountBinding State + Audit"}),
"SOC-01-BTN-CANDIDATE-DECIDE":("decideCandidate",{"payload_schema":"DecideSocialContentCandidateRequest"}),
"SOC-01-BTN-CONTENT-SAVE":("saveDraft",{"payload_schema":"SaveSocialContentDraftRequest","persistence_owner":"Existing Draft/Candidate Owner / Social Content Draft + Audit"}),
"SOC-01-BTN-CREDENTIAL-REVEAL":("revealSocialCredential",{"payload_schema":"RevealSocialCredentialRequest","persistence_owner":"Secret/Vault Owner + Credential Reveal Audit"}),
"SOC-01-BTN-KILL":("setKillSwitch",{"payload_schema":"SetKillSwitchRequest","persistence_owner":"Existing Operations Governance / Kill Switch State + Audit"}),
"SOC-01-BTN-MANUAL-COMPLETE":("completeSocialManualAction",{"payload_schema":"CompleteSocialManualActionRequest","persistence_owner":"System 07 / Social Manual Action + Audit"}),
"SOC-01-BTN-PLATFORM-CONFIG":("configureGovernedResource",{"payload_schema":"ConfigureGovernedResourceRequest","persistence_owner":"Resolved Governed Resource Canonical Owner + Governance Audit"}),
"SOC-01-BTN-POLICY-CONFIG":("configureSocialTargetPolicy",{"payload_schema":"ConfigureSocialTargetPolicyRequest","persistence_owner":"System 07 / Social Target Posting Policy + Audit"}),
"SOC-01-BTN-PUBLISH":("requestSocialTargetPublish",{"payload_schema":"SocialTargetPublishRequest","persistence_owner":"System 07 / PublishRequest + Target Publish State + Audit"}),
"SOC-01-BTN-REFRESH":("refreshProjection",{"payload_schema":"RefreshSocialProjectionRequest","persistence_owner":"System 07 / Social Projection + Projection Audit"}),
"SOC-01-BTN-TARGET-DISCOVERY":("createSocialTargetDiscovery",{"payload_schema":"CreateSocialTargetDiscoveryRequest","persistence_owner":"System 07 / SocialTargetDiscoveryJob + Discovery Evidence"}),
"SOC-01-BTN-TARGET-JOIN":("requestSocialTargetJoin",{"payload_schema":"SocialTargetJoinRequest","persistence_owner":"System 07 / Social Target Join Request + Audit"}),
"SOC-01-INP-RECORD-SEARCH":("searchProjection",{"payload_schema":"SearchSocialProjectionRequest","persistence_owner":"System 07 / Search Projection Result + Search Audit"}),
"SOC-01-INP-TARGET-SEARCH":("searchProjection",{"payload_schema":"SearchSocialProjectionRequest","persistence_owner":"System 07 / Search Projection Result + Search Audit"}),
}
def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(h,*ns):
    hs=[norm(x).lower() for x in h]
    for n in ns:
        for i,x in enumerate(hs):
            if n.lower() in x:return i
def exact(h,n):
    for i,x in enumerate(h):
        if norm(x).lower()==n.lower():return i
def parse(path):
    d=Document(path);out=[]
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        ix={"control":ci,"action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),"permission":hfind(h,"permission","auth resource"),
        "payload_schema":hfind(h,"payload / schema","payload","schema"),"operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
        "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),"runtime_status":hfind(h,"runtime status")}
        for row in t.rows[1:]:
            v=[norm(c.text) for c in row.cells];uid=v[ci] if ci<len(v) else ""
            if not uid or uid in {"—","-"}:continue
            out.append({k:(v[i] if i is not None and i<len(v) else "") for k,i in ix.items()})
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        b=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(b.get(f,"")) and len(vals)==1:b[f]=vals[0]
        out[uid]=b
    return out
def blob(p):
    b=Path(p).read_bytes();return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()
for f,s in EXPECTED_SHA.items():assert blob(f)==s,(f,blob(f),s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""
preP,preO=compose(parse(SOC)),compose(parse(S07))
for uid,(op,fs) in T.items():
    for m,label in [(preP,"PAGE"),(preO,"OWNER")]:
        assert uid in m,(label,uid);r=m[uid];assert r["operation"]==op,(label,uid,r["operation"],op);assert r["runtime_status"]=="EFFECTFUL_EXACT"
        for f,nv in fs.items():assert missing(r[f]),(label,uid,f,r[f],nv)
    assert not missing(preP[uid]["method_path"]),(uid,"METHOD_MISSING")
assert preO["SOC-01-BTN-CANDIDATE-DECIDE"]["persistence_owner"]=="public.context_candidates"
def patch(path):
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact(h,"Control UID")
        if ci is None:continue
        ix={"payload_schema":hfind(h,"payload / schema","payload","schema"),"persistence_owner":hfind(h,"persistence owner")}
        for row in t.rows[1:]:
            v=[norm(c.text) for c in row.cells];uid=v[ci] if ci<len(v) else ""
            if uid not in T:continue
            for f,nv in T[uid][1].items():
                i=ix[f]
                if i is None or i>=len(row.cells):continue
                old=norm(row.cells[i].text)
                if old==nv:hits[(uid,f)]+=1;continue
                assert missing(old),(path,uid,f,old,nv)
                row.cells[i].text=nv;hits[(uid,f)]+=1
    exp={(u,f) for u,(_,fs) in T.items() for f in fs}
    assert set(hits)==exp,(path,exp-set(hits),dict(hits));d.save(path);Document(path);return {f"{u}:{f}":n for (u,f),n in hits.items()}
patches={"SYSTEM07_OWNER":patch(S07),"SOC_PAGE":patch(SOC)}
def landscape(d):
    s=d.add_section(WD_SECTION.NEW_PAGE);s.orientation=WD_ORIENT.LANDSCAPE;s.page_width,s.page_height=s.page_height,s.page_width
    s.top_margin=Inches(.33);s.bottom_margin=Inches(.33);s.left_margin=Inches(.28);s.right_margin=Inches(.28)
def table(d,h,rows):
    t=d.add_table(rows=1,cols=len(h));t.style="Table Grid"
    tr=t.rows[0]._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");tr.append(e)
    for i,x in enumerate(h):t.rows[0].cells[i].text=x
    for rr in rows:
        c=t.add_row().cells
        for i,x in enumerate(rr):c[i].text=str(x)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(3.7)
for path,title in [(S07,"Batch 47 · System 07 Social Canonical Contracts"),(SOC,"Batch 47 · SOC-01 Social Contract Binding Ledger")]:
    d=Document(path);landscape(d);d.add_heading(title,1)
    d.add_paragraph(f"[{MARK}] Social effectful contracts preserve external-provider truth, secret isolation, candidate governance and derived-projection boundaries. Credential reveal never persists plaintext secret; local publish request acceptance never equals external publish success.")
    rows=[[u,op,fs.get("payload_schema","UNCHANGED"),fs.get("persistence_owner","UNCHANGED")] for u,(op,fs) in T.items()]
    table(d,["Control UID","Operation","Payload / Schema","Persistence Owner"],rows);d.save(path);Document(path)
postP,postO=compose(parse(SOC)),compose(parse(S07))
for uid,(op,fs) in T.items():
    for m in [postP,postO]:
        for f,v in fs.items():assert m[uid][f]==v,(uid,f,m[uid][f],v)
allm={p:compose(parse(fn)) for p,fn in PAGES.items()};rem=[]
for p,m in allm.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if missing(op) or st not in {"EFFECTFUL_EXACT","READ_EXACT"}:continue
        if missing(r.get("method_path","")):rem.append(("method_path",p,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("payload_schema","")):rem.append(("payload_schema",p,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("persistence_owner","")):rem.append(("persistence_owner",p,uid))
rc=collections.Counter(x[0] for x in rem);assert len(rem)==239,(len(rem),rc);assert rc=={"method_path":62,"payload_schema":98,"persistence_owner":79},rc
d=Document(LOGIC);landscape(d);d.add_heading("Batch 47 · System 07 Social Contract Closure",1)
table(d,["Metric","Value"],[["Pre denominator","268"],["Cells closed","29 = Payload 15 + Persistence 14"],["Post denominator","239 = Method 62 + Payload 98 + Persistence 79"],["Plaintext secret persisted","0"],["Local request treated as external success","0"],["Runtime execution claimed","False"]])
machine={"marker":MARK,"base_head":BASE_HEAD,"pre_contract_cells":268,"system07_cells_closed":29,"payload_schema_closed":15,"persistence_owner_closed":14,
"post_contract_cells":239,"remaining_method_path":62,"remaining_payload_schema":98,"remaining_persistence_owner":79,"plaintext_secret_persisted":0,"local_request_as_external_success":0,"runtime_execution_claimed":False}
d.add_paragraph("BATCH47_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")));d.save(LOGIC);Document(LOGIC)
changed=[S07,SOC,LOGIC]
Path("__batch47_remediation_report.json").write_text(json.dumps({"machine":machine,"targets":T,"patches":patches,"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed}},ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH47="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
