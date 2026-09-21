from docx import Document
import json,re
FILES=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx",
]
TERMS=[
"LOCAL_WORKING_DRAFT_EXACT",
"UI_LOCAL_OR_READ_SOURCE",
"ORCHESTRATES_EXISTING_EXACT_OPERATIONS",
"owner-orchestrated local command envelope",
"NO_PUBLIC_API_ID_IN_CURRENT_AUTHORITY",
"working draft",
"local draft",
"local-only",
"read source",
"existing exact operations",
"orchestrates existing",
"single operation",
"single runtime owner",
]
def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
hits=[]
for fn in FILES:
    d=Document(fn)
    for i,p in enumerate(d.paragraphs,1):
        t=norm(p.text)
        if t and any(x.lower() in t.lower() for x in TERMS):
            hits.append({"file":fn,"kind":"paragraph","index":i,"text":t})
    for ti,tb in enumerate(d.tables,1):
        for ri,row in enumerate(tb.rows,1):
            t=" | ".join(norm(c.text) for c in row.cells)
            if t and any(x.lower() in t.lower() for x in TERMS):
                hits.append({"file":fn,"kind":"table","table":ti,"row":ri,"text":t})
print("BATCH14_STATUS_CONTEXT="+json.dumps(hits,ensure_ascii=False,sort_keys=True))
