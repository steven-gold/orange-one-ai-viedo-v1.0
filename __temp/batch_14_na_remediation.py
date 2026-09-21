from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib

MARK="ACPOS-20260921-BATCH-14-LOCAL-ORCHESTRATION-NA-SEMANTICS-V1"
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S09="09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
EDIT="ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
IAM="ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
 S05:"4bcaef86cff9a59db1b5f312832d7e6152b1b7ba",
 S09:"fd6498f43a8b629a2fdb43c738a5bd33d99979f1",
 EDIT:"09d45fce05a0c3ca71f240a225ef5135d86c7a50",
 IAM:"1cff72bf4dca24b61e1efc8626c106c5fbb45da9",
 LOGIC:"67eada556abc54794059cb46b90e16f567f9f0e0",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":EDIT,
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":IAM,
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()
for f,s in EXPECTED_SHA.items():
    assert blob(f)==s,(f,blob(f),s)

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def display_only_type(t):
    u=(t or "").upper()
    return u in {"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"} or any(x in u for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(h,"control uid")
        if ci is None:continue
        idx={
          "control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),
          "method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),
          "persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status"),
        }
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"table":ti+1,"row":ri,"headers":h})
            out.append(rec)
    return out

def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out

# Batch-13 rules.
def is_wb_read_projection_action_na(r):
    return (
        missing(r.get("action","")) and r.get("type","")=="SECTION_OPEN" and
        r.get("gate","")=="PAGE_READ" and r.get("permission","")=="workspace.dashboard.view" and
        r.get("payload_schema","")=="N/A_READ_PROJECTION" and
        r.get("operation","")=="getDashboardReadModel" and
        r.get("method_path","")=="GET /v1/dashboard/read-model" and
        r.get("runtime_owner","")=="DASHBOARD_READ_MODEL" and r.get("runtime_status","")=="READ_EXACT"
    )
def is_aiapi_direct_operation_action_na(r):
    return (
        missing(r.get("action","")) and r.get("type","")=="BUTTON_OR_ROW_ACTION" and
        r.get("payload_schema","")=="AIAPI Page Operation-specific Form / Provider Profile Field Contract" and
        not missing(r.get("permission","")) and not missing(r.get("operation","")) and
        not missing(r.get("method_path","")) and not missing(r.get("runtime_owner","")) and
        r.get("runtime_status","") in {"EFFECTFUL_EXACT","READ_EXACT"}
    )

# Batch-14 exact structural applicability.
def is_edit_local_working_draft_operation_na(r):
    return (
        missing(r.get("operation",""))
        and r.get("runtime_status","")=="LOCAL_WORKING_DRAFT_EXACT"
        and r.get("method_path","")=="NO_PUBLIC_API_BY_AUTHORITY"
        and not missing(r.get("action",""))
        and not missing(r.get("gate",""))
        and not missing(r.get("permission",""))
        and not missing(r.get("runtime_owner",""))
    )
def is_iam_orchestration_operation_na(r):
    return (
        missing(r.get("operation",""))
        and r.get("runtime_status","")=="ORCHESTRATES_EXISTING_EXACT_OPERATIONS"
        and not missing(r.get("action",""))
        and not missing(r.get("gate",""))
        and not missing(r.get("permission",""))
    )
def is_iam_orchestration_runtime_owner_na(r):
    return (
        missing(r.get("runtime_owner",""))
        and r.get("runtime_status","")=="ORCHESTRATES_EXISTING_EXACT_OPERATIONS"
        and not missing(r.get("action",""))
        and not missing(r.get("gate",""))
        and not missing(r.get("permission",""))
    )

def classify_pre(page,field,r):
    v=r.get(field,"");st=r.get("runtime_status","")
    if not missing(v):return "BOUND"
    if "UI_LOCAL_EXACT" in st:
        if field in {"operation","runtime_owner"}:return "LEGITIMATE_NA_UI_LOCAL"
        if field=="action" and display_only_type(r.get("type","")):return "LEGITIMATE_NA_UI_LOCAL"
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in st and field=="operation":return "CLOSED_BY_OWNER_LOCAL_COMMAND"
    if "READ_EXACT" in st and display_only_type(r.get("type","")):
        if field=="action":return "LEGITIMATE_NA_READ_PRESENTATION"
        if field=="operation" and not missing(r.get("runtime_owner","")):return "READ_OWNER_BOUND_OPERATION_UNSPECIFIED"
    if field=="runtime_owner" and any(x in st for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return "KNOWN_RUNTIME_BLOCKER"
    if v=="SOURCE_NOT_DEFINED":return "SOURCE_RESOLUTION_REQUIRED"
    if field=="action" and page=="WB-01" and is_wb_read_projection_action_na(r):return "LEGITIMATE_NA_READ_PROJECTION_ACTION"
    if field=="action" and page=="AIAPI-01" and is_aiapi_direct_operation_action_na(r):return "LEGITIMATE_NA_DIRECT_OPERATION_ACTION"
    return "DEFINITION_BINDING_GAP"

def classify_post(page,field,r):
    pre=classify_pre(page,field,r)
    if pre!="DEFINITION_BINDING_GAP":return pre
    if page=="EDIT-01" and field=="operation" and is_edit_local_working_draft_operation_na(r):
        return "LEGITIMATE_NA_LOCAL_WORKING_DRAFT_OPERATION"
    if page=="IAM-01" and field=="operation" and is_iam_orchestration_operation_na(r):
        return "LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_OPERATION"
    if page=="IAM-01" and field=="runtime_owner" and is_iam_orchestration_runtime_owner_na(r):
        return "LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_RUNTIME_OWNER"
    return pre

def inventory(classifier):
    rows=[];counts=collections.Counter()
    for page,fn in PAGES.items():
        bm=compose(parse_controls(fn))
        for uid,r in bm.items():
            for field in ["action","gate","permission","operation","runtime_owner"]:
                cls=classifier(page,field,r)
                counts[cls]+=1
                if cls=="DEFINITION_BINDING_GAP":rows.append((page,uid,field))
    return rows,counts

pre_rows,pre_counts=inventory(classify_pre)
assert len(pre_rows)==209,len(pre_rows)
edit_map=compose(parse_controls(EDIT))
iam_map=compose(parse_controls(IAM))
edit_targets=sorted(uid for uid,r in edit_map.items() if is_edit_local_working_draft_operation_na(r))
iam_op_targets=sorted(uid for uid,r in iam_map.items() if is_iam_orchestration_operation_na(r))
iam_owner_targets=sorted(uid for uid,r in iam_map.items() if is_iam_orchestration_runtime_owner_na(r))
assert len(edit_targets)==43,(len(edit_targets),edit_targets)
assert iam_op_targets==["IAM-01-BTN-COMPLETE"],iam_op_targets
assert iam_owner_targets==["IAM-01-BTN-COMPLETE"],iam_owner_targets

expected_removed={("EDIT-01",uid,"operation") for uid in edit_targets}|{
    ("IAM-01","IAM-01-BTN-COMPLETE","operation"),
    ("IAM-01","IAM-01-BTN-COMPLETE","runtime_owner"),
}
assert expected_removed.issubset(set(pre_rows)),sorted(expected_removed-set(pre_rows))

# Same structural truth must exist in canonical owner rows.
s05=compose(parse_controls(S05));s09=compose(parse_controls(S09))
for uid in edit_targets:
    assert uid in s05,uid
    assert is_edit_local_working_draft_operation_na(s05[uid]),(uid,s05[uid])
assert "IAM-01-BTN-COMPLETE" in s09
assert is_iam_orchestration_operation_na(s09["IAM-01-BTN-COMPLETE"]),s09["IAM-01-BTN-COMPLETE"]
assert is_iam_orchestration_runtime_owner_na(s09["IAM-01-BTN-COMPLETE"]),s09["IAM-01-BTN-COMPLETE"]

post_rows,post_counts=inventory(classify_post)
removed=set(pre_rows)-set(post_rows);added=set(post_rows)-set(pre_rows)
assert removed==expected_removed,{"missing":sorted(expected_removed-removed),"unexpected":sorted(removed-expected_removed)}
assert not added,sorted(added)
assert len(post_rows)==164,len(post_rows)

# Helpers.
def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.34);sec.bottom_margin=Inches(.34);sec.left_margin=Inches(.30);sec.right_margin=Inches(.30)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=4.7):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row):c[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for r in p.runs:r.font.size=Pt(fs)
    return t

# Canonical System 05 rule.
d=Document(S05)
txt="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
assert MARK not in txt,"S05_BATCH14_ALREADY_APPLIED"
add_landscape(d)
d.add_heading("Batch 14 · Local Working-Draft Operation Applicability Rule",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::SYSTEM05] ").bold=True
p.add_run("For EDIT controls with Runtime Status=LOCAL_WORKING_DRAFT_EXACT, Method/Path=NO_PUBLIC_API_BY_AUTHORITY, an existing exact Action UID/Gate/Permission, and a non-empty local Runtime Owner, the interaction is a local working-draft command handled inside that owner boundary. No independent registered Operation is applicable. Operation MUST remain '—'; no public or synthetic Operation ID may be invented.")
add_table(d,["Target UID","Field","Runtime Status","Method/Path","Local Runtime Owner","Applicability"],[
    [uid,"Operation","LOCAL_WORKING_DRAFT_EXACT","NO_PUBLIC_API_BY_AUTHORITY",s05[uid]["runtime_owner"],"LEGITIMATE_NA_LOCAL_WORKING_DRAFT_OPERATION"]
    for uid in edit_targets
],4.1)
d.save(S05);Document(S05)

# Canonical System 09 orchestration rule.
d=Document(S09)
txt="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
assert MARK not in txt,"S09_BATCH14_ALREADY_APPLIED"
add_landscape(d)
d.add_heading("Batch 14 · Aggregate Orchestration Operation / Owner Applicability Rule",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::SYSTEM09] ").bold=True
p.add_run("IAM-01-BTN-COMPLETE is explicitly classified ORCHESTRATES_EXISTING_EXACT_OPERATIONS. The aggregate UI control delegates to already-defined exact IAM operations and their respective runtime owners. Therefore the aggregate row has no single Operation and no single Runtime Owner; both cells MUST remain '—'. This rule does not alter constituent operation contracts.")
add_table(d,["Target UID","Field","Runtime Status","Applicability","Handling"],[
["IAM-01-BTN-COMPLETE","Operation","ORCHESTRATES_EXISTING_EXACT_OPERATIONS","LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_OPERATION","Delegate to constituent exact operations"],
["IAM-01-BTN-COMPLETE","Runtime Owner","ORCHESTRATES_EXISTING_EXACT_OPERATIONS","LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_RUNTIME_OWNER","Delegate to constituent operation owners"],
],5.0)
d.save(S09);Document(S09)

# Page-level applicability ledgers, not product controls.
d=Document(EDIT)
txt="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
assert MARK not in txt,"EDIT_BATCH14_ALREADY_APPLIED"
add_landscape(d)
d.add_heading("Batch 14 · Local Working-Draft Operation Applicability Ledger",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::PAGE::EDIT-01] ").bold=True
p.add_run("Applicability-only ledger. Existing product rows remain unchanged; Operation stays '—'. Target UID is used deliberately so this table cannot become a second Control denominator.")
add_table(d,["Target UID","Field","Current Cell","Runtime Status","Method/Path","Runtime Owner","Handling"],[
    [uid,"Operation","—","LOCAL_WORKING_DRAFT_EXACT","NO_PUBLIC_API_BY_AUTHORITY",edit_map[uid]["runtime_owner"],"DO_NOT_INVENT_OPERATION"]
    for uid in edit_targets
],4.0)
d.save(EDIT);Document(EDIT)

d=Document(IAM)
txt="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
assert MARK not in txt,"IAM_BATCH14_ALREADY_APPLIED"
add_landscape(d)
d.add_heading("Batch 14 · Aggregate Orchestration Applicability Ledger",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::PAGE::IAM-01] ").bold=True
p.add_run("Applicability-only ledger. IAM-01-BTN-COMPLETE remains an orchestration aggregate; its product row is unchanged.")
add_table(d,["Target UID","Field","Current Cell","Runtime Status","Handling"],[
["IAM-01-BTN-COMPLETE","Operation","—","ORCHESTRATES_EXISTING_EXACT_OPERATIONS","DELEGATE_TO_EXISTING_EXACT_OPERATIONS"],
["IAM-01-BTN-COMPLETE","Runtime Owner","—","ORCHESTRATES_EXISTING_EXACT_OPERATIONS","DELEGATE_TO_CONSTITUENT_RUNTIME_OWNERS"],
],5.0)
d.save(IAM);Document(IAM)

# Central logic.
logic=Document(LOGIC)
txt="\n".join([p.text for p in logic.paragraphs]+[c.text for t in logic.tables for row in t.rows for c in row.cells])
assert MARK not in txt,"LOGIC_BATCH14_ALREADY_APPLIED"
add_landscape(logic)
logic.add_heading("Batch 14 · Local / Aggregate Operation Applicability Closure",level=1)
p=logic.add_paragraph();p.add_run("["+MARK+"] ").bold=True
p.add_run("Fresh Batch-14 audit examined 189 Operation/Runtime Owner gaps. It found 5 weak single-value relations, 66 conflicting relations, and 118 with no exact relation. Weak relations based only on shared Gate/Permission are not accepted as construction authority. Batch 14 closes only structural N/A classes proven by exact Runtime Status plus exact row boundary fields.")

add_table(logic,["Class","Count","Decision"],[
["Pre-batch Definition Binding Gap",209,"Current Batch-13 denominator."],
["EDIT LOCAL_WORKING_DRAFT_EXACT Operation",43,"N/A only when Method/Path=NO_PUBLIC_API_BY_AUTHORITY and local Runtime Owner is already exact."],
["IAM aggregate Operation",1,"N/A because row explicitly ORCHESTRATES_EXISTING_EXACT_OPERATIONS."],
["IAM aggregate Runtime Owner",1,"N/A because aggregate delegates to constituent exact operation owners."],
["Total Batch-14 N/A reclassification",45,"No Operation or Runtime Owner value written."],
["Weak unique relation candidates",5,"Rejected for this batch; Gate/Permission-only uniqueness is not operation-specific authority."],
["Exact relation conflicts",66,"Remain fail-closed."],
["No exact relation",118,"Remain canonical-spec remediation."],
["UI_LOCAL_OR_READ_SOURCE subset",42,"Remain unresolved; status is intentionally ambiguous between UI-local and read-source behavior."],
["Post-batch Definition Binding Gap",164,"163 unique-owner canonical-spec gaps + 1 preserved SYS Gate conflict."],
["Business binding mutations",0,"Product registry cells remain unchanged."],
],4.7)

logic.add_heading("45 Explicit N/A Rows",level=2)
add_table(logic,["Page","Target UID","Field","Applicability Class"],[
    *[["EDIT-01",uid,"Operation","LEGITIMATE_NA_LOCAL_WORKING_DRAFT_OPERATION"] for uid in edit_targets],
    ["IAM-01","IAM-01-BTN-COMPLETE","Operation","LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_OPERATION"],
    ["IAM-01","IAM-01-BTN-COMPLETE","Runtime Owner","LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_RUNTIME_OWNER"],
],3.9)

logic.add_heading("Batch 14 Fail-Closed Boundary",level=2)
add_table(logic,["Rule","Requirement"],[
["B14-01","LOCAL_WORKING_DRAFT_EXACT alone is insufficient; Method/Path must be NO_PUBLIC_API_BY_AUTHORITY and a local Runtime Owner must already be explicit."],
["B14-02","ORCHESTRATES_EXISTING_EXACT_OPERATIONS means the aggregate row delegates to constituent contracts; it does not authorize inventing an aggregate Operation or Runtime Owner."],
["B14-03","Shared Gate or Permission yielding one candidate Operation is not accepted unless the relation is operation-specific in current authority."],
["B14-04","UI_LOCAL_OR_READ_SOURCE remains unresolved until each owner defines local-vs-read behavior or an exact binding."],
["B14-05","The preserved SYS-01-BTN-NAV-OPEN Gate conflict remains fail-closed."],
],4.8)

machine={
 "marker":MARK,
 "pre_definition_gaps":209,
 "edit_local_working_draft_operation_na":43,
 "iam_aggregate_operation_na":1,
 "iam_aggregate_runtime_owner_na":1,
 "na_reclassified":45,
 "post_definition_gaps":164,
 "remaining_unique_owner_remediation":163,
 "preserved_sys_gate_conflict":1,
 "weak_unique_relation_rejected":5,
 "exact_relation_conflicts":66,
 "no_exact_relation":118,
 "ui_local_or_read_source_unresolved":42,
 "business_binding_mutations":0,
}
logic.add_paragraph("BATCH14_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

# Final product rows still blank in target fields.
edit_after=compose(parse_controls(EDIT));iam_after=compose(parse_controls(IAM))
for uid in edit_targets:assert missing(edit_after[uid]["operation"]),(uid,edit_after[uid]["operation"])
assert missing(iam_after["IAM-01-BTN-COMPLETE"]["operation"])
assert missing(iam_after["IAM-01-BTN-COMPLETE"]["runtime_owner"])

report={
 "machine":machine,
 "edit_targets":edit_targets,
 "iam_targets":[
   {"uid":"IAM-01-BTN-COMPLETE","field":"operation"},
   {"uid":"IAM-01-BTN-COMPLETE","field":"runtime_owner"},
 ],
 "removed_from_definition_gap":[{"page":p,"uid":u,"field":f} for p,u,f in sorted(removed)],
 "pre_classification_counts":dict(pre_counts),
 "post_classification_counts":dict(post_counts),
 "output_hashes":{f:blob(f) for f in [S05,S09,EDIT,IAM,LOGIC]},
}
Path("__batch14_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH14_REMEDIATION="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
