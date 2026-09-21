from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-45-SYSTEM09-IAM-SYS-CONTRACT-REMEDIATION-V1"
BASE_HEAD="ec5ed099f9ab70d5498aaedf96726bbf6f53686f"
S09="09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
IAM="ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx"
SYS="ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S09:"2ae1d6462c2c6338a0f437721b0c62dd938749a3",
IAM:"575bb09781ae825ab48f20f2bde524165aab9fc3",
SYS:"84deebe31d6d8d66fad294490920581654657412",
LOGIC:"3ee5b096ef9adcf6338eb7466c2ccecb969dbcbd",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":IAM,
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":SYS,
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]
TARGETS={
"IAM-01-BTN-SAVE-DRAFT":{
 "page":"IAM-01","operation":"saveDraft",
 "fields":{"payload_schema":"SaveIAMAuthorizationDraftRequest","persistence_owner":"IAM Runtime / Authorization Draft + Audit"},
 "basis":"Save Draft persists only editable authorization draft state and audit/correlation identity. Department preset/select-all remain draft conveniences and do not grant runtime authorization."
},
"IAM-01-BTN-VALIDATE":{
 "page":"IAM-01","operation":"validateDraft",
 "fields":{"payload_schema":"ValidateIAMAuthorizationDraftRequest","persistence_owner":"IAM Runtime / Draft Validation Evidence"},
 "basis":"Validation evaluates the exact governed draft and records validation evidence/result. It does not materialize effective assignments or approve the draft."
},
"IAM-01-BTN-PREVIEW":{
 "page":"IAM-01","operation":"previewAuthorizationImpact",
 "fields":{"payload_schema":"PreviewAuthorizationImpactRequest","persistence_owner":"IAM Runtime / Authorization Impact Preview + Audit"},
 "basis":"Preview persists derived impact-preview evidence for an exact draft/version. Preview state never becomes effective permission state and stale preview must be rejected or refreshed."
},
"SYS-01-BTN-CANDIDATE-CREATE":{
 "page":"SYS-01","operation":"createCandidate",
 "fields":{"payload_schema":"CreateSystemChangeCandidateRequest","persistence_owner":"System 09 / ChangeCandidate + Audit"},
 "basis":"Create Candidate persists an append-only bounded ChangeCandidate under System 09. It is not Current Authority and cannot mutate Production."
},
"SYS-01-BTN-CR-CREATE":{
 "page":"SYS-01","operation":"createChangeRequest",
 "fields":{"payload_schema":"CreateSystemChangeRequestRequest","persistence_owner":"System 09 / System Change Request + Audit"},
 "basis":"Change Request persists governed scope/reason/impact/risk/validation references and audit lineage. It remains subject to permission/gate and later authorization."
},
"SYS-01-BTN-SANDBOX-TEST":{
 "page":"SYS-01","operation":"runSandboxTest",
 "fields":{"persistence_owner":"System 09 / ValidationRun + Sandbox Evidence/Audit"},
 "basis":"Sandbox persists exact candidate/version test run and evidence only. Passing sandbox evidence is not Production acceptance and must be invalidated when candidate version changes."
},
}

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out
def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():
    got=blob(f);assert got==s,(f,got,s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

page_files={"IAM-01":IAM,"SYS-01":SYS}
pre_pages={p:compose(parse_controls(fn)) for p,fn in page_files.items()}
pre_owner=compose(parse_controls(S09))
for uid,s in TARGETS.items():
    for m,label in [(pre_pages[s["page"]],"PAGE"),(pre_owner,"OWNER")]:
        assert uid in m,(label,uid,"MISSING")
        r=m[uid]
        assert r["operation"]==s["operation"],(label,uid,r["operation"],s["operation"])
        assert r["runtime_status"]=="EFFECTFUL_EXACT",(label,uid,r["runtime_status"])
        for field,new in s["fields"].items():
            assert missing(r[field]),(label,uid,field,r[field],new)
        assert not missing(r["method_path"]),(label,uid,"METHOD_MISSING")
    if uid=="SYS-01-BTN-SANDBOX-TEST":
        assert pre_pages["SYS-01"][uid]["payload_schema"]=="AIAPI Page Operation-specific Form / Provider Profile Field Contract"

def patch(path,target_uids):
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact_index(h,"Control UID")
        if ci is None:continue
        idx={"payload_schema":hfind(h,"payload / schema","payload","schema"),"persistence_owner":hfind(h,"persistence owner")}
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals):continue
            uid=vals[ci]
            if uid not in target_uids:continue
            for field,new in TARGETS[uid]["fields"].items():
                fi=idx[field]
                if fi is None or fi>=len(row.cells):continue
                old=norm(row.cells[fi].text)
                if old==new:hits[(uid,field)]+=1;continue
                assert missing(old),(path,uid,field,old,new)
                row.cells[fi].text=new;hits[(uid,field)]+=1
    expected={(u,f) for u in target_uids for f in TARGETS[u]["fields"]}
    assert set(hits)==expected,(path,dict(hits),expected-set(hits))
    d.save(path);Document(path)
    return {f"{u}:{f}":n for (u,f),n in hits.items()}

patches={
"IAM_PAGE":patch(IAM,[u for u,s in TARGETS.items() if s["page"]=="IAM-01"]),
"SYS_PAGE":patch(SYS,[u for u,s in TARGETS.items() if s["page"]=="SYS-01"]),
"SYSTEM09_OWNER":patch(S09,list(TARGETS)),
}

def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.33);sec.bottom_margin=Inches(.33);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=3.9):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for rr in rows:
        cells=t.add_row().cells
        for i,v in enumerate(rr):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
    return t

for path,title,prefix,uids in [
(S09,"Batch 45 · System 09 IAM/SYS Canonical Contracts","SYSTEM09",list(TARGETS)),
(IAM,"Batch 45 · IAM Draft / Validation / Preview Binding Ledger","PAGE::IAM-01",[u for u,s in TARGETS.items() if s["page"]=="IAM-01"]),
(SYS,"Batch 45 · SYS Change / Sandbox Binding Ledger","PAGE::SYS-01",[u for u,s in TARGETS.items() if s["page"]=="SYS-01"]),
]:
    d=Document(path);add_landscape(d);d.add_heading(title,level=1)
    p=d.add_paragraph();p.add_run(f"[{MARK}::{prefix}] ").bold=True
    p.add_run("These bindings preserve the authorization boundary: drafts, validations, previews, candidates, change requests and sandbox runs remain non-authoritative until their later governed gates are satisfied. No self-approval or direct Production mutation is introduced.")
    rr=[]
    for uid in uids:
        s=TARGETS[uid]
        rr.append([uid,s["operation"],s["fields"].get("payload_schema","UNCHANGED"),s["fields"].get("persistence_owner","UNCHANGED"),s["basis"]])
    add_table(d,["Control UID","Operation","Payload / Schema","Persistence Owner","Semantics"],rr,3.6)
    d.save(path);Document(path)

post_pages={p:compose(parse_controls(fn)) for p,fn in page_files.items()}
post_owner=compose(parse_controls(S09))
for uid,s in TARGETS.items():
    for m,label in [(post_pages[s["page"]],"PAGE"),(post_owner,"OWNER")]:
        for f,v in s["fields"].items():
            assert m[uid][f]==v,(label,uid,f,m[uid][f],v)

all_maps={p:compose(parse_controls(fn)) for p,fn in PAGES.items()}
remaining=[]
for page,m in all_maps.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if missing(op) or st not in {"EFFECTFUL_EXACT","READ_EXACT"}:continue
        if missing(r.get("method_path","")):remaining.append(("method_path",page,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("payload_schema","")):remaining.append(("payload_schema",page,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("persistence_owner","")):remaining.append(("persistence_owner",page,uid))
rc=collections.Counter(f for f,_,_ in remaining)
assert len(remaining)==268,(len(remaining),rc)
assert rc=={"method_path":62,"payload_schema":113,"persistence_owner":93},rc

logic=Document(LOGIC);add_landscape(logic);logic.add_heading("Batch 45 · System 09 IAM/SYS Contract Closure",level=1)
p=logic.add_paragraph();p.add_run(f"[{MARK}] ").bold=True
p.add_run("System 09 closes eleven IAM/SYS cells while preserving authorization and simulation boundaries. IAM draft/validation/preview do not create effective permissions; ChangeCandidate/ChangeRequest/Sandbox remain pre-authorization artifacts/evidence and do not mutate Current Authority or Production.")
add_table(logic,["Metric","Value","Result"],[
["Pre contract denominator",279,"Method 62 + Payload 118 + Persistence 99"],
["System 09 cells closed",11,"Payload 5 + Persistence 6"],
["Post contract denominator",268,"Method 62 + Payload 113 + Persistence 93"],
["Self-approval introduced",0,"Forbidden"],
["Production mutation from sandbox/candidate",0,"None"],
["Runtime execution claimed","False","Design-contract closure only"],
],4.1)
machine={"marker":MARK,"base_head":BASE_HEAD,"pre_contract_cells":279,"system09_cells_closed":11,
"payload_schema_closed":5,"persistence_owner_closed":6,"post_contract_cells":268,
"remaining_method_path":62,"remaining_payload_schema":113,"remaining_persistence_owner":93,
"self_approval_introduced":0,"production_mutation_from_sandbox_candidate":0,"runtime_execution_claimed":False}
logic.add_paragraph("BATCH45_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

changed=[S09,IAM,SYS,LOGIC]
report={"machine":machine,"targets":TARGETS,"patches":patches,"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed}}
Path("__batch45_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH45="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
