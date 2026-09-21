from pathlib import Path
from docx import Document
import json,re,subprocess

BASE_HEAD="ff050d80da5495ed70a4e4a6380d89226e5f8e00"
S07="07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
SOC="ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx"
OPS=[
"bindSocialAccount","unbindSocialAccount","decideCandidate","saveDraft","revealSocialCredential",
"setKillSwitch","completeSocialManualAction","configureGovernedResource","configureSocialTargetPolicy",
"requestSocialTargetPublish","refreshProjection","createSocialTargetDiscovery","requestSocialTargetJoin","searchProjection"
]
KW=["social","account","credential","secret","target","discovery","join","publish","publishing","policy","kill switch","manual action",
"projection","search","candidate","draft","provider","adapter","evidence","audit","external","token","oauth","platform","community"]

assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
rx=re.compile("|".join(re.escape(x) for x in OPS+KW),re.I)

def collect(path):
    d=Document(path);out={"paragraphs":[],"rows":[]}
    for i,p in enumerate(d.paragraphs):
        t=norm(p.text)
        if t and rx.search(t):out["paragraphs"].append({"index":i,"text":t[:4000]})
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];joined=" | ".join(vals)
            if joined and rx.search(joined):
                out["rows"].append({"table":ti+1,"row":ri,"headers":h,"values":vals,"joined":joined[:7000]})
    return out

owner=collect(S07);page=collect(SOC)
summary={"operations":len(OPS),"owner_paragraph_hits":len(owner["paragraphs"]),"owner_table_hits":len(owner["rows"]),
"page_paragraph_hits":len(page["paragraphs"]),"page_table_hits":len(page["rows"])}
report={"marker":"ACPOS-20260922-BATCH-46-SYSTEM07-SOCIAL-CONTRACT-DESIGN-AUDIT-V1","base_head":BASE_HEAD,
"summary":summary,"operations":OPS,"owner_evidence":owner,"page_evidence":page}
Path("__batch46_system07_social_contract_audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
Path("__batch46_summary.txt").write_text(json.dumps(summary,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print("BATCH46_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
