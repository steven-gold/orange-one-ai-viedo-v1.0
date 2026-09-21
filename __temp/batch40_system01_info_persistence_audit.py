from pathlib import Path
from docx import Document
import json,re,subprocess

BASE_HEAD="1be8067f103dc2a1402771cfd4635b622a6dbe26"
S01="01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx"
INFO="ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
OPS=["refreshProjection","searchProjection","exportProjection"]
KW=["projection","refresh","search","export","information","knowledge","persistence","persist","materialized","cache","index","audit","evidence","snapshot","read model"]

assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
rx=re.compile("|".join(re.escape(x) for x in OPS+KW),re.I)

def collect(path):
    d=Document(path);out={"paragraphs":[],"rows":[]}
    for i,p in enumerate(d.paragraphs):
        t=norm(p.text)
        if t and rx.search(t):out["paragraphs"].append({"index":i,"text":t[:3500]})
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];joined=" | ".join(vals)
            if joined and rx.search(joined):
                out["rows"].append({"table":ti+1,"row":ri,"headers":h,"values":vals,"joined":joined[:6000]})
    return out

owner=collect(S01);page=collect(INFO)
summary={
"operations":3,
"owner_paragraph_hits":len(owner["paragraphs"]),"owner_table_hits":len(owner["rows"]),
"page_paragraph_hits":len(page["paragraphs"]),"page_table_hits":len(page["rows"]),
}
report={"marker":"ACPOS-20260922-BATCH-40-SYSTEM01-INFO-PERSISTENCE-DESIGN-AUDIT-V1",
"base_head":BASE_HEAD,"summary":summary,"operations":OPS,"owner_evidence":owner,"page_evidence":page}
Path("__batch40_system01_info_persistence_audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
Path("__batch40_summary.txt").write_text(json.dumps(summary,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print("BATCH40_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
