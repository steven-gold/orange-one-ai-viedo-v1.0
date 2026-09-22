from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-54-QA-CONTRACT-REMEDIATION-V1"
BASE_HEAD="fc451684389519dfbb0ff540816e7b48f85f38ed"
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
QA="ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S05:"9087fb79fb54fe976d3897d1d70e5b791112990d",
QA:"bcf92170271fb4676396bb33d38f69d597c23b86",
LOGIC:"2ac2ae31b3d5a37fb246ac905bcd4e650dc636a8",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx","ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx","EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx","IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":QA,"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx","STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx","VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx"}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

TARGETS={
"QA-01-BTN-START-REVIEW":{
 "operation":"startQaReview","route":"POST /v1/qa/reviews","payload":"StartQaReviewRequest",
 "persistence":"QAService / QAReviewCase + QAScorecard + Audit",
 "basis":"Starts QA only from ExactOutputVersion + QAValidationScriptView + approved CriteriaVersion/QualityGatePolicy + Evidence + Rights/Policy + checksum/provenance. Missing required criterion/evidence is BLOCK."
},
"QA-01-BTN-FINDING-CREATE":{
 "operation":"createFinding","route":"POST /v1/findings","payload":"CreateFindingRequest",
 "persistence":"public.findings",
 "basis":"Finding must bind owning domain, exact source version, evidence, severity and expected correction; route/persistence are existing Current authority."
},
"QA-01-BTN-CORRECTION":{
 "operation":"createCorrectionRequest","route":"POST /v1/correction-requests","payload":"CreateCorrectionRequestPayload",
 "persistence":"public.correction_requests",
 "basis":"CorrectionRequest routes to the original canonical owner with exact source + required delta/evidence. QA does not directly edit upstream output."
},
"QA-01-BTN-RECHECK":{
 "operation":"startRecheck","route":"POST /v1/qa/reviews/{reviewId}/rechecks","payload":"StartQaRecheckRequest",
 "persistence":"QAService / QARecheckCase + QAScorecard + Audit",
 "basis":"Recheck is only against a new exact corrected version; historical PASS cannot close a finding on a newer revision."
},
"QA-01-BTN-RELEASE-CREATE":{
 "operation":"createReleasePackage","route":"POST /v1/qa/release-candidates","payload":"CreateReleaseCandidateRequest",
 "persistence":"ReleaseService / ReleaseCandidate + Provenance/Audit",
 "basis":"Creates ReleaseCandidate only from exact QA PASS/version/checksum/provenance. ReleaseCandidate is not published/released-to-channel truth."
},
"QA-01-BTN-AUTO-START":{
 "operation":"startAutoTopicQaReview","route":"POST /v1/qa/reviews/auto","payload":"StartAutoTopicQaReviewRequest",
 "persistence":"QAService / AutoQAReviewCase + QAScorecard + Evaluation Evidence",
 "basis":"AUTO review evaluates exact output against approved criteria/evidence. Machine score cannot bypass required manual review or auto-release."
},
"QA-01-BTN-FINDING-POINT":{
 "operation":"setQAFindingPoint","route":"LOCAL LOCAL","payload":"SetQAFindingPointCommand",
 "persistence":"CLIENT/QA / EvidenceCaptureDraft (session-local, non-canonical)",
 "basis":"Client-only evidence marker mutation. It records a local draft point bound to the current exact output; it does not create a Finding until the governed Finding create command is submitted."
},
"QA-01-BTN-EVIDENCE-IN":{
 "operation":"setQAEvidenceRangeIn","route":"LOCAL LOCAL","payload":"SetQAEvidenceRangeInCommand",
 "persistence":"CLIENT/QA / EvidenceCaptureDraft (session-local, non-canonical)",
 "basis":"Client-only local range-start capture. No server API or canonical QA decision is implied."
},
"QA-01-BTN-EVIDENCE-OUT":{
 "operation":"setQAEvidenceRangeOut","route":"LOCAL LOCAL","payload":"SetQAEvidenceRangeOutCommand",
 "persistence":"CLIENT/QA / EvidenceCaptureDraft (session-local, non-canonical)",
 "basis":"Client-only local range-end capture. Range remains a draft until bound into a governed Finding/evidence submission."
},
"QA-01-BTN-MANUAL-MODIFY":{
 "operation":"modifyManualReview","route":"POST /v1/qa/manual-reviews/{reviewCaseId}/modify","payload":"ModifyManualReviewRequest",
 "persistence":"ManualReviewService / ManualReviewCase + Human Decision/Audit",
 "basis":"Modifies the exact manual review case under an authorized reviewer. Required manual review cannot be bypassed; unresolved case/reviewer/state fails closed."
},
"QA-01-BTN-MANUAL-PASS":{
 "operation":"passManualReview","route":"POST /v1/qa/manual-reviews/{reviewCaseId}/pass","payload":"PassManualReviewRequest",
 "persistence":"ManualReviewService / ManualReviewCase + Human Decision/Audit",
 "basis":"Records authorized manual-review PASS for exact case/output/version. Machine score alone is insufficient and cannot substitute reviewer evidence."
},
}

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),"permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),"operation":hfind(h,"operation"),
          "method_path":hfind(h,"method / path","method","path"),"runtime_owner":hfind(h,"runtime owner"),
          "persistence_owner":hfind(h,"persistence owner"),"runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out
def collect_uid_operations(path):
    d=Document(path);out=collections.defaultdict(set)
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ui=hfind(h,"control uid","target uid");oi=hfind(h,"operation")
        if ui is None or oi is None:continue
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ui>=len(vals) or oi>=len(vals):continue
            uid,op=vals[ui],vals[oi]
            if uid and uid not in {"—","-"} and op and op not in {"—","-"}:out[uid].add(op)
    return out
def blob(path):
    b=Path(path).read_bytes();return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():assert blob(f)==s,(f,blob(f),s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""
assert len(TARGETS)==11

pre_page=compose(parse_controls(QA));pre_owner=compose(parse_controls(S05));owner_ops=collect_uid_operations(S05)
for uid,s in TARGETS.items():
    assert pre_page[uid]["operation"]==s["operation"],("PAGE",uid,pre_page[uid]["operation"],s["operation"])
    assert s["operation"] in owner_ops.get(uid,set()),("OWNER_OPERATION_MISSING",uid,s["operation"],sorted(owner_ops.get(uid,set())))
    for m,label in [(pre_page,"PAGE"),(pre_owner,"OWNER")]:
        r=m[uid];assert r["runtime_status"]=="EFFECTFUL_EXACT",(label,uid,r["runtime_status"])
        assert missing(r["payload_schema"]),(label,uid,"PAYLOAD_ALREADY",r["payload_schema"])
        if uid not in {"QA-01-BTN-FINDING-CREATE","QA-01-BTN-CORRECTION"}:
            assert missing(r["method_path"]),(label,uid,"METHOD_ALREADY",r["method_path"])
            assert missing(r["persistence_owner"]),(label,uid,"PERSIST_ALREADY",r["persistence_owner"])
        else:
            assert r["method_path"]==s["route"],(label,uid,r["method_path"],s["route"])
            assert r["persistence_owner"]==s["persistence"],(label,uid,r["persistence_owner"],s["persistence"])

def patch(path):
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact_index(h,"Control UID")
        if ci is None:continue
        idx={"method_path":hfind(h,"method / path","method","path"),"payload_schema":hfind(h,"payload / schema","payload","schema"),
             "persistence_owner":hfind(h,"persistence owner")}
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if uid not in TARGETS:continue
            s=TARGETS[uid]
            assignments={"payload_schema":s["payload"]}
            if uid not in {"QA-01-BTN-FINDING-CREATE","QA-01-BTN-CORRECTION"}:
                assignments["method_path"]=s["route"];assignments["persistence_owner"]=s["persistence"]
            for f,new in assignments.items():
                i=idx[f]
                if i is None or i>=len(row.cells):continue
                old=norm(row.cells[i].text)
                if old==new:hits[(uid,f)]+=1;continue
                assert missing(old),(path,uid,f,old,new)
                row.cells[i].text=new;hits[(uid,f)]+=1
    expected={(u,"payload_schema") for u in TARGETS}
    expected|={(u,"method_path") for u in TARGETS if u not in {"QA-01-BTN-FINDING-CREATE","QA-01-BTN-CORRECTION"}}
    expected|={(u,"persistence_owner") for u in TARGETS if u not in {"QA-01-BTN-FINDING-CREATE","QA-01-BTN-CORRECTION"}}
    assert set(hits)==expected,(path,sorted(expected-set(hits)),dict(hits))
    d.save(path);Document(path)
    return {f"{u}:{f}":n for (u,f),n in hits.items()}

patches={"SYSTEM05_OWNER":patch(S05),"QA_PAGE":patch(QA)}

def landscape(d):
    s=d.add_section(WD_SECTION.NEW_PAGE);s.orientation=WD_ORIENT.LANDSCAPE;s.page_width,s.page_height=s.page_height,s.page_width
    s.top_margin=Inches(.33);s.bottom_margin=Inches(.33);s.left_margin=Inches(.28);s.right_margin=Inches(.28)
def table(d,h,rows,fs=3.45):
    t=d.add_table(rows=1,cols=len(h));t.style="Table Grid"
    tr=t.rows[0]._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");tr.append(e)
    for i,x in enumerate(h):
        t.rows[0].cells[i].text=str(x);sh=OxmlElement("w:shd");sh.set(qn("w:fill"),"EDE9FE");t.rows[0].cells[i]._tc.get_or_add_tcPr().append(sh)
    for rr in rows:
        c=t.add_row().cells
        for i,x in enumerate(rr):c[i].text=str(x)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
for path,title,prefix in [(S05,"Batch 54 · QA Canonical Contract Registry","SYSTEM05"),(QA,"Batch 54 · QA-01 Contract Binding Ledger","PAGE::QA-01")]:
    d=Document(path);landscape(d);d.add_heading(title,1)
    d.add_paragraph(f"[{MARK}::{prefix}] QA contracts preserve exact source/version/criteria/evidence boundaries. Machine score cannot bypass manual review. ReleaseCandidate is not published truth. CLIENT/QA evidence-point/range mutations are LOCAL LOCAL session-draft operations and do not fabricate server APIs or canonical Findings.")
    rows=[[u,s["operation"],s["route"],s["payload"],s["persistence"],s["basis"]] for u,s in TARGETS.items()]
    table(d,["Control UID","Operation","Method / Path","Payload / Schema","Persistence Owner","Semantics"],rows)
    d.save(path);Document(path)

post_page=compose(parse_controls(QA));post_owner=compose(parse_controls(S05))
for uid,s in TARGETS.items():
    for m,label in [(post_page,"PAGE"),(post_owner,"OWNER")]:
        assert m[uid]["payload_schema"]==s["payload"],(label,uid,"PAYLOAD",m[uid]["payload_schema"])
        assert m[uid]["method_path"]==s["route"],(label,uid,"METHOD",m[uid]["method_path"])
        assert m[uid]["persistence_owner"]==s["persistence"],(label,uid,"PERSIST",m[uid]["persistence_owner"])

allm={p:compose(parse_controls(fn)) for p,fn in PAGES.items()};rem=[]
for p,m in allm.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if missing(op) or st not in {"EFFECTFUL_EXACT","READ_EXACT"}:continue
        if missing(r.get("method_path","")):rem.append(("method_path",p,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("payload_schema","")):rem.append(("payload_schema",p,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("persistence_owner","")):rem.append(("persistence_owner",p,uid))
rc=collections.Counter(x[0] for x in rem);assert len(rem)==86,(len(rem),rc);assert rc=={"method_path":25,"payload_schema":33,"persistence_owner":28},rc

d=Document(LOGIC);landscape(d);d.add_heading("Batch 54 · QA Contract Closure",1)
machine={"marker":MARK,"base_head":BASE_HEAD,"pre_contract_cells":115,"qa_cells_closed":29,"method_path_closed":9,"payload_schema_closed":11,"persistence_owner_closed":9,
"post_contract_cells":86,"remaining_method_path":25,"remaining_payload_schema":33,"remaining_persistence_owner":28,
"client_local_routes":3,"server_apis_fabricated_for_client_local":0,"manual_review_bypass_introduced":0,"release_candidate_as_publish_truth":0,"runtime_execution_claimed":False}
table(d,["Metric","Value"],[["Pre denominator","115"],["QA cells closed","29 = Method 9 + Payload 11 + Persistence 9"],["Post denominator","86 = Method 25 + Payload 33 + Persistence 28"],["CLIENT/QA LOCAL LOCAL commands","3"],["Fabricated server APIs for local evidence capture","0"],["Manual review bypass introduced","0"],["Runtime execution claimed","False"]],4.0)
d.add_paragraph("BATCH54_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")));d.save(LOGIC);Document(LOGIC)
changed=[S05,QA,LOGIC]
Path("__batch54_remediation_report.json").write_text(json.dumps({"machine":machine,"targets":TARGETS,"patches":patches,"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed}},ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH54="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
