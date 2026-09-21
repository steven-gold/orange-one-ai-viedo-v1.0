from pathlib import Path
from docx import Document
import json,re,collections

SYSTEMS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def display_only_type(t):
    u=(t or "").upper()
    return u in {"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"} or any(x in u for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]; ci=hfind(h,"control uid")
        if ci is None:continue
        idx={
          "control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),
          "method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),
          "persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status"),
        }
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"source":path,"table":ti+1,"row":ri,"headers":h,"values":vals})
            out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out

def is_wb_na(r):
    return missing(r.get("action","")) and r.get("type")=="SECTION_OPEN" and r.get("gate")=="PAGE_READ" and r.get("permission")=="workspace.dashboard.view" and r.get("payload_schema")=="N/A_READ_PROJECTION" and r.get("operation")=="getDashboardReadModel" and r.get("method_path")=="GET /v1/dashboard/read-model" and r.get("runtime_owner")=="DASHBOARD_READ_MODEL" and r.get("runtime_status")=="READ_EXACT"
def is_aiapi_na(r):
    return missing(r.get("action","")) and r.get("type")=="BUTTON_OR_ROW_ACTION" and r.get("payload_schema")=="AIAPI Page Operation-specific Form / Provider Profile Field Contract" and not missing(r.get("permission","")) and not missing(r.get("operation","")) and not missing(r.get("method_path","")) and not missing(r.get("runtime_owner","")) and r.get("runtime_status") in {"EFFECTFUL_EXACT","READ_EXACT"}
def is_edit_na(r):
    return missing(r.get("operation","")) and r.get("runtime_status")=="LOCAL_WORKING_DRAFT_EXACT" and r.get("method_path")=="NO_PUBLIC_API_BY_AUTHORITY" and not missing(r.get("action","")) and not missing(r.get("gate","")) and not missing(r.get("permission","")) and not missing(r.get("runtime_owner",""))
def is_iam_op_na(r):
    return missing(r.get("operation","")) and r.get("runtime_status")=="ORCHESTRATES_EXISTING_EXACT_OPERATIONS" and not missing(r.get("action","")) and not missing(r.get("gate","")) and not missing(r.get("permission",""))
def is_iam_owner_na(r):
    return missing(r.get("runtime_owner","")) and r.get("runtime_status")=="ORCHESTRATES_EXISTING_EXACT_OPERATIONS" and not missing(r.get("action","")) and not missing(r.get("gate","")) and not missing(r.get("permission",""))
def classify(page,field,r):
    v=r.get(field,"");st=r.get("runtime_status","")
    if not missing(v):return "BOUND"
    if "UI_LOCAL_EXACT" in st:
        if field in {"operation","runtime_owner"}:return "LEGITIMATE_NA_UI_LOCAL"
        if field=="action" and display_only_type(r.get("type","")):return "LEGITIMATE_NA_UI_LOCAL"
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in st and field=="operation":return "CLOSED_BY_OWNER_LOCAL_COMMAND"
    if "READ_EXACT" in st and display_only_type(r.get("type","")):
        if field=="action":return "LEGITIMATE_NA_READ_PRESENTATION"
        if field=="operation" and not missing(r.get("runtime_owner","")):return "READ_OWNER_BOUND_OPERATION_UNSPECIFIED"
    if field=="runtime_owner" and any(x in st for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return "KNOWN_RUNTIME_BLOCKER"
    if v=="SOURCE_NOT_DEFINED":return "SOURCE_RESOLUTION_REQUIRED"
    if page=="WB-01" and field=="action" and is_wb_na(r):return "LEGITIMATE_NA_READ_PROJECTION_ACTION"
    if page=="AIAPI-01" and field=="action" and is_aiapi_na(r):return "LEGITIMATE_NA_DIRECT_OPERATION_ACTION"
    if page=="EDIT-01" and field=="operation" and is_edit_na(r):return "LEGITIMATE_NA_LOCAL_WORKING_DRAFT_OPERATION"
    if page=="IAM-01" and field=="operation" and is_iam_op_na(r):return "LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_OPERATION"
    if page=="IAM-01" and field=="runtime_owner" and is_iam_owner_na(r):return "LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_RUNTIME_OWNER"
    return "DEFINITION_BINDING_GAP"

# Canonical owner rows and reference-only precedence
owner_rows=collections.defaultdict(list);reference_only=collections.defaultdict(lambda:collections.defaultdict(list));generic=collections.defaultdict(list)
for sf in SYSTEMS:
    d=Document(sf)
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        rows=[]
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if any(vals):
                rec={"source":sf,"table":ti,"row":ri,"headers":headers,"values":vals,"joined":" | ".join(vals)}
                rows.append(rec);generic[sf].append(rec)
        ci=hfind(headers,"control uid")
        if ci is not None:
            for rec in rows:
                uid=rec["values"][ci] if ci<len(rec["values"]) else ""
                if uid and uid not in {"—","-"}:owner_rows[uid].append(rec)
        rui=next((i for i,h in enumerate(headers) if h.lower()=="reference uid"),None)
        coi=next((i for i,h in enumerate(headers) if h.lower()=="canonical owner"),None)
        mfi=next((i for i,h in enumerate(headers) if h.lower()=="missing field"),None)
        if rui is not None and coi is not None:
            for rec in rows:
                vals=rec["values"];uid=vals[rui] if rui<len(vals) else "";owner=vals[coi] if coi<len(vals) else "";field=vals[mfi] if mfi is not None and mfi<len(vals) else ""
                if uid and owner:reference_only[uid][field].append({"secondary":sf,"canonical_name":owner})

current=[]
for page,fn in PAGES.items():
    for uid,r in compose(parse_controls(fn)).items():
        for field in ["action","gate","permission","operation","runtime_owner"]:
            if classify(page,field,r)=="DEFINITION_BINDING_GAP":
                current.append({"page":page,"file":fn,"uid":uid,"field":field,**{k:r.get(k,"") for k in FIELDS[1:]}})
assert len(current)==164,len(current)
ordinary=[x for x in current if not (x["page"]=="SYS-01" and x["uid"]=="SYS-01-BTN-NAV-OPEN" and x["field"]=="gate")]
assert len(ordinary)==163
focus=[x for x in ordinary if x["field"] in {"operation","runtime_owner"}]
assert len(focus)==144,len(focus)
ui=[x for x in focus if x["runtime_status"]=="UI_LOCAL_OR_READ_SOURCE"]
assert len(ui)==42,len(ui)

def effective_owner(uid,field):
    physical=sorted(set(r["source"] for r in owner_rows.get(uid,[])))
    secondary=set()
    for rr in reference_only.get(uid,{}).values():
        for item in rr:secondary.add(item["secondary"])
    eff=[s for s in physical if s not in secondary]
    assert len(eff)==1,(uid,field,physical,eff)
    return eff[0]

# exact same-action/gate/permission candidate sets
TARGET_HEADERS={
"operation":["Operation","Operation ID","operationId","API Operation","Service Operation"],
"runtime_owner":["Runtime Owner","Runtime Service Owner","Service Owner"],
}
ANCHORS={
"action":["Action UID","Action ID"],
"gate":["Gate","Gate UID"],
"permission":["Permission","Permission / Auth Resource","Auth Resource","Authorization Resource"],
}
def cols(headers,names):
    hs=[norm(h).lower() for h in headers]
    return [i for i,h in enumerate(hs) if any(h==n.lower() for n in names)]
def target_candidates(owner,target,anchor_name,anchor_value):
    vals=collections.defaultdict(list)
    if missing(anchor_value):return vals
    for rec in generic[owner]:
        ac=cols(rec["headers"],ANCHORS[anchor_name]);tc=cols(rec["headers"],TARGET_HEADERS[target])
        if not ac or not tc:continue
        if not any(i<len(rec["values"]) and rec["values"][i]==anchor_value for i in ac):continue
        for j in tc:
            if j<len(rec["values"]):
                v=rec["values"][j]
                if not missing(v):vals[v].append({"table":rec["table"],"row":rec["row"],"anchor":anchor_name,"anchor_value":anchor_value,"target_header":rec["headers"][j]})
    return vals

result=[]
for q in ui:
    owner=effective_owner(q["uid"],q["field"]);q["canonical_owner"]=owner
    evidence={}
    for a in ["action","gate","permission"]:
        vals=target_candidates(owner,q["field"],a,q.get(a,""))
        evidence[a]={k:v for k,v in vals.items()}
    q["anchor_candidates"]=evidence
    # exact owner UID rows
    q["owner_uid_rows"]=[{"table":r["table"],"row":r["row"],"text":r["joined"]} for r in owner_rows.get(q["uid"],[]) if r["source"]==owner]
    result.append(q)

by_page_field=collections.Counter((x["page"],x["field"]) for x in result)
by_owner_field=collections.Counter((x["canonical_owner"],x["field"]) for x in result)
by_action=collections.Counter((x["page"],x["action"],x["field"]) for x in result)
payload={
"denominator":42,
"page_field":{" | ".join(k):v for k,v in by_page_field.items()},
"owner_field":{" | ".join(k):v for k,v in by_owner_field.items()},
"action_field":{" | ".join(k):v for k,v in by_action.items()},
"rows":result,
}
Path("__batch15_ui_local_read_source_audit.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH15_AUDIT="+json.dumps({
"denominator":42,
"page_field":{" | ".join(k):v for k,v in by_page_field.items()},
"owner_field":{" | ".join(k):v for k,v in by_owner_field.items()},
"action_field":{" | ".join(k):v for k,v in by_action.items()},
},ensure_ascii=False,sort_keys=True))
print("BATCH15_ROWS_BEGIN")
for x in result:
    print(json.dumps(x,ensure_ascii=False,sort_keys=True))
print("BATCH15_ROWS_END")
