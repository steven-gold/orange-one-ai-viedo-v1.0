from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
import hashlib,json

LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED="29a48e5b67bed457ec4261c0cedfb3477a57dde0"
MARK="ACPOS-20260921-BATCH-12-EXACT-RELATION-REMEDIATION-V1"

def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

def snapshot(doc):
    return "\n".join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])

assert blob(LOGIC)==EXPECTED,(blob(LOGIC),EXPECTED)
d=Document(LOGIC)
before=snapshot(d)

found=False
for p in d.paragraphs:
    if p.text.startswith("BATCH12_MACHINE_JSON="):
        found=True
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=0.85
        for r in p.runs:
            r.font.size=Pt(3.2)
        break
assert found,"BATCH12_MACHINE_JSON_NOT_FOUND"

sec=d.sections[-1]
sec.top_margin=Inches(.22)
sec.bottom_margin=Inches(.20)
sec.left_margin=Inches(.24)
sec.right_margin=Inches(.24)

# Tighten only Batch-12 headings/paragraph spacing in the final section.
for p in d.paragraphs:
    if p.text.startswith("Batch 12 · Exact-Relation Canonical Spec Remediation"):
        p.paragraph_format.space_after=Pt(2)
    elif p.text.startswith("["+MARK+"]"):
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(2)
        p.paragraph_format.line_spacing=0.95

d.save(LOGIC)
d2=Document(LOGIC)
assert snapshot(d2)==before,"TEXT_CHANGED_IN_LAYOUT_ONLY_FIX"
report={"layout_only":True,"text_unchanged":True,"marker":MARK,"logic_blob":blob(LOGIC)}
Path("__batch12_layout_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH12_LAYOUT="+json.dumps(report,ensure_ascii=False,sort_keys=True))
