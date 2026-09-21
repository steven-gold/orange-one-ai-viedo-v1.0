from pathlib import Path
from docx import Document
import json,re,collections

SYSTEMS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}

def norm(x): return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def get(vals,i): return "" if i is None or i>=len(vals) else norm(vals[i])
def missing(v): return v in {"","—","-","SOURCE_NOT_DEFINED"}

DISPLAY_ONLY_TYPES={"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"}
def display_only_type(t):
    typ=(t or "").upper()
    return typ in DISPLAY_ONLY_TYPES or any(x in typ for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]

def parse_page_rows(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(h,"control uid")
        if ci is None:continue
        idx={
          "control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),"operation":hfind(h,"operation"),
          "runtime_owner":hfind(h,"runtime owner"),"runtime_status":hfind(h,"runtime status")
        }
        for ri,r in enumerate(t.rows[1:],2):
            vals=[c.text for c in r.cells]
            uid=get(vals,ci)
            if not uid or uid in {"—","-"} or not re.search(r"[A-Za-z0-9]",uid):continue
            rec={k:get(vals,idx[k]) for k in FIELDS}
            rec.update({"source":str(path),"table":ti+1,"row":ri,"headers":h,"indices":idx})
            out.append(rec)
    return out

def compose(rows):
    groups=collections.defaultdict(list)
    for r in rows:groups[r["control"]].append(r)
    out={}
    for uid,rs in groups.items():
        ranked=sorted(rs,key=lambda r:sum(1 for k in FIELDS[1:] if not missing(r.get(k,""))),reverse=True)
        base=dict(ranked[0])
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        base["same_uid_rows"]=rs
        out[uid]=base
    return out

def classify_gap(field,row):
    val=row.get(field,"");status=row.get("runtime_status","")
    if not missing(val):return "BOUND"
    if "UI_LOCAL_EXACT" in status:
        if field in {"operation","runtime_owner"}:return "LEGITIMATE_NA_UI_LOCAL"
        if field=="action" and display_only_type(row.get("type","")):return "LEGITIMATE_NA_UI_LOCAL"
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in status and field=="operation":return "CLOSED_BY_OWNER_LOCAL_COMMAND"
    if "READ_EXACT" in status and display_only_type(row.get("type","")):
        if field=="action":return "LEGITIMATE_NA_READ_PRESENTATION"
        if field=="operation" and not missing(row.get("runtime_owner","")):return "READ_OWNER_BOUND_OPERATION_UNSPECIFIED"
    if field=="runtime_owner" and any(x in status for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return "KNOWN_RUNTIME_BLOCKER"
    if val=="SOURCE_NOT_DEFINED":return "SOURCE_RESOLUTION_REQUIRED"
    return "DEFINITION_BINDING_GAP"

# Generic system rows with exact control UID ownership.
sys_rows=[]
for sf in SYSTEMS:
    d=Document(sf)
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(headers,"control uid")
        if ci is None:continue
        idx={
          "control":ci,"type":hfind(headers,"type"),"label":hfind(headers,"label","顯示名稱"),
          "action":hfind(headers,"action uid"),"gate":hfind(headers,"gate uid","gate"),
          "permission":hfind(headers,"permission","auth resource"),"operation":hfind(headers,"operation"),
          "runtime_owner":hfind(headers,"runtime owner"),"runtime_status":hfind(headers,"runtime status")
        }
        for ri,row in enumerate(t.rows[1:],2):
            vals=[c.text for c in row.cells]
            uid=get(vals,ci)
            if not uid or uid in {"—","-"}:continue
            rec={k:get(vals,idx[k]) for k in FIELDS}
            rec.update({"source":sf,"table":ti+1,"row":ri,"headers":headers})
            sys_rows.append(rec)

owner_index=collections.defaultdict(list)
for r in sys_rows: owner_index[r["control"]].append(r)

queue=[]
for page,fn in PAGES.items():
    rows=parse_page_rows(fn);bm=compose(rows)
    for uid,t in sorted(bm.items()):
        for field in ["action","gate","permission","operation","runtime_owner"]:
            if classify_gap(field,t)!="DEFINITION_BINDING_GAP":continue
            if page=="SYS-01" and uid=="SYS-01-BTN-NAV-OPEN" and field=="gate":continue
            owners=owner_index.get(uid,[])
            if owners:
                owner_state="CANONICAL_OWNER_ROW_FIELD_UNDEFINED"
            else:
                owner_state="CANONICAL_OWNER_ROW_ABSENT"
            queue.append({
              "page":page,"file":fn,"uid":uid,"field":field,
              "type":t.get("type",""),"label":t.get("label",""),
              "action":t.get("action",""),"gate":t.get("gate",""),
              "permission":t.get("permission",""),"operation":t.get("operation",""),
              "runtime_owner":t.get("runtime_owner",""),"runtime_status":t.get("runtime_status",""),
              "owner_state":owner_state,
              "owners":owners,
              "field_column_present":any(r["indices"].get(field) is not None for r in t["same_uid_rows"])
            })
assert len(queue)==310,len(queue)
assert sum(1 for q in queue if q["owner_state"]=="CANONICAL_OWNER_ROW_FIELD_UNDEFINED")==282
assert sum(1 for q in queue if q["owner_state"]=="CANONICAL_OWNER_ROW_ABSENT")==28

# Evidence-grounded classification for the 282 owner-row gaps.
def owner_signals(q):
    owners=q["owners"]
    statuses=sorted(set(r.get("runtime_status","") for r in owners if not missing(r.get("runtime_status",""))))
    types=sorted(set(r.get("type","") for r in owners if not missing(r.get("type",""))))
    actions=sorted(set(r.get("action","") for r in owners if not missing(r.get("action",""))))
    gates=sorted(set(r.get("gate","") for r in owners if not missing(r.get("gate",""))))
    perms=sorted(set(r.get("permission","") for r in owners if not missing(r.get("permission",""))))
    ops=sorted(set(r.get("operation","") for r in owners if not missing(r.get("operation",""))))
    ros=sorted(set(r.get("runtime_owner","") for r in owners if not missing(r.get("runtime_owner",""))))
    return {"statuses":statuses,"types":types,"actions":actions,"gates":gates,"permissions":perms,"operations":ops,"runtime_owners":ros}

def classify_owner_gap(q,s):
    field=q["field"];status=" | ".join(s["statuses"]+[q.get("runtime_status","")]).upper()
    typ=" | ".join(s["types"]+[q.get("type","")]).upper()

    # Explicit runtime-not-ready evidence takes precedence only for Runtime Owner.
    if field=="runtime_owner" and any(k in status for k in [
        "SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED",
        "NOT_EXECUTED","IMPLEMENTATION_REQUIRED","RUNTIME_BLOCKED"
    ]):
        return "FORMAL_RUNTIME_BLOCKED"

    # Explicit UI-local/read-presentation evidence can classify structural N/A.
    if any(k in status for k in ["UI_LOCAL_EXACT","UI_LOCAL_ONLY"]):
        if field in {"operation","runtime_owner"}:
            return "FORMAL_NOT_APPLICABLE"
        if field=="action" and display_only_type(typ):
            return "FORMAL_NOT_APPLICABLE"
    if "READ_EXACT" in status and display_only_type(typ) and field=="action":
        return "FORMAL_NOT_APPLICABLE"

    # No explicit N/A or blocked marker: owner contract itself is incomplete.
    return "CANONICAL_SPEC_REMEDIATION_REQUIRED"

owner_rows=[]
owner_counts=collections.Counter();by_field=collections.defaultdict(collections.Counter);by_page=collections.defaultdict(collections.Counter)
for q in [x for x in queue if x["owner_state"]=="CANONICAL_OWNER_ROW_FIELD_UNDEFINED"]:
    s=owner_signals(q);cls=classify_owner_gap(q,s)
    owner_sources=sorted(set(r["source"] for r in q["owners"]))
    rec={k:v for k,v in q.items() if k!="owners"}|{"signals":s,"owner_sources":owner_sources,"remediation_class":cls}
    owner_rows.append(rec);owner_counts[cls]+=1;by_field[q["field"]][cls]+=1;by_page[q["page"]][cls]+=1

# WB 28 mapping audit. Exact owner is absent; identify only exact existing anchors from the WB control itself.
wb_rows=[]
wb_classes=collections.Counter()
for q in [x for x in queue if x["owner_state"]=="CANONICAL_OWNER_ROW_ABSENT"]:
    assert q["page"]=="WB-01",q
    anchors=[]
    for key in ["action","gate","permission","operation","runtime_owner"]:
        v=q.get(key,"")
        if not missing(v):anchors.append({"field":key,"value":v})
    # Synthetic/compact UIDs are visibly abbreviated; do not expand them.
    uid_form="SYNTHETIC_OR_COMPACT_UID" if ("…" in q["uid"] or "..." in q["uid"]) else "UNREGISTERED_EXACT_UID"
    cls="WB_CANONICAL_OWNER_MAPPING_REQUIRED"
    if anchors:
        cls="WB_CANONICAL_OWNER_MAPPING_REQUIRED_WITH_EXISTING_ANCHOR"
    rec={k:v for k,v in q.items() if k!="owners"}|{"uid_form":uid_form,"existing_exact_anchors":anchors,"mapping_class":cls}
    wb_rows.append(rec);wb_classes[cls]+=1

owner_source_counts=collections.Counter()
owner_source_cardinality=collections.Counter()
for r in owner_rows:
    owner_source_cardinality[len(r["owner_sources"])]+=1
    for s in r["owner_sources"]:owner_source_counts[s]+=1

payload={
 "denominator":310,
 "owner_source_counts":dict(owner_source_counts),
 "owner_source_cardinality":{str(k):v for k,v in owner_source_cardinality.items()},
 "owner_field_undefined_count":282,
 "owner_row_absent_count":28,
 "owner_classification":dict(owner_counts),
 "owner_source_counts":dict(owner_source_counts),
 "owner_source_cardinality":{str(k):v for k,v in owner_source_cardinality.items()},
 "owner_by_field":{k:dict(v) for k,v in by_field.items()},
 "owner_by_page":{k:dict(v) for k,v in by_page.items()},
 "wb_mapping_classification":dict(wb_classes),
 "owner_rows":owner_rows,
 "wb_rows":wb_rows
}
Path("__batch08_authority_remediation_audit.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH08_AUDIT="+json.dumps({
 "denominator":310,
 "owner_field_undefined_count":282,
 "owner_row_absent_count":28,
 "owner_classification":dict(owner_counts),
 "owner_by_field":{k:dict(v) for k,v in by_field.items()},
 "owner_by_page":{k:dict(v) for k,v in by_page.items()},
 "wb_mapping_classification":dict(wb_classes)
},ensure_ascii=False,sort_keys=True))
