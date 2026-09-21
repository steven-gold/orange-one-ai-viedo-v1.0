from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib

MARK="ACPOS-20260921-BATCH-09-GATE-SCOPED-RUNTIME-OWNER-V1"
OWNER05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
OWNER06="06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
TARGET_PAGES={
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
}
EXPECTED_SHA={
OWNER05:"8e3767e2194e11bbc1971172bc2e6404ffd7797b",
OWNER06:"0a6cfbfdc3b28b95a29d8b0608489342b4470314",
TARGET_PAGES["ASSET-01"]:"71c6c4b8b734a51db5f83b73fc76d455973b8775",
TARGET_PAGES["DEV-01"]:"a98a7eedabcf27af1ecb7015591b2a7c85ecb8de",
TARGET_PAGES["EDIT-01"]:"f478bbe5e80b4b26b73e131536db740724364003",
TARGET_PAGES["ERP-01"]:"87e3086adc0f705a8e13ac9c73dabb2b30cafa33",
LOGIC:"5a61addbbe9975b3ef9cc7ef9c6f35d756dc3290",
}
ALL_PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}

def blob(path):
    b=Path(path).read_bytes();return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()
for f,s in EXPECTED_SHA.items():assert blob(f)==s,(f,blob(f),s)

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def header_exact(headers,name):
    n=norm(name).lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for n in needles:
        n=n.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None

DISPLAY={"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"}
def display_only(t):
    u=(t or "").upper();return u in DISPLAY or any(x in u for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])
FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]

def parse_page(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),"action":hfind(h,"action uid"),"gate":hfind(h,"gate"),"permission":hfind(h,"permission","auth resource"),"operation":hfind(h,"operation"),"runtime_owner":hfind(h,"runtime owner"),"runtime_status":hfind(h,"runtime status")}
        for ri,r in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in r.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"table":ti+1,"row":ri,"indices":idx});out.append(rec)
    return out

def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        base["same_uid_rows"]=rs;out[uid]=base
    return out

def gap(field,r):
    v=r.get(field,"");st=r.get("runtime_status","")
    if not missing(v):return False
    if "UI_LOCAL_EXACT" in st and (field in {"operation","runtime_owner"} or (field=="action" and display_only(r.get("type","")))):return False
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in st and field=="operation":return False
    if "READ_EXACT" in st and display_only(r.get("type","")):
        if field=="action":return False
        if field=="operation" and not missing(r.get("runtime_owner","")):return False
    if field=="runtime_owner" and any(x in st for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return False
    if v=="SOURCE_NOT_DEFINED":return False
    return True

# Gate -> unanimous Runtime Owner map, strictly inside owner contract tables with exact headers.
def gate_runtime_map(owner_file):
    d=Document(owner_file);m=collections.defaultdict(list)
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        gi=header_exact(h,"Gate");ri=header_exact(h,"Runtime Owner")
        if gi is None or ri is None:continue
        for rowno,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if gi>=len(vals) or ri>=len(vals):continue
            g=vals[gi];ro=vals[ri]
            if missing(g) or missing(ro):continue
            m[g].append({"value":ro,"table":ti+1,"row":rowno})
    out={}
    for g,rows in m.items():
        vals=sorted(set(x["value"] for x in rows))
        if len(vals)==1:out[g]={"value":vals[0],"evidence":rows}
        else:out[g]={"conflict":vals,"evidence":rows}
    return out

maps={OWNER05:gate_runtime_map(OWNER05),OWNER06:gate_runtime_map(OWNER06)}

# Reconstruct exactly the 22 page Runtime Owner gaps allowed by the Batch-09 policy.
OWNER_BY_PAGE={"ASSET-01":OWNER05,"EDIT-01":OWNER05,"DEV-01":OWNER06,"ERP-01":OWNER06}
accepted=[]
for page,fn in TARGET_PAGES.items():
    bm=compose(parse_page(fn));owner=OWNER_BY_PAGE[page]
    for uid,r in sorted(bm.items()):
        if not gap("runtime_owner",r):continue
        g=r.get("gate","")
        if missing(g):continue
        gm=maps[owner].get(g)
        if not gm or "value" not in gm:continue
        accepted.append({"page":page,"file":fn,"uid":uid,"gate":g,"runtime_owner":gm["value"],"owner_doc":owner,"evidence":gm["evidence"]})

expected_counts=collections.Counter(x["page"] for x in accepted)
assert len(accepted)==22,(len(accepted),accepted)
assert expected_counts==collections.Counter({"EDIT-01":13,"ERP-01":6,"DEV-01":2,"ASSET-01":1}),expected_counts

# Guard the 6 rejected semantic candidates from being smuggled into this batch.
assert all(x["runtime_owner"]!="Production secret/output forbidden" for x in accepted)

# Patch exact page Runtime Owner cells.
page_changed=collections.Counter()
for page,fn in TARGET_PAGES.items():
    targets={x["uid"]:x for x in accepted if x["page"]==page}
    if not targets:continue
    d=Document(fn)
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(h,"control uid");ri=hfind(h,"runtime owner")
        if ci is None or ri is None:continue
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals):continue
            uid=vals[ci]
            if uid not in targets:continue
            old=norm(row.cells[ri].text)
            val=targets[uid]["runtime_owner"]
            if missing(old):
                row.cells[ri].text=val;page_changed[page]+=1
            elif old!=val:
                raise AssertionError(("PAGE_RUNTIME_OWNER_CONFLICT",page,uid,old,val))
    # each target must now appear in at least one exact row with value
    d.save(fn);Document(fn)
    rows=parse_page(fn);comp=compose(rows)
    for uid,x in targets.items():assert comp[uid]["runtime_owner"]==x["runtime_owner"],(page,uid,comp[uid].get("runtime_owner"),x["runtime_owner"])

# Patch owner exact Control UID rows; if no exact Runtime Owner column, append exact UID addendum.
owner_changed=collections.Counter();owner_addendum=collections.defaultdict(list)
for owner in [OWNER05,OWNER06]:
    targets={x["uid"]:x for x in accepted if x["owner_doc"]==owner}
    d=Document(owner)
    touched=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=header_exact(h,"Control UID");ri=header_exact(h,"Runtime Owner")
        if ci is None or ri is None:continue
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals):continue
            uid=vals[ci]
            if uid not in targets:continue
            old=norm(row.cells[ri].text);val=targets[uid]["runtime_owner"]
            if missing(old):
                row.cells[ri].text=val;touched[uid]+=1;owner_changed[owner]+=1
            elif old!=val:
                raise AssertionError(("OWNER_RUNTIME_CONFLICT",owner,uid,old,val))
    for uid,x in targets.items():
        if touched[uid]==0:owner_addendum[owner].append(x)
    if owner_addendum[owner]:
        sec=d.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width,sec.page_height=sec.page_height,sec.page_width
        sec.top_margin=Inches(.4);sec.bottom_margin=Inches(.4);sec.left_margin=Inches(.4);sec.right_margin=Inches(.4)
        d.add_heading("Batch 09 · Exact Control Runtime Owner Closure Addendum",level=1)
        p=d.add_paragraph();p.add_run("["+MARK+"] ").bold=True
        p.add_run("The rows below close Runtime Owner only. Each value is inherited from a unanimous exact Gate -> Runtime Owner contract inside this same Canonical Owner document. No Action, Gate, Permission, Operation, API route, or new owner identity is created.")
        t=d.add_table(rows=1,cols=4);t.style="Table Grid"
        for i,h in enumerate(["Control UID","Gate","Runtime Owner","Authority Rule"]):t.rows[0].cells[i].text=h
        for x in owner_addendum[owner]:
            r=t.add_row().cells;r[0].text=x["uid"];r[1].text=x["gate"];r[2].text=x["runtime_owner"];r[3].text="UNANIMOUS_SAME_OWNER_GATE_RUNTIME_BINDING"
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:run.font.size=Pt(7)
    # append closure ledger even if all direct
    sec=d.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.4);sec.bottom_margin=Inches(.4);sec.left_margin=Inches(.4);sec.right_margin=Inches(.4)
    d.add_heading("Batch 09 · Gate-scoped Runtime Owner Closure Ledger",level=1)
    p=d.add_paragraph();p.add_run("["+MARK+"::"+Path(owner).name+"] ").bold=True
    p.add_run("This ledger supersedes the listed Batch-08 Runtime Owner remediation rows for current-state verification.")
    t=d.add_table(rows=1,cols=5);t.style="Table Grid"
    for i,h in enumerate(["Page","Control UID","Gate","Runtime Owner","Status"]):t.rows[0].cells[i].text=h
    for x in sorted(targets.values(),key=lambda z:(z["page"],z["uid"])):
        r=t.add_row().cells;r[0].text=x["page"];r[1].text=x["uid"];r[2].text=x["gate"];r[3].text=x["runtime_owner"];r[4].text="CLOSED_BY_BATCH_09"
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for run in p.runs:run.font.size=Pt(6)
    d.save(owner);Document(owner)

# Fresh 18-page Definition Binding Gap count using existing classification rules.
def total_definition_gaps():
    total=0;by=collections.Counter()
    for page,fn in ALL_PAGES.items():
        bm=compose(parse_page(fn))
        for uid,r in bm.items():
            for f in ["action","gate","permission","operation","runtime_owner"]:
                if gap(f,r):
                    total+=1;by[page]+=1
    return total,by
post_total,post_by=total_definition_gaps()
assert post_total==289,(post_total,post_by)

# Central evidence.
logic=Document(LOGIC)
txt="\n".join([p.text for p in logic.paragraphs]+[c.text for t in logic.tables for row in t.rows for c in row.cells])
assert MARK not in txt,"BATCH09_ALREADY_PRESENT"
sec=logic.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width,sec.page_height=sec.page_height,sec.page_width
sec.top_margin=Inches(.35);sec.bottom_margin=Inches(.35);sec.left_margin=Inches(.3);sec.right_margin=Inches(.3)
logic.add_heading("Batch 09 · Gate-scoped Runtime Owner Closure",level=1)
p=logic.add_paragraph();p.add_run("["+MARK+"] ").bold=True
p.add_run("Batch 09 closes only Runtime Owner fields when the control has a unique Canonical Owner, the page already binds an exact Gate, and every non-empty Runtime Owner contract row for that exact Gate inside the same owner document agrees on one value. Gate->Operation propagation is explicitly forbidden because a Gate may protect multiple operations.")

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def table(headers,rows,fs=4.6):
    t=logic.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row):c[i].text=str(v)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:r.font.size=Pt(fs)
    return t

logic.add_heading("Batch Result",level=2)
table(["Item","Count / State"],[
["Pre-batch Definition Binding Gap",311],
["Gate-scoped Runtime Owner closed",22],
["Post-batch Definition Binding Gap",289],
["Remaining unresolved excluding preserved SYS Gate conflict",288],
["Preserved SYS-01-BTN-NAV-OPEN Gate conflict",1],
["Rejected semantic Operation candidates",5],
["Rejected semantic Gate policy-text candidate",1],
["Batch-09 Action/Gate/Permission/Operation mutation",0],
],5.1)

logic.add_heading("22 Closed Bindings",level=2)
table(["Page","Control UID","Gate","Runtime Owner","Canonical Owner"],[
[x["page"],x["uid"],x["gate"],x["runtime_owner"],Path(x["owner_doc"]).name] for x in accepted
],4.0)

logic.add_heading("Inference Boundary",level=2)
table(["Rule","Decision"],[
["Gate -> Runtime Owner","Allowed only under unique owner + exact Gate + unanimous same-owner Runtime Owner."],
["Gate -> Operation","FORBIDDEN. One Gate can protect multiple operations; the 5 apparent unique candidates remain unresolved."],
["Operation -> Gate policy text","Not used when target Gate schema expects canonical binding semantics; SYS sandbox policy text remains unresolved."],
["Cross-system transfer","FORBIDDEN."],
],5.0)

page_hashes={p:blob(f) for p,f in TARGET_PAGES.items()}
owner_hashes={OWNER05:blob(OWNER05),OWNER06:blob(OWNER06)}
machine={"marker":MARK,"pre_definition_gaps":311,"closed_runtime_owner":22,"post_definition_gaps":289,"remaining_unresolved":288,"preserved_conflict":1,"page_counts":dict(expected_counts),"page_hashes":page_hashes,"owner_hashes":owner_hashes}
logic.add_paragraph("BATCH09_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

Path("__batch09_remediation_report.json").write_text(json.dumps({"machine":machine,"accepted":accepted,"post_by_page":dict(post_by),"owner_direct_cells":dict(owner_changed),"owner_addenda":{k:len(v) for k,v in owner_addendum.items()}},ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH09_REMEDIATION="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
