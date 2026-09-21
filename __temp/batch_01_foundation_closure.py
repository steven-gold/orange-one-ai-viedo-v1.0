from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import hashlib, json, re, sys

MARK="ACPOS-20260921-BATCH-01-FOUNDATION-CLOSURE-V1"
BRANCH="0921acpos"
FILES={
"index":("00_INDEX_ACPOS_Mother_Compliant_Basic_Design_Master.docx","230c4313fc07038a043a17676dcfa4ac77613e47"),
"shell":("ACPOS_GLOBAL_HOME_SHELL_NAVIGATION_Mother_Basic_Design_OPTIMIZED.docx","9ee35b10c2cdbd3cfa3b537678a809b08aa08498"),
"logic":("ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx","97f86daef306722dbc9be9e63a154297ce31f469"),
"core":("ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx","8bc3752adfec1a443c1aee4935c1aee9067e5f50"),
"sys":("ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx","042dc9e71220dfe38d428ff513cd4ba6aa4bfca0"),
"dev":("ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx","5853224aa34f2217440349fbd009972a94e7a3c1"),
}

def git_blob_sha(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for key,(fn,expected) in FILES.items():
    actual=git_blob_sha(fn)
    if actual!=expected:
        raise SystemExit(f"INPUT_CHANGED:{fn}:{actual}!={expected}")

def all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p

def text(doc):
    return "\n".join(p.text for p in all_paragraphs(doc))

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr()
    e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)

def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr()
    e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)

def add_table(doc, headers, rows, fs=6.2):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:r.font.size=Pt(fs)
    return t

def landscape(doc):
    s=doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation=WD_ORIENT.LANDSCAPE
    s.page_width,s.page_height=s.page_height,s.page_width
    s.top_margin=Inches(.42);s.bottom_margin=Inches(.42);s.left_margin=Inches(.42);s.right_margin=Inches(.42)

def add_marker_heading(doc,title):
    landscape(doc)
    doc.add_heading(title,level=1)
    p=doc.add_paragraph()
    p.add_run(f"[{MARK}] ").bold=True
    p.add_run("本節為 0921acpos Current Word package 的正式缺口閉合內容。只定義 Basic Design / authority contract；若缺 Runtime / Production evidence，狀態仍必須維持 NOT_EXECUTED / IMPLEMENTATION_REQUIRED，不得因文件定義完成而宣告 Runtime 完成。")

def replace_para(p,new_text):
    if p.text==new_text:return
    for r in p.runs:r.text=""
    if p.runs:p.runs[0].text=new_text
    else:p.add_run(new_text)

def patch_index(path):
    d=Document(path)
    if MARK in text(d): return "already"
    changed=[]
    for p in all_paragraphs(d):
        s=p.text.strip()
        if ("20 Basic Design" in s and "21" in s and "Word" in s) or ("20" in s and "Basic Design" in s and "Index" in s and "21" in s):
            replace_para(p,"Current Physical Package Denominator：31 DOCX = 1 Index + 9 System Normative + 18 Page Design + 1 Login + 1 System Logic + 1 Global Shell。Denominator MUST be re-resolved from the current branch physical scan before every formal audit.")
            changed.append("denominator")
        elif "GitHub branch" in s and re.search(r"\bnew\b",s):
            replace_para(p,re.sub(r"\bnew\b",BRANCH,s))
            changed.append("branch")
        elif "165c6440" in s:
            replace_para(p,"Current HEAD：MUST be resolved fresh from GitHub branch 0921acpos at audit/execution time. Historical SHA references are evidence only and MUST NOT be treated as Current Truth.")
            changed.append("head")
    add_marker_heading(d,"Current Package Identity / Dynamic Denominator Closure")
    add_table(d,["Item","Current rule"],[
        ["Branch","0921acpos"],
        ["Physical package denominator","31 DOCX: 1 Index + 9 System Normative + 18 Page Design + Login + System Logic + Global Shell"],
        ["HEAD identity","Dynamic. Resolve fresh from GitHub before audit, batch execution, validation and completion declaration; do not hard-code a historical SHA as Current Truth."],
        ["File applicability","Every physical DOCX in the package is in denominator unless an explicit Current Authority exclusion exists."],
        ["Evidence","Current branch root scan + exact Git blob SHA for each target file + fresh batch validation run."],
        ["Stale metadata rule","Any embedded branch/SHA/count that conflicts with the fresh physical scan is historical evidence only, not Current package authority."],
    ],6.2)
    d.save(path)
    return changed

TYPO=[
("TY01","target_uid","Every governed text target has stable target UID; page-only representative samples are insufficient."),
("TY02","page_or_scope_uid","Bind target to exact page/shell/shared scope."),
("TY03","viewport_breakpoint","Resolve desktop/tablet/mobile or explicit NOT_APPLICABLE."),
("TY04","language_theme","Bind zh-TW/zh-CN/en and theme variant when applicable."),
("TY05","typography_authority_ref","Every target resolves to this Global Typography Authority or an approved page override."),
("TY06","font_family","Use approved UI font stack; CJK must have deterministic fallback."),
("TY07","font_size","Explicit token/value required for each target class."),
("TY08","font_weight","Explicit semantic weight; do not infer from screenshot."),
("TY09","line_height","Explicit line-height token/value."),
("TY10","letter_spacing","Explicit tracking value/token."),
("TY11","container_width","Text metrics must be evaluated against actual container width."),
("TY12","wrap_rule","Define wrap/no-wrap behavior and max lines."),
("TY13","truncation_rule","Define ellipsis/clamp/overflow behavior."),
("TY14","expected_line_count_rule","Define expected line-count envelope when layout-sensitive."),
("TY15","tolerance","Define measurable tolerance for computed metric/render comparison."),
]
TOKENS=[
("display-xl","32px","700","1.20","-0.02em","Page title / major dashboard heading"),
("heading-lg","24px","700","1.25","-0.01em","Primary section heading"),
("heading-md","20px","600","1.30","0","Card / panel heading"),
("body-md","14px","400","1.55","0","Primary body / form content"),
("body-sm","12px","400","1.50","0","Secondary text / helper / metadata"),
("label-md","13px","600","1.35","0","Form label / table header / compact control label"),
("button-md","14px","600","1.00","0","Primary / secondary button text"),
("mono-sm","12px","400","1.45","0","UID / hash / technical reference"),
]
def patch_shell(path):
    d=Document(path)
    if MARK in text(d): return "already"
    add_marker_heading(d,"Global Typography Authority / TY01–TY15")
    d.add_paragraph("Canonical font stack：Inter / system-ui / -apple-system / BlinkMacSystemFont / 'Segoe UI' / 'Noto Sans TC' / 'PingFang TC' / sans-serif。若實際產品環境無 Inter，必須落到可量測且可重現的 system/CJK fallback；不得以不同字型截圖替代 computed metrics。")
    add_table(d,["Token","Font size","Weight","Line height","Letter spacing","Applicability"],TOKENS,6.0)
    add_table(d,["ID","Field / Check","Closure rule"],TYPO,5.9)
    d.add_heading("Typography Denominator / Computed Metrics",level=2)
    add_table(d,["Rule","Requirement"],[
        ["Target denominator","Every visible governed text target in Global Shell + each page must be inventoried. Screenshot samples do not close the denominator."],
        ["Computed metrics","Validation must compare resolved font-family, size, weight, line-height, letter-spacing, width, wrapping and truncation at the declared viewport/language."],
        ["Tolerance","Default typography metric tolerance = ±1 CSS px for box/line-height geometry and ±1 rendered line unless page authority defines stricter tolerance; semantic overflow/clipping = zero tolerance."],
        ["Responsive behavior","If a target changes token/wrap at a breakpoint, each applicable breakpoint is a separate denominator row."],
        ["Override","Page-level override must name target_uid + authority_ref + reason; unbound local typography is BLOCK."],
    ],6.0)
    d.save(path)
    return "updated"

L=[
("L01","Previous Context","Resolve predecessor state / exact source refs."),
("L02","Trigger","Named control/event that begins the function."),
("L03","Input","Typed user/system inputs and exact refs."),
("L04","Validation","Pre-action validation and fail-closed checks."),
("L05","Action","User/system intent mapped to one governed action."),
("L06","Operation","Canonical operation / command name."),
("L07","Payload","Typed payload fields + constraints."),
("L08","Owner","Canonical owner of the operation."),
("L09","Runtime","Runtime/port/worker/provider binding or explicit IMPLEMENTATION_REQUIRED."),
("L10","Authorization","Permission / role / scope check."),
("L11","Mutation_or_Read","Exact read/mutation boundary."),
("L12","State","Pre/post state transition."),
("L13","Output","Exact output/ref/version produced."),
("L14","UI_Effect","Visible page effect / updated region."),
("L15","Post_Action_Validation","Fresh validation after operation."),
("L16","Success_Next_State","Named successful state."),
("L17","Next_Function_or_Page","Explicit downstream consumer / route."),
("L18","Failure","Canonical failure taxonomy / condition."),
("L19","Recovery","Recovery owner/action."),
("L20","Retry","Retry eligibility / idempotency relationship."),
("L21","Cancel","Cancellation semantics when applicable."),
("L22","Resume_Back","Persisted resume/back behavior."),
("L23","Downstream_Readiness","Consumer-required refs/schema/evidence are present."),
("L24","Audit","Audit event/evidence/correlation refs persisted."),
]
S=[
("S01","Entity Contract"),("S02","Operation Contract"),("S03","Input/Payload Schema"),("S04","Validation Contract"),
("S05","Authorization/Permission"),("S06","Mutation Contract"),("S07","Read Contract"),("S08","State Machine"),
("S09","Transaction Boundary"),("S10","Idempotency"),("S11","Concurrency/Race"),("S12","Consistency"),
("S13","Async/Job Semantics"),("S14","Queue/Worker"),("S15","Timeout"),("S16","Retry"),
("S17","Cancel"),("S18","Error Taxonomy"),("S19","Recovery"),("S20","Resume"),
("S21","Version/Checksum/Exact Ref"),("S22","Dependency"),("S23","Provider/External Boundary"),("S24","Permission/RLS Boundary"),
("S25","Audit/Correlation"),("S26","Observability"),("S27","Producer/Consumer Schema"),("S28","Evidence Contract"),
("S29","Terminal Closure"),("S30","Applicability Disposition"),
]
H=[
("H01","Predecessor output exact-ref ready"),("H02","Predecessor state allowed"),("H03","Version/checksum bound"),
("H04","Authority/profile/policy refs bound"),("H05","Required evidence present"),("H06","Permission context portable"),
("H07","Producer schema matches consumer input"),("H08","Consumer can reject stale/superseded ref"),("H09","Failure routes back to canonical owner"),
("H10","Retry/resume preserves lineage"),("H11","Audit correlation continues across boundary"),("H12","Successor readiness explicitly PASS/BLOCK"),
]
def patch_logic(path):
    d=Document(path)
    if MARK in text(d): return "already"
    add_marker_heading(d,"Exhaustive Functional / System Logic / Cross-stage Closure Contract")
    d.add_paragraph("對每個 REQUIRED Function / Operation / system-logic concern，必須逐項標示 REQUIRED 或 NOT_APPLICABLE；NOT_APPLICABLE 必須有 Authority reason。缺欄、UNKNOWN、以其他頁存在作替代、或只寫 UI 名稱，都不能取得 closure credit。")
    d.add_heading("L01–L24 Functional Forward/Backward Chain",level=2)
    add_table(d,["ID","Node","Required closure"],L,5.7)
    d.add_heading("S01–S30 System Logic Applicability + Closure",level=2)
    add_table(d,["ID","System logic concern","Disposition rule"],[(i,n,"REQUIRED→must bind exact contract/evidence; NOT_APPLICABLE→must state authority-backed reason.") for i,n in S],5.6)
    d.add_heading("H01–H12 Cross-stage Consumer Readiness",level=2)
    add_table(d,["ID","Cross-stage check","Required result"],[(i,n,"PASS with exact refs/evidence, or BLOCK with owner/recovery; no implicit readiness.") for i,n in H],5.7)
    d.add_heading("Exhaustive Denominator Rule",level=2)
    add_table(d,["Denominator","Rule"],[
        ["Functions","Every visible/effectful function + system-triggered function in the page contract."],
        ["Operations","Every operation reachable from those functions, including cancel/retry/recheck/submit/approve."],
        ["System logic","S01–S30 evaluated per applicable operation/domain; absence is not implicit N/A."],
        ["Cross-stage","Every predecessor→successor handoff in the package, not representative examples."],
        ["Completion","Only when all REQUIRED rows close and all N/A rows carry authority reason may Stage-02 functional closure be claimed."],
    ],6.0)
    d.save(path)
    return "updated"

def patch_core(path):
    d=Document(path)
    if MARK in text(d): return "already"
    hits=0
    for p in all_paragraphs(d):
        s=p.text
        if "unresolved LockReview contract stays BLOCK" in s:
            replace_para(p,s.replace("unresolved LockReview contract stays BLOCK","LockReview contract is resolved by the current Basic Design authority; BLOCK only when required LockReview inputs, validation, approval evidence, permission, or exact binding are missing/failed"))
            hits+=1
    add_marker_heading(d,"CORE-01 LockReview Contract Truth Reconciliation")
    add_table(d,["Item","Current contract"],[
        ["LockReview authority","Defined in CORE-01 Basic Design. The former generic statement that the LockReview contract itself is unresolved is superseded."],
        ["When BLOCK remains valid","Required inputs missing; authority binding missing; permission denied; validation failed; approval evidence absent/invalid; stale version/checksum; downstream readiness failed."],
        ["Runtime truth","Design contract closure does not assert Runtime implementation. If no fresh runtime evidence exists, state remains NOT_EXECUTED / IMPLEMENTATION_REQUIRED."],
        ["Downstream handoff","Only a successfully locked exact Blueprint/Canonical Script version with required evidence can satisfy downstream readiness."],
    ],6.1)
    d.save(path)
    return {"replaced_stale_phrase":hits}

VALIDATIONS=[
("CONTENT_INCOMPLETE","Required static catalog/profile/checklist/dimension exists only as schema/container or lacks canonical rows.","BLOCK completion declaration; route to canonical content owner.","Concrete canonical rows + render contract + lifecycle/owner evidence."),
("STALE_QA_EVIDENCE","Artifact/output/version/checksum changed after scorecard/review evidence.","BLOCK release/delivery/publish; invalidate current QA eligibility.","Fresh evaluation/recheck on exact new version."),
("CRITERIA_LIFECYCLE_REACHABILITY","Required governed criteria resource cannot traverse required lifecycle/control path.","BLOCK criteria activation/consumer readiness.","Reachable Configure→Submit Review→Approve/Activate path with evidence."),
("DENOMINATOR_ALIGNMENT","Current physical/control/function denominator differs from embedded/historical count.","BLOCK formal closure.","Fresh physical/registry scan + reconciled denominator."),
("QUALITY_CONTEXT_UNBOUND","Consumer lacks exact profile/version/policy/scorecard/finding/evidence refs.","BLOCK governed consumption.","Bind exact quality context before action."),
("HARD_BLOCK_BYPASS","Total score/manual decision attempts to bypass rights/policy/evidence/checksum/provenance hard block.","BLOCK action regardless of score.","Hard-block owner resolves required check; fresh gate evaluation."),
("RUNTIME_CLAIM","Word/design authority exists but no current runtime/production evidence supports completion claim.","BLOCK runtime-complete declaration only; design may remain defined.","Fresh implementation + validation evidence on current revision."),
]
def patch_sys(path):
    d=Document(path)
    if MARK in text(d): return "already"
    add_marker_heading(d,"System Inspection / Validation Catalog Synchronization")
    add_table(d,["Validation UID","Trigger","Enforcement","Recovery evidence"],VALIDATIONS,5.5)
    d.add_heading("Inspection Execution Rule",level=2)
    add_table(d,["Rule","Requirement"],[
        ["Applicability","Every validation is REQUIRED or NOT_APPLICABLE per inspected scope; silent omission forbidden."],
        ["Freshness","Evidence must bind current branch/revision/exact artifact refs. Historical PASS cannot substitute."],
        ["Owner","Finding routes to canonical owner; SYS may detect/orchestrate but must not silently take over product ownership."],
        ["False-completion guard","Document existence, UI presence, route name, schema presence or previous workflow success alone cannot satisfy runtime/product closure."],
        ["Terminal state","PASS only after fresh revalidation; otherwise BLOCKED / IMPLEMENTATION_REQUIRED / REVIEW_REQUIRED as applicable."],
    ],6.0)
    d.save(path)
    return "updated"

DELIVERY=[
("exact_output_ref","Exact deliverable artifact/content ref; latest/alias-only prohibited."),
("output_version_ref","Exact reviewed version."),
("artifact_checksum","Checksum/manifest must match reviewed output."),
("criteria_profile_uid","Applicable governed QA profile."),
("criteria_version_id","Exact approved criteria version."),
("quality_gate_policy_version","Exact gate policy version."),
("scorecard_id","Fresh scorecard for exact output/version."),
("qa_gate_status","Must be PASS for governed external delivery."),
("required_checks","All applicable hard checks PASS, including Rights/Policy/Provenance/Checksum where required."),
("open_blocking_findings","Must equal 0."),
("evidence_refs","Review/recheck/manual-case/required-check evidence lineage."),
("delivery_correlation_id","Audit correlation linking approval→delivery attempt→provider result."),
]
def patch_dev(path):
    d=Document(path)
    if MARK in text(d): return "already"
    add_marker_heading(d,"DEV-01 External Delivery QA Eligibility / Exact Binding")
    add_table(d,["Required field","Contract"],DELIVERY,5.8)
    d.add_heading("Delivery Gate / Failure Semantics",level=2)
    add_table(d,["Condition","Result"],[
        ["Missing exact output/version/checksum","BLOCK; do not deliver."],
        ["Missing/STALE scorecard or changed artifact after QA","BLOCK; require fresh evaluation/recheck."],
        ["Criteria profile/version/policy not exact or not approved","BLOCK."],
        ["Hard required check FAIL or open blocking finding > 0","BLOCK regardless of aggregate score or human business approval."],
        ["Eligible","Only after all required bindings validate may DEV create/send the external delivery command."],
        ["Provider/runtime not implemented","Keep implementation state NOT_EXECUTED / IMPLEMENTATION_REQUIRED; Word contract alone is not Production evidence."],
    ],5.9)
    d.save(path)
    return "updated"

results={}
results["index"]=patch_index(FILES["index"][0])
results["shell"]=patch_shell(FILES["shell"][0])
results["logic"]=patch_logic(FILES["logic"][0])
results["core"]=patch_core(FILES["core"][0])
results["sys"]=patch_sys(FILES["sys"][0])
results["dev"]=patch_dev(FILES["dev"][0])

# Semantic verification, fail closed.
checks={
FILES["index"][0]:[MARK,"31 DOCX","0921acpos","MUST be resolved fresh"],
FILES["shell"][0]:[MARK,"TY01","TY15","computed metrics","font_family"],
FILES["logic"][0]:[MARK,"L01","L24","S01","S30","H01","H12","NOT_APPLICABLE"],
FILES["core"][0]:[MARK,"LockReview contract is resolved","NOT_EXECUTED / IMPLEMENTATION_REQUIRED"],
FILES["sys"][0]:[MARK,"CONTENT_INCOMPLETE","STALE_QA_EVIDENCE","CRITERIA_LIFECYCLE_REACHABILITY","DENOMINATOR_ALIGNMENT","QUALITY_CONTEXT_UNBOUND","HARD_BLOCK_BYPASS","RUNTIME_CLAIM"],
FILES["dev"][0]:[MARK,"artifact_checksum","criteria_version_id","scorecard_id","open_blocking_findings","delivery_correlation_id"],
}
for fn,needles in checks.items():
    d=Document(fn);t=text(d)
    for n in needles:
        if n not in t: raise SystemExit(f"SEMANTIC_MISSING:{fn}:{n}")
    if fn==FILES["core"][0] and "unresolved LockReview contract stays BLOCK" in t:
        raise SystemExit("STALE_CORE_CONTRADICTION_REMAINS")

print(json.dumps({"marker":MARK,"results":results,"files":[v[0] for v in FILES.values()]},ensure_ascii=False,indent=2))
