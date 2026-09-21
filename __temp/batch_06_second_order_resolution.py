from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib

BASE_HEAD="22761e902112a39fdf0853022e56704c725b1ad8"
MARK="ACPOS-20260921-BATCH-06-SECOND-ORDER-CANONICAL-RESOLUTION-V1"
SYS_DOC="ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx"
SYSTEM_LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
SYS_SHA="964f5cd5983bb169b659f21030a336c7555f642c"
LOGIC_SHA="4ec3536925eba83dabb632cd602a1738b672681e"

SYSTEMS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}

def git_blob_sha(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

assert git_blob_sha(SYS_DOC)==SYS_SHA,(git_blob_sha(SYS_DOC),SYS_SHA)
assert git_blob_sha(SYSTEM_LOGIC)==LOGIC_SHA,(git_blob_sha(SYSTEM_LOGIC),LOGIC_SHA)

def norm(x):
    return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def get(vals,i):
    return "" if i is None or i>=len(vals) else norm(vals[i])
def is_missing(v):
    return v in {"","—","-","SOURCE_NOT_DEFINED"}

DISPLAY_ONLY_TYPES={"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"}
def display_only(r):
    typ=r.get("type","").upper()
    return typ in DISPLAY_ONLY_TYPES or any(x in typ for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]

def parse_rows(path,require_control=False):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows: continue
        h=[norm(c.text) for c in t.rows[0].cells]
        idx={
          "control":hfind(h,"control uid"),
          "type":hfind(h,"type"),
          "label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),
          "gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "operation":hfind(h,"operation"),
          "runtime_owner":hfind(h,"runtime owner"),
          "runtime_status":hfind(h,"runtime status")
        }
        if require_control and idx["control"] is None: continue
        if not require_control and all(idx[k] is None for k in ["control","action","gate","permission","operation","runtime_owner"]): continue
        for ri,r in enumerate(t.rows[1:],2):
            vals=[c.text for c in r.cells]
            rec={k:get(vals,idx[k]) for k in FIELDS}
            if require_control:
                uid=rec["control"]
                if not uid or uid in {"—","-"} or not re.search(r"[A-Za-z0-9]",uid): continue
            else:
                if not any(not is_missing(rec[k]) for k in ["control","action","gate","permission","operation","runtime_owner"]): continue
            rec.update({"source":str(path),"table":ti+1,"row":ri,"indices":idx})
            out.append(rec)
    return out

def classify(field,row):
    val=row.get(field,"");status=row.get("runtime_status","")
    if not is_missing(val): return "BOUND"
    if "UI_LOCAL_EXACT" in status:
        if field in {"operation","runtime_owner"}: return "LEGITIMATE_NA_UI_LOCAL"
        if field=="action" and display_only(row): return "LEGITIMATE_NA_UI_LOCAL"
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in status and field=="operation":
        return "CLOSED_BY_OWNER_LOCAL_COMMAND"
    if "READ_EXACT" in status and display_only(row):
        if field=="action": return "LEGITIMATE_NA_READ_PRESENTATION"
        if field=="operation" and not is_missing(row.get("runtime_owner","")):
            return "READ_OWNER_BOUND_OPERATION_UNSPECIFIED"
    if field=="runtime_owner" and any(x in status for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):
        return "KNOWN_RUNTIME_BLOCKER"
    if val=="SOURCE_NOT_DEFINED": return "SOURCE_RESOLUTION_REQUIRED"
    return "DEFINITION_BINDING_GAP"

def compose_by_uid(rows):
    groups=collections.defaultdict(list)
    for r in rows: groups[r["control"]].append(r)
    out={}
    for uid,rs in groups.items():
        ranked=sorted(rs,key=lambda r:sum(1 for k in FIELDS[1:] if not is_missing(r.get(k,""))),reverse=True)
        base=dict(ranked[0])
        for field in ["type","label","action","gate","permission","operation","runtime_owner","runtime_status"]:
            vals=sorted(set(r.get(field,"") for r in rs if not is_missing(r.get(field,""))))
            if is_missing(base.get(field,"")) and len(vals)==1:
                base[field]=vals[0]
            if len(vals)>1:
                base[field+"_same_uid_conflict"]=vals
        base["same_uid_rows"]=rs
        out[uid]=base
    return out

SRC_ROWS=[]
SYSTEM_TEXT={}
for f in SYSTEMS:
    SRC_ROWS += parse_rows(f,False)
    d=Document(f)
    txt="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for rr in t.rows for c in rr.cells])
    SYSTEM_TEXT[f]=txt

def candidate_records(target,field,page_rows):
    candidates=[]
    keys=[]
    if not is_missing(target.get("control","")): keys.append(("control",target["control"]))
    if field in {"gate","permission","operation","runtime_owner"} and not is_missing(target.get("action","")):
        keys.append(("action",target["action"]))
    if field=="permission" and not is_missing(target.get("gate","")):
        keys.append(("gate",target["gate"]))
    if field=="runtime_owner" and not is_missing(target.get("operation","")):
        keys.append(("operation",target["operation"]))
    seen=set()
    for key,val in keys:
        for r in SRC_ROWS+page_rows:
            if r.get(key,"")!=val: continue
            fv=r.get(field,"")
            if is_missing(fv): continue
            sig=(fv,r["source"],r["table"],r["row"],key,val)
            if sig not in seen:
                seen.add(sig)
                candidates.append({"value":fv,"source":r["source"],"table":r["table"],"row":r["row"],"matched_by":key,"matched_value":val})
    return candidates

def compute_graph():
    rep={"total_definition_gaps":0,"resolvable":0,"unresolved":0,"conflict":0,"rows":[],"conflicts":[],"page_rows":{},"composed":{}}
    for page,fn in PAGES.items():
        rows=parse_rows(fn,True);bm=compose_by_uid(rows)
        rep["page_rows"][page]=rows;rep["composed"][page]=bm
        for uid,t in sorted(bm.items()):
            for field in ["action","gate","permission","operation","runtime_owner"]:
                if classify(field,t)!="DEFINITION_BINDING_GAP": continue
                rep["total_definition_gaps"]+=1
                cs=candidate_records(t,field,rows)
                vals=collections.defaultdict(list)
                for c in cs: vals[c["value"]].append(c)
                base={"page":page,"file":fn,"uid":uid,"field":field,"type":t.get("type",""),"label":t.get("label",""),"action":t.get("action",""),"gate":t.get("gate",""),"permission":t.get("permission",""),"operation":t.get("operation",""),"runtime_owner":t.get("runtime_owner",""),"runtime_status":t.get("runtime_status","")}
                if len(vals)==1:
                    v=next(iter(vals));rep["resolvable"]+=1;rep["rows"].append(base|{"value":v,"sources":vals[v]})
                elif len(vals)==0:
                    rep["unresolved"]+=1;rep["rows"].append(base|{"value":None,"sources":[]})
                else:
                    rep["conflict"]+=1;rep["conflicts"].append(base|{"candidate_values":dict(vals)})
    return rep

pre=compute_graph()
assert pre["total_definition_gaps"]==312,pre["total_definition_gaps"]
assert pre["resolvable"]==1,pre["resolvable"]
assert pre["unresolved"]==310,pre["unresolved"]
assert pre["conflict"]==1,pre["conflict"]
one=[r for r in pre["rows"] if r.get("value") is not None]
assert len(one)==1,one
r0=one[0]
assert (r0["page"],r0["uid"],r0["field"],r0["value"])==("SYS-01","SYS-01-BTN-NAV-OPEN","runtime_owner","Projection / source page read owner"),r0

# Apply only the one second-order binding directly to existing SYS table cell(s).
doc=Document(SYS_DOC)
changed=0
for t in doc.tables:
    if not t.rows: continue
    h=[norm(c.text) for c in t.rows[0].cells]
    ci=hfind(h,"control uid");ri=hfind(h,"runtime owner")
    if ci is None or ri is None: continue
    for row in t.rows[1:]:
        vals=[c.text for c in row.cells]
        if get(vals,ci)=="SYS-01-BTN-NAV-OPEN":
            old=norm(row.cells[ri].text)
            if is_missing(old):
                row.cells[ri].text="Projection / source page read owner";changed+=1
            elif old!="Projection / source page read owner":
                raise AssertionError(("SYS_RUNTIME_OWNER_CONFLICT",old))
assert changed>0,changed
doc.save(SYS_DOC)
Document(SYS_DOC)

post=compute_graph()
assert post["total_definition_gaps"]==311,post["total_definition_gaps"]
assert post["resolvable"]==0,post["resolvable"]
assert post["unresolved"]==310,post["unresolved"]
assert post["conflict"]==1,post["conflict"]

# Deep exact-key coverage classification for the 310 unresolved rows.
queue=[]
coverage_counts=collections.Counter()
field_counts=collections.Counter()
page_counts=collections.Counter()
for r in [x for x in post["rows"] if x.get("value") is None]:
    page=r["page"];uid=r["uid"];field=r["field"]
    target=post["composed"][page][uid]
    same=target.get("same_uid_rows",[])
    field_column_present=any(x["indices"].get(field) is not None for x in same)
    keys=[]
    for key in ["control","action","gate","operation"]:
        v=target.get(key,"")
        if not is_missing(v):
            keys.append((key,v))
    hits=[]
    for key,val in keys:
        for sf,txt in SYSTEM_TEXT.items():
            if val and val in txt:
                hits.append({"key":key,"value":val,"source":sf})
    # De-duplicate exact source hits.
    uniq=[];seen=set()
    for h in hits:
        sig=(h["key"],h["value"],h["source"])
        if sig not in seen:seen.add(sig);uniq.append(h)
    if uniq and field_column_present:
        cls="EXACT_KEY_PRESENT_FIELD_CELL_UNBOUND"
    elif uniq and not field_column_present:
        cls="EXACT_KEY_PRESENT_PAGE_SCHEMA_FIELD_ABSENT"
    elif (not uniq) and field_column_present:
        cls="NO_SYSTEM_AUTHORITY_KEY_REFERENCE_FIELD_CELL_UNBOUND"
    else:
        cls="NO_SYSTEM_AUTHORITY_KEY_REFERENCE_PAGE_SCHEMA_FIELD_ABSENT"
    rec={
      "page":page,"uid":uid,"field":field,"type":target.get("type",""),"label":target.get("label",""),
      "action":target.get("action",""),"gate":target.get("gate",""),"permission":target.get("permission",""),
      "operation":target.get("operation",""),"runtime_status":target.get("runtime_status",""),
      "coverage_class":cls,"field_column_present":field_column_present,
      "exact_key_hits":uniq[:12],"exact_key_hit_count":len(uniq)
    }
    queue.append(rec);coverage_counts[cls]+=1;field_counts[field]+=1;page_counts[page]+=1

assert len(queue)==310,len(queue)

# Preserve the known conflict exactly.
conf=post["conflicts"][0]
assert (conf["page"],conf["uid"],conf["field"])==("SYS-01","SYS-01-BTN-NAV-OPEN","gate"),conf
conflict_values=sorted(conf["candidate_values"].keys())
assert conflict_values==["CURRENT_VIEW_AND_SOURCE_ACTION_GATE","SOURCE_PAGE_GATE"],conflict_values

# Append authoritative Batch 06 resolution ledger.
logic=Document(SYSTEM_LOGIC)
txt="\n".join([p.text for p in logic.paragraphs]+[c.text for t in logic.tables for rr in t.rows for c in rr.cells])
assert MARK not in txt,"BATCH06_ALREADY_PRESENT"

sec=logic.add_section(WD_SECTION.NEW_PAGE)
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width,sec.page_height=sec.page_height,sec.page_width
sec.top_margin=Inches(.35);sec.bottom_margin=Inches(.35);sec.left_margin=Inches(.3);sec.right_margin=Inches(.3)
logic.add_heading("Batch 06 · Second-order Exact Binding / Canonical Resolution Queue",level=1)
p=logic.add_paragraph()
p.add_run("["+MARK+"] ").bold=True
p.add_run("Batch 06 closes the single second-order binding unlocked by Batch 05, then classifies every remaining unresolved Definition Binding Gap by exact canonical-key coverage and page-schema materialization. No label similarity, fuzzy inference, owner guessing, API invention, or conflict auto-selection is allowed.")

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def table(headers,rows,fs=4.5):
    t=logic.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row):c[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for pp in cell.paragraphs:
                for rr in pp.runs:rr.font.size=Pt(fs)
    return t

logic.add_heading("Batch Result",level=2)
table(["Item","Count / State"],[
["Pre-batch Definition Binding Gap",312],
["Second-order exact binding applied",1],
["Post-batch Definition Binding Gap",311],
["Remaining unresolved queue",310],
["Preserved authority conflict",1],
["Post-batch exact-graph resolvable",0],
],5.4)

logic.add_heading("Second-order Binding Closed",level=2)
table(["Page","Control UID","Field","Exact Value","Authority Chain"],[
["SYS-01","SYS-01-BTN-NAV-OPEN","Runtime Owner","Projection / source page read owner","Operation getUiProjection -> exact registered projection/source-page read owner"]
],5.1)

logic.add_heading("Canonical Resolution Queue Classification",level=2)
table(["Coverage Class","Count"],[[k,coverage_counts[k]] for k in sorted(coverage_counts)],5.2)
table(["Missing Field","Count"],[[k,field_counts[k]] for k in sorted(field_counts)],5.2)
table(["Page","Remaining Unresolved"],[[k,page_counts[k]] for k in sorted(page_counts)],5.0)

logic.add_heading("310-row Canonical Resolution Queue",level=2)
table(["Page","Control UID","Missing Field","Type","Action","Gate","Operation","Runtime Status","Page Field Schema","Authority Coverage","Exact Authority Sources"],[
[
 q["page"],q["uid"],q["field"],q["type"],q["action"],q["gate"],q["operation"],q["runtime_status"],
 "COLUMN_PRESENT" if q["field_column_present"] else "COLUMN_ABSENT",
 q["coverage_class"],
 "; ".join(sorted(set(h["source"] for h in q["exact_key_hits"]))) if q["exact_key_hits"] else "NONE"
] for q in queue
],3.55)

logic.add_heading("Preserved Conflict / Fail-closed",level=2)
table(["Page","Control UID","Field","Conflicting Values","Disposition"],[
["SYS-01","SYS-01-BTN-NAV-OPEN","Gate"," | ".join(conflict_values),"PRESERVED_UNRESOLVED. Same ACT-NAV-OPEN has different Gate authority by source-page context; Batch 06 does not select a value."]
],5.0)

machine={
 "marker":MARK,
 "pre":{"definition_gaps":312,"resolvable":1,"unresolved":310,"conflict":1},
 "applied":{"page":"SYS-01","uid":"SYS-01-BTN-NAV-OPEN","field":"runtime_owner","value":"Projection / source page read owner","changed_cells":changed},
 "post":{"definition_gaps":311,"resolvable":0,"unresolved":310,"conflict":1},
 "coverage_counts":dict(coverage_counts),
 "field_counts":dict(field_counts),
 "page_counts":dict(page_counts),
 "preserved_conflict":{"page":"SYS-01","uid":"SYS-01-BTN-NAV-OPEN","field":"gate","values":conflict_values}
}
logic.add_paragraph("BATCH06_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(SYSTEM_LOGIC)
Document(SYSTEM_LOGIC)

report={
 "machine":machine,
 "queue":queue,
 "conflict":conf,
 "sys_blob_sha":git_blob_sha(SYS_DOC),
 "system_logic_blob_sha":git_blob_sha(SYSTEM_LOGIC)
}
Path("__batch06_resolution_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH06_RESULT="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
