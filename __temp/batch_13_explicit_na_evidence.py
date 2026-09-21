from docx import Document
import json,re

FILES={
"S01":"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"S03":"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"WB":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"AIAPI":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"LOGIC":"ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx",
}
def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())

out={}
# exact WB System01 rows
for key,fn in [("S01",FILES["S01"]),("WB",FILES["WB"])]:
    d=Document(fn);hits=[]
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if vals and vals[0].startswith("CTRL-WORKSPACE-WB-01-"):
                hits.append({"table":ti,"row":ri,"headers":headers,"values":vals})
    out[key+"_WB_ROWS"]=hits

# AIAPI READ_EXACT rows with missing Action in owner/page
for key,fn in [("S03",FILES["S03"]),("AIAPI",FILES["AIAPI"])]:
    d=Document(fn);hits=[]
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        hl=[h.lower() for h in headers]
        ci=next((i for i,h in enumerate(hl) if "control uid" in h),None)
        ai=next((i for i,h in enumerate(hl) if "action uid" in h),None)
        si=next((i for i,h in enumerate(hl) if "runtime status" in h),None)
        if ci is None:continue
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            uid=vals[ci] if ci<len(vals) else ""
            if not uid.startswith("CTRL-ADMIN-AIAPI-"):continue
            action=vals[ai] if ai is not None and ai<len(vals) else ""
            status=vals[si] if si is not None and si<len(vals) else ""
            if action in {"","—","-"} and status=="READ_EXACT":
                hits.append({"table":ti,"row":ri,"headers":headers,"values":vals})
    out[key+"_AIAPI_READ_NO_ACTION"]=hits

# explicit N/A/projection/read-action language
for key,fn in FILES.items():
    d=Document(fn);ctx=[]
    terms=["N/A_READ_PROJECTION","read projection","read-only projection","no effectful action","action uid","direct operation","operation-specific"]
    for i,p in enumerate(d.paragraphs,1):
        t=norm(p.text)
        if t and any(x.lower() in t.lower() for x in terms):
            ctx.append({"kind":"paragraph","index":i,"text":t})
    for ti,tb in enumerate(d.tables,1):
        for ri,row in enumerate(tb.rows,1):
            t=" | ".join(norm(c.text) for c in row.cells)
            if t and any(x.lower() in t.lower() for x in terms):
                ctx.append({"kind":"table","table":ti,"row":ri,"text":t})
    out[key+"_CONTEXT"]=ctx[:200]

print("BATCH13_EXPLICIT="+json.dumps(out,ensure_ascii=False,sort_keys=True))
