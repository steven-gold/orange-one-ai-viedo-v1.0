from pathlib import Path
from docx import Document
import json,re,subprocess,collections

BASE_HEAD="c17e7bbf4ddf9208c6a80cea659627f84ede87a9"
S02="02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"

assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def collect(path,keywords):
    d=Document(path);out={"paragraphs":[],"table_rows":[]}
    rx=re.compile("|".join(re.escape(x) for x in keywords),re.I)
    for i,p in enumerate(d.paragraphs):
        t=norm(p.text)
        if t and rx.search(t):out["paragraphs"].append({"index":i,"text":t[:2000]})
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            joined=" | ".join(vals)
            if rx.search(joined):
                out["table_rows"].append({"table":ti+1,"row":ri,"headers":h,"values":vals,"joined":joined[:4000]})
    return out

s05=collect(S05,["ManualReviewService","manual review","MANUAL-MODIFY","MANUAL-PASS","manual modify","manual pass","QA_USE","QA-01-GATE-MANUAL-DECISION"])
s02=collect(S02,["Shared Conversation Core","attach","attachment","context","CONVERSATION-ATTACH","conversation composer","system.ai.use"])
summary={
"s05_paragraph_hits":len(s05["paragraphs"]),"s05_table_row_hits":len(s05["table_rows"]),
"s02_paragraph_hits":len(s02["paragraphs"]),"s02_table_row_hits":len(s02["table_rows"]),
}
report={"marker":"ACPOS-20260922-BATCH-31-OWNER-OPERATION-VOCABULARY-AUDIT-V1","base_head":BASE_HEAD,"summary":summary,
"manual_review_owner_evidence":s05,"conversation_attach_owner_evidence":s02}
Path("__batch31_operation_vocabulary_audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
Path("__batch31_summary.txt").write_text(json.dumps(summary,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print("BATCH31_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
