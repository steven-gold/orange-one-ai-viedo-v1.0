from pathlib import Path
from docx import Document
import json,re,subprocess

BASE_HEAD="01343f341271c65d45405c4fb6de0dac96de5339"
S03="03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
AIAPI="ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx"
OPS=["configureGovernedResource","approveGovernedResource","runSandboxTest","setKillSwitch"]
KW=["governed resource","sandbox","kill switch","provider profile","persistence","persist","audit","approval","configuration","quarantine","runtime","secret","test"]

assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

def norm(x):
    return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())

rx=re.compile("|".join(re.escape(x) for x in OPS+KW),re.I)

def collect(path):
    d=Document(path);out={"paragraphs":[],"rows":[]}
    for i,p in enumerate(d.paragraphs):
        t=norm(p.text)
        if t and rx.search(t):
            out["paragraphs"].append({"index":i,"text":t[:3000]})
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            joined=" | ".join(vals)
            if joined and rx.search(joined):
                out["rows"].append({"table":ti+1,"row":ri,"headers":h,"values":vals,"joined":joined[:5000]})
    return out

owner=collect(S03)
page=collect(AIAPI)
summary={
    "operations":4,
    "owner_paragraph_hits":len(owner["paragraphs"]),
    "owner_table_hits":len(owner["rows"]),
    "page_paragraph_hits":len(page["paragraphs"]),
    "page_table_hits":len(page["rows"]),
}
report={
    "marker":"ACPOS-20260922-BATCH-37-SYSTEM03-PERSISTENCE-DESIGN-AUDIT-V1",
    "base_head":BASE_HEAD,
    "summary":summary,
    "operations":OPS,
    "owner_evidence":owner,
    "page_evidence":page,
}
Path("__batch37_system03_persistence_audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
Path("__batch37_summary.txt").write_text(json.dumps(summary,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print("BATCH37_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
