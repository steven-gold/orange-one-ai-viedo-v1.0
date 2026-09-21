from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,os

MARK="ACPOS-20260921-BATCH-08-AUTHORITY-REMEDIATION-DENOMINATOR-V1"
WB_MARK="ACPOS-20260921-BATCH-08-WB-CANONICAL-OWNER-MAPPING-V1"

SYSTEMS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
TARGET_OWNER_DOCS={
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx":14,
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx":19,
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx":133,
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx":54,
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx":6,
}
WB_FILE="ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
LOGIC_FILE="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"

EXPECTED_SHA={
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx":"aa3ddacece36fc55400f98396867c31bb1c3a12b",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx":"ef69737bb09ea510638e1843359cc75c937d0758",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx":"77cf27afbb2cb4572a006e8dab5cfa52206525d8",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx":"e86a0f2118e885680d98f1c39c2d4cf07e498947",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx":"dc52068889ed6d1f959d7cadc7048644957960a3",
WB_FILE:"75842c122d81823c8b3e4a696c56a035e606477a",
LOGIC_FILE:"b52b05687fd6425e69ed360351b371b9d331e9bc",
}

PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":WB_FILE,
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}

def blob_sha(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,sha in EXPECTED_SHA.items():
    assert blob_sha(f)==sha,(f,blob_sha(f),sha)

def norm(x): return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def get(vals,i): return "" if i is None or i>=len(vals) else norm(vals[i])
def missing(v): return v in {"","—","-","SOURCE_NOT_DEFINED"}

DISPLAY_ONLY_TYPES={"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"}
def display_only_type(t):
    typ=(t or "").upper()
    return typ in DISPLAY_ONLY_TYPES or any(x in typ for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]

def parse_page_rows(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(h,"control uid")
        if ci is None:continue
        idx={
          "control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),"operation":hfind(h,"operation"),
          "runtime_owner":hfind(h,"runtime owner"),"runtime_status":hfind(h,"runtime status")
        }
        for ri,r in enumerate(t.rows[1:],2):
            vals=[c.text for c in r.cells]
            uid=get(vals,ci)
            if not uid or uid in {"—","-"} or not re.search(r"[A-Za-z0-9]",uid):continue
            rec={k:get(vals,idx[k]) for k in FIELDS}
            rec.update({"source":str(path),"table":ti+1,"row":ri,"headers":h,"indices":idx})
            out.append(rec)
    return out

def compose(rows):
    groups=collections.defaultdict(list)
    for r in rows:groups[r["control"]].append(r)
    out={}
    for uid,rs in groups.items():
        ranked=sorted(rs,key=lambda r:sum(1 for k in FIELDS[1:] if not missing(r.get(k,""))),reverse=True)
        base=dict(ranked[0])
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        base["same_uid_rows"]=rs
        out[uid]=base
    return out

def classify_gap(field,row):
    val=row.get(field,"");status=row.get("runtime_status","")
    if not missing(val):return "BOUND"
    if "UI_LOCAL_EXACT" in status:
        if field in {"operation","runtime_owner"}:return "LEGITIMATE_NA_UI_LOCAL"
        if field=="action" and display_only_type(row.get("type","")):return "LEGITIMATE_NA_UI_LOCAL"
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in status and field=="operation":return "CLOSED_BY_OWNER_LOCAL_COMMAND"
    if "READ_EXACT" in status and display_only_type(row.get("type","")):
        if field=="action":return "LEGITIMATE_NA_READ_PRESENTATION"
        if field=="operation" and not missing(row.get("runtime_owner","")):return "READ_OWNER_BOUND_OPERATION_UNSPECIFIED"
    if field=="runtime_owner" and any(x in status for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return "KNOWN_RUNTIME_BLOCKER"
    if val=="SOURCE_NOT_DEFINED":return "SOURCE_RESOLUTION_REQUIRED"
    return "DEFINITION_BINDING_GAP"

# Exact owner rows from 01-09.
owner_index=collections.defaultdict(list)
for sf in SYSTEMS:
    d=Document(sf)
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(headers,"control uid")
        if ci is None:continue
        idx={
          "control":ci,"type":hfind(headers,"type"),"label":hfind(headers,"label","顯示名稱"),
          "action":hfind(headers,"action uid"),"gate":hfind(headers,"gate uid","gate"),
          "permission":hfind(headers,"permission","auth resource"),"operation":hfind(headers,"operation"),
          "runtime_owner":hfind(headers,"runtime owner"),"runtime_status":hfind(headers,"runtime status")
        }
        for ri,row in enumerate(t.rows[1:],2):
            vals=[c.text for c in row.cells]
            uid=get(vals,ci)
            if not uid or uid in {"—","-"}:continue
            rec={k:get(vals,idx[k]) for k in FIELDS}
            rec.update({"source":sf,"table":ti+1,"row":ri})
            owner_index[uid].append(rec)

queue=[]
for page,fn in PAGES.items():
    rows=parse_page_rows(fn);bm=compose(rows)
    for uid,t in sorted(bm.items()):
        for field in ["action","gate","permission","operation","runtime_owner"]:
            if classify_gap(field,t)!="DEFINITION_BINDING_GAP":continue
            if page=="SYS-01" and uid=="SYS-01-BTN-NAV-OPEN" and field=="gate":continue
            owners=owner_index.get(uid,[])
            sources=sorted(set(r["source"] for r in owners))
            queue.append({
              "page":page,"file":fn,"uid":uid,"field":field,"type":t.get("type",""),"label":t.get("label",""),
              "action":t.get("action",""),"gate":t.get("gate",""),"permission":t.get("permission",""),
              "operation":t.get("operation",""),"runtime_owner":t.get("runtime_owner",""),"runtime_status":t.get("runtime_status",""),
              "owner_sources":sources,"field_column_present":any(r["indices"].get(field) is not None for r in t["same_uid_rows"])
            })

assert len(queue)==310,len(queue)
owner_defined=[q for q in queue if q["owner_sources"]]
wb_absent=[q for q in queue if not q["owner_sources"]]
assert len(owner_defined)==282,len(owner_defined)
assert len(wb_absent)==28,len(wb_absent)
assert all(q["page"]=="WB-01" for q in wb_absent)

unique=[q for q in owner_defined if len(q["owner_sources"])==1]
multi=[q for q in owner_defined if len(q["owner_sources"])>1]
assert len(unique)==226,len(unique)
assert len(multi)==56,len(multi)

by_owner=collections.defaultdict(list)
for q in unique:by_owner[q["owner_sources"][0]].append(q)
assert {k:len(v) for k,v in by_owner.items()}==TARGET_OWNER_DOCS,{k:len(v) for k,v in by_owner.items()}

collision_pairs=collections.Counter(tuple(q["owner_sources"]) for q in multi)
expected_pairs={
 tuple(sorted(["03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx","09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"])):18,
 tuple(sorted(["01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx","09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"])):18,
 tuple(sorted(["05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx","09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"])):8,
 tuple(sorted(["02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx","04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"])):7,
 tuple(sorted(["02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx","06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"])):2,
 tuple(sorted(["02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx","08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"])):2,
 tuple(sorted(["02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx","09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"])):1,
}
assert collision_pairs==collections.Counter(expected_pairs),collision_pairs

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=4.7):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row):c[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:r.font.size=Pt(fs)
    return t

def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.35);sec.bottom_margin=Inches(.35);sec.left_margin=Inches(.3);sec.right_margin=Inches(.3)

# Write unique-owner remediation registers into the exact canonical owner docs.
owner_hashes={}
for sf,expected_count in TARGET_OWNER_DOCS.items():
    rows=by_owner[sf]
    assert len(rows)==expected_count
    d=Document(sf)
    text="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for rr in t.rows for c in rr.cells])
    marker=MARK+"::"+os.path.basename(sf)
    assert marker not in text,(sf,"BATCH08_ALREADY_PRESENT")
    add_landscape(d)
    d.add_heading("Batch 08 · Canonical Spec Remediation Register",level=1)
    p=d.add_paragraph();p.add_run("["+marker+"] ").bold=True
    p.add_run("Every row below is owned uniquely by this Canonical System document. Current evidence does not authorize a concrete missing value, a LEGITIMATE_N/A classification, or a Runtime-blocked reclassification. The canonical owner must explicitly define the field or explicitly declare its non-applicability/blocking semantics before any page binding may close.")
    add_table(d,["State","Count"],[
      ["CANONICAL_SPEC_REMEDIATION_REQUIRED",len(rows)],
      ["FORMAL_NOT_APPLICABLE",0],
      ["FORMAL_RUNTIME_BLOCKED",0],
      ["PAGE_BINDING_MUTATION_AUTHORIZED",0],
    ],5.2)
    add_table(d,["Page","Control UID","Missing Field","Type","Existing Action","Existing Gate","Existing Permission","Existing Operation","Existing Runtime Owner","Runtime Status","Required Owner Action"],[
      [q["page"],q["uid"],q["field"],q["type"],q["action"],q["gate"],q["permission"],q["operation"],q["runtime_owner"],q["runtime_status"],
       "DEFINE_EXACT_VALUE_OR_EXPLICIT_NA/BLOCKED_IN_CANONICAL_OWNER"]
      for q in rows
    ],3.6)
    d.save(sf);Document(sf);owner_hashes[sf]=blob_sha(sf)

# WB owner-mapping register; no Action/Gate values are created.
wb=Document(WB_FILE)
wtxt="\n".join([p.text for p in wb.paragraphs]+[c.text for t in wb.tables for rr in t.rows for c in rr.cells])
assert WB_MARK not in wtxt,"WB_BATCH08_ALREADY_PRESENT"
group=collections.defaultdict(list)
for q in wb_absent:group[q["uid"]].append(q)
assert len(group)==14,len(group)
add_landscape(wb)
wb.add_heading("Batch 08 · WB Canonical Control UID / Owner Mapping Register",level=1)
p=wb.add_paragraph();p.add_run("["+WB_MARK+"] ").bold=True
p.add_run("The 14 WB compact/synthetic control UIDs below account for 28 unresolved fields (Action + Gate). They already carry exact Permission / Operation / Runtime Owner anchors in the page design, but no exact Control UID owner row exists in System 01–09. Action and Gate MUST remain unbound until a canonical owner mapping is established.")
wb_rows=[]
for uid,qs in sorted(group.items()):
    q=qs[0]
    missing_fields=" + ".join(sorted(x["field"].upper() for x in qs))
    wb_rows.append([uid,missing_fields,q["permission"],q["operation"],q["runtime_owner"],q["runtime_status"],"CANONICAL_CONTROL_UID_OWNER_MAPPING_REQUIRED"])
add_table(wb,["WB Control UID","Missing Fields","Existing Permission","Existing Operation","Existing Runtime Owner","Runtime Status","Mapping State"],wb_rows,4.5)
wb.save(WB_FILE);Document(WB_FILE)
wb_hash=blob_sha(WB_FILE)

# Central System Logic ledger.
logic=Document(LOGIC_FILE)
ltxt="\n".join([p.text for p in logic.paragraphs]+[c.text for t in logic.tables for rr in t.rows for c in rr.cells])
assert MARK not in ltxt,"LOGIC_BATCH08_ALREADY_PRESENT"
add_landscape(logic)
logic.add_heading("Batch 08 · Authority Remediation Denominator / Owner Mapping",level=1)
p=logic.add_paragraph();p.add_run("["+MARK+"] ").bold=True
p.add_run("Batch 08 resolves the remediation ownership of all 310 unresolved Definition Binding Gap rows without inventing missing contract values. Exact owner evidence is authoritative; page labels, shared operations, and cross-domain similarities are non-authoritative.")

add_table(logic,["Denominator Class","Count","Batch 08 Disposition"],[
["CANONICAL_OWNER_ROW_FIELD_UNDEFINED",282,"All require canonical specification remediation; current evidence supports 0 formal N/A and 0 Runtime-blocked reclassifications."],
["Unique canonical owner",226,"Remediation register persisted in the exact owner document."],
["Multiple canonical owners",56,"OWNER_COLLISION; do not write a missing value until canonical ownership/contract precedence is resolved."],
["WB CANONICAL_OWNER_ROW_ABSENT",28,"14 WB controls × Action/Gate. Canonical Control UID / Owner mapping required before binding."],
["Admissible page binding mutation",0,"No page Action/Gate/Permission/Operation/Runtime Owner value is fabricated in Batch 08."],
],5.0)

logic.add_heading("Unique-owner Remediation Distribution",level=2)
add_table(logic,["Canonical Owner Document","Rows","Current Page-binding Effect"],[
 [os.path.basename(sf),len(by_owner[sf]),"NONE — canonical spec remediation register only"]
 for sf in TARGET_OWNER_DOCS
],4.6)

logic.add_heading("Multi-owner Collision Distribution",level=2)
add_table(logic,["Owner A","Owner B","Rows","Disposition"],[
 [os.path.basename(pair[0]),os.path.basename(pair[1]),count,"OWNER_COLLISION / FAIL-CLOSED"]
 for pair,count in sorted(collision_pairs.items(),key=lambda x:(-x[1],x[0]))
],4.3)

logic.add_heading("56-row Multi-owner Collision Ledger",level=2)
add_table(logic,["Page","Control UID","Missing Field","Owner Documents","Existing Action","Existing Gate","Existing Operation","Runtime Status","Disposition"],[
 [q["page"],q["uid"],q["field"],"; ".join(os.path.basename(x) for x in q["owner_sources"]),q["action"],q["gate"],q["operation"],q["runtime_status"],"OWNER_COLLISION"]
 for q in multi
],3.7)

logic.add_heading("28-row WB Canonical Owner Mapping Ledger",level=2)
add_table(logic,["Control UID","Missing Field","Permission","Operation","Runtime Owner","Disposition"],[
 [q["uid"],q["field"],q["permission"],q["operation"],q["runtime_owner"],"CANONICAL_CONTROL_UID_OWNER_MAPPING_REQUIRED"]
 for q in wb_absent
],4.0)

logic.add_heading("Batch 08 Classification Decision",level=2)
logic.add_paragraph("For the 282 exact-owner rows, no current canonical Runtime Status / Control Type evidence supports converting any row to FORMAL_NOT_APPLICABLE or FORMAL_RUNTIME_BLOCKED. Therefore all 282 remain CANONICAL_SPEC_REMEDIATION_REQUIRED. This is a statement about current evidence only; a later canonical owner amendment may explicitly define N/A or blocked semantics.")
logic.add_paragraph("For the 28 WB rows, the existing Permission / Operation / Runtime Owner anchors are retained as evidence only. They do not authorize construction of Action, Gate, or a canonical owner identity.")

machine={
 "marker":MARK,
 "denominator":310,
 "owner_field_undefined":282,
 "classification":{"CANONICAL_SPEC_REMEDIATION_REQUIRED":282,"FORMAL_NOT_APPLICABLE":0,"FORMAL_RUNTIME_BLOCKED":0},
 "unique_owner_rows":226,
 "unique_owner_distribution":{os.path.basename(k):len(v) for k,v in by_owner.items()},
 "multi_owner_rows":56,
 "collision_pairs":{" + ".join(os.path.basename(x) for x in k):v for k,v in collision_pairs.items()},
 "wb_owner_absent_rows":28,
 "wb_controls":14,
 "page_binding_mutations":0
}
logic.add_paragraph("BATCH08_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC_FILE);Document(LOGIC_FILE)
logic_hash=blob_sha(LOGIC_FILE)

report={
 "machine":machine,
 "unique_owner_rows":unique,
 "multi_owner_rows":multi,
 "wb_owner_absent_rows":wb_absent,
 "output_hashes":{"owner_docs":owner_hashes,"wb":wb_hash,"system_logic":logic_hash}
}
Path("__batch08_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH08_REMEDIATION="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
