from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json, re

MARK="ACPOS-20260921-CONCRETE-BUSINESS-CONTENT-CLOSURE-V1"
STATUS="NEW_AUTHORITY_AFTER_COMPLETE_SOURCE_SCAN"

profiles={
"QCP-CORE-BASELINE-V1":[
("QA-DIM-CORE-GOAL","目標／Blueprint 意圖符合","Goal / Blueprint Intent Compliance",15,False,"核心故事、Topic、Blueprint 與明確目標一致"),
("QA-DIM-CORE-LOGIC","故事／邏輯結構","Story / Logic Structure",20,False,"因果、前後條件、敘事結構無矛盾"),
("QA-DIM-CORE-DNA-PLAN","角色／DNA 規劃一致性","Character / DNA Plan Consistency",15,False,"角色身份、外觀、行為、DNA 規劃不互相衝突"),
("QA-DIM-CORE-SCENE-OBJECT","場景／物件連續規劃","Scene / Object Continuity Plan",10,False,"場景、物件、出現條件與連續性規劃完整"),
("QA-DIM-CORE-COMPLETENESS","製作需求完整性","Production Requirement Completeness",15,False,"ASSET/VIDEO/EDIT/VOICE/Subtitle 所需資訊可施工"),
("QA-DIM-CORE-AMBIGUITY","歧義／衝突控制","Ambiguity / Contradiction Control",10,False,"無未裁決歧義、互斥指令或不明 owner"),
("QA-DIM-EVIDENCE-PROVENANCE","證據／來源可追溯","Evidence / Provenance",10,False,"來源、版本、引用、決策證據完整"),
("QA-DIM-RIGHTS-POLICY","權利／政策","Rights / Policy",5,True,"權利、政策、禁止內容與使用限制必須 PASS"),
],
"QCP-ASSET-IMAGE-V1":[
("QA-DIM-ASSET-SCRIPT","腳本／Manifest 符合","Script / Manifest Compliance",15,False,"符合 ASSET Script、Manifest Item 與用途"),
("QA-DIM-ASSET-IDENTITY","人物／角色 DNA 一致性","Character / Identity DNA Fidelity",20,False,"臉、身形、服裝、標誌特徵與 DNA 一致"),
("QA-DIM-ASSET-OBJECT","物件一致性","Object Fidelity",10,False,"物件形狀、材質、尺寸比例與指定版本一致"),
("QA-DIM-ASSET-SCENE","場景／環境一致性","Scene / Environment Fidelity",10,False,"環境、空間、背景元素與 Scene DNA 一致"),
("QA-DIM-ASSET-COMPOSITION","構圖／取景","Composition / Framing",10,False,"主體位置、鏡位意圖、留白與視覺重心合理"),
("QA-DIM-ASSET-GEOMETRY","解剖／幾何／生成瑕疵","Anatomy / Geometry / Artifact Integrity",15,False,"手指、肢體、透視、破圖、重影、錯字等無明顯瑕疵"),
("QA-DIM-ASSET-STYLE","風格／色彩／光線一致性","Style / Color / Lighting Consistency",10,False,"風格、色調、光線與系列基準一致"),
("QA-DIM-ASSET-TECH","技術品質","Technical Image Quality",5,False,"尺寸、解析度、格式、清晰度與檔案完整"),
("QA-DIM-RIGHTS-POLICY","權利／政策","Rights / Policy",5,True,"權利與政策要求必須 PASS"),
],
"QCP-ASSET-AUDIO-V1":[
("QA-DIM-AUDIO-SCRIPT","腳本／Manifest 符合","Script / Manifest Compliance",15,False,"符合 Voice/Music/SFX Script 與用途"),
("QA-DIM-AUDIO-IDENTITY","聲音身份／語意一致","Voice Identity / Semantic Fit",20,False,"指定角色聲線、語氣或音樂/SFX 意圖一致"),
("QA-DIM-AUDIO-CONTENT","發音／內容正確","Pronunciation / Content Accuracy",15,False,"台詞、名稱、語言、發音與文字來源正確"),
("QA-DIM-AUDIO-TIMING","時間／長度","Timing / Duration",10,False,"開始、結束、節奏與所屬片段長度合理"),
("QA-DIM-AUDIO-NOISE","雜訊／爆音／失真","Noise / Clipping / Distortion",15,False,"無不可接受雜訊、爆音、削波、失真"),
("QA-DIM-AUDIO-LEVEL","音量／動態","Level / Dynamics",10,False,"音量、峰值、動態與可懂度符合用途"),
("QA-DIM-AUDIO-TECH","技術格式","Technical Audio Format",10,False,"取樣率、聲道、檔案格式與完整性符合輸出契約"),
("QA-DIM-RIGHTS-POLICY","權利／政策","Rights / Policy",5,True,"授權、素材來源與政策要求必須 PASS"),
],
"QCP-VIDEO-BASELINE-V1":[
("VIDEO-01-SCORE-SCRIPT","腳本符合","Script Compliance",15,False,"符合 Blueprint/VIDEO Script 與 shot beats"),
("VIDEO-01-SCORE-IDENTITY","人物／身份穩定","Character / Identity Stability",15,False,"人物臉、身形、服裝、DNA 穩定"),
("VIDEO-01-SCORE-MOTION","動作／身體連續性","Motion / Body Continuity",15,False,"動作連續、物理合理、無跳變"),
("VIDEO-01-SCORE-CAMERA","鏡頭／視線","Camera / Gaze",10,False,"鏡位、構圖、視線與行動方向符合腳本"),
("VIDEO-01-SCORE-SCENE","場景／連續性","Scene / Continuity",10,False,"場景、物件、空間與幀間狀態連續"),
("VIDEO-01-SCORE-TIMING","時間／片長","Timing / Duration",10,False,"片段長度、節奏、事件時序符合要求"),
("VIDEO-01-SCORE-TECH","技術品質","Technical Quality",10,False,"無破幀、閃爍、嚴重壓縮或不可用輸出"),
("VIDEO-01-SCORE-ASSET","素材／DNA 忠實度","Asset / DNA Fidelity",10,False,"綁定素材與 DNA 在影片中保持一致"),
("VIDEO-01-SCORE-RIGHTS","權利／政策","Rights / Policy",5,True,"權利與政策 Hard Block 必須 PASS"),
],
"QCP-EDIT-FINAL-V1":[
("QA-DIM-EDIT-ASSEMBLY","鏡頭順序／組接符合","Assembly / Shot Order Compliance",12,False,"Shot order、cut point 與 Script/Assembly 規格一致"),
("QA-DIM-EDIT-PACING","節奏／剪接／轉場","Pacing / Cuts / Transitions",10,False,"節奏、剪接點、轉場自然且符合情緒"),
("QA-DIM-EDIT-VISUAL","畫面／色彩／幀連續","Visual / Color / Frame Continuity",8,False,"色彩、亮度、構圖與幀間連續無異常"),
("QA-DIM-EDIT-VOICE","配音／對白品質","Voice / Dialogue Quality",8,False,"對白清楚、聲線正確、無截斷或不合理停頓"),
("QA-DIM-EDIT-BGM","BGM 音量／Duck","BGM Balance / Ducking",6,False,"BGM 與對白平衡、ducking 合理"),
("QA-DIM-EDIT-SFX","音效時間／層次","SFX Timing / Layering",6,False,"SFX 時點、層次與事件一致"),
("QA-DIM-EDIT-AVSYNC","音畫同步","Audio / Video Sync",8,False,"聲音事件與畫面事件同步"),
("QA-DIM-EDIT-LIPSYNC","嘴型／口型同步","Lip Sync",10,False,"嘴型與 DialogueTimingBinding／Voice Segment 對齊"),
("QA-DIM-EDIT-SUB-CONTENT","字幕內容正確","Subtitle Content Accuracy",8,False,"字幕文字、語言、專有名詞與台詞一致"),
("QA-DIM-EDIT-SUB-TIMING","字幕時間／可讀性／安全區","Subtitle Timing / Readability / Safe Area",8,False,"In/Out、閱讀時間、換行、遮擋、安全區符合規範"),
("QA-DIM-EDIT-TECH","最終輸出技術完整","Technical Output Integrity",8,False,"Resolution/FPS/Codec/Duration/Tracks/Manifest 可驗證"),
("QA-DIM-EDIT-RIGHTS","權利／來源／政策","Rights / Provenance / Policy",8,True,"Rights、Provenance、Policy Hard Block 必須 PASS"),
],
"QCP-QA-FINAL-OUTPUT-V1":[
("QA-DIM-FINAL-SCRIPT","腳本／目標符合","Script / Goal Compliance",6,False,"最終輸出符合 Blueprint、Script 與目標"),
("QA-DIM-FINAL-IDENTITY","人物／角色一致性","Character Identity Consistency",9,False,"人物臉、身形、服裝、DNA 跨鏡頭一致"),
("QA-DIM-FINAL-SCENE","場景／物件一致性","Scene / Object Consistency",7,False,"場景、物件、空間關係與版本一致"),
("QA-DIM-FINAL-COMPOSITION","構圖／取景","Composition / Framing",6,False,"構圖、主體位置、視覺重心符合意圖"),
("QA-DIM-FINAL-MOTION","動作／身體連續性","Motion / Body Continuity",7,False,"動作連續、姿態合理、無瞬移"),
("QA-DIM-FINAL-FRAME","幀間連續性","Frame-to-Frame Continuity",8,False,"相鄰幀無閃爍、跳變、物體突變或不合理 morph"),
("QA-DIM-FINAL-CAMERA","鏡頭／視線連續","Camera / Gaze Continuity",5,False,"鏡位、軸線、視線與運動方向連續"),
("QA-DIM-FINAL-LIPSYNC","嘴型／口型同步","Lip Sync",7,False,"嘴型與實際語音/DialogueTimingBinding 對齊"),
("QA-DIM-FINAL-VOICE","配音／對白品質","Voice / Dialogue Quality",6,False,"聲線、語氣、清晰度與角色一致"),
("QA-DIM-FINAL-AUDIO","音訊技術品質","Audio Technical Quality",5,False,"無 clipping、noise、dropout、失真"),
("QA-DIM-FINAL-MIX","BGM／音效平衡與時間","BGM / SFX Balance & Timing",5,False,"BGM、SFX 與對白音量/時點合理"),
("QA-DIM-FINAL-SUBCONTENT","字幕內容正確","Subtitle Content Accuracy",6,False,"字幕與台詞、語言、專有名詞一致"),
("QA-DIM-FINAL-SUBTIMING","字幕時間／可讀性／安全區","Subtitle Timing / Readability / Safe Area",6,False,"字幕 in/out、閱讀時間、換行、遮擋、安全區正確"),
("QA-DIM-FINAL-EDIT","剪輯節奏／轉場","Editing Pacing / Transition",6,False,"節奏、cut、transition 與敘事情緒一致"),
("QA-DIM-FINAL-AVSYNC","音畫同步","Audio / Video Sync",5,False,"對白、SFX、音樂事件與畫面同步"),
("QA-DIM-FINAL-TECH","輸出技術完整","Technical Output Integrity",4,False,"Resolution/FPS/Codec/Duration/Manifest/Decode 正常"),
("QA-DIM-FINAL-RIGHTS","權利／政策／來源","Rights / Policy / Provenance",2,True,"Rights/Policy/Provenance Hard Block 必須 PASS"),
],
}
for pid,rows in profiles.items():
    assert sum(x[3] for x in rows)==100,(pid,sum(x[3] for x in rows))

profile_meta=[
("QCP-CORE-BASELINE-V1","CORE","CORE / Blueprint / Canonical Script",95,"NEW 95 threshold authority; criteria rows new authority"),
("QCP-ASSET-IMAGE-V1","ASSET","IMAGE asset outputs",95,"NEW 95 threshold authority; concrete dimensions new authority"),
("QCP-ASSET-AUDIO-V1","ASSET","VOICE / MUSIC / SFX asset outputs",95,"NEW 95 threshold authority; concrete dimensions new authority"),
("QCP-VIDEO-BASELINE-V1","VIDEO","VIDEO candidate outputs",95,"Threshold 95 and 9 dimensions derived from existing VIDEO authority; weights new authority"),
("QCP-EDIT-FINAL-V1","EDITING","EDIT stage/final outputs",95,"Threshold 95 derived from existing EDIT authority; concrete dimensions/weights new authority"),
("QCP-QA-FINAL-OUTPUT-V1","QA","Final output / release-candidate review",95,"Concrete final QA catalog and weights new authority"),
]
required_checks=[
("QA-CHK-OUTPUT-IDENTITY","Exact Output Identity","Exact output_id/version/artifact ref resolved; no latest/default alias"),
("QA-CHK-CHECKSUM","Checksum / Manifest","Artifact checksum matches exact manifest/output bytes"),
("QA-CHK-PROVENANCE","Provenance / Lineage","Blueprint/Script/Asset/Video/Edit lineage and source refs are complete"),
("QA-CHK-CRITERIA","Approved Criteria Binding","Exact CriteriaVersion + QualityGatePolicyVersion is APPROVED and applicable"),
("QA-CHK-RIGHTS","Rights / Policy","Rights, content policy, channel/use restrictions PASS"),
("QA-CHK-BOUND-INPUTS","Bound Inputs","Required input asset/script/DNA refs exist and are not stale/superseded"),
("QA-CHK-DECODE","Decode / Technical Manifest","Output opens/decodes; duration/resolution/fps/track manifest matches contract"),
("QA-CHK-BLOCKING-FINDING","Blocking Findings","No unresolved blocking Finding/Hard Block"),
("QA-CHK-MANUAL-REVIEW","Manual Review","Required ManualReviewCase is closed/passed"),
("QA-CHK-RECHECK","Correction/Recheck Freshness","After correction, QA uses a verified NEW exact output version before closure"),
]
issue_categories=[
("QA-ISSUE-SCRIPT","SCRIPT_COMPLIANCE","腳本／目標不符"),
("QA-ISSUE-IDENTITY","IDENTITY_CONSISTENCY","人物／角色／DNA 一致性"),
("QA-ISSUE-SCENE","SCENE_OBJECT_CONTINUITY","場景／物件連續性"),
("QA-ISSUE-COMPOSITION","COMPOSITION_FRAMING","構圖／取景"),
("QA-ISSUE-MOTION","MOTION_CONTINUITY","動作／身體連續性"),
("QA-ISSUE-FRAME","FRAME_CONTINUITY","幀間連續性"),
("QA-ISSUE-CAMERA","CAMERA_GAZE","鏡頭／視線"),
("QA-ISSUE-LIPSYNC","LIP_SYNC","嘴型／口型同步"),
("QA-ISSUE-VOICE","VOICE_DIALOGUE","配音／對白"),
("QA-ISSUE-AUDIO","AUDIO_TECHNICAL","音訊技術品質"),
("QA-ISSUE-MIX","BGM_SFX","BGM／音效"),
("QA-ISSUE-SUBCONTENT","SUBTITLE_CONTENT","字幕內容"),
("QA-ISSUE-SUBTIMING","SUBTITLE_TIMING_READABILITY","字幕時間／可讀性"),
("QA-ISSUE-EDIT","EDITING_PACING","剪輯節奏／轉場"),
("QA-ISSUE-AVSYNC","AV_SYNC","音畫同步"),
("QA-ISSUE-TECH","TECHNICAL_OUTPUT","技術輸出"),
("QA-ISSUE-RIGHTS","RIGHTS_POLICY_PROVENANCE","權利／政策／來源"),
("QA-ISSUE-EVIDENCE","EVIDENCE_PROVENANCE","證據／追溯"),
]

FILES={
"INDEX":"00_INDEX_ACPOS_Mother_Compliant_Basic_Design_Master.docx",
"GLOBAL":"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"CORELOGIC":"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"CREATIVE":"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"SYSLOGIC":"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"CORE":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ASSET":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"VIDEO":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"EDIT":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"QA":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
}

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr()
    e=OxmlElement('w:tblHeader'); e.set(qn('w:val'),'true'); trPr.append(e)
def shade(cell,fill='EDE9FE'):
    tcPr=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:shd'); e.set(qn('w:fill'),fill); tcPr.append(e)
def table(doc,headers,rows,fs=6.2):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; repeat_header(t.rows[0])
    for i,h in enumerate(headers): t.rows[0].cells[i].text=str(h); shade(t.rows[0].cells[i])
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row): c[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs: r.font.size=Pt(fs)
    return t
def landscape(doc):
    s=doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation=WD_ORIENT.LANDSCAPE
    s.page_width,s.page_height=s.page_height,s.page_width
    s.top_margin=Inches(.38); s.bottom_margin=Inches(.38); s.left_margin=Inches(.35); s.right_margin=Inches(.35)
def has_mark(doc):
    return any(MARK in p.text for p in doc.paragraphs)
def title(doc,name):
    doc.add_heading(name,level=1)
    p=doc.add_paragraph(); p.add_run(f'[{MARK}] ').bold=True
    p.add_run('這一段是完整來源掃描後補入的 Concrete Business Content Authority。既有來源已明確的內容沿用；原來源只有抽象 Criteria/Dimension 容器而缺實際明細者，以下項目標示為本次新 Authority。')
def profile_rows(pid):
    return profiles[pid]
def add_profile(doc,pid,show_status=True):
    meta=next(x for x in profile_meta if x[0]==pid)
    doc.add_heading(f'{pid} · {meta[1]} · {meta[2]}',level=2)
    table(doc,["Profile UID","Department","Applicability","Pass Threshold","Authority Basis"],[meta],6.3)
    rows=[]
    for order,(uid,zh,en,w,hard,definition) in enumerate(profiles[pid],1):
        rows.append([order,uid,zh,en,w,"HARD BLOCK" if hard else "NO",definition])
    table(doc,["Order","Dimension UID","審查項目","English","Weight","Hard Block","Definition"],rows,5.7)
    doc.add_paragraph('Scoring: 每個 Dimension 0–100；Weighted Score = Score × Weight / 100；Total = 全部適用 Dimension 加總。PASS = Total ≥ 95 且所有 Hard Block / Required Check PASS。頁面不得自行新增第二套 threshold。')

def add_global_gate(doc):
    landscape(doc); title(doc,'Business Content Completeness Gate / 業務內容完整性 Gate')
    doc.add_paragraph('完成判定不得只看 Section、Control、Action、API 或空的 Registry/Library 容器。任何頁面只要宣稱管理 Canonical Catalog / Criteria / Dimension / Checklist / Profile / Preset / Stage Definition，就必須存在可施工的具體項目與顯示規則；否則 status = CONTENT_INCOMPLETE。')
    table(doc,["Class","Completion Rule","Examples"],[
        ["STATIC_CANONICAL_CATALOG","必須列出實際 UID/名稱/語意/適用範圍/必要參數或權重；空 Dimension Library 不得 PASS","QA Criteria、Score Dimension、Asset Type、Stage Definition、Role/Permission fixed group"],
        ["DYNAMIC_RUNTIME_RECORD","不得在文件 seed 假資料；完整度以 Schema/Owner/Operation/Empty-State/Render Contract 判定","Account、Project、Finance amount、Provider profile、Knowledge source、Strategy fact"],
        ["HYBRID","固定 Catalog 必須列明；實際 Instance 由 Runtime 投影","QA criteria catalog + runtime CriteriaVersion、Provider capability types + real provider profiles"],
    ],6.4)
    table(doc,["Invariant","Rule"],[
        ["CONTENT-01","Abstract container ≠ content. Registry/Library/Table header without canonical rows is incomplete when the product requires a baseline catalog."],
        ["CONTENT-02","Page UI MUST render concrete catalog rows, not only JSON blob / count / placeholder."],
        ["CONTENT-03","If Current source lacks the required static catalog after complete source scan, create a clearly marked New Authority rather than pretending existing authority is complete."],
        ["CONTENT-04","Dynamic business records remain empty until real data exists; do not confuse intentional empty runtime data with missing static authority."],
        ["CONTENT-05","QA/Review/Score pages require actual review dimensions/checks plus their UI rendering contract before completion may be declared."],
    ],6.4)
    table(doc,["18-page audit result","Classification"],[
        ["CORE-01","FIXED HERE: concrete CORE evaluation profile was missing."],
        ["ASSET-01","FIXED HERE: concrete ASSET image/audio evaluation dimensions were missing."],
        ["VIDEO-01","EXISTING CONTENT OK: 9 concrete dimensions existed; this closure adds governed weights/profile mapping only."],
        ["EDIT-01","FIXED HERE: Stage Evaluation used dimension_scores but did not name concrete dimensions."],
        ["QA-01 / SG-02","FIXED HERE: concrete review catalog + page-row rendering contract were missing."],
        ["WB/DB/INFO/IAM/DEV/SOC/ERP/AIAPI/STR/Admin-STR/KB/SYS","NO SAME DEFECT FOUND: concrete schema/stages/entities exist; runtime instances are intentionally real-data-only and may be empty."],
    ],6.3)

def add_sg02(doc):
    landscape(doc); title(doc,'SG-02 Concrete QA Criteria Catalog / QA 審查項目正式明細')
    doc.add_paragraph('SG-02 不再只管理抽象 Dimension。以下 6 個 baseline profiles 與其 dimensions 是本次新增的正式施工 Authority；Runtime 仍需依 DRAFT→REVIEW→APPROVED lifecycle materialize CriteriaVersion，文件定義不等於 Production 已核准。')
    for pid,_,_,_,_ in profile_meta: add_profile(doc,pid)
    doc.add_heading('Canonical Required Checks',level=2)
    table(doc,["Check UID","Name","Blocking Rule"],required_checks,6.0)
    doc.add_heading('Canonical QA Issue Categories',level=2)
    table(doc,["Category UID","Enum","Display"],issue_categories,6.0)
    doc.add_heading('SG-02 Page Rendering Contract',level=2)
    table(doc,["Surface","Must Display"],[
        ["Criteria Table","Profile UID / Criteria Version / Department / Artifact Applicability / Status / Threshold / Dimension Count / Required Check Count / Approval Ref"],
        ["Dimension Library Drawer","逐項 Dimension UID / 中文名 / English / Definition；禁止只顯示 dimensions JSON/count"],
        ["Threshold Drawer","每個 Profile 的 overall pass threshold=95；Hard Block override；VIDEO/EDIT 95 沿用既有 authority，其餘 95 為本次新 Authority"],
        ["Required Checks Drawer","逐項顯示上述 10 個 checks、PASS/BLOCK 狀態與 Evidence refs"],
        ["Department Mapping","CORE→CORE profile；ASSET→IMAGE/AUDIO profile by artifact type；VIDEO→VIDEO profile；EDITING→EDIT profile；QA→FINAL OUTPUT profile"],
        ["Approval / Impact","顯示 DRAFT/REVIEW/APPROVED/ACTIVE/SUPERSEDED、content hash、affected tasks、revalidation requirement"],
    ],6.1)
    doc.add_heading('SG-02 Lifecycle Closure',level=2)
    table(doc,["State","Control / Operation","Rule"],[
        ["EMPTY","Existing Configure control = Create Draft mode","有 quality.criteria.configure 權限時可從 canonical baseline profile 建立第一個 DRAFT；不可因 EMPTY 永久 disable。"],
        ["DRAFT","Configure","只可修改 DRAFT；APPROVED/ACTIVE immutable。"],
        ["DRAFT → REVIEW","NEW: Submit Review","新增 Control UID CTRL-ADMIN-SG-02-ACT-04-ACT-SUBMIT-REVIEW；Action UID SG-02-ACT-SUBMIT-REVIEW；new operation submitGovernedResourceReview；POST /v1/governance/resources/{id}/submit-review；permission quality.criteria.configure；產生 audit event governance.review_submitted。"],
        ["REVIEW","Approve","沿用 approveGovernedResource；permission governance.approve；通過後 materialize APPROVED quality_criteria_version。"],
        ["APPROVED / ACTIVE","Read-only","修改必須建立新 version；不得原地覆寫。"],
    ],6.0)
    doc.add_paragraph('New denominator: SG-02 Current Control Denominator = 12（原 11 + Submit Review）。任何仍宣稱 11 的舊 closure row 被此 V1 business-content closure supersede。')
    doc.add_heading('Production Acceptance Truth (2026-09-21 validation)',level=2)
    table(doc,["Evidence","Result"],[
        ["Source/runtime release tests","6/6 PASS"],
        ["Production release SHA","e28f528eab7c739fc8d585d962d74d3a689fa7da"],
        ["Production read-only controls","9/9 PASS · getUiProjection · HTTP 200 · drawers visible"],
        ["Production criteria data at validation","EMPTY · criteria_versions=[] · dimensions=null · required_checks=null · gate_policy=null"],
        ["Concrete QA detail in Production","MISSING at validation"],
        ["Configure / Approve","Both disabled because no materialized Criteria lifecycle data"],
        ["Effectful acceptance","BLOCKED until new lifecycle closure is implemented and at least one formal CriteriaVersion is DRAFT→REVIEW→APPROVED"],
    ],6.2)

def add_qa(doc):
    landscape(doc); title(doc,'QA-01 Concrete Review Item Display / 實際審查項目顯示')
    add_profile(doc,"QCP-QA-FINAL-OUTPUT-V1")
    doc.add_heading('QA-01 Scorecard Row Contract',level=2)
    table(doc,["Column","Rule"],[
        ["Order","沿 approved profile 固定排序；不可 frontend 自排"],
        ["Dimension UID","Exact UID"],
        ["審查項目","顯示具體中文名稱，例如 人物／角色一致性、幀間連續性、嘴型／口型同步、字幕內容、字幕時間／可讀性、音訊品質、構圖等"],
        ["Weight","Profile-bound；當前適用 rows 總和必須 =100"],
        ["Score","System/AI 0–100；人工不可直接輸入"],
        ["Weighted Score","score×weight/100"],
        ["Status","PASS / FINDING / HARD_BLOCK / N/A"],
        ["Hard Block","是/否；任何 Hard Block FAIL 覆蓋 total score"],
        ["Evidence Refs","Viewer exact frame/time/range/output/checksum/provenance refs"],
        ["Finding Count","該 dimension 的 open/closed findings"],
    ],6.0)
    table(doc,["UI behavior","Exact Rule"],[
        ["Criteria header","顯示 Profile UID、Criteria Version、Policy Version、Department、Artifact Type、Threshold=95、Approval status"],
        ["Score rows","QA-01-LIST-SCORE-DIMENSIONS 必須逐項 render approved criteria rows；不得只顯示 Dimension count/JSON blob"],
        ["Row click","聚焦 Exact Output Viewer 對應 Evidence/frame/time/range；不修改上游輸出"],
        ["Missing Criteria","BLOCK QA scoring；顯示 QA-01-ERR-CRITERIA-001；禁止空 Scorecard 或 generic fallback"],
        ["Required Checks","固定列出 10 個 canonical checks 與 Evidence / PASS/BLOCK"],
        ["Finding Category","必須從 canonical issue categories 選；缺 category = BLOCK + REPORT_AUTHORITY_GAP"],
        ["Manual Review","仍只有 MODIFY / PASS；不新增人工 score input；Manual PASS 不能繞過 QA Gate"],
    ],6.0)

def add_core(doc):
    landscape(doc); title(doc,'CORE-01 Concrete Evaluation Dimensions')
    add_profile(doc,"QCP-CORE-BASELINE-V1")
    doc.add_paragraph('CORE-01-CMP-EVALUATION / CORE-01-FLD-EVALUATION 必須逐項顯示上述 8 個 dimensions、weight、score、hard-block、evidence；不得只顯示「CORE Evaluation / Criteria Version / Evidence」單一摘要。')

def add_asset(doc):
    landscape(doc); title(doc,'ASSET-01 Concrete Evaluation Dimensions')
    add_profile(doc,"QCP-ASSET-IMAGE-V1")
    add_profile(doc,"QCP-ASSET-AUDIO-V1")
    doc.add_paragraph('ASSET-01-FLD-DIMENSIONS 必須依 Asset Type 選擇 IMAGE 或 AUDIO profile，逐項顯示 Dimension UID / Label / Weight / Score / Hard Block / Evidence。CHARACTER/SCENE/OBJECT/VFX/THUMBNAIL 等 image-compatible 類使用 IMAGE profile；VOICE/MUSIC/SFX 使用 AUDIO profile。不得再只顯示抽象欄位名稱。')

def add_video(doc):
    landscape(doc); title(doc,'VIDEO-01 Governed Dimension Weights')
    add_profile(doc,"QCP-VIDEO-BASELINE-V1")
    doc.add_paragraph('VIDEO 已有 9 個 concrete Score UID；本 closure 不建立第二套 UID，只補治理權重與 SG-02 profile binding。既有 VIDEO-01-FLD-SCORE-* 控制逐項顯示。')

def add_edit(doc):
    landscape(doc); title(doc,'EDIT-01 Concrete Stage Evaluation Dimensions')
    add_profile(doc,"QCP-EDIT-FINAL-V1")
    table(doc,["Stage","Primary Applicable Dimensions"],[
        ["Stage 01 Assembly","Assembly/Shot Order; Pacing/Cuts/Transitions; Visual/Frame Continuity"],
        ["Stage 02 Audio","Voice/Dialogue; BGM Balance; SFX Timing; Audio/Video Sync"],
        ["Stage 03 Sync","Audio/Video Sync; Lip Sync; Subtitle Content; Subtitle Timing/Readability/Safe Area"],
        ["Stage 04 Finalize","All applicable dimensions + Technical Output + Rights/Provenance Hard Block"],
    ],6.2)
    doc.add_paragraph('EDIT-01-CMP-STAGE-EVALUATION 必須逐項 render current stage applicable dimensions，而不是只有 overall_score / dimension_scores JSON/summary。Stage 01–03 confirmation uses the same concrete profile rows; Finalize performs full-profile evaluation.')

def add_logic_summary(doc,profiles_to_add,title_text):
    landscape(doc); title(doc,title_text)
    for pid in profiles_to_add:
        meta=next(x for x in profile_meta if x[0]==pid)
        table(doc,["Profile UID","Department","Applicability","Pass Threshold","Dimension Count","Authority Basis"],[[pid,meta[1],meta[2],meta[3],len(profiles[pid]),meta[4]]],6.2)
    doc.add_paragraph('Cross-document invariant: the owning page must render concrete dimension rows from the exact approved profile. An abstract CriteriaVersion/Dimension container alone is CONTENT_INCOMPLETE.')

def add_index(doc):
    landscape(doc); title(doc,'Concrete Business Content Closure Index')
    table(doc,["Page","Previous defect","Closure"],[
        ["CORE-01","Approved CORE Criteria referenced, no concrete dimensions","8-dimension CORE baseline profile added + page row-render requirement"],
        ["ASSET-01","Scorecard fields existed, no concrete dimension catalog","9-dimension IMAGE + 8-dimension AUDIO profiles added"],
        ["VIDEO-01","9 dimensions existed but no governed weights/profile","Existing 9 UIDs preserved; weights + SG-02 mapping added"],
        ["EDIT-01","dimension_scores existed, no concrete stage dimension names","12-dimension EDIT profile + stage applicability added"],
        ["QA-01","Generated dimension table existed, but actual review items absent","17 concrete final-output review items + exact scorecard display contract added"],
        ["SG-02","Dimension/Criteria containers empty; lifecycle lacked DRAFT→REVIEW action","6 baseline profiles + dimension/check/category catalogs + Submit Review authority; denominator 12"],
    ],6.0)
    table(doc,["Other pages audited","Result"],[
        ["WB-01 / DB-01 / INFO-01 / IAM-01","No same defect: fixed schemas/sections/entities are concrete; runtime records are intentionally dynamic."],
        ["DEV-01 / SOC-01 / ERP-01 / AIAPI-01","No same defect: business entities, stages/contracts/fields are concrete; values/providers/amounts remain real-data-only."],
        ["STR-01 / Admin STR-01 / KB-01 / SYS-01","No same defect: workflow/entity/review/source contracts are concrete; dynamic facts/candidates/sources must not be seeded."],
    ],6.1)

actions={
"GLOBAL":add_global_gate,
"CORE":add_core,
"ASSET":add_asset,
"VIDEO":add_video,
"EDIT":add_edit,
"QA":add_qa,
"SG02":add_sg02,
"INDEX":add_index,
}
updated=[]
for key,fn in FILES.items():
    p=Path(fn); d=Document(p)
    if has_mark(d):
        print("SKIP_ALREADY_APPLIED",fn); continue
    if key in actions:
        actions[key](d)
    elif key=="CORELOGIC":
        add_logic_summary(d,["QCP-CORE-BASELINE-V1"],'CORE Criteria Concrete Content Closure')
    elif key=="CREATIVE":
        add_logic_summary(d,["QCP-ASSET-IMAGE-V1","QCP-ASSET-AUDIO-V1","QCP-VIDEO-BASELINE-V1","QCP-EDIT-FINAL-V1","QCP-QA-FINAL-OUTPUT-V1"],'Creative Production Criteria Concrete Content Closure')
    elif key=="SYSLOGIC":
        landscape(d); title(d,'Quality Governance / Content Completeness Closure')
        table(d,["Rule","Requirement"],[
            ["Concrete Criteria","QA/CORE/ASSET/VIDEO/EDIT review surfaces require actual dimension rows, not empty catalogs."],
            ["SG-02 lifecycle","EMPTY→DRAFT→REVIEW→APPROVED must be reachable; Submit Review is now required New Authority."],
            ["Runtime acceptance","Production read-only SG-02 9/9 PASS; effectful acceptance remains blocked until lifecycle implementation and real CriteriaVersion materialization."],
            ["Completion declaration","No page may be declared complete if a required static business catalog is absent, even when UID/API bindings are complete."],
        ],6.2)
    d.save(p); Document(p); updated.append(fn)

Path("__temp").mkdir(exist_ok=True)
Path("__temp/concrete_content_patch_report.json").write_text(json.dumps({
 "marker":MARK,
 "updated":updated,
 "profiles":{k:{"dimensions":len(v),"weight_sum":sum(x[3] for x in v)} for k,v in profiles.items()},
 "sg02_new_control_denominator":12,
 "new_submit_review":{"control_uid":"CTRL-ADMIN-SG-02-ACT-04-ACT-SUBMIT-REVIEW","operation":"submitGovernedResourceReview","path":"POST /v1/governance/resources/{id}/submit-review","permission":"quality.criteria.configure"},
},ensure_ascii=False,indent=2),encoding="utf-8")
print("UPDATED",len(updated),updated)
