from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-20-AUDIT-SEMANTICS-IAM-CONFLICT-CLOSURE-V1"
BASE_HEAD="9248e011612f72b7ee8c8b04f71bbe68db9ede63"
S09="09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
IAM="ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
IAM_UID="IAM-01-BTN-COMPLETE"
IAM_PERMISSION="iam.user.create/configure + iam.permission.configure"

PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":IAM,
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
SYSTEM_DOCS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
S09,
]
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]
COMPARE_FIELDS=["action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]
EXPECTED_SHA={
S09:"b55eb7ca56ae8c19339b6b8f6443f0f925937690",
IAM:"2fdb636a2b72380e0d8f51e42660b2085c4c4201",
LOGIC:"70d66128cd1faf5be2f63de903992adbaf6c4c31",
}

def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()
for f,s in EXPECTED_SHA.items():
    got=blob(f);assert got==s,(f,got,s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        fv={}
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            fv[f]=vals
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        base["_field_values"]=fv;base["_occurrences"]=len(rs);out[uid]=base
    return out

# Precondition: exact same-UID conflict exists only in IAM page and System09 is unanimous.
iam_pre=[r for r in parse_controls(IAM) if r["control"]==IAM_UID]
pre_vals=sorted(set(r["permission"] for r in iam_pre if not missing(r["permission"])))
assert pre_vals==["iam.user.create/configure + iam.permission.configure","user manage + permission configure"],pre_vals
s09_rows=[r for r in parse_controls(S09) if r["control"]==IAM_UID]
s09_vals=sorted(set(r["permission"] for r in s09_rows if not missing(r["permission"])))
assert s09_vals==[IAM_PERMISSION],s09_vals

# Normalize all physical page occurrences of the same control UID to exact Canonical Owner permission.
d=Document(IAM);patched=0;seen=0
for t in d.tables:
    if not t.rows:continue
    h=[norm(c.text) for c in t.rows[0].cells]
    ci=exact_index(h,"Control UID");pi=exact_index(h,"Permission / Auth Resource")
    if pi is None:pi=hfind(h,"permission","auth resource")
    if ci is None or pi is None:continue
    for row in t.rows[1:]:
        vals=[norm(c.text) for c in row.cells]
        if ci>=len(vals) or vals[ci]!=IAM_UID:continue
        seen+=1
        old=norm(row.cells[pi].text)
        if old!=IAM_PERMISSION:
            row.cells[pi].text=IAM_PERMISSION;patched+=1
assert seen>=2,(seen,patched)
assert patched>=1,(seen,patched)
d.save(IAM);Document(IAM)
post_vals=sorted(set(r["permission"] for r in parse_controls(IAM) if r["control"]==IAM_UID and not missing(r["permission"])))
assert post_vals==[IAM_PERMISSION],post_vals

def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.33);sec.bottom_margin=Inches(.33);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=4.4):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for rr in rows:
        cells=t.add_row().cells
        for i,v in enumerate(rr):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
    return t

# Canonical owner evidence ledger.
d=Document(S09);txt="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
assert MARK not in txt
add_landscape(d);d.add_heading("Batch 20 · IAM Permission Conflict Closure / Audit Identity Semantics",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::SYSTEM09] ").bold=True
p.add_run("IAM-01-BTN-COMPLETE is one control identity. System 09 Canonical Owner occurrences unanimously bind Permission/Auth Resource to 'iam.user.create/configure + iam.permission.configure'. The page-local legacy phrase 'user manage + permission configure' is not a second permission identity and is superseded for Current construction binding. This closure changes no runtime implementation claim.")
add_table(d,["Control UID","Field","Canonical Value","Superseded Page Alias","Result"],[[IAM_UID,"Permission / Auth Resource",IAM_PERMISSION,"user manage + permission configure","NORMALIZE_TO_CANONICAL_OWNER_VALUE"]],4.3)
d.save(S09);Document(S09)

# Page evidence ledger.
d=Document(IAM);add_landscape(d);d.add_heading("Batch 20 · IAM Complete Permission Normalization",level=1)
p=d.add_paragraph();p.add_run("["+MARK+"::PAGE::IAM-01] ").bold=True
p.add_run("All physical occurrences of IAM-01-BTN-COMPLETE now use the same exact Permission/Auth Resource from System 09. No Action, Gate, Operation, Method/Path, Runtime Owner or runtime status is changed.")
add_table(d,["Control UID","Canonical Permission","Physical Occurrences Seen","Cells Changed","Status"],[[IAM_UID,IAM_PERMISSION,seen,patched,"SAME_UID_FIELD_CONFLICT_CLOSED"]],4.5)
d.save(IAM);Document(IAM)

# Formalize audit semantics in System Logic; this is audit interpretation, not a product runtime mutation.
logic=Document(LOGIC);lt="\n".join([p.text for p in logic.paragraphs]+[c.text for t in logic.tables for row in t.rows for c in row.cells])
assert MARK not in lt
add_landscape(logic);logic.add_heading("Batch 20 · Cross-System Audit Identity Semantics Hardening",level=1)
p=logic.add_paragraph();p.add_run("["+MARK+"] ").bold=True
p.add_run("Batch 18/19 proved that a scanner which treats the lexical Operation cell as a globally unique key creates false conflicts across independently owned domains. Current remediation history already requires exact relations to be resolved inside the same Canonical Owner and forbids cross-system transfer. Therefore cross-system consistency audit MUST resolve Canonical Owner and exact Control/contract context before declaring an Operation conflict.")
add_table(logic,["Audit Rule","Required Semantics","Fail-Closed Boundary"],[
["Operation identity comparison","Resolve Canonical Owner first. Compare exact owner-bound contract using Control UID + Operation + Method/Path + Gate/Permission/Runtime context. Lexical Operation equality alone is not a global identity collision.","Only contradictory exact bindings inside the same Current owner/contract context may block."],
["Cross-owner same Operation text","Treat as domain-scoped lexical reuse unless Current Canonical Authority explicitly declares a shared operation identity.","Do not rename or transfer owners/routes from another domain by name similarity."],
["LOCAL LOCAL","Transport sentinel for UI-local/no-public-route behavior; not a network Method/Path identity.","Never count two LOCAL LOCAL rows as an API route collision."],
["NO_PUBLIC_API_BY_AUTHORITY","Explicit no-public-route sentinel; not a missing or colliding public API route.","Preserve owner-local command semantics."],
["Same Control UID field conflict","Two different non-empty values for the same field of the same Current Control UID remain a true blocker unless supersession/precedence resolves one value.","IAM-01-BTN-COMPLETE resolved by System 09 canonical permission in this batch."],
["Action fan-out","A repeated Action UID is review-only unless Current owner semantics prove the action identity is one-to-one with one Operation.","STR-01-ACT-NOOP remains legitimate shared read/no-op; CORE fan-out remains review evidence, not auto-mutation."],
],4.0)

# Corrected full-package audit.
page_rows={p:parse_controls(f) for p,f in PAGES.items()}
page_maps={p:compose(rows) for p,rows in page_rows.items()}
findings=[]
uid_pages=collections.defaultdict(list)
for page,m in page_maps.items():
    for uid in m:uid_pages[uid].append(page)
for uid,pages in uid_pages.items():
    if len(pages)>1:findings.append({"class":"CONTROL_UID_CROSS_PAGE_COLLISION","uid":uid,"pages":pages})
for page,m in page_maps.items():
    for uid,r in m.items():
        for f in COMPARE_FIELDS:
            vals=r["_field_values"].get(f,[])
            if len(vals)>1:findings.append({"class":"PAGE_INTERNAL_FIELD_CONFLICT","page":page,"uid":uid,"field":f,"values":vals})
# Only same-page operation consistency is authoritative without a resolved shared Canonical Owner.
for page,m in page_maps.items():
    ops=collections.defaultdict(list)
    routes=collections.defaultdict(list)
    for uid,r in m.items():
        op=r.get("operation","");mp=r.get("method_path","")
        if not missing(op):ops[op].append((uid,r))
        if not missing(mp) and mp not in {"LOCAL LOCAL","NO_PUBLIC_API_BY_AUTHORITY"}:routes[mp].append((uid,r))
    for op,rows in ops.items():
        methods=sorted(set(r.get("method_path","") for _,r in rows if not missing(r.get("method_path","")) and r.get("method_path","") not in {"LOCAL LOCAL","NO_PUBLIC_API_BY_AUTHORITY"}))
        owners=sorted(set(r.get("runtime_owner","") for _,r in rows if not missing(r.get("runtime_owner",""))))
        if len(methods)>1:findings.append({"class":"SAME_PAGE_OPERATION_METHOD_CONFLICT","page":page,"operation":op,"methods":methods})
        if len(owners)>1:findings.append({"class":"SAME_PAGE_OPERATION_OWNER_CONFLICT","page":page,"operation":op,"owners":owners})
    for mp,rows in routes.items():
        operations=sorted(set(r.get("operation","") for _,r in rows if not missing(r.get("operation",""))))
        if len(operations)>1:findings.append({"class":"SAME_PAGE_ROUTE_OPERATION_COLLISION","page":page,"method_path":mp,"operations":operations})

transport=[];payload=[];persistence=[]
for page,m in page_maps.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if st in {"EFFECTFUL_EXACT","READ_EXACT"} and not missing(op):
            if missing(r.get("method_path","")):transport.append((page,uid,st,op))
            if st=="EFFECTFUL_EXACT" and missing(r.get("payload_schema","")):payload.append((page,uid,op))
            if st=="EFFECTFUL_EXACT" and missing(r.get("persistence_owner","")):persistence.append((page,uid,op))

summary={
"root_docx":31,"page_contracts":18,"unique_page_control_uids":len(uid_pages),
"cross_page_uid_collisions":sum(1 for f in findings if f["class"]=="CONTROL_UID_CROSS_PAGE_COLLISION"),
"page_internal_field_conflicts":sum(1 for f in findings if f["class"]=="PAGE_INTERNAL_FIELD_CONFLICT"),
"same_page_operation_method_conflicts":sum(1 for f in findings if f["class"]=="SAME_PAGE_OPERATION_METHOD_CONFLICT"),
"same_page_operation_owner_conflicts":sum(1 for f in findings if f["class"]=="SAME_PAGE_OPERATION_OWNER_CONFLICT"),
"same_page_route_operation_collisions":sum(1 for f in findings if f["class"]=="SAME_PAGE_ROUTE_OPERATION_COLLISION"),
"corrected_blocker_count":len(findings),
"transport_unbound":len(transport),"effectful_payload_unbound":len(payload),"effectful_persistence_owner_unbound":len(persistence),
"iam_conflict_closed":1,"global_operation_name_false_blockers_reclassified":13,"local_local_false_positive_reclassified":1,
}
assert summary["cross_page_uid_collisions"]==0,summary
assert summary["page_internal_field_conflicts"]==0,summary
assert summary["same_page_operation_method_conflicts"]==0,summary
assert summary["same_page_operation_owner_conflicts"]==0,summary
assert summary["same_page_route_operation_collisions"]==0,summary
assert summary["corrected_blocker_count"]==0,summary
assert (summary["transport_unbound"],summary["effectful_payload_unbound"],summary["effectful_persistence_owner_unbound"])==(80,137,124),summary

add_table(logic,["Corrected Fresh Audit","Count","Status"],[
["Cross-page Control UID collision",summary["cross_page_uid_collisions"],"PASS"],
["Same-page field conflict",summary["page_internal_field_conflicts"],"PASS"],
["Same-page Operation→Method conflict",summary["same_page_operation_method_conflicts"],"PASS"],
["Same-page Operation→Runtime Owner conflict",summary["same_page_operation_owner_conflicts"],"PASS"],
["Same-page real route collision",summary["same_page_route_operation_collisions"],"PASS"],
["Corrected blocker count",summary["corrected_blocker_count"],"0"],
["Transport binding inventory",summary["transport_unbound"],"NEXT_DENOMINATOR / NOT AUTO-BLOCKER"],
["Effectful Payload/Schema inventory",summary["effectful_payload_unbound"],"NEXT_DENOMINATOR / APPLICABILITY REQUIRED"],
["Effectful Persistence Owner inventory",summary["effectful_persistence_owner_unbound"],"NEXT_DENOMINATOR / APPLICABILITY REQUIRED"],
],4.2)
machine={
"marker":MARK,"base_head":BASE_HEAD,"iam_same_uid_permission_conflict_closed":1,
"global_operation_name_false_blockers_reclassified":13,"local_local_false_positive_reclassified":1,
"corrected_blocker_count":0,"transport_unbound":80,"effectful_payload_unbound":137,
"effectful_persistence_owner_unbound":124,"runtime_execution_claimed":False,
}
logic.add_paragraph("BATCH20_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

changed=[S09,IAM,LOGIC]
report={"machine":machine,"summary":summary,"findings":findings,"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed},"iam_pre_values":pre_vals,"iam_post_values":post_vals}
Path("__batch20_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH20="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
