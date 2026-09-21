from pathlib import Path
from docx import Document
import json,re,collections

SYSTEMS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
WB="ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"

def norm(x): return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v): return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None

FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]

def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows: continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(h,"control uid")
        if ci is None: continue
        idx={
          "control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate"),
          "permission":hfind(h,"permission","auth resource"),
          "operation":hfind(h,"operation"),"runtime_owner":hfind(h,"runtime owner"),
          "runtime_status":hfind(h,"runtime status")
        }
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}: continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"source":path,"table":ti+1,"row":ri})
            out.append(rec)
    return out

# WB exact 14 controls: Action and Gate both missing, anchors present.
wb_rows=parse_controls(WB)
groups=collections.defaultdict(list)
for r in wb_rows:
    if missing(r["action"]) and missing(r["gate"]) and not missing(r["permission"]) and not missing(r["operation"]) and not missing(r["runtime_owner"]):
        groups[r["control"]].append(r)
# Filter the Batch-08 owner-absent set: compact/synthetic dashboard read controls.
candidates=[]
for uid,rs in groups.items():
    # choose richest row
    r=max(rs,key=lambda z:sum(not missing(z.get(f,"")) for f in FIELDS[1:]))
    if r["permission"]=="workspace.dashboard.view" and r["operation"]=="getDashboardReadModel" and r["runtime_owner"]=="DASHBOARD_READ_MODEL":
        candidates.append(r)
assert len(candidates)==14,(len(candidates),[(r["control"],r["permission"],r["operation"],r["runtime_owner"]) for r in candidates])

# Generic rows + document authority text.
system_rows=[]
doc_text={}
for sf in SYSTEMS:
    d=Document(sf)
    texts=[norm(p.text) for p in d.paragraphs if norm(p.text)]
    generic=[]
    for ti,t in enumerate(d.tables):
        headers=[norm(c.text) for c in t.rows[0].cells] if t.rows else []
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if any(vals):
                generic.append({"source":sf,"table":ti+1,"row":ri,"headers":headers,"values":vals,"joined":" | ".join(vals)})
                system_rows.append(generic[-1])
    doc_text[sf]={"paragraphs":texts,"rows":generic}

def exact_anchor_hits(q):
    p,o,ro=q["permission"],q["operation"],q["runtime_owner"]
    triple=[];op_owner=[];op=[];owner=[];perm=[]
    for r in system_rows:
        vals=r["values"]
        hs=[h.lower() for h in r["headers"]]
        # exact value equality somewhere in row; strong match is all anchors in same row.
        if p in vals and o in vals and ro in vals: triple.append(r)
        if o in vals and ro in vals: op_owner.append(r)
        if o in vals: op.append(r)
        if ro in vals: owner.append(r)
        if p in vals: perm.append(r)
    return triple,op_owner,op,owner,perm

def by_source(rows):
    c=collections.Counter(r["source"] for r in rows)
    return dict(c)

def authority_context(sf,anchors):
    pats=[
      r"canonical owner",r"canonical.*authority",r"source of truth",r"authoritative",
      r"dashboard",r"read model",r"projection",r"workspace"
    ]
    out=[]
    for i,p in enumerate(doc_text[sf]["paragraphs"],1):
        if any(a and a in p for a in anchors) or any(re.search(pt,p,re.I) for pt in pats):
            out.append({"kind":"paragraph","index":i,"text":p})
    for rr in doc_text[sf]["rows"]:
        t=rr["joined"]
        if any(a and a in t for a in anchors):
            out.append({"kind":"table","table":rr["table"],"row":rr["row"],"text":t})
    return out[:60]

results=[]
for q in sorted(candidates,key=lambda x:x["control"]):
    triple,op_owner,op,owner,perm=exact_anchor_hits(q)
    sources_union=sorted(set(r["source"] for r in triple+op_owner+op+owner+perm))
    results.append({
      "control":q["control"],"label":q["label"],"type":q["type"],
      "permission":q["permission"],"operation":q["operation"],"runtime_owner":q["runtime_owner"],"runtime_status":q["runtime_status"],
      "triple_by_source":by_source(triple),
      "operation_owner_by_source":by_source(op_owner),
      "operation_by_source":by_source(op),
      "runtime_owner_by_source":by_source(owner),
      "permission_by_source":by_source(perm),
      "contexts":{sf:authority_context(sf,[q["permission"],q["operation"],q["runtime_owner"]]) for sf in sources_union}
    })

# Since all 14 share the same anchor triple, verify identical source profiles.
profiles=[(tuple(sorted(r["triple_by_source"].items())),tuple(sorted(r["operation_owner_by_source"].items())),tuple(sorted(r["operation_by_source"].items())),tuple(sorted(r["runtime_owner_by_source"].items())),tuple(sorted(r["permission_by_source"].items()))) for r in results]
assert len(set(profiles))==1,len(set(profiles))
summary={
 "denominator_controls":14,
 "unresolved_fields":28,
 "anchor":{"permission":"workspace.dashboard.view","operation":"getDashboardReadModel","runtime_owner":"DASHBOARD_READ_MODEL"},
 "source_profile":{
   "triple_by_source":results[0]["triple_by_source"],
   "operation_owner_by_source":results[0]["operation_owner_by_source"],
   "operation_by_source":results[0]["operation_by_source"],
   "runtime_owner_by_source":results[0]["runtime_owner_by_source"],
   "permission_by_source":results[0]["permission_by_source"],
 }
}
Path("__batch11_wb_owner_mapping_audit.json").write_text(json.dumps({"summary":summary,"rows":results},ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH11_WB_AUDIT="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
print("BATCH11_CONTEXT_BEGIN")
# Print one representative control with all authority contexts because the 14 anchor profiles are identical.
print(json.dumps(results[0],ensure_ascii=False,sort_keys=True))
print("BATCH11_CONTEXT_END")
