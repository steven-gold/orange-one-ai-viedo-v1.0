from pathlib import Path
from docx import Document
import json,re,collections,subprocess

BASE_HEAD="0440a385d21d29833846104fcad01197531fc22d"
TARGET_UID="EDIT-01-BTN-FLOW-START"
OPS=["createEditingRuntimeRun","completeAssembly"]
KEY_FIELDS={"operation","method / path","method","path","payload / schema","payload","schema","persistence owner","runtime owner","action uid","gate uid","permission"}
DOCS=sorted(Path(".").glob("*.docx"))
assert len(DOCS)==31
assert subprocess.check_output(["git","-c","core.quotePath=false","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def header_index(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None

edit_evidence=[]
source_not_defined_rows=[]
placeholder_context=[]
legacy_context=[]
operation_evidence=[]

for path in DOCS:
    d=Document(path)
    # paragraphs
    for i,p in enumerate(d.paragraphs):
        t=norm(p.text)
        if not t:continue
        if TARGET_UID in t or any(op in t for op in OPS):
            edit_evidence.append({"file":path.name,"kind":"paragraph","index":i,"text":t[:5000]})
        if re.search(r"\bPLACEHOLDER\b|\bTODO\b|\bTBD\b|\bFIXME\b",t,re.I):
            placeholder_context.append({"file":path.name,"kind":"paragraph","index":i,"text":t[:2500]})
        if re.search(r"\bR9(?:\.0\.1)?\b|\bR5\b",t,re.I):
            legacy_context.append({"file":path.name,"kind":"paragraph","index":i,"text":t[:2500]})
    # tables
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        ci=header_index(headers,"control uid","target uid")
        ri=header_index(headers,"runtime status")
        oi=header_index(headers,"operation")
        for rno,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            joined=" | ".join(vals)
            uid=vals[ci] if ci is not None and ci<len(vals) else ""
            runtime=vals[ri] if ri is not None and ri<len(vals) else ""
            op=vals[oi] if oi is not None and oi<len(vals) else ""
            if uid==TARGET_UID or TARGET_UID in joined or any(x in joined for x in OPS):
                edit_evidence.append({"file":path.name,"kind":"table_row","table":ti,"row":rno,"headers":headers,"values":vals,"uid":uid,"runtime_status":runtime,"operation":op})
            if any(x in joined for x in OPS):
                operation_evidence.append({"file":path.name,"table":ti,"row":rno,"headers":headers,"values":vals,"uid":uid,"runtime_status":runtime,"operation":op})
            if "SOURCE_NOT_DEFINED" in joined:
                locations=[]
                for idx,v in enumerate(vals):
                    if "SOURCE_NOT_DEFINED" in v:
                        locations.append({"column":headers[idx] if idx<len(headers) else str(idx),"value":v})
                status="NARRATIVE_OR_NONRUNTIME"
                if runtime in {"EFFECTFUL_EXACT","READ_EXACT"}:status="CURRENT_EXACT_RUNTIME_ROW"
                elif "BLOCKED" in runtime.upper():status="BLOCKED_RUNTIME_ROW"
                elif runtime:status="OTHER_RUNTIME_STATUS"
                source_not_defined_rows.append({"file":path.name,"table":ti,"row":rno,"uid":uid,"runtime_status":runtime,"operation":op,"classification":status,"locations":locations,"headers":headers,"values":vals})
            if re.search(r"\bPLACEHOLDER\b|\bTODO\b|\bTBD\b|\bFIXME\b",joined,re.I):
                placeholder_context.append({"file":path.name,"kind":"table_row","table":ti,"row":rno,"headers":headers,"values":vals})
            if re.search(r"\bR9(?:\.0\.1)?\b|\bR5\b",joined,re.I):
                legacy_context.append({"file":path.name,"kind":"table_row","table":ti,"row":rno,"headers":headers,"values":vals})

# Deduce edit target row semantics.
target_rows=[x for x in edit_evidence if x.get("kind")=="table_row" and x.get("uid")==TARGET_UID]
target_current=[]
for x in target_rows:
    if x.get("runtime_status") in {"EFFECTFUL_EXACT","READ_EXACT"}:
        target_current.append(x)

snd_by_class=collections.Counter(x["classification"] for x in source_not_defined_rows)
snd_exact_binding=[]
for x in source_not_defined_rows:
    if x["classification"]!="CURRENT_EXACT_RUNTIME_ROW":continue
    for loc in x["locations"]:
        if any(k in loc["column"].lower() for k in KEY_FIELDS):
            snd_exact_binding.append(x);break

# Legacy heuristic classification.
legacy_classified=[]
for x in legacy_context:
    txt=x.get("text","") if x["kind"]=="paragraph" else " | ".join(x["values"])
    if re.search(r"舊\s*R9|legacy.*R9|歷史.*R9|不可壓過 Current",txt,re.I):
        cls="EXPLICIT_LEGACY_WARNING"
    elif re.search(r"\bt\d+/r(?:5|9)\b",txt,re.I):
        cls="TABLE_ROW_REFERENCE_FALSE_POSITIVE"
    else:
        cls="REVIEW_CANDIDATE"
    legacy_classified.append(x|{"classification":cls})

# Placeholder heuristic classification.
placeholder_classified=[]
for x in placeholder_context:
    txt=x.get("text","") if x["kind"]=="paragraph" else " | ".join(x["values"])
    if re.search(r"OPERATION-PLACEHOLDER-REMEDIATION|Operation Placeholder Closure|canonical instruction placeholder|schema/ref placeholder|Localization / Placeholder|placeholder missing|placeholder;",txt,re.I):
        cls="NORMATIVE_OR_HISTORY_TERM"
    elif re.search(r"\bTODO\b",txt,re.I) and re.search(r"List todo|FOLLOW_UPS",txt,re.I):
        cls="SCHEMA_LITERAL_TERM"
    else:
        cls="REVIEW_CANDIDATE"
    placeholder_classified.append(x|{"classification":cls})

summary={
"target_uid":TARGET_UID,
"target_exact_rows":len(target_current),
"target_all_rows":len(target_rows),
"operation_evidence_rows":len(operation_evidence),
"source_not_defined_rows":len(source_not_defined_rows),
"source_not_defined_by_class":dict(snd_by_class),
"source_not_defined_exact_binding_rows":len(snd_exact_binding),
"placeholder_hits":len(placeholder_classified),
"placeholder_review_candidates":sum(x["classification"]=="REVIEW_CANDIDATE" for x in placeholder_classified),
"legacy_hits":len(legacy_classified),
"legacy_review_candidates":sum(x["classification"]=="REVIEW_CANDIDATE" for x in legacy_classified),
}
report={"marker":"ACPOS-20260922-BATCH-59-BATCH58-FINDING-CLASSIFICATION-AUDIT-V1","base_head":BASE_HEAD,"summary":summary,
"edit_target_exact_rows":target_current,"edit_target_all_rows":target_rows,"operation_evidence":operation_evidence,
"source_not_defined_rows":source_not_defined_rows,"source_not_defined_exact_binding_rows":snd_exact_binding,
"placeholder_classification":placeholder_classified,"legacy_classification":legacy_classified}
Path("__batch59_finding_classification.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
Path("__batch59_summary.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH59_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
