from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-51-SYSTEM04-CORE-CONTRACT-REMEDIATION-V1"
BASE_HEAD="7770fb3cd0466d14b971a3354d9bd35f1a45cd15"
S04="04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
CORE="ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S04:"f9523a78a8ca212b143166b7bea932be76dce446",
CORE:"31f0f25b6c7292ab5d1e98986708264b5842de29",
LOGIC:"017cb2969b8cd197c58f3a38a33cebb3c2e6813e",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":CORE,
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

TARGETS={
"CORE-01-BTN-PROJECT-CREATE":{
 "operation":"createProjectDraft","route":"POST /v1/core/projects/drafts","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"CreateProjectDraftRequest","persistence":"System 04 / Project + ProjectDraftVersion + Audit",
 "basis":"Creates a governed editable ProjectDraftVersion under one Project identity. Raw AI output is not canonical authority; idempotency/version/audit identity are required."
},
"CORE-01-BTN-TOPIC-CREATE":{
 "operation":"createTopicDraft","route":"POST /v1/core/projects/{projectId}/topics/drafts","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"CreateTopicDraftRequest","persistence":"System 04 / Topic + TopicProductionScope Draft + Audit",
 "basis":"Creates a production-scoped Topic draft only from valid Project lineage. Topic cannot recreate Project Story/World/Base DNA/Project Blueprint authority."
},
"CORE-01-FLD-EVALUATION":{
 "operation":"evaluateCoreCandidate","route":"POST /v1/core/candidates/{candidateId}/evaluation","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"EvaluateCoreCandidateRequest","persistence":"System 04 / CandidateVersion + Evaluation Evidence + QualityCriteriaRef",
 "basis":"Formal evaluation requires exact approved criteria and evidence. Score/AI consensus cannot auto-adopt, auto-approve or auto-lock a candidate."
},
"CORE-01-FLD-HUMAN-DECISION":{
 "operation":"recordCoreHumanDecision","route":"POST /v1/core/candidates/{candidateId}/human-decisions","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"RecordCoreHumanDecisionRequest","persistence":"System 04 / CandidateVersion + Human Decision Evidence/Audit",
 "basis":"Records explicit authorized human decision/rationale bound to exact candidate/version/evidence. Human decision is distinct from later lock approval."
},
"CORE-01-BTN-CANDIDATE-CREATE":{
 "operation":"createCoreCandidate","route":"POST /v1/core/candidates","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"CreateCoreCandidateRequest","persistence":"System 04 / CandidateVersion + Source/Context Lineage",
 "basis":"CandidateVersion is immutable and non-authoritative until the governed human decision/adoption path completes."
},
"CORE-01-BTN-CANDIDATE-CONFIRM":{
 "operation":"acceptCoreCandidate","route":"POST /v1/core/candidates/{candidateId}/accept","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"AcceptCoreCandidateRequest","persistence":"System 04 / CandidateVersion + Adopted Canon Version Lineage + Decision Evidence",
 "basis":"Acceptance requires exact human decision evidence and creates/binds a new canonical version lineage; prior candidate/version evidence remains immutable."
},
"CORE-01-BTN-RETURN-MODIFY":{
 "operation":"requestCoreCandidateRevision","route":"POST /v1/core/candidates/{candidateId}/revision-requests","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"RequestCoreCandidateRevisionRequest","persistence":"System 04 / CandidateVersion Revision Request + Audit",
 "basis":"Revision request never overwrites the candidate in place. Any material correction produces a new candidate/version with predecessor lineage."
},
"CORE-01-BTN-PROJECT-VALIDATE":{
 "operation":"validateProjectDraft","route":"POST /v1/core/projects/drafts/{draftId}/validate","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"ValidateProjectDraftRequest","persistence":"System 04 / ProjectDraftVersion + Validation Evidence",
 "basis":"Validates the exact draft/version against required structured contracts; validation evidence does not confirm or lock the Project."
},
"CORE-01-BTN-PROJECT-CONFIRM":{
 "operation":"confirmProjectDraft","route":"POST /v1/core/projects/drafts/{draftId}/confirm","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"ConfirmProjectDraftRequest","persistence":"System 04 / Project + ProjectDraftVersion Confirmed State + Audit",
 "basis":"Confirmation binds exact draft/version and human evidence into Project lineage without silently overwriting prior versions."
},
"CORE-01-BTN-STORY-CANDIDATE":{
 "operation":"createStoryCandidateSet","route":"POST /v1/core/projects/{projectId}/story/candidate-sets","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"CreateStoryCandidateSetRequest","persistence":"System 04 / StoryCandidateSet + CandidateVersion + Audit",
 "basis":"Creates immutable Story alternatives for human comparison/decision. Raw model output cannot directly write Story canonical authority."
},
"CORE-01-BTN-DNA-LOCK":{
 "operation":"requestDNALock","route":"POST /v1/state-commands/dna/requestdnalock","route_basis":"OWNER_DEFINED_REGISTERED_STATE_COMMAND",
 "payload":"RequestDNALockRequest","persistence":"VersionLockService / LockReviewRef + DNAResolution Target/Audit",
 "basis":"Materializes the required registered state-command route. DNA candidate, resolution and lock review remain separate; a request never implies approval."
},
"CORE-01-BTN-CORE-REVIEW":{
 "operation":"submitCoreReview","route":"POST /v1/core/projects/{projectId}/core-review-requests","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"SubmitCoreReviewRequest","persistence":"System 04 / ProjectBlueprintPackage Review Request + QualityCriteriaRef + Audit",
 "basis":"Submits exact Project/Blueprint/candidate/evidence refs for governed review. Missing approved criteria/evidence blocks; review submission is not approval."
},
"CORE-01-BTN-PROJECT-LOCK":{
 "operation":"requestMotherLock","route":"POST /v1/locks/mother-requests","route_basis":"SOURCE_EXACT_REUSED",
 "payload":"RequestMotherLockRequest","persistence":"VersionLockService / LockReviewRef + ProjectStructureVersion Target/Audit",
 "basis":"Reuses the exact Current endpoint. Request binds scope, expected version/hash, criteria/evidence and reviewer path; requester cannot decide own lock."
},
"CORE-01-BTN-BLUEPRINT-CREATE":{
 "operation":"createProjectBlueprintCandidate","route":"POST /v1/core/projects/{projectId}/blueprint-candidates","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"CreateProjectBlueprintCandidateRequest","persistence":"System 04 / ProjectBlueprintPackage Candidate + CandidateVersion/Audit",
 "basis":"Project Blueprint belongs only to Project Core. Candidate construction binds exact upstream Canon/DNA/Chapter/Rights/Quality/dependency refs and remains non-authoritative."
},
"CORE-01-BTN-BLUEPRINT-VALIDATE":{
 "operation":"validateProjectBlueprint","route":"POST /v1/core/projects/{projectId}/blueprint-candidates/{candidateId}/validate","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"ValidateProjectBlueprintRequest","persistence":"System 04 / ProjectBlueprintPackage Validation Evidence + QualityCriteriaRef",
 "basis":"Validation checks exact required upstream refs and approved criteria. Any missing/stale required ref fails closed."
},
"CORE-01-BTN-BLUEPRINT-APPROVE":{
 "operation":"approveProjectBlueprint","route":"POST /v1/core/projects/{projectId}/blueprint-candidates/{candidateId}/approve","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":"ApproveProjectBlueprintRequest","persistence":"System 04 / ProjectBlueprintPackage Approved Version + Human Decision/Audit",
 "basis":"Approval requires authorized human/reviewer evidence and creates/binds immutable approved version lineage; AI cannot self-approve."
},
"CORE-01-BTN-CHILD-LOCK":{
 "operation":"requestChildLock","route":"POST /v1/locks/child-requests","route_basis":"SOURCE_EXACT_REUSED",
 "payload":"RequestChildLockRequest","persistence":"VersionLockService / LockReviewRef + TopicBlueprintVersion Target/Audit",
 "basis":"Reuses the exact Current endpoint. Child Lock binds exact Topic Blueprint Version/hash and required review evidence; request is not approval."
},
"CORE-01-BTN-CANONICAL-SCRIPT":{
 "operation":"getCanonicalScript","route":"GET /v1/core/canonical-script-versions/{canonicalScriptVersionId}","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":None,"persistence":None,
 "basis":"READ_EXACT lookup of one provider-neutral canonical script version. No provider-specific prompt syntax is persisted or synthesized by this read."
},
"CORE-01-BTN-CANDIDATE-COMPARE":{
 "operation":"compareCoreCandidates","route":"GET /v1/core/candidates/compare","route_basis":"OWNER_DEFINED_CURRENT_CONTRACT",
 "payload":None,"persistence":None,
 "basis":"READ_EXACT comparison over explicit candidate/version refs. All candidates and decision rationale remain inspectable; comparison cannot auto-adopt or auto-lock."
},
}

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out
def blob(path):
    b=Path(path).read_bytes();return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():
    got=blob(f);assert got==s,(f,got,s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""
assert len(TARGETS)==19
assert sum(1 for s in TARGETS.values() if s["payload"] is not None)==17
assert sum(1 for s in TARGETS.values() if s["persistence"] is not None)==17
assert sum(1 for s in TARGETS.values() if s["route_basis"]=="SOURCE_EXACT_REUSED")==2

def collect_uid_operations(path):
    d=Document(path);out=collections.defaultdict(set)
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ui=hfind(h,"control uid","target uid")
        oi=hfind(h,"operation")
        if ui is None or oi is None:continue
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ui>=len(vals) or oi>=len(vals):continue
            uid,op=vals[ui],vals[oi]
            if uid and uid not in {"—","-"} and op and op not in {"—","-"}:
                out[uid].add(op)
    return out

pre_page=compose(parse_controls(CORE));pre_owner=compose(parse_controls(S04))
owner_uid_ops=collect_uid_operations(S04)
for uid,s in TARGETS.items():
    assert uid in pre_page,("PAGE",uid,"MISSING")
    assert uid in pre_owner,("OWNER",uid,"MISSING")
    pr=pre_page[uid];orr=pre_owner[uid]
    assert pr["operation"]==s["operation"],("PAGE",uid,pr["operation"],s["operation"])
    assert s["operation"] in owner_uid_ops.get(uid,set()),("OWNER_OPERATION_AUTHORITY_MISSING",uid,s["operation"],sorted(owner_uid_ops.get(uid,set())))
    for r,label in [(pr,"PAGE"),(orr,"OWNER")]:
        assert missing(r["method_path"]),(label,uid,"METHOD_ALREADY_DEFINED",r["method_path"])
        if s["payload"] is not None:
            assert r["runtime_status"]=="EFFECTFUL_EXACT",(label,uid,r["runtime_status"])
            assert missing(r["payload_schema"]),(label,uid,"PAYLOAD_ALREADY_DEFINED",r["payload_schema"])
            assert missing(r["persistence_owner"]),(label,uid,"PERSIST_ALREADY_DEFINED",r["persistence_owner"])
        else:
            assert r["runtime_status"]=="READ_EXACT",(label,uid,r["runtime_status"])

def patch(path):
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact_index(h,"Control UID")
        if ci is None:continue
        idx={"method_path":hfind(h,"method / path","method","path"),
             "payload_schema":hfind(h,"payload / schema","payload","schema"),
             "persistence_owner":hfind(h,"persistence owner")}
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals):continue
            uid=vals[ci]
            if uid not in TARGETS:continue
            s=TARGETS[uid]
            assignments={"method_path":s["route"]}
            if s["payload"] is not None:
                assignments["payload_schema"]=s["payload"]
                assignments["persistence_owner"]=s["persistence"]
            for field,new in assignments.items():
                fi=idx[field]
                if fi is None or fi>=len(row.cells):continue
                old=norm(row.cells[fi].text)
                if old==new:hits[(uid,field)]+=1;continue
                assert missing(old),(path,uid,field,old,new)
                row.cells[fi].text=new;hits[(uid,field)]+=1
    expected=set()
    for uid,s in TARGETS.items():
        expected.add((uid,"method_path"))
        if s["payload"] is not None:
            expected.add((uid,"payload_schema"));expected.add((uid,"persistence_owner"))
    assert set(hits)==expected,(path,"PATCH_MISMATCH",sorted(expected-set(hits)),dict(hits))
    d.save(path);Document(path)
    return {f"{u}:{f}":n for (u,f),n in hits.items()}

patches={"SYSTEM04_OWNER":patch(S04),"CORE_PAGE":patch(CORE)}

def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.33);sec.bottom_margin=Inches(.33);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=3.45):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for rr in rows:
        cells=t.add_row().cells
        for i,v in enumerate(rr):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
    return t

for path,title,prefix in [
(S04,"Batch 51 · System 04 CORE Canonical Contract Registry","SYSTEM04"),
(CORE,"Batch 51 · CORE-01 Canonical Contract Binding Ledger","PAGE::CORE-01"),
]:
    d=Document(path);add_landscape(d);d.add_heading(title,level=1)
    p=d.add_paragraph();p.add_run(f"[{MARK}::{prefix}] ").bold=True
    p.add_run("This registry closes the CORE Method/Path, Payload/Schema and Persistence Owner gaps at the design-authority layer. Two exact lock request routes are reused; seventeen previously undefined routes are now formally defined by System 04. No route definition is runtime/deployment evidence. Candidate≠Authority, request≠approval, AI self-approval is forbidden, immutable version lineage is mandatory, and missing exact refs/evidence fail closed.")
    rr=[]
    for uid,s in TARGETS.items():
        rr.append([uid,s["operation"],s["route"],s["route_basis"],s["payload"] or "N/A_READ",s["persistence"] or "N/A_READ_ONLY",s["basis"]])
    add_table(d,["Control UID","Operation","Method / Path","Route Authority","Payload / Schema","Persistence Owner","Required Semantics"],rr)
    d.save(path);Document(path)

post_page=compose(parse_controls(CORE));post_owner=compose(parse_controls(S04))
for uid,s in TARGETS.items():
    for m,label in [(post_page,"PAGE"),(post_owner,"OWNER")]:
        r=m[uid]
        assert r["method_path"]==s["route"],(label,uid,"METHOD",r["method_path"],s["route"])
        if s["payload"] is not None:
            assert r["payload_schema"]==s["payload"],(label,uid,"PAYLOAD",r["payload_schema"],s["payload"])
            assert r["persistence_owner"]==s["persistence"],(label,uid,"PERSIST",r["persistence_owner"],s["persistence"])

# Route collision check within CORE page.
route_ops=collections.defaultdict(set)
for uid,r in post_page.items():
    mp=r.get("method_path","");op=r.get("operation","")
    if missing(mp) or mp in {"LOCAL LOCAL","NO_PUBLIC_API_BY_AUTHORITY"}:continue
    route_ops[mp].add(op)
collisions={k:sorted(v) for k,v in route_ops.items() if len(v)>1}
assert not collisions,collisions

all_maps={p:compose(parse_controls(fn)) for p,fn in PAGES.items()}
remaining=[]
for page,m in all_maps.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if missing(op) or st not in {"EFFECTFUL_EXACT","READ_EXACT"}:continue
        if missing(r.get("method_path","")):remaining.append(("method_path",page,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("payload_schema","")):remaining.append(("payload_schema",page,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("persistence_owner","")):remaining.append(("persistence_owner",page,uid))
rc=collections.Counter(f for f,_,_ in remaining)
assert len(remaining)==137,(len(remaining),rc)
assert rc=={"method_path":43,"payload_schema":56,"persistence_owner":38},rc

logic=Document(LOGIC);add_landscape(logic);logic.add_heading("Batch 51 · System 04 CORE Contract Closure",level=1)
p=logic.add_paragraph();p.add_run(f"[{MARK}] ").bold=True
p.add_run("System 04 closes all 53 CORE-01 contract cells. Mother/Child lock request routes reuse exact Current authority; the other seventeen routes are formal System 04 design-contract definitions, not claims of deployed APIs. Candidate/decision/review/lock boundaries remain separate and all canonical versions preserve immutable lineage.")
add_table(logic,["Metric","Value","Result"],[
["Pre contract denominator",190,"Method 62 + Payload 73 + Persistence 55"],
["System 04 cells closed",53,"Method 19 + Payload 17 + Persistence 17"],
["Exact Current lock routes reused",2,"Mother + Child"],
["New Current design route definitions",17,"System 04 owner authority only"],
["Post contract denominator",137,"Method 43 + Payload 56 + Persistence 38"],
["Same-page Method/Path collisions",0,"PASS"],
["AI self-approval introduced",0,"Forbidden"],
["Runtime / API deployment claimed","False","Design-contract closure only"],
],4.0)
machine={
"marker":MARK,"base_head":BASE_HEAD,"pre_contract_cells":190,"system04_cells_closed":53,
"method_path_closed":19,"payload_schema_closed":17,"persistence_owner_closed":17,
"exact_lock_routes_reused":2,"new_current_design_route_definitions":17,
"post_contract_cells":137,"remaining_method_path":43,"remaining_payload_schema":56,"remaining_persistence_owner":38,
"same_page_route_collisions":0,"ai_self_approval_introduced":0,"runtime_execution_claimed":False,
}
logic.add_paragraph("BATCH51_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

changed=[S04,CORE,LOGIC]
report={"machine":machine,"targets":TARGETS,"patches":patches,"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed}}
Path("__batch51_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH51="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
