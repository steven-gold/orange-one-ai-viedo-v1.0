from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-53-EDIT-CONTRACT-REMEDIATION-V1"
BASE_HEAD="3eab30d574f9401402bf179d380aa46eb896c2f2"
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
EDIT="ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S05:"82244e2c0465df58fd5b46eae043c704f0f19bc2",
EDIT:"5f48109884fe062c9cd60f116212b4dc53f0d67c",
LOGIC:"4f940b4dc3e2a00d926ab8839b472f0f739bd01d",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":EDIT,
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

EFFECTFUL={
"EDIT-01-BTN-FLOW-START":{
 "operation":"createEditingRuntimeRun | completeAssembly",
 "payload":"CreateEditingRuntimeRunRequest | CompleteEditingAssemblyRequest",
 "persistence":"acpos_runtime.editing_runtime_runs_runtime; acpos_runtime.editing_runtime_steps_runtime; public.editing_timelines",
 "basis":"Start/continue creates the exact editing runtime run and completes assembly only against the bound input fingerprint/source versions. Runtime step lineage and timeline version remain reconstructable."
},
"EDIT-01-BTN-HANDOFF-2":{"operation":"handoffVoiceToQA","payload":"HandoffVoiceToQARequest","persistence":None,
 "basis":"Handoff packages the exact locked/final edit+voice refs for QA; existing public.handoffs owner remains unchanged."},
"EDIT-01-BTN-LIPSYNC-EXECUTE":{"operation":"completeLipSync","payload":"CompleteLipSyncRequest","persistence":None,
 "basis":"Lip-sync completion binds exact visual/voice segment versions and alignment evidence; existing editing runtime persistence remains unchanged."},
"EDIT-01-BTN-MIX-EXECUTE":{"operation":"completeAudioMix","payload":"CompleteAudioMixRequest","persistence":None,
 "basis":"Audio mix completion binds exact voice/music/SFX sources and mix version/evidence; existing editing runtime persistence remains unchanged."},
"EDIT-01-BTN-OUTPUT-SAVE":{"operation":"saveEditOutputVersion","payload":"SaveEditOutputVersionRequest","persistence":None,
 "basis":"Saves one immutable output version from the exact render job; existing render/task-output owners remain unchanged."},
"EDIT-01-BTN-RENDER-CANCEL":{"operation":"cancelEditRender","payload":"CancelEditRenderRequest","persistence":None,
 "basis":"Cancels only the exact render job if cancellation remains legal; historical attempts/evidence are preserved."},
"EDIT-01-BTN-RENDER-EXECUTE":{"operation":"startEditRender","payload":"StartEditRenderRequest","persistence":None,
 "basis":"Starts a render from exact editing timeline/source versions with idempotency/correlation identity; job creation is not render completion."},
"EDIT-01-BTN-STAGE-CONFIRM":{"operation":"decideOutputCandidate","payload":"DecideOutputCandidateRequest","persistence":None,
 "basis":"Decision binds exact task/output version and authorized decision evidence; it cannot silently rewrite the output."},
"EDIT-01-BTN-SUB-API-SYNC":{"operation":"completeSubtitle","payload":"CompleteSubtitleRequest","persistence":None,
 "basis":"Subtitle sync/completion binds exact subtitle cue/version and dialogue-sync context; existing runtime-step persistence remains unchanged."},
"EDIT-01-BTN-VERSION-LOCK":{"operation":"lockEditVersion","payload":"LockEditVersionRequest","persistence":None,
 "basis":"Locks the exact edit version under the existing production output version lock owner. Lock request/decision evidence must remain auditable."},
"EDIT-01-BTN-VERSION-RESTORE":{"operation":"restoreEditVersionAsDraft","payload":"RestoreEditVersionAsDraftRequest","persistence":None,
 "basis":"Restore creates a new editable draft/timeline lineage from an exact prior version; no in-place overwrite of historical edit versions."},
"EDIT-01-BTN-VERSION-SAVE":{"operation":"saveEditVersion","payload":"SaveEditVersionRequest","persistence":None,
 "basis":"Saves a new immutable editing timeline/version snapshot from the exact working state; predecessor lineage remains intact."},
}

READS={
"EDIT-01-LBL-API-JOB":{"operation":"getEditCorrectionJobStatus","route":"GET /v1/editing-runtime-runs/{runId}/correction/job-status",
 "basis":"READ_EXACT projection of the correction job state for the current editing run."},
"EDIT-01-LBL-BINDING-FINGERPRINT":{"operation":"getEditCorrectionBindingFingerprint","route":"GET /v1/editing-runtime-runs/{runId}/correction/binding-fingerprint",
 "basis":"READ_EXACT projection of the exact input/source binding fingerprint; no mutation."},
"EDIT-01-LBL-CURRENT-SCRIPT-SECTION":{"operation":"getEditCorrectionScriptSection","route":"GET /v1/editing-runtime-runs/{runId}/correction/script-section",
 "basis":"READ_EXACT projection of the correction script section bound to the current run/version."},
"EDIT-01-LBL-LIPSYNC-SYNC-BINDING":{"operation":"getEditLipSyncBinding","route":"GET /v1/editing-runtime-runs/{runId}/dialogue-sync/lip-sync-binding",
 "basis":"READ_EXACT lip-sync binding projection over exact voice/visual refs."},
"EDIT-01-LBL-SUB-SYNC-BINDING":{"operation":"getEditSubtitleSyncBinding","route":"GET /v1/editing-runtime-runs/{runId}/dialogue-sync/subtitle-binding",
 "basis":"READ_EXACT subtitle synchronization binding projection."},
"EDIT-01-LST-IMPORT-QUEUE":{"operation":"getEditImportQueue","route":"GET /v1/editing-runtime-runs/{runId}/imports/queue",
 "basis":"READ_EXACT import queue projection for the current editing run; display does not create queue work."},
"EDIT-01-LST-VOICE-TAKES":{"operation":"getEditVoiceTakes","route":"GET /v1/editing-runtime-runs/{runId}/voice/takes",
 "basis":"READ_EXACT voice-take/source projection with exact version refs."},
"EDIT-01-PNL-API-CANDIDATE":{"operation":"getEditCorrectionCandidatePreview","route":"GET /v1/editing-runtime-runs/{runId}/correction/candidate-preview",
 "basis":"READ_EXACT correction candidate preview. Preview action remains legitimately N/A and cannot apply the correction."},
"EDIT-01-PNL-FINAL-PREVIEW":{"operation":"getEditFinalPreview","route":"GET /v1/editing-runtime-runs/{runId}/preview/final",
 "basis":"READ_EXACT final preview projection. Preview action remains legitimately N/A and cannot finalize/lock the output."},
}

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),"permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),"operation":hfind(h,"operation"),
          "method_path":hfind(h,"method / path","method","path"),"runtime_owner":hfind(h,"runtime owner"),
          "persistence_owner":hfind(h,"persistence owner"),"runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out
def collect_uid_operations(path):
    d=Document(path);out=collections.defaultdict(set)
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ui=hfind(h,"control uid","target uid");oi=hfind(h,"operation")
        if ui is None or oi is None:continue
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ui>=len(vals) or oi>=len(vals):continue
            uid,op=vals[ui],vals[oi]
            if uid and uid not in {"—","-"} and op and op not in {"—","-"}:out[uid].add(op)
    return out
def blob(path):
    b=Path(path).read_bytes();return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():assert blob(f)==s,(f,blob(f),s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""
assert len(EFFECTFUL)==12 and len(READS)==9

pre_page=compose(parse_controls(EDIT));pre_owner=compose(parse_controls(S05));owner_ops=collect_uid_operations(S05)
for uid,s in EFFECTFUL.items():
    for m,label in [(pre_page,"PAGE"),(pre_owner,"OWNER")]:
        assert uid in m,(label,uid)
        r=m[uid]
        if label=="PAGE": assert r["operation"]==s["operation"],(uid,r["operation"],s["operation"])
        else: assert s["operation"] in owner_ops.get(uid,set()),("OWNER_OPERATION_MISSING",uid,s["operation"],sorted(owner_ops.get(uid,set())))
        assert r["runtime_status"]=="EFFECTFUL_EXACT",(label,uid,r["runtime_status"])
        assert missing(r["payload_schema"]),(label,uid,"PAYLOAD_ALREADY",r["payload_schema"])
    if s["persistence"] is not None:
        assert missing(pre_page[uid]["persistence_owner"]) and missing(pre_owner[uid]["persistence_owner"])
for uid,s in READS.items():
    assert pre_page[uid]["operation"]==s["operation"],(uid,pre_page[uid]["operation"],s["operation"])
    assert s["operation"] in owner_ops.get(uid,set()),("OWNER_OPERATION_MISSING",uid,s["operation"])
    assert pre_page[uid]["runtime_status"]=="READ_EXACT"
    assert missing(pre_page[uid]["method_path"]) and missing(pre_owner[uid]["method_path"])

def patch(path):
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact_index(h,"Control UID")
        if ci is None:continue
        idx={"method_path":hfind(h,"method / path","method","path"),"payload_schema":hfind(h,"payload / schema","payload","schema"),
             "persistence_owner":hfind(h,"persistence owner")}
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals):continue
            uid=vals[ci]
            assignments={}
            if uid in EFFECTFUL:
                assignments["payload_schema"]=EFFECTFUL[uid]["payload"]
                if EFFECTFUL[uid]["persistence"] is not None:assignments["persistence_owner"]=EFFECTFUL[uid]["persistence"]
            elif uid in READS:
                assignments["method_path"]=READS[uid]["route"]
            else:continue
            for f,new in assignments.items():
                i=idx[f]
                if i is None or i>=len(row.cells):continue
                old=norm(row.cells[i].text)
                if old==new:hits[(uid,f)]+=1;continue
                assert missing(old),(path,uid,f,old,new)
                row.cells[i].text=new;hits[(uid,f)]+=1
    expected={(u,"payload_schema") for u in EFFECTFUL}|{(u,"method_path") for u in READS}|{(u,"persistence_owner") for u,s in EFFECTFUL.items() if s["persistence"] is not None}
    assert set(hits)==expected,(path,sorted(expected-set(hits)),dict(hits))
    d.save(path);Document(path)
    return {f"{u}:{f}":n for (u,f),n in hits.items()}

patches={"SYSTEM05_OWNER":patch(S05),"EDIT_PAGE":patch(EDIT)}

def landscape(d):
    s=d.add_section(WD_SECTION.NEW_PAGE);s.orientation=WD_ORIENT.LANDSCAPE;s.page_width,s.page_height=s.page_height,s.page_width
    s.top_margin=Inches(.33);s.bottom_margin=Inches(.33);s.left_margin=Inches(.28);s.right_margin=Inches(.28)
def table(d,h,rows,fs=3.55):
    t=d.add_table(rows=1,cols=len(h));t.style="Table Grid"
    tr=t.rows[0]._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");tr.append(e)
    for i,x in enumerate(h):
        t.rows[0].cells[i].text=str(x);sh=OxmlElement("w:shd");sh.set(qn("w:fill"),"EDE9FE");t.rows[0].cells[i]._tc.get_or_add_tcPr().append(sh)
    for rr in rows:
        c=t.add_row().cells
        for i,x in enumerate(rr):c[i].text=str(x)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
for path,title,prefix in [(S05,"Batch 53 · EDIT Canonical Contract Registry","SYSTEM05"),(EDIT,"Batch 53 · EDIT-01 Contract Binding Ledger","PAGE::EDIT-01")]:
    d=Document(path);landscape(d);d.add_heading(title,1)
    d.add_paragraph(f"[{MARK}::{prefix}] Existing effectful EDIT routes and persistence owners are preserved. This batch defines missing payload schemas, one FLOW persistence binding, and nine read-only editing-runtime projection routes. Read projections do not mutate state; preview remains non-effectful.")
    rows=[]
    for u,s in EFFECTFUL.items(): rows.append([u,s["operation"],"UNCHANGED",s["payload"],s["persistence"] or "UNCHANGED",s["basis"]])
    for u,s in READS.items(): rows.append([u,s["operation"],s["route"],"N/A_READ","N/A_READ_ONLY",s["basis"]])
    table(d,["Control UID","Operation","Method / Path","Payload / Schema","Persistence Owner","Semantics"],rows)
    d.save(path);Document(path)

post_page=compose(parse_controls(EDIT));post_owner=compose(parse_controls(S05))
for uid,s in EFFECTFUL.items():
    assert post_page[uid]["payload_schema"]==s["payload"],(uid,post_page[uid]["payload_schema"])
    assert post_owner[uid]["payload_schema"]==s["payload"],(uid,post_owner[uid]["payload_schema"])
    if s["persistence"] is not None:
        assert post_page[uid]["persistence_owner"]==s["persistence"]
        assert post_owner[uid]["persistence_owner"]==s["persistence"]
for uid,s in READS.items():
    assert post_page[uid]["method_path"]==s["route"],(uid,post_page[uid]["method_path"])
    assert post_owner[uid]["method_path"]==s["route"],(uid,post_owner[uid]["method_path"])

allm={p:compose(parse_controls(fn)) for p,fn in PAGES.items()};rem=[]
for p,m in allm.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if missing(op) or st not in {"EFFECTFUL_EXACT","READ_EXACT"}:continue
        if missing(r.get("method_path","")):rem.append(("method_path",p,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("payload_schema","")):rem.append(("payload_schema",p,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("persistence_owner","")):rem.append(("persistence_owner",p,uid))
rc=collections.Counter(x[0] for x in rem);assert len(rem)==115,(len(rem),rc);assert rc=={"method_path":34,"payload_schema":44,"persistence_owner":37},rc

d=Document(LOGIC);landscape(d);d.add_heading("Batch 53 · EDIT Contract Closure",1)
machine={"marker":MARK,"base_head":BASE_HEAD,"pre_contract_cells":137,"edit_cells_closed":22,"method_path_closed":9,"payload_schema_closed":12,"persistence_owner_closed":1,
"post_contract_cells":115,"remaining_method_path":34,"remaining_payload_schema":44,"remaining_persistence_owner":37,
"existing_effectful_routes_changed":0,"read_routes_defined":9,"runtime_execution_claimed":False}
table(d,["Metric","Value"],[["Pre denominator","137"],["EDIT cells closed","22 = Method 9 + Payload 12 + Persistence 1"],["Post denominator","115 = Method 34 + Payload 44 + Persistence 37"],["Existing effectful routes changed","0"],["Runtime execution claimed","False"]],4.0)
d.add_paragraph("BATCH53_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")));d.save(LOGIC);Document(LOGIC)
changed=[S05,EDIT,LOGIC]
Path("__batch53_remediation_report.json").write_text(json.dumps({"machine":machine,"effectful":EFFECTFUL,"reads":READS,"patches":patches,"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed}},ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH53="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
