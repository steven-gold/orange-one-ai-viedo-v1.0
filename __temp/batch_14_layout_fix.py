from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
import hashlib,json

LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED="4392870d4eeec39ddd9ee2982f3d04f06c482619"
MARK="ACPOS-20260921-BATCH-14-LOCAL-ORCHESTRATION-NA-SEMANTICS-V1"

def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

def snapshot(doc):
    return "\n".join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])

assert blob(LOGIC)==EXPECTED,(blob(LOGIC),EXPECTED)
d=Document(LOGIC)
before=snapshot(d)

# Tighten final Batch-14 section only.
sec=d.sections[-1]
sec.top_margin=Inches(.20)
sec.bottom_margin=Inches(.18)
sec.left_margin=Inches(.24)
sec.right_margin=Inches(.24)

for p in d.paragraphs:
    if p.text.startswith("Batch 14 · Local / Aggregate Operation Applicability Closure"):
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(1)
        p.paragraph_format.line_spacing=.9
        for r in p.runs:r.font.size=Pt(15)
    elif p.text.startswith("["+MARK+"]"):
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(1)
        p.paragraph_format.line_spacing=.88
        for r in p.runs:r.font.size=Pt(7.2)
    elif p.text.startswith("BATCH14_MACHINE_JSON="):
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=.75
        for r in p.runs:r.font.size=Pt(3.0)

# The Batch-14 section appended the last four tables in System Logic.
for t in d.tables[-4:]:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0)
                p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=.85
                for r in p.runs:
                    if r.font.size is None or r.font.size.pt>3.5:
                        r.font.size=Pt(3.5)

d.save(LOGIC)
d2=Document(LOGIC)
assert snapshot(d2)==before,"TEXT_CHANGED_IN_LAYOUT_ONLY_FIX"
report={"layout_only":True,"text_unchanged":True,"marker":MARK,"logic_blob":blob(LOGIC)}
Path("__batch14_layout_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH14_LAYOUT="+json.dumps(report,ensure_ascii=False,sort_keys=True))
