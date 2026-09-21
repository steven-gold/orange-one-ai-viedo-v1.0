from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,os

MARK="ACPOS-20260921-BATCH-10-CANONICAL-OWNER-PRECEDENCE-V1"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"

S01="01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx"
S02="02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S03="03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S04="04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S06="06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S07="07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S08="08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S09="09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
SYSTEMS=[S01,S02,S03,S04,S05,S06,S07,S08,S09]

EXPECTED_SHA={
S02:"37a6e1e10a9f2ff1cd6b28e60634d0b59146e155",
S04:"ad51ec663d36ce9fe9d444dfc4d70030c97d2b2a",
S08:"78e94b76bbc3697e3375b49b4590e8724418cd58",
S09:"be6c0436b3682a4e8e17d1977ebf9501ec830451",
LOGIC:"d23f6a0e873546511665355a34bdbd2607e9f9a1",
}

PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}

# Pair precedence is not inferred from row richness. It is grounded in existing owner-boundary text:
# S03 owns AI Execution/AIAPI runtime; S09 resolves/inspects owners but forbids parallel owner.
# S02 is the unique Shared Conversation Core; domain systems consume it.
# S01 owns information lifecycle/projection/knowledge concerns represented by the DB/KB rows here.
# S06 owns DEV domain commands; S02 explicitly does not replace Domain Authority.
# S05 owns creative-production/QA concerns; S09 is governance inspection.
PRECEDENCE={
tuple(sorted((S03,S09))):S03,
tuple(sorted((S02,S04))):S02,
tuple(sorted((S01,S09))):S01,
tuple(sorted((S02,S06))):S06,
tuple(sorted((S05,S09))):S05,
tuple(sorted((S02,S08))):S02,
tuple(sorted((S02,S09))):S02,
}
AUTHORITY_BASIS={
tuple(sorted((S03,S09))):"System 03 declares AI Execution / Script Compiler / Tool / AIAPI Runtime as Canonical Owner; System 09 is owner-resolution/governance inspection and forbids parallel owner.",
tuple(sorted((S02,S04))):"System 02 declares the whole-site unique Shared Conversation Core and forbids a second conversation core; CORE is a Domain Extension/consumer for these conversation controls.",
tuple(sorted((S01,S09))):"System 01 owner scope covers whole-site state projection, context/information retrieval and knowledge/experience admission; System 09 resolves owner/duplicate findings and forbids parallel authority.",
tuple(sorted((S02,S06))):"System 02 explicitly does not replace Domain Authority; DEV discovery/email dispatch are domain commands, therefore System 06 is the canonical domain owner.",
tuple(sorted((S05,S09))):"System 05 owns creative-production/QA domain concerns; System 09 performs owner-resolution/governance inspection and forbids parallel owner.",
tuple(sorted((S02,S08))):"System 02 is the unique Shared Conversation Core; STR conversation view/attachment controls are conversation-core concerns consumed by Strategy.",
tuple(sorted((S02,S09))):"System 02 is the unique Shared Conversation Core; SYS attachment is a conversation-core concern while System 09 remains governance/engineering consumer.",
}

def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()
for f,s in EXPECTED_SHA.items():
    assert blob(f)==s,(f,blob(f),s)

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def display_only_type(t):
    u=(t or "").upper()
    return u in {"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"} or any(x in u for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]

def parse_control_rows(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),"action":hfind(h,"action uid"),"gate":hfind(h,"gate"),"permission":hfind(h,"permission","auth resource"),"operation":hfind(h,"operation"),"runtime_owner":hfind(h,"runtime owner"),"runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"source":path,"table":ti+1,"row":ri})
            out.append(rec)
    return out

def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        base["same_uid_rows"]=rs;out[uid]=base
    return out

def gap(field,r):
    v=r.get(field,"");st=r.get("runtime_status","")
    if not missing(v):return False
    if "UI_LOCAL_EXACT" in st and (field in {"operation","runtime_owner"} or (field=="action" and display_only_type(r.get("type","")))):return False
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in st and field=="operation":return False
    if "READ_EXACT" in st and display_only_type(r.get("type","")):
        if field=="action":return False
        if field=="operation" and not missing(r.get("runtime_owner","")):return False
    if field=="runtime_owner" and any(x in st for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return False
    if v=="SOURCE_NOT_DEFINED":return False
    return True

# Physical owner registrations from 9 System docs.
owner_index=collections.defaultdict(list)
for sf in SYSTEMS:
    for r in parse_control_rows(sf):
        owner_index[r["control"]].append(r)

# Reconstruct current gaps and collision rows.
current=[]
for page,fn in PAGES.items():
    bm=compose(parse_control_rows(fn))
    for uid,r in bm.items():
        for field in ["action","gate","permission","operation","runtime_owner"]:
            if not gap(field,r):continue
            sources=sorted(set(x["source"] for x in owner_index.get(uid,[])))
            current.append({"page":page,"uid":uid,"field":field,"owner_sources":sources})
assert len(current)==277,len(current)

sys_conf=[x for x in current if x["page"]=="SYS-01" and x["uid"]=="SYS-01-BTN-NAV-OPEN" and x["field"]=="gate"]
assert len(sys_conf)==1,sys_conf

collisions=[x for x in current if len(x["owner_sources"])>1]
assert len(collisions)==56,len(collisions)
raw_pair_counts=collections.Counter(tuple(x["owner_sources"]) for x in collisions)
assert set(raw_pair_counts)==set(PRECEDENCE),raw_pair_counts

resolved=[]
reference_rows=collections.defaultdict(list)
canonical_counts=collections.Counter()
reference_counts=collections.Counter()
for x in collisions:
    pair=tuple(sorted(x["owner_sources"]))
    canonical=PRECEDENCE[pair]
    refs=[s for s in pair if s!=canonical]
    assert len(refs)==1
    ref=refs[0]
    rec=x|{"canonical_owner":canonical,"reference_only_owner":ref,"authority_basis":AUTHORITY_BASIS[pair],"status":"OWNER_RESOLVED_REFERENCE_ONLY_SECONDARY"}
    resolved.append(rec)
    reference_rows[ref].append(rec)
    canonical_counts[canonical]+=1
    reference_counts[ref]+=1

assert canonical_counts==collections.Counter({S03:18,S02:10,S01:18,S06:2,S05:8}),canonical_counts
assert reference_counts==collections.Counter({S09:45,S04:7,S02:2,S08:2}),reference_counts

# Effective post-Batch10 ownership classification.
effective_unique=0;effective_absent=0;effective_collision=0
for x in current:
    if x in collisions:
        effective_unique+=1
    elif len(x["owner_sources"])==0:
        effective_absent+=1
    elif len(x["owner_sources"])==1:
        effective_unique+=1
    else:
        effective_collision+=1
# isolate preserved SYS Gate conflict from ordinary remediation denominator
assert effective_unique==249,(effective_unique,effective_absent,effective_collision)
assert effective_absent==28
assert effective_collision==0
effective_unique_remediation=effective_unique-1
assert effective_unique_remediation==248

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.35);sec.bottom_margin=Inches(.35);sec.left_margin=Inches(.3);sec.right_margin=Inches(.3)
def add_table(doc,headers,rows,fs=4.5):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row):c[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:r.font.size=Pt(fs)
    return t

# Secondary occurrence supersession registers.
for sf in [S02,S04,S08,S09]:
    rows=reference_rows[sf]
    assert rows,(sf,"NO_REFERENCE_ROWS")
    d=Document(sf)
    txt="\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
    local_marker=MARK+"::REFERENCE_ONLY::"+Path(sf).name
    assert local_marker not in txt,(sf,"BATCH10_ALREADY_PRESENT")
    add_landscape(d)
    d.add_heading("Batch 10 · Cross-System Reference-Only Supersession Register",level=1)
    p=d.add_paragraph();p.add_run("["+local_marker+"] ").bold=True
    p.add_run("The exact UIDs below remain in this document for integration, inspection, or domain-consumer context only. Their physical presence MUST NOT be interpreted as Canonical Owner registration. Batch 10 resolves the canonical owner from pre-existing owner-boundary authority; no business binding value is changed.")
    add_table(d,["Reference UID","Page","Missing Field","Canonical Owner","Local Role","Authority Basis"],[
      [r["uid"],r["page"],r["field"],Path(r["canonical_owner"]).name,"CROSS_SYSTEM_REFERENCE_ONLY",r["authority_basis"]]
      for r in sorted(rows,key=lambda z:(z["page"],z["uid"],z["field"]))
    ],3.9)
    d.add_paragraph("Scanner rule: tables in this supersession register deliberately use 'Reference UID', not 'Control UID'. Listed UIDs are reference-only in this document and MUST be excluded from canonical-owner denominator calculations.")
    d.save(sf);Document(sf)

# Central System Logic precedence ledger.
logic=Document(LOGIC)
ltxt="\n".join([p.text for p in logic.paragraphs]+[c.text for t in logic.tables for row in t.rows for c in row.cells])
assert MARK not in ltxt,"LOGIC_BATCH10_ALREADY_PRESENT"
add_landscape(logic)
logic.add_heading("Batch 10 · Canonical Owner Precedence / Collision Closure",level=1)
p=logic.add_paragraph();p.add_run("["+MARK+"] ").bold=True
p.add_run("All 56 prior OWNER_COLLISION rows were physically duplicated but had no divergent non-empty binding values. Batch 10 resolves ownership only. The secondary occurrence is retained as CROSS_SYSTEM_REFERENCE_ONLY and cannot supply or override Action/Gate/Permission/Operation/Runtime Owner values.")

add_table(logic,["Pair","Rows","Canonical Owner","Reference-only Owner","Authority Basis"],[
 [" + ".join(Path(s).stem for s in pair),raw_pair_counts[pair],Path(PRECEDENCE[pair]).name,Path([s for s in pair if s!=PRECEDENCE[pair]][0]).name,AUTHORITY_BASIS[pair]]
 for pair in sorted(raw_pair_counts)
],3.8)

logic.add_heading("56-row Ownership Resolution Ledger",level=2)
add_table(logic,["Page","Affected UID","Missing Field","Canonical Owner","Reference-only Source","Status"],[
 [r["page"],r["uid"],r["field"],Path(r["canonical_owner"]).name,Path(r["reference_only_owner"]).name,r["status"]]
 for r in sorted(resolved,key=lambda z:(z["page"],z["uid"],z["field"]))
],3.8)

logic.add_heading("Post-Batch 10 Denominator",level=2)
add_table(logic,["Class","Count","Meaning"],[
 ["Definition Binding Gap total",277,"UNCHANGED; Batch 10 changes ownership classification only."],
 ["Unique-owner unresolved remediation",248,"Includes the 56 collision rows now mapped to one canonical owner; excludes preserved SYS Gate conflict."],
 ["Owner collision",0,"56 -> 0 after explicit precedence + reference-only supersession."],
 ["WB owner absent",28,"Unchanged; requires WB Canonical Control UID / Owner mapping."],
 ["SYS-01-BTN-NAV-OPEN Gate conflict",1,"Unchanged and preserved fail-closed."],
 ["Business binding values mutated",0,"No Action/Gate/Permission/Operation/Runtime Owner value changed in Batch 10."],
],5.0)

logic.add_heading("Owner Precedence Rules",level=2)
add_table(logic,["Rule","Requirement"],[
 ["CP-01","Physical duplicate Control rows do not create multiple canonical owners when an explicit owner-boundary contract identifies one owner."],
 ["CP-02","The non-canonical occurrence remains reference-only; it may support integration/inspection context but cannot close a missing binding."],
 ["CP-03","Row richness, file recency, lexical similarity, and first-match order never determine owner precedence."],
 ["CP-04","If an explicit owner-boundary rule is absent or contradictory, collision remains fail-closed."],
 ["CP-05","Ownership resolution does not authorize missing business values; those remain separate canonical-spec remediation work."],
],5.0)

machine={
 "marker":MARK,
 "definition_gap_total":277,
 "raw_owner_collision":56,
 "resolved_owner_collision":56,
 "post_owner_collision":0,
 "unique_owner_unresolved_remediation":248,
 "wb_owner_absent":28,
 "preserved_sys_gate_conflict":1,
 "business_binding_mutations":0,
 "canonical_owner_counts":{Path(k).name:v for k,v in canonical_counts.items()},
 "reference_only_counts":{Path(k).name:v for k,v in reference_counts.items()},
}
logic.add_paragraph("BATCH10_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

report={
 "machine":machine,
 "resolved_rows":resolved,
 "raw_pair_counts":{" + ".join(k):v for k,v in raw_pair_counts.items()},
 "modified_reference_docs":[S02,S04,S08,S09],
 "output_hashes":{f:blob(f) for f in [S02,S04,S08,S09,LOGIC]}
}
Path("__batch10_owner_precedence_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH10_REMEDIATION="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
