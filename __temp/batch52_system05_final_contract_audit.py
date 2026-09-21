from pathlib import Path
from docx import Document
import json,re,collections,subprocess

BASE_HEAD="ed9689931041b59cea3cfc0b985b096455fd405a"
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
SYSTEM_DOCS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
S05,
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
S02=SYSTEM_DOCS[1];S03=SYSTEM_DOCS[2];S04=SYSTEM_DOCS[3];S06=SYSTEM_DOCS[5];S08=SYSTEM_DOCS[7];S09=SYSTEM_DOCS[8]
PAIR_PRECEDENCE={
 tuple(sorted((S03,S09))):S03,
 tuple(sorted((S02,S04))):S02,
 tuple(sorted((S02,S06))):S06,
 tuple(sorted((S02,S08))):S02,
 tuple(sorted((S02,S09))):S02,
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out

assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

page_maps={p:compose(parse_controls(fn)) for p,fn in PAGES.items()}
sys_maps={sf:compose(parse_controls(sf)) for sf in SYSTEM_DOCS}

remaining=[]
for page,m in page_maps.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if missing(op) or st not in {"EFFECTFUL_EXACT","READ_EXACT"}:continue
        fs=[]
        if missing(r.get("method_path","")):fs.append("method_path")
        if st=="EFFECTFUL_EXACT" and missing(r.get("payload_schema","")):fs.append("payload_schema")
        if st=="EFFECTFUL_EXACT" and missing(r.get("persistence_owner","")):fs.append("persistence_owner")
        if not fs:continue
        sources=sorted(sf for sf,sm in sys_maps.items() if uid in sm)
        if len(sources)==1:owner=sources[0]
        else:
            pair=tuple(sources);assert pair in PAIR_PRECEDENCE,(page,uid,fs,sources)
            owner=PAIR_PRECEDENCE[pair]
        for f in fs:
            remaining.append({"field":f,"page":page,"uid":uid,"label":r.get("label",""),"type":r.get("type",""),
              "operation":op,"runtime_status":st,"runtime_owner":r.get("runtime_owner",""),
              "method_path":r.get("method_path",""),"payload_schema":r.get("payload_schema",""),
              "persistence_owner":r.get("persistence_owner",""),"action":r.get("action",""),
              "gate":r.get("gate",""),"permission":r.get("permission",""),"canonical_owner":owner})

assert len(remaining)==137,len(remaining)
not_s05=[x for x in remaining if x["canonical_owner"]!=S05]
assert not not_s05,not_s05[:20]

controls={}
for x in remaining:
    key=(x["page"],x["uid"])
    controls.setdefault(key,{"page":x["page"],"uid":x["uid"],"label":x["label"],"type":x["type"],"operation":x["operation"],
      "runtime_status":x["runtime_status"],"runtime_owner":x["runtime_owner"],"method_path":x["method_path"],
      "payload_schema":x["payload_schema"],"persistence_owner":x["persistence_owner"],"action":x["action"],
      "gate":x["gate"],"permission":x["permission"],"missing_fields":[]})
    controls[key]["missing_fields"].append(x["field"])

ops=sorted(set(v["operation"] for v in controls.values()))
keywords=ops+[
"asset","layer","patch","mask","candidate","image","production","handoff","video","generation","clip","shot","retry",
"edit","timeline","correction","subtitle","lip sync","voice take","import queue","final preview","qa","finding",
"manual review","evidence range","review point","release","criteria","quality","review","correction","task","job",
"request","runtime","persistence","artifact","version","audit","provider","generation job","production package"
]
rx=re.compile("|".join(re.escape(x) for x in sorted(set(keywords),key=len,reverse=True)),re.I)
d=Document(S05)
evidence={"paragraphs":[],"rows":[]}
for i,p in enumerate(d.paragraphs):
    t=norm(p.text)
    if t and rx.search(t):evidence["paragraphs"].append({"index":i,"text":t[:5000]})
for ti,t in enumerate(d.tables):
    if not t.rows:continue
    h=[norm(c.text) for c in t.rows[0].cells]
    for ri,row in enumerate(t.rows[1:],2):
        vals=[norm(c.text) for c in row.cells];joined=" | ".join(vals)
        if joined and rx.search(joined):
            evidence["rows"].append({"table":ti+1,"row":ri,"headers":h,"values":vals,"joined":joined[:8500]})

owner_rows=parse_controls(S05)
owner_exact={}
for key,c in controls.items():
    owner_exact[f"{c['page']}::{c['uid']}"]=[r for r in owner_rows if r["control"]==c["uid"]]

by_page={}
for p in sorted(set(x["page"] for x in remaining)):
    rr=[x for x in remaining if x["page"]==p]
    by_page[p]={
      "cells":len(rr),
      "unique_controls":len(set(x["uid"] for x in rr)),
      "fields":dict(collections.Counter(x["field"] for x in rr)),
      "runtime_status":dict(collections.Counter(x["runtime_status"] for x in rr)),
      "runtime_owners":dict(collections.Counter(x["runtime_owner"] for x in rr)),
      "operations":dict(collections.Counter(x["operation"] for x in rr)),
    }

summary={
"system05_cells":137,
"unique_controls":len(controls),
"operations":len(ops),
"by_page":by_page,
"by_field":dict(collections.Counter(x["field"] for x in remaining)),
"by_runtime_status":dict(collections.Counter(x["runtime_status"] for x in remaining)),
"by_runtime_owner":dict(collections.Counter(x["runtime_owner"] for x in remaining)),
"owner_paragraph_hits":len(evidence["paragraphs"]),
"owner_table_hits":len(evidence["rows"]),
}
report={"marker":"ACPOS-20260922-BATCH-52-SYSTEM05-FINAL-REMAINING-CONTRACT-AUDIT-V1","base_head":BASE_HEAD,
"summary":summary,"operations":ops,"controls":list(controls.values()),"owner_exact_rows":owner_exact,
"owner_evidence":evidence,"remaining_cells":remaining}
Path("__batch52_system05_contract_audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
Path("__batch52_summary.txt").write_text(json.dumps({"summary":summary,"controls":list(controls.values())},ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print("BATCH52_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
