from pathlib import Path
from docx import Document
import json,re,collections,subprocess

BASE_HEAD="ec3d35324f751693f41aad1a7d94ab4474929065"
S02="02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
PAGES=[
"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
]
OPS=[
"sendConversationMessage","stopConversationGeneration","createConversationThread","analyzeConversationMessage",
"createConversationBranch","createAssistantSummary","createAssistantStructuredDecision","attachConversationContext",
"getStrategicConversationProjection","getConversationAttachmentContext",
]
KEYWORDS=[
"conversation","thread","message","branch","summary","decision","attachment","context","generation","memory",
"outbox","inbox","persist","database","table","conversation_id","thread_id","message_id","attachment_id",
]

assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
rx=re.compile("|".join(re.escape(x) for x in OPS+KEYWORDS),re.I)

def collect(path):
    d=Document(path);out={"paragraphs":[],"rows":[]}
    for i,p in enumerate(d.paragraphs):
        t=norm(p.text)
        if t and rx.search(t):
            out["paragraphs"].append({"index":i,"text":t[:3000]})
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            joined=" | ".join(vals)
            if joined and rx.search(joined):
                out["rows"].append({"table":ti+1,"row":ri,"headers":h,"values":vals,"joined":joined[:5000]})
    return out

owner=collect(S02)
pages={p:collect(p) for p in PAGES}

# Focused structured extraction from owner tables containing Operation.
structured=[]
d=Document(S02)
for ti,t in enumerate(d.tables):
    if not t.rows:continue
    h=[norm(c.text) for c in t.rows[0].cells]
    hs=[x.lower() for x in h]
    oi=None
    for i,x in enumerate(hs):
        if "operation" in x: oi=i;break
    if oi is None:continue
    for ri,row in enumerate(t.rows[1:],2):
        vals=[norm(c.text) for c in row.cells]
        if oi>=len(vals):continue
        op=vals[oi]
        if op in OPS:
            structured.append({"operation":op,"table":ti+1,"row":ri,"headers":h,"values":vals})

summary={
"operations":len(OPS),
"owner_paragraph_hits":len(owner["paragraphs"]),
"owner_table_hits":len(owner["rows"]),
"structured_operation_rows":len(structured),
"page_paragraph_hits":sum(len(v["paragraphs"]) for v in pages.values()),
"page_table_hits":sum(len(v["rows"]) for v in pages.values()),
}
report={"marker":"ACPOS-20260922-BATCH-35-SYSTEM02-CONTRACT-DESIGN-AUDIT-V1","base_head":BASE_HEAD,
"summary":summary,"operations":OPS,"owner_evidence":owner,"page_evidence":pages,"structured_operation_rows":structured}
Path("__batch35_system02_contract_design_audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
Path("__batch35_summary.txt").write_text(json.dumps(summary,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print("BATCH35_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
