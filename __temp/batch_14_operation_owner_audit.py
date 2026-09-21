from pathlib import Path
from docx import Document
import json,re,collections

SYSTEMS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_cols(headers,names):
    hs=[norm(x).lower() for x in headers]
    return [i for i,h in enumerate(hs) if any(h==n.lower() for n in names)]
def display_only_type(t):
    u=(t or "").upper()
    return u in {"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"} or any(x in u for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(h,"control uid")
        if ci is None:continue
        idx={
          "control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),
          "method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),
          "persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status"),
        }
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"source":path,"table":ti+1,"row":ri,"headers":h,"values":vals})
            out.append(rec)
    return out

def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out

# Batch-13 current classifier.
def is_wb_read_projection_action_na(r):
    return (
        missing(r.get("action","")) and r.get("type","")=="SECTION_OPEN" and
        r.get("gate","")=="PAGE_READ" and r.get("permission","")=="workspace.dashboard.view" and
        r.get("payload_schema","")=="N/A_READ_PROJECTION" and
        r.get("operation","")=="getDashboardReadModel" and
        r.get("method_path","")=="GET /v1/dashboard/read-model" and
        r.get("runtime_owner","")=="DASHBOARD_READ_MODEL" and r.get("runtime_status","")=="READ_EXACT"
    )
def is_aiapi_direct_operation_action_na(r):
    return (
        missing(r.get("action","")) and r.get("type","")=="BUTTON_OR_ROW_ACTION" and
        r.get("payload_schema","")=="AIAPI Page Operation-specific Form / Provider Profile Field Contract" and
        not missing(r.get("permission","")) and not missing(r.get("operation","")) and
        not missing(r.get("method_path","")) and not missing(r.get("runtime_owner","")) and
        r.get("runtime_status","") in {"EFFECTFUL_EXACT","READ_EXACT"}
    )
def classify(page,field,r):
    v=r.get(field,"");st=r.get("runtime_status","")
    if not missing(v):return "BOUND"
    if "UI_LOCAL_EXACT" in st:
        if field in {"operation","runtime_owner"}:return "LEGITIMATE_NA_UI_LOCAL"
        if field=="action" and display_only_type(r.get("type","")):return "LEGITIMATE_NA_UI_LOCAL"
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in st and field=="operation":return "CLOSED_BY_OWNER_LOCAL_COMMAND"
    if "READ_EXACT" in st and display_only_type(r.get("type","")):
        if field=="action":return "LEGITIMATE_NA_READ_PRESENTATION"
        if field=="operation" and not missing(r.get("runtime_owner","")):return "READ_OWNER_BOUND_OPERATION_UNSPECIFIED"
    if field=="runtime_owner" and any(x in st for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return "KNOWN_RUNTIME_BLOCKER"
    if v=="SOURCE_NOT_DEFINED":return "SOURCE_RESOLUTION_REQUIRED"
    if field=="action" and page=="WB-01" and is_wb_read_projection_action_na(r):return "LEGITIMATE_NA_READ_PROJECTION_ACTION"
    if field=="action" and page=="AIAPI-01" and is_aiapi_direct_operation_action_na(r):return "LEGITIMATE_NA_DIRECT_OPERATION_ACTION"
    return "DEFINITION_BINDING_GAP"

# Parse all generic owner rows and Batch-10 reference-only precedence.
owner_rows=collections.defaultdict(list)
generic=collections.defaultdict(list)
reference_only=collections.defaultdict(lambda:collections.defaultdict(list))
doc_context={}
for sf in SYSTEMS:
    d=Document(sf)
    paras=[norm(p.text) for p in d.paragraphs if norm(p.text)]
    allrows=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if not any(vals):continue
            rec={"source":sf,"table":ti+1,"row":ri,"headers":headers,"values":vals,"joined":" | ".join(vals)}
            generic[sf].append(rec);allrows.append(rec)
        ci=hfind(headers,"control uid")
        if ci is not None:
            for rec in [x for x in allrows if x["table"]==ti+1]:
                if ci<len(rec["values"]):
                    uid=rec["values"][ci]
                    if uid and uid not in {"—","-"}:owner_rows[uid].append(rec)
        rui=next((i for i,h in enumerate(headers) if h.lower()=="reference uid"),None)
        coi=next((i for i,h in enumerate(headers) if h.lower()=="canonical owner"),None)
        mfi=next((i for i,h in enumerate(headers) if h.lower()=="missing field"),None)
        if rui is not None and coi is not None:
            for rec in [x for x in allrows if x["table"]==ti+1]:
                vals=rec["values"]
                uid=vals[rui] if rui<len(vals) else ""
                owner=vals[coi] if coi<len(vals) else ""
                field=vals[mfi] if mfi is not None and mfi<len(vals) else ""
                if uid and owner:
                    reference_only[uid][field].append({"secondary":sf,"canonical_name":owner,"table":rec["table"],"row":rec["row"]})
    doc_context[sf]={"paragraphs":paras,"rows":generic[sf]}

# Current 209 gaps and 208 ordinary unique-owner rows.
current=[]
page_maps={}
for page,fn in PAGES.items():
    bm=compose(parse_controls(fn));page_maps[page]=bm
    for uid,r in bm.items():
        for field in ["action","gate","permission","operation","runtime_owner"]:
            if classify(page,field,r)=="DEFINITION_BINDING_GAP":
                current.append({"page":page,"file":fn,"uid":uid,"field":field,**{k:r.get(k,"") for k in FIELDS[1:]}})
assert len(current)==209,len(current)
ordinary=[x for x in current if not (x["page"]=="SYS-01" and x["uid"]=="SYS-01-BTN-NAV-OPEN" and x["field"]=="gate")]
assert len(ordinary)==208
targets=[x for x in ordinary if x["field"] in {"operation","runtime_owner"}]
assert len(targets)==189,len(targets)

# Effective owner resolution using Batch-10 reference-only supersession.
for q in targets:
    physical=sorted(set(r["source"] for r in owner_rows.get(q["uid"],[])))
    secondary=set()
    for rr in reference_only.get(q["uid"],{}).values():
        for item in rr:secondary.add(item["secondary"])
    effective=[s for s in physical if s not in secondary]
    assert len(effective)==1,(q["page"],q["uid"],q["field"],physical,effective)
    q["canonical_owner"]=effective[0]

# Relation candidate extraction.
TARGET_HEADERS={
 "operation":["Operation","Operation ID","operationId","API Operation","Service Operation"],
 "runtime_owner":["Runtime Owner","Runtime Service Owner","Service Owner"],
}
ANCHORS={
 "operation":[
   ("action",["Action UID","Action ID"]),
   ("method_path",["Method / Path","Method/Path","HTTP Method / Path","API / Method","API/Method"]),
   ("permission",["Permission","Permission / Auth Resource","Auth Resource","Authorization Resource"]),
   ("gate",["Gate","Gate UID"]),
 ],
 "runtime_owner":[
   ("operation",["Operation","Operation ID","operationId","API Operation","Service Operation"]),
   ("method_path",["Method / Path","Method/Path","HTTP Method / Path","API / Method","API/Method"]),
   ("persistence_owner",["Persistence Owner","Persistence"]),
   ("action",["Action UID","Action ID"]),
 ]
}

def exact_relation(q):
    owner=q["canonical_owner"];target=q["field"]
    vals=collections.defaultdict(list)
    for anchor_field,headers in ANCHORS[target]:
        anchor=q.get(anchor_field,"")
        if missing(anchor):continue
        for rec in generic[owner]:
            acols=exact_cols(rec["headers"],headers)
            if not acols or not any(i<len(rec["values"]) and rec["values"][i]==anchor for i in acols):continue
            tcols=exact_cols(rec["headers"],TARGET_HEADERS[target])
            for ti in tcols:
                if ti>=len(rec["values"]):continue
                value=rec["values"][ti]
                if missing(value):continue
                vals[value].append({"anchor_field":anchor_field,"anchor":anchor,"table":rec["table"],"row":rec["row"],"target_header":rec["headers"][ti]})
    if len(vals)==1:
        value=next(iter(vals))
        return "UNIQUE_EXACT_RELATION",value,vals[value]
    if len(vals)>1:
        return "EXACT_RELATION_CONFLICT",None,dict(vals)
    return "NO_EXACT_RELATION",None,[]

# Semantic N/A / boundary contexts for exact UID and runtime status.
TERMS=[
"LOCAL_WORKING_DRAFT_EXACT","UI_LOCAL_OR_READ_SOURCE","ORCHESTRATES_EXISTING_EXACT_OPERATIONS",
"owner-orchestrated","working draft","local draft","local-only","local only","no public route",
"no independent operation","no separate operation","no separate runtime owner","read source","projection",
"existing exact operations","orchestrates existing","runtime owner","operation"
]
def contexts_for(q):
    sf=q["canonical_owner"];uid=q["uid"];status=q["runtime_status"]
    out=[]
    for i,p in enumerate(doc_context[sf]["paragraphs"],1):
        low=p.lower()
        if uid in p or (status and status in p) or any(t.lower() in low for t in TERMS):
            if uid in p or (status and status in p) or any(x in low for x in ["local working","working draft","local-only","local only","orchestrat","no public route","read source","projection","no separate operation","no independent operation"]):
                out.append({"kind":"paragraph","index":i,"text":p})
    for rec in doc_context[sf]["rows"]:
        text=rec["joined"];low=text.lower()
        if uid in text or (status and status in text) or any(t.lower() in low for t in TERMS):
            if uid in text or (status and status in text) or any(x in low for x in ["local working","working draft","local-only","local only","orchestrat","no public route","read source","projection","no separate operation","no independent operation"]):
                out.append({"kind":"table","table":rec["table"],"row":rec["row"],"text":text})
    return out[:80]

relation_counts=collections.Counter();status_field=collections.Counter();owner_field=collections.Counter();page_field=collections.Counter()
rows=[]
for q in targets:
    st,val,evidence=exact_relation(q)
    q["relation_status"]=st
    if val is not None:q["relation_value"]=val
    q["relation_evidence"]=evidence
    q["semantic_contexts"]=contexts_for(q)
    relation_counts[st]+=1
    status_field[(q["runtime_status"],q["field"])]+=1
    owner_field[(q["canonical_owner"],q["field"])]+=1
    page_field[(q["page"],q["field"])]+=1
    rows.append(q)

payload={
 "denominator":189,
 "relation_counts":dict(relation_counts),
 "status_field":{" | ".join(k):v for k,v in status_field.items()},
 "owner_field":{" | ".join(k):v for k,v in owner_field.items()},
 "page_field":{" | ".join(k):v for k,v in page_field.items()},
 "rows":rows,
}
Path("__batch14_operation_owner_audit.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH14_AUDIT="+json.dumps({
 "denominator":189,
 "relation_counts":dict(relation_counts),
 "status_field":{" | ".join(k):v for k,v in status_field.items()},
 "owner_field":{" | ".join(k):v for k,v in owner_field.items()},
 "page_field":{" | ".join(k):v for k,v in page_field.items()},
},ensure_ascii=False,sort_keys=True))
print("BATCH14_CANDIDATES_BEGIN")
for r in rows:
    if r["relation_status"]!="NO_EXACT_RELATION" or r["runtime_status"] in {"LOCAL_WORKING_DRAFT_EXACT","UI_LOCAL_OR_READ_SOURCE","ORCHESTRATES_EXISTING_EXACT_OPERATIONS"}:
        print(json.dumps({
          "page":r["page"],"uid":r["uid"],"field":r["field"],"type":r["type"],
          "status":r["runtime_status"],"owner":r["canonical_owner"],
          "action":r["action"],"gate":r["gate"],"permission":r["permission"],
          "payload_schema":r["payload_schema"],"operation":r["operation"],
          "method_path":r["method_path"],"runtime_owner":r["runtime_owner"],
          "persistence_owner":r["persistence_owner"],
          "relation_status":r["relation_status"],"relation_value":r.get("relation_value"),
          "relation_evidence":r["relation_evidence"],
          "semantic_contexts":r["semantic_contexts"][:25],
        },ensure_ascii=False,sort_keys=True))
print("BATCH14_CANDIDATES_END")
