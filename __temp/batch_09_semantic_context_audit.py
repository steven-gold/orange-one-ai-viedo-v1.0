from pathlib import Path
from docx import Document
import json,re,collections

SYSTEMS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
TARGET_OWNERS=set([
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
])

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for n in needles:
        n=n.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None

DISPLAY={"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"}
def display_only(t):
    u=(t or "").upper()
    return u in DISPLAY or any(x in u for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]
def page_rows(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),"action":hfind(h,"action uid"),"gate":hfind(h,"gate"),"permission":hfind(h,"permission","auth resource"),"operation":hfind(h,"operation"),"runtime_owner":hfind(h,"runtime owner"),"runtime_status":hfind(h,"runtime status")}
        for ri,r in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in r.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"source":path,"table":ti+1,"row":ri,"indices":idx});out.append(rec)
    return out

def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        base["same_uid_rows"]=rs;out[uid]=base
    return out

def gap(field,r):
    v=r.get(field,"");st=r.get("runtime_status","")
    if not missing(v):return False
    if "UI_LOCAL_EXACT" in st and (field in {"operation","runtime_owner"} or (field=="action" and display_only(r.get("type","")))):return False
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in st and field=="operation":return False
    if "READ_EXACT" in st and display_only(r.get("type","")):
        if field=="action":return False
        if field=="operation" and not missing(r.get("runtime_owner","")):return False
    if field=="runtime_owner" and any(x in st for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return False
    if v=="SOURCE_NOT_DEFINED":return False
    return True

# Exact owner source index, using standard Control UID headers only.
owner_idx=collections.defaultdict(set)
for sf in SYSTEMS:
    d=Document(sf)
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        for r in t.rows[1:]:
            vals=[norm(c.text) for c in r.cells]
            if ci<len(vals) and vals[ci] and vals[ci] not in {"—","-"}:owner_idx[vals[ci]].add(sf)

allq=[]
for page,fn in PAGES.items():
    bm=compose(page_rows(fn))
    for uid,r in bm.items():
        for f in ["action","gate","permission","operation","runtime_owner"]:
            if gap(f,r):
                if page=="SYS-01" and uid=="SYS-01-BTN-NAV-OPEN" and f=="gate":continue
                owners=sorted(owner_idx.get(uid,set()))
                allq.append({"page":page,"file":fn,"uid":uid,"field":f,"type":r.get("type",""),"action":r.get("action",""),"gate":r.get("gate",""),"permission":r.get("permission",""),"operation":r.get("operation",""),"runtime_owner":r.get("runtime_owner",""),"runtime_status":r.get("runtime_status",""),"owners":owners})
unique=[q for q in allq if len(q["owners"])==1]
assert len(unique)==226,len(unique)

# Build richer same-owner contexts: all table rows + paragraphs.
docs={}
for sf in TARGET_OWNERS:
    d=Document(sf)
    tables=[]
    header_catalog=collections.Counter()
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        header_catalog[tuple(headers)]+=1
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if any(vals):tables.append({"table":ti+1,"row":ri,"headers":headers,"values":vals})
    paras=[norm(p.text) for p in d.paragraphs if norm(p.text)]
    docs[sf]={"tables":tables,"paras":paras,"headers":[{"headers":list(k),"count":v} for k,v in header_catalog.most_common()]}

FIELD_HEADER_PATTERNS={
"action":[r"action"],
"gate":[r"gate",r"policy gate"],
"permission":[r"permission",r"auth(?:orization)? resource",r"rbac"],
"operation":[r"operation",r"operationid",r"operation id",r"api operation",r"service operation"],
"runtime_owner":[r"runtime.*owner",r"owner.*runtime",r"service.*owner",r"execution.*owner",r"runtime service"],
}
EXCLUDE_HEADER=[r"status",r"description",r"desc",r"note",r"evidence",r"count",r"gap",r"missing",r"required owner action"]

def field_cols(headers,field):
    out=[]
    for i,h in enumerate(headers):
        hl=norm(h).lower()
        if any(re.search(p,hl) for p in FIELD_HEADER_PATTERNS[field]) and not any(re.search(p,hl) for p in EXCLUDE_HEADER):
            out.append(i)
    return out

LABELS={
"action":["Action UID","Action ID","Action"],
"gate":["Gate UID","Gate ID","Gate"],
"permission":["Permission Resource","Permission","Auth Resource","Authorization Resource"],
"operation":["operationId","Operation ID","Operation","API Operation","Service Operation"],
"runtime_owner":["Runtime Owner","Runtime Service Owner","Service Owner","Execution Owner"],
}
def parse_labeled(text,field):
    vals=[]
    for lab in LABELS[field]:
        # Delimit to semicolon/pipe/newline; require explicit label separator.
        pat=re.compile(r"(?i)(?:^|[;|。；])\s*"+re.escape(lab)+r"\s*[:=→]\s*([^;|。；]+)")
        for m in pat.finditer(text):
            v=norm(m.group(1))
            if v and len(v)<=180 and not missing(v):vals.append((v,lab))
    return vals

results=[];counts=collections.Counter();byfield=collections.defaultdict(collections.Counter);byowner=collections.defaultdict(collections.Counter)
for q in unique:
    sf=q["owners"][0];ctx=docs[sf];anchors=[]
    for k,v in [("control",q["uid"]),("action",q["action"]),("gate",q["gate"]),("operation",q["operation"])]:
        if not missing(v):anchors.append((k,v))
    cand=collections.defaultdict(list)
    for tr in ctx["tables"]:
        # row must contain an exact anchor cell, not substring.
        matches=[(k,v) for k,v in anchors if v in tr["values"]]
        if not matches:continue
        for ci in field_cols(tr["headers"],q["field"]):
            if ci>=len(tr["values"]):continue
            v=tr["values"][ci]
            if missing(v):continue
            cand[v].append({"kind":"TABLE_EXPLICIT_HEADER","table":tr["table"],"row":tr["row"],"header":tr["headers"][ci],"matches":matches})
        joined=" | ".join(tr["values"])
        for v,lab in parse_labeled(joined,q["field"]):
            cand[v].append({"kind":"TABLE_LABELED_TEXT","table":tr["table"],"row":tr["row"],"label":lab,"matches":matches})
    for pi,p in enumerate(ctx["paras"],1):
        matches=[(k,v) for k,v in anchors if v and v in p]
        if not matches:continue
        for v,lab in parse_labeled(p,q["field"]):
            cand[v].append({"kind":"PARAGRAPH_LABELED_TEXT","paragraph":pi,"label":lab,"matches":matches})
    if len(cand)==1:
        v=next(iter(cand));status="UNIQUE_EXPLICIT_SEMANTIC_BINDING";rec=q|{"status":status,"value":v,"evidence":cand[v]}
    elif len(cand)>1:
        status="EXPLICIT_SEMANTIC_CONFLICT";rec=q|{"status":status,"values":dict(cand)}
    else:
        status="NO_EXPLICIT_SEMANTIC_BINDING";rec=q|{"status":status}
    results.append(rec);counts[status]+=1;byfield[q["field"]][status]+=1;byowner[sf][status]+=1

Path("__batch09_semantic_context_audit.json").write_text(json.dumps({
 "denominator":226,"counts":dict(counts),"by_field":{k:dict(v) for k,v in byfield.items()},"by_owner":{k:dict(v) for k,v in byowner.items()},"rows":results,
 "header_catalog":{sf:docs[sf]["headers"] for sf in TARGET_OWNERS}
},ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH09_SEMANTIC="+json.dumps({"denominator":226,"counts":dict(counts),"by_field":{k:dict(v) for k,v in byfield.items()},"by_owner":{k:dict(v) for k,v in byowner.items()}},ensure_ascii=False,sort_keys=True))
print("SEMANTIC_RESOLVABLE_BEGIN")
for r in results:
    if r["status"]=="UNIQUE_EXPLICIT_SEMANTIC_BINDING":print(json.dumps(r,ensure_ascii=False,sort_keys=True))
print("SEMANTIC_RESOLVABLE_END")
print("SEMANTIC_CONFLICT_BEGIN")
for r in results:
    if r["status"]=="EXPLICIT_SEMANTIC_CONFLICT":print(json.dumps(r,ensure_ascii=False,sort_keys=True))
print("SEMANTIC_CONFLICT_END")
