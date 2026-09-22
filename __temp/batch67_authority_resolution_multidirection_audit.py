from pathlib import Path
from docx import Document
import json,re,collections,subprocess

BASE_HEAD="c8223a3886cc0a493bb37876651df0d06cb201be"
DOCS=sorted(Path(".").glob("*.docx"))
assert len(DOCS)==31
assert subprocess.check_output(["git","-c","core.quotePath=false","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

TARGET_UIDS=[
"SYS-01-BTN-ATTACH",
"CORE-01-BTN-CANDIDATE-CREATE","CORE-01-FLD-ASSISTANT-SUMMARY","CORE-01-FLD-EVALUATION","CORE-01-FLD-HUMAN-DECISION","CORE-01-FLD-STRUCTURED-DECISION",
"STR-01-BTN-ATTACH","STR-01-VIEW-CONVERSATION",
"AIAPI-01-BTN-CONFIGURE-GOVERNED-RESOURCE","AIAPI-01-BTN-APPROVE-GOVERNED-RESOURCE",
"CORE-01-BTN-SEND","CORE-01-FLD-MESSAGE","STR-01-BTN-SEND","STR-01-BTN-STOP","STR-01-INP-MESSAGE",
"IAM-01-BTN-AUDIT","SYS-01-BTN-SANDBOX-TEST","SYS-01-BTN-SEND","SYS-01-BTN-STOP",
"CTRL-ADMIN-STR-01-ACT-03-ACT-NAV-OPEN","CTRL-ADMIN-STR-02-ACT-03-ACT-NAV-OPEN","CTRL-ADMIN-STR-03-ACT-03-ACT-NAV-OPEN",
"CTRL-ADMIN-STR-04-ACT-02-ACT-NAV-OPEN","CTRL-ADMIN-STR-05-ACT-03-ACT-NAV-OPEN","CTRL-ADMIN-STR-06-ACT-04-ACT-NAV-OPEN",
"QA-01-BTN-EVIDENCE-IN","QA-01-BTN-EVIDENCE-OUT","QA-01-BTN-FINDING-POINT",
"KB-01-CTL-CITATION-OPEN","KB-01-CTL-REPLAY-OPEN",
"CTRL-ADMIN-SG-02-CRITERIA-TABLE-OPEN","CTRL-ADMIN-SG-02-DIMENSION-LIBRARY-OPEN","CTRL-ADMIN-SG-02-THRESHOLDS-OPEN",
"CTRL-ADMIN-SG-02-DEPARTMENT-MAPPING-OPEN","CTRL-ADMIN-SG-02-REQUIRED-CHECKS-OPEN","CTRL-ADMIN-SG-02-GATE-POLICY-OPEN",
"CTRL-ADMIN-SG-02-APPROVAL-OPEN","CTRL-ADMIN-SG-02-IMPACT-OPEN","CTRL-ADMIN-SG-02-ACT-03-ACT-NAV-OPEN",
]
OPS=[
"attachConversationContext","createCoreCandidate","createAssistantSummary","evaluateCoreCandidate","recordCoreHumanDecision","createAssistantStructuredDecision",
"getConversationAttachmentContext","getStrategicConversationProjection","configureGovernedResource","approveGovernedResource","sendConversationMessage",
"stopConversationGeneration","getUiProjection","runSandboxTest","setQAEvidenceRangeIn","setQAEvidenceRangeOut","setQAFindingPoint"
]
KEYWORDS=TARGET_UIDS+OPS+["canonical owner","owner","conversation core","assistant summary","structured decision","human decision","evaluation",
"citation","replay","quality criteria","SG-02","QA review","source page read owner","Projection / source page read owner","editing.task.execute"]

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None

uid_rows=collections.defaultdict(list)
op_rows=collections.defaultdict(list)
narrative=[]
rx=re.compile("|".join(re.escape(x) for x in sorted(set(KEYWORDS),key=len,reverse=True)),re.I)
for p in DOCS:
    d=Document(p)
    for i,para in enumerate(d.paragraphs):
        t=norm(para.text)
        if t and rx.search(t):
            narrative.append({"file":p.name,"kind":"paragraph","index":i,"text":t[:5000]})
    for ti,tbl in enumerate(d.tables,1):
        if not tbl.rows:continue
        h=[norm(c.text) for c in tbl.rows[0].cells]
        ci=hfind(h,"control uid","target uid");oi=hfind(h,"operation")
        for ri,row in enumerate(tbl.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];joined=" | ".join(vals)
            uid=vals[ci] if ci is not None and ci<len(vals) else ""
            op=vals[oi] if oi is not None and oi<len(vals) else ""
            rec={"file":p.name,"table":ti,"row":ri,"headers":h,"values":vals,"uid":uid,"operation":op}
            if uid in TARGET_UIDS:uid_rows[uid].append(rec)
            if op in OPS:op_rows[op].append(rec)
            if joined and rx.search(joined):
                narrative.append(rec|{"kind":"table_row"})

# Compact exact row projection for target UIDs.
def field(rec,*names):
    h=rec["headers"];v=rec["values"]
    i=hfind(h,*names)
    return v[i] if i is not None and i<len(v) else ""
target_compact={}
for uid,rows in uid_rows.items():
    target_compact[uid]=[]
    for r in rows:
        target_compact[uid].append({
          "file":r["file"],"table":r["table"],"row":r["row"],
          "uid_status":field(r,"uid status"),"type":field(r,"type"),"label":field(r,"label"),
          "action":field(r,"action uid"),"gate":field(r,"gate uid","gate"),"permission":field(r,"permission","auth resource"),
          "payload":field(r,"payload / schema","payload","schema"),"operation":field(r,"operation"),
          "method":field(r,"method / path","method","path"),"runtime_owner":field(r,"runtime owner"),
          "persistence":field(r,"persistence owner"),"runtime_status":field(r,"runtime status")
        })

# Candidate owner-resolution evidence: rows for KB and SG across system docs.
owner_candidates={}
for uid in ["KB-01-CTL-CITATION-OPEN","KB-01-CTL-REPLAY-OPEN",
            "CTRL-ADMIN-SG-02-CRITERIA-TABLE-OPEN","CTRL-ADMIN-SG-02-DIMENSION-LIBRARY-OPEN","CTRL-ADMIN-SG-02-THRESHOLDS-OPEN",
            "CTRL-ADMIN-SG-02-DEPARTMENT-MAPPING-OPEN","CTRL-ADMIN-SG-02-REQUIRED-CHECKS-OPEN","CTRL-ADMIN-SG-02-GATE-POLICY-OPEN",
            "CTRL-ADMIN-SG-02-APPROVAL-OPEN","CTRL-ADMIN-SG-02-IMPACT-OPEN","CTRL-ADMIN-SG-02-ACT-03-ACT-NAV-OPEN"]:
    owner_candidates[uid]=[r for r in target_compact.get(uid,[]) if re.match(r"0[1-9]_ACPOS_",r["file"])]

# Action reuse map only exact runtime rows.
action_map=collections.defaultdict(list)
for uid,rows in target_compact.items():
    for r in rows:
        if r["runtime_status"] in {"EFFECTFUL_EXACT","READ_EXACT"} and r["action"] not in {"","—","-"}:
            action_map[r["action"]].append(r)

summary={
 "target_uid_count":len(TARGET_UIDS),
 "target_uids_found":sum(1 for u in TARGET_UIDS if uid_rows.get(u)),
 "operation_count":len(OPS),
 "operations_found":sum(1 for o in OPS if op_rows.get(o)),
 "narrative_hits":len(narrative),
 "sys_attach_rows":len(uid_rows["SYS-01-BTN-ATTACH"]),
 "core_candidate_action_exact_rows":len(action_map["CORE-01-ACT-CANDIDATE-CREATE"]),
 "str_noop_action_exact_rows":len(action_map["STR-01-ACT-NOOP"]),
 "kb_owner_candidate_rows":sum(len(v) for k,v in owner_candidates.items() if k.startswith("KB-01")),
 "sg_owner_candidate_rows":sum(len(v) for k,v in owner_candidates.items() if k.startswith("CTRL-ADMIN-SG-02")),
}
report={"marker":"ACPOS-20260922-BATCH-67-AUTHORITY-RESOLUTION-MULTIDIRECTION-AUDIT-V1","base_head":BASE_HEAD,
"summary":summary,"target_rows":target_compact,"operation_rows":dict(op_rows),"action_map":dict(action_map),
"owner_candidates":owner_candidates,"narrative":narrative}
Path("__batch67_authority_resolution_audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
Path("__batch67_summary.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH67_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
