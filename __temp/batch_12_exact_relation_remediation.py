from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib

MARK="ACPOS-20260921-BATCH-12-EXACT-RELATION-REMEDIATION-V1"
S03="03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
AIAPI="ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx"
SG02="ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S03:"c48b94ebffc71c65416a45a2c71b688616a73f9c",
S05:"878e1b204907c36f77383e56ac6ec0fc20cc7c70",
AIAPI:"ef610d7f6f2bcc72890f1ba563c063f737cae60e",
SG02:"db3eadbfb0173b22c14c41044718bf68d644c0fe",
LOGIC:"86d2e77e2a66b7bde677e469cf8998e33e1ecc0a",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]

TARGETS=[
{"page":"AIAPI-01","file":AIAPI,"uid":"AIAPI-01-BTN-CONFIGURE-GOVERNED-RESOURCE","field":"permission","operation":"configureGovernedResource","value":"quality.criteria.configure","owner":S03},
{"page":"AIAPI-01","file":AIAPI,"uid":"AIAPI-01-BTN-APPROVE-GOVERNED-RESOURCE","field":"permission","operation":"approveGovernedResource","value":"governance.approve","owner":S03},
{"page":"SG-02","file":SG02,"uid":"CTRL-ADMIN-SG-02-CRITERIA-TABLE-OPEN","field":"action","operation":"getUiProjection","value":"ACT-NAV-OPEN","owner":S05},
{"page":"SG-02","file":SG02,"uid":"CTRL-ADMIN-SG-02-DIMENSION-LIBRARY-OPEN","field":"action","operation":"getUiProjection","value":"ACT-NAV-OPEN","owner":S05},
{"page":"SG-02","file":SG02,"uid":"CTRL-ADMIN-SG-02-THRESHOLDS-OPEN","field":"action","operation":"getUiProjection","value":"ACT-NAV-OPEN","owner":S05},
{"page":"SG-02","file":SG02,"uid":"CTRL-ADMIN-SG-02-DEPARTMENT-MAPPING-OPEN","field":"action","operation":"getUiProjection","value":"ACT-NAV-OPEN","owner":S05},
{"page":"SG-02","file":SG02,"uid":"CTRL-ADMIN-SG-02-REQUIRED-CHECKS-OPEN","field":"action","operation":"getUiProjection","value":"ACT-NAV-OPEN","owner":S05},
{"page":"SG-02","file":SG02,"uid":"CTRL-ADMIN-SG-02-GATE-POLICY-OPEN","field":"action","operation":"getUiProjection","value":"ACT-NAV-OPEN","owner":S05},
{"page":"SG-02","file":SG02,"uid":"CTRL-ADMIN-SG-02-APPROVAL-OPEN","field":"action","operation":"getUiProjection","value":"ACT-NAV-OPEN","owner":S05},
{"page":"SG-02","file":SG02,"uid":"CTRL-ADMIN-SG-02-IMPACT-OPEN","field":"action","operation":"getUiProjection","value":"ACT-NAV-OPEN","owner":S05},
]

def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():
    assert blob(f)==s,(f,blob(f),s)

def norm(x):
    return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())

def missing(v):
    return v in {"","—","-","SOURCE_NOT_DEFINED"}

def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None

def display_only_type(t):
    u=(t or "").upper()
    return u in {"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"} or any(x in u for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(headers,"control uid")
        if ci is None:continue
        idx={
          "control":ci,"type":hfind(headers,"type"),"label":hfind(headers,"label","顯示名稱"),
          "action":hfind(headers,"action uid"),"gate":hfind(headers,"gate uid","gate"),
          "permission":hfind(headers,"permission","auth resource"),"operation":hfind(headers,"operation"),
          "runtime_owner":hfind(headers,"runtime owner"),"runtime_status":hfind(headers,"runtime status")
        }
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"table":ti+1,"row":ri,"headers":headers})
            out.append(rec)
    return out

def compose(rows):
    groups=collections.defaultdict(list)
    for r in rows:groups[r["control"]].append(r)
    out={}
    for uid,rs in groups.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out

def is_gap(field,r):
    v=r.get(field,"");st=r.get("runtime_status","")
    if not missing(v):return False
    if "UI_LOCAL_EXACT" in st:
        if field in {"operation","runtime_owner"}:return False
        if field=="action" and display_only_type(r.get("type","")):return False
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in st and field=="operation":return False
    if "READ_EXACT" in st and display_only_type(r.get("type","")):
        if field=="action":return False
        if field=="operation" and not missing(r.get("runtime_owner","")):return False
    if field=="runtime_owner" and any(x in st for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return False
    if v=="SOURCE_NOT_DEFINED":return False
    return True

def gap_inventory():
    out=[]
    for page,fn in PAGES.items():
        bm=compose(parse_controls(fn))
        for uid,r in bm.items():
            for field in ["action","gate","permission","operation","runtime_owner"]:
                if is_gap(field,r):
                    out.append((page,uid,field))
    return out

pre=gap_inventory()
assert len(pre)==249,len(pre)
pre_set=set(pre)
for t in TARGETS:
    assert (t["page"],t["uid"],t["field"]) in pre_set,(t,"NOT_CURRENT_GAP")

# Verify each exact relation directly in current canonical owner tables.
def exact_cols(headers,names):
    hs=[norm(x).lower() for x in headers]
    return [i for i,h in enumerate(hs) if any(h==n.lower() for n in names)]

def verify_relation(target):
    d=Document(target["owner"])
    vals=collections.defaultdict(list)
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        op_cols=exact_cols(headers,["Operation","Operation ID","operationId","API Operation","Service Operation"])
        target_headers=["Permission","Permission / Auth Resource","Auth Resource","Authorization Resource"] if target["field"]=="permission" else ["Action UID","Action ID"]
        f_cols=exact_cols(headers,target_headers)
        if not op_cols or not f_cols:continue
        for ri,row in enumerate(t.rows[1:],2):
            cells=[norm(c.text) for c in row.cells]
            if not any(i<len(cells) and cells[i]==target["operation"] for i in op_cols):continue
            for fi in f_cols:
                if fi<len(cells) and not missing(cells[fi]):
                    vals[cells[fi]].append({"table":ti+1,"row":ri,"header":headers[fi]})
    assert list(vals.keys())==[target["value"]] or set(vals.keys())=={target["value"]},(target,dict(vals))
    return vals[target["value"]]

for t in TARGETS:
    t["relation_evidence"]=verify_relation(t)

# Patch page product rows only.
page_patch_counts=collections.Counter()
for fn in sorted(set(t["file"] for t in TARGETS)):
    d=Document(fn)
    file_targets=[t for t in TARGETS if t["file"]==fn]
    seen=collections.Counter()
    for table in d.tables:
        if not table.rows:continue
        headers=[norm(c.text) for c in table.rows[0].cells]
        ci=hfind(headers,"control uid")
        oi=hfind(headers,"operation")
        fi=hfind(headers,"permission","auth resource") if file_targets[0]["field"]=="permission" else hfind(headers,"action uid")
        if ci is None or oi is None or fi is None:continue
        for row in table.rows[1:]:
            cells=[norm(c.text) for c in row.cells]
            if ci>=len(cells):continue
            uid=cells[ci]
            matches=[t for t in file_targets if t["uid"]==uid]
            if not matches:continue
            t=matches[0]
            op=cells[oi] if oi<len(cells) else ""
            if op!=t["operation"]:continue
            old=norm(row.cells[fi].text)
            if missing(old):
                row.cells[fi].text=t["value"]
                seen[uid]+=1
                page_patch_counts[t["page"]]+=1
            elif old!=t["value"]:
                raise AssertionError(("PAGE_FIELD_CONFLICT",fn,uid,t["field"],old,t["value"]))
    d.save(fn);Document(fn)
    for t in file_targets:
        assert seen[t["uid"]]>=1,(fn,t["uid"],"TARGET_ROW_NOT_PATCHED")

# Add canonical owner closure ledgers; do not mutate historical registries.
def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.35);sec.bottom_margin=Inches(.35);sec.left_margin=Inches(.3);sec.right_margin=Inches(.3)

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)

def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)

def add_table(doc,headers,rows,fs=5.0):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row):c[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:r.font.size=Pt(fs)
    return t

for owner in [S03,S05]:
    rows=[t for t in TARGETS if t["owner"]==owner]
    d=Document(owner)
    text_all="\n".join([p.text for p in d.paragraphs]+[c.text for tb in d.tables for row in tb.rows for c in row.cells])
    local=MARK+"::"+Path(owner).name
    assert local not in text_all,(owner,"ALREADY_APPLIED")
    add_landscape(d)
    d.add_heading("Batch 12 · Exact-Relation Canonical Field Closure",level=1)
    p=d.add_paragraph();p.add_run("["+local+"] ").bold=True
    p.add_run("The rows below close only fields supported by a unique exact Operation relation already present in this same Canonical Owner document. No fuzzy label matching, cross-system transfer, first-match inference, or new API/owner identity is used.")
    if owner==S03:
        headers=["Control UID","Operation","Permission","Authority Relation","Status"]
        data=[[x["uid"],x["operation"],x["value"],"EXACT_OPERATION_TO_PERMISSION","CLOSED_BY_BATCH_12"] for x in rows]
    else:
        headers=["Control UID","Operation","Action UID","Authority Relation","Status"]
        data=[[x["uid"],x["operation"],x["value"],"EXACT_OPERATION_TO_ACTION","CLOSED_BY_BATCH_12"] for x in rows]
    add_table(d,headers,data,5.0)
    d.save(owner);Document(owner)

post=gap_inventory()
assert len(post)==239,(len(post),post[-30:])
removed=set(pre)-set(post)
added=set(post)-set(pre)
expected_removed={(t["page"],t["uid"],t["field"]) for t in TARGETS}
assert removed==expected_removed,{"missing_removed":sorted(expected_removed-removed),"unexpected_removed":sorted(removed-expected_removed)}
assert not added,sorted(added)

# Central evidence and explicit rejected candidate.
logic=Document(LOGIC)
logic_text="\n".join([p.text for p in logic.paragraphs]+[c.text for tb in logic.tables for row in tb.rows for c in row.cells])
assert MARK not in logic_text,"BATCH12_ALREADY_PRESENT"
add_landscape(logic)
logic.add_heading("Batch 12 · Exact-Relation Canonical Spec Remediation",level=1)
p=logic.add_paragraph();p.add_run("["+MARK+"] ").bold=True
p.add_run("Fresh audit denominator: 249 current Definition Binding Gaps = 248 unique-owner remediation + 1 preserved SYS Gate conflict. Batch 12 closes only 10 fields supported by unique exact same-owner Operation relations.")

add_table(logic,["Class","Count","Decision"],[
["Pre-batch Definition Binding Gap",249,"Fresh current denominator."],
["Unique-owner remediation",248,"Canonical owner resolved."],
["Unique exact relations found",11,"Field-format validation applied before mutation."],
["Accepted exact relations",10,"2 Permission + 8 Action UID."],
["Rejected policy-text Gate candidate",1,"SYS-01-BTN-SANDBOX-TEST: policy text is not a canonical Gate UID."],
["Exact relation conflicts",15,"Remain fail-closed."],
["No exact relation",222,"Remain canonical-spec remediation."],
["Post-batch Definition Binding Gap",239,"10 exact fields closed."],
["Business values invented",0,"All values pre-existed in the same Canonical Owner."],
],5.0)

logic.add_heading("10 Closed Fields",level=2)
add_table(logic,["Page","Control UID","Field","Exact Operation","Resolved Value","Canonical Owner"],[
[x["page"],x["uid"],x["field"],x["operation"],x["value"],Path(x["owner"]).name] for x in TARGETS
],4.2)

logic.add_heading("Rejected Candidate Boundary",level=2)
add_table(logic,["Control UID","Field","Observed Same-owner Value","Decision"],[
["SYS-01-BTN-SANDBOX-TEST","Gate","Production secret/output forbidden","REJECTED: policy prose is not a canonical Gate UID; remains unresolved."],
],5.0)

logic.add_heading("Batch 12 Invariants",level=2)
add_table(logic,["Rule","Requirement"],[
["B12-01","Exact Operation equality and one unanimous target value inside the same Canonical Owner are required."],
["B12-02","A semantically relevant natural-language policy must not be written into a UID field."],
["B12-03","Conflict, ambiguity, or absent relation remains fail-closed and cannot receive completion credit."],
["B12-04","Page binding and Canonical Owner closure evidence must agree exactly."],
],5.0)

machine={
"marker":MARK,
"pre_definition_gaps":249,
"accepted_exact_relations":10,
"accepted_permission":2,
"accepted_action":8,
"rejected_policy_text_gate":1,
"exact_relation_conflicts":15,
"no_exact_relation":222,
"post_definition_gaps":239,
"remaining_unique_owner_remediation":238,
"preserved_sys_gate_conflict":1,
"invented_values":0,
}
logic.add_paragraph("BATCH12_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

report={
"machine":machine,
"targets":TARGETS,
"removed":[{"page":a,"uid":b,"field":field} for a,b,field in sorted(removed)],
"page_patch_counts":dict(page_patch_counts),
"output_hashes":{f:blob(f) for f in [S03,S05,AIAPI,SG02,LOGIC]},
}
Path("__batch12_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH12_REMEDIATION="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
