from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib

BATCH04_HEAD="bcb2d0db630ad2ad88727b6be548e216a4fe4129"
MARK="ACPOS-20260921-BATCH-05-AUTHORITY-BACKED-REMEDIATION-V1"

SYSTEMS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]

PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
TARGET_PAGES={"IAM-01","SG-02","SYS-01","WB-01","ADMIN-STR-01"}
SYSTEM_LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"

def norm(x):
    return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())

def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None

def get(vals,i):
    return "" if i is None or i>=len(vals) else norm(vals[i])

def is_missing(v):
    return v in {"","—","-","SOURCE_NOT_DEFINED"}

DISPLAY_ONLY_TYPES={"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"}
def display_only(r):
    typ=r.get("type","").upper()
    return typ in DISPLAY_ONLY_TYPES or any(x in typ for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]

def parse_rows(path,require_control=False,with_cells=False):
    d=Document(path)
    out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        idx={
          "control":hfind(h,"control uid"),
          "type":hfind(h,"type"),
          "label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),
          "gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "operation":hfind(h,"operation"),
          "runtime_owner":hfind(h,"runtime owner"),
          "runtime_status":hfind(h,"runtime status")
        }
        if require_control and idx["control"] is None:continue
        if not require_control and all(idx[k] is None for k in ["control","action","gate","permission","operation","runtime_owner"]):continue
        for ri,r in enumerate(t.rows[1:],2):
            vals=[c.text for c in r.cells]
            rec={k:get(vals,idx[k]) for k in FIELDS}
            if require_control:
                uid=rec["control"]
                if not uid or uid in {"—","-"} or not re.search(r"[A-Za-z0-9]",uid):continue
            else:
                if not any(not is_missing(rec[k]) for k in ["control","action","gate","permission","operation","runtime_owner"]):continue
            rec.update({"source":str(path),"table":ti+1,"row":ri,"_indices":idx})
            if with_cells:
                rec["_doc"]=d
                rec["_table_obj"]=t
                rec["_row_obj"]=r
            out.append(rec)
    return d,out

def classify(field,row):
    val=row.get(field,"");status=row.get("runtime_status","")
    if not is_missing(val):return "BOUND"
    if "UI_LOCAL_EXACT" in status:
        if field in {"operation","runtime_owner"}:return "LEGITIMATE_NA_UI_LOCAL"
        if field=="action" and display_only(row):return "LEGITIMATE_NA_UI_LOCAL"
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in status and field=="operation":
        return "CLOSED_BY_OWNER_LOCAL_COMMAND"
    if "READ_EXACT" in status and display_only(row):
        if field=="action":return "LEGITIMATE_NA_READ_PRESENTATION"
        if field=="operation" and not is_missing(row.get("runtime_owner","")):
            return "READ_OWNER_BOUND_OPERATION_UNSPECIFIED"
    if field=="runtime_owner" and any(x in status for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):
        return "KNOWN_RUNTIME_BLOCKER"
    if val=="SOURCE_NOT_DEFINED":return "SOURCE_RESOLUTION_REQUIRED"
    return "DEFINITION_BINDING_GAP"

def best(rows):
    b={}
    for r in rows:
        uid=r["control"]
        score=sum(1 for k in FIELDS[1:] if not is_missing(r.get(k,"")))
        if uid not in b or score>b[uid][0]:
            b[uid]=(score,r)
    return {k:v[1] for k,v in b.items()}

def build_source_pool():
    src=[]
    for f in SYSTEMS:
        assert Path(f).exists(),f
        _,rows=parse_rows(f,False)
        src += rows
    return src

SRC=build_source_pool()

def candidate_records(target,field,page_rows):
    candidates=[]
    keys=[]
    if not is_missing(target.get("control","")):
        keys.append(("control",target["control"]))
    if field in {"gate","permission","operation","runtime_owner"} and not is_missing(target.get("action","")):
        keys.append(("action",target["action"]))
    if field=="permission" and not is_missing(target.get("gate","")):
        keys.append(("gate",target["gate"]))
    if field=="runtime_owner" and not is_missing(target.get("operation","")):
        keys.append(("operation",target["operation"]))
    seen=set()
    for key,val in keys:
        for r in SRC+page_rows:
            if r.get(key,"")!=val:continue
            fv=r.get(field,"")
            if is_missing(fv):continue
            sig=(fv,r["source"],r["table"],r["row"],key,val)
            if sig not in seen:
                seen.add(sig)
                candidates.append({
                  "value":fv,"source":r["source"],"table":r["table"],"row":r["row"],
                  "matched_by":key,"matched_value":val
                })
    return candidates

def compute_graph():
    report={"total_definition_gaps":0,"resolvable":0,"unresolved":0,"conflict":0,"pages":{},"rows":[],"conflicts":[]}
    for page,fn in PAGES.items():
        _,rows=parse_rows(fn,True)
        bm=best(rows)
        pc={"definition_gaps":0,"resolvable":0,"unresolved":0,"conflict":0}
        for uid,t in sorted(bm.items()):
            for field in ["action","gate","permission","operation","runtime_owner"]:
                if classify(field,t)!="DEFINITION_BINDING_GAP":continue
                report["total_definition_gaps"]+=1;pc["definition_gaps"]+=1
                cs=candidate_records(t,field,rows)
                vals=collections.defaultdict(list)
                for c in cs:vals[c["value"]].append(c)
                base={
                  "page":page,"file":fn,"uid":uid,"field":field,"type":t.get("type",""),
                  "label":t.get("label",""),"action":t.get("action",""),"gate":t.get("gate",""),
                  "permission":t.get("permission",""),"operation":t.get("operation",""),
                  "runtime_owner":t.get("runtime_owner",""),"runtime_status":t.get("runtime_status","")
                }
                if len(vals)==1:
                    value=next(iter(vals))
                    rec=base|{"value":value,"sources":vals[value]}
                    report["resolvable"]+=1;pc["resolvable"]+=1;report["rows"].append(rec)
                elif len(vals)==0:
                    report["unresolved"]+=1;pc["unresolved"]+=1
                else:
                    rec=base|{"candidate_values":dict(vals)}
                    report["conflict"]+=1;pc["conflict"]+=1;report["conflicts"].append(rec)
        report["pages"][page]=pc
    return report

pre=compute_graph()
assert pre["total_definition_gaps"]==346,pre["total_definition_gaps"]
assert pre["resolvable"]==34,pre["resolvable"]
assert pre["unresolved"]==311,pre["unresolved"]
assert pre["conflict"]==1,pre["conflict"]
assert pre["conflicts"][0]["page"]=="SYS-01" and pre["conflicts"][0]["uid"]=="SYS-01-BTN-NAV-OPEN" and pre["conflicts"][0]["field"]=="gate",pre["conflicts"]

applied=[]
page_cell_updates=collections.Counter()
page_binding_updates=collections.Counter()

for page in sorted(TARGET_PAGES):
    fn=PAGES[page]
    doc=Document(fn)
    # Build writable table row records from this exact document instance.
    writable=[]
    for ti,t in enumerate(doc.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        idx={
          "control":hfind(h,"control uid"),
          "type":hfind(h,"type"),
          "label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),
          "gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "operation":hfind(h,"operation"),
          "runtime_owner":hfind(h,"runtime owner"),
          "runtime_status":hfind(h,"runtime status")
        }
        if idx["control"] is None:continue
        for ri,r in enumerate(t.rows[1:],2):
            vals=[c.text for c in r.cells]
            uid=get(vals,idx["control"])
            if not uid or uid in {"—","-"} or not re.search(r"[A-Za-z0-9]",uid):continue
            writable.append({"uid":uid,"table":ti+1,"row":ri,"row_obj":r,"idx":idx})

    for rec in [x for x in pre["rows"] if x["page"]==page]:
        field=rec["field"];uid=rec["uid"];value=rec["value"]
        changed=0
        for wr in writable:
            if wr["uid"]!=uid:continue
            ci=wr["idx"].get(field)
            if ci is None:continue
            cell=wr["row_obj"].cells[ci]
            old=norm(cell.text)
            if is_missing(old):
                cell.text=value
                changed+=1
            elif old!=value:
                raise AssertionError(("TARGET_CELL_CONFLICT",page,uid,field,old,value,wr["table"],wr["row"]))
        assert changed>0,("NO_WRITABLE_TARGET_CELL",page,uid,field,value)
        applied.append({
          "page":page,"file":fn,"uid":uid,"field":field,"value":value,
          "changed_cells":changed,
          "matched_by":sorted(set(s["matched_by"] for s in rec["sources"])),
          "source_files":sorted(set(s["source"] for s in rec["sources"]))
        })
        page_cell_updates[page]+=changed
        page_binding_updates[page]+=1
    doc.save(fn)
    Document(fn)

assert len(applied)==34,len(applied)
assert sum(page_binding_updates.values())==34,page_binding_updates

post=compute_graph()
assert post["total_definition_gaps"]==312,post["total_definition_gaps"]

def git_blob_sha(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

page_hashes={page:git_blob_sha(PAGES[page]) for page in sorted(TARGET_PAGES)}

# Persist current evidence in System Logic.
doc=Document(SYSTEM_LOGIC)
alltxt="\n".join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
assert MARK not in alltxt,"BATCH05_ALREADY_PRESENT"

s=doc.add_section(WD_SECTION.NEW_PAGE)
s.orientation=WD_ORIENT.LANDSCAPE
s.page_width,s.page_height=s.page_height,s.page_width
s.top_margin=Inches(.35);s.bottom_margin=Inches(.35);s.left_margin=Inches(.3);s.right_margin=Inches(.3)

doc.add_heading("Batch 05 · Authority-backed Direct Binding Remediation",level=1)
p=doc.add_paragraph()
p.add_run("["+MARK+"] ").bold=True
p.add_run("This batch remediates only bindings proven by exact canonical graph keys. Label similarity, fuzzy matching, inferred API routes, invented owners, and cross-domain semantic guessing are forbidden.")

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr()
    e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)

def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr()
    e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)

def table(headers,rows,fs=4.6):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row):c[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for pp in cell.paragraphs:
                for rr in pp.runs:rr.font.size=Pt(fs)
    return t

doc.add_heading("Batch Result",level=2)
table(["Item","Count / State"],[
["Pre-batch direct Definition Binding Gap",346],
["Exact graph uniquely resolvable",34],
["Applied bindings",34],
["Pre-batch unresolved with no exact authority",311],
["Pre-batch explicit authority conflict",1],
["Post-batch direct Definition Binding Gap",post["total_definition_gaps"]],
["Post-batch exact-graph resolvable for a later batch",post["resolvable"]],
["Post-batch unresolved",post["unresolved"]],
["Post-batch conflicts",post["conflict"]],
],5.4)

doc.add_heading("Modified Page Evidence",level=2)
table(["Page","Applied bindings","Changed physical cells","Current Git blob SHA"],[
 [page,page_binding_updates[page],page_cell_updates[page],page_hashes[page]] for page in sorted(TARGET_PAGES)
],5.0)

doc.add_heading("Applied Exact Bindings",level=2)
table(["Page","Control UID","Field","Exact Value","Matched By","Authority Source"],[
 [r["page"],r["uid"],r["field"],r["value"]," / ".join(r["matched_by"]),"; ".join(r["source_files"])]
 for r in applied
],3.8)

doc.add_heading("Preserved Conflict / No Guessing",level=2)
conf=pre["conflicts"][0]
vals=sorted(conf["candidate_values"].keys())
table(["Page","Control UID","Field","Conflict Values","Disposition"],[
 [conf["page"],conf["uid"],conf["field"]," | ".join(vals),"PRESERVED_UNRESOLVED. Do not choose a Gate until canonical owner/context authority disambiguates the intended source page domain."]
],5.0)

doc.add_paragraph("The Batch-03/04 page hash snapshots remain historical evidence. For the five modified pages above, the Batch-05 hashes are the current evidence and supersede the older page hashes for current-state verification.")

machine={
 "marker":MARK,
 "pre":{"definition_gaps":346,"resolvable":34,"unresolved":311,"conflict":1},
 "applied_bindings":34,
 "post":{"definition_gaps":post["total_definition_gaps"],"resolvable":post["resolvable"],"unresolved":post["unresolved"],"conflict":post["conflict"]},
 "modified_pages":{page:{"bindings":page_binding_updates[page],"changed_cells":page_cell_updates[page],"blob_sha":page_hashes[page]} for page in sorted(TARGET_PAGES)},
 "preserved_conflict":{"page":conf["page"],"uid":conf["uid"],"field":conf["field"],"values":vals}
}
doc.add_paragraph("BATCH05_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
doc.save(SYSTEM_LOGIC)
Document(SYSTEM_LOGIC)

Path("__batch05_remediation_report.json").write_text(json.dumps({
 "pre":pre,
 "applied":applied,
 "post":post,
 "page_hashes":page_hashes,
 "system_logic_blob_sha":git_blob_sha(SYSTEM_LOGIC)
},ensure_ascii=False,indent=2),encoding="utf-8")

print("BATCH05_REMEDIATION="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
for r in applied:
    print("APPLIED="+json.dumps(r,ensure_ascii=False,sort_keys=True))
