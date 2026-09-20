from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

MARK="ACPOS-20260921-NINE-SYSTEM-QA-SYNC-V1"
FILES={
"01":"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02":"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03":"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04":"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05":"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06":"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07":"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08":"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09":"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
}

PROFILES=[
["QCP-CORE-BASELINE-V1","CORE","CORE / Blueprint / Canonical Script",8,95],
["QCP-ASSET-IMAGE-V1","ASSET","IMAGE-compatible asset outputs",9,95],
["QCP-ASSET-AUDIO-V1","ASSET","VOICE / MUSIC / SFX outputs",8,95],
["QCP-VIDEO-BASELINE-V1","VIDEO","VIDEO candidate outputs",9,95],
["QCP-EDIT-FINAL-V1","EDITING","EDIT stage / final outputs",12,95],
["QCP-QA-FINAL-OUTPUT-V1","QA","Final output / release-candidate review",17,95],
]

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr()
    e=OxmlElement('w:tblHeader');e.set(qn('w:val'),'true');trPr.append(e)

def shade(cell,fill='EDE9FE'):
    tcPr=cell._tc.get_or_add_tcPr()
    e=OxmlElement('w:shd');e.set(qn('w:fill'),fill);tcPr.append(e)

def table(doc,headers,rows,fs=6.3):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';repeat_header(t.rows[0])
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:r.font.size=Pt(fs)
    return t

def landscape(doc):
    s=doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation=WD_ORIENT.LANDSCAPE
    s.page_width,s.page_height=s.page_height,s.page_width
    s.top_margin=Inches(.4);s.bottom_margin=Inches(.4);s.left_margin=Inches(.35);s.right_margin=Inches(.35)

def has_mark(doc):
    return any(MARK in p.text for p in doc.paragraphs)

def heading(doc,title):
    landscape(doc)
    doc.add_heading(title,level=1)
    p=doc.add_paragraph()
    p.add_run(f'[{MARK}] ').bold=True
    p.add_run('本節同步目前九份系統文件對 QA Criteria、Business Content Completeness、SG-02 lifecycle 與跨系統 evidence lineage 的共同理解。此同步只更新 Word authority，不代表任何外部 Runtime/Production 已施工。')

def add_common_truth(doc):
    doc.add_heading('Canonical Quality Governance Baseline',level=2)
    table(doc,["Profile UID","Owner","Applicability","Dimensions","Pass Threshold"],PROFILES,6.2)
    table(doc,["Invariant","System-wide rule"],[
        ["QA-SYNC-01","任何需要 Quality Review / Score / Evaluation 的系統，MUST resolve exact approved Criteria Profile + CriteriaVersion + QualityGatePolicyVersion；不得用 generic/default/latest 取代。"],
        ["QA-SYNC-02","Static canonical catalog 必須有具體 rows；只有 Registry/Library/Dimension 容器而沒有 canonical content = CONTENT_INCOMPLETE。"],
        ["QA-SYNC-03","Dynamic runtime records 可以為空；禁止用 mock/sample/fake business rows 讓頁面看起來完整。"],
        ["QA-SYNC-04","QA score 不能覆蓋 Hard Block / Required Check；Rights / Policy / Evidence / Checksum / Provenance 等 blocker FAIL 時不得因總分高而 PASS。"],
        ["QA-SYNC-05","Correction 必須產生 NEW exact output/version 後才可 Recheck；不得在原輸出上直接關閉 Finding。"],
        ["QA-SYNC-06","Human Review 完成不等於 QA Gate PASS；人工不能繞過 Hard Block / Evidence / Policy / Criteria。"],
        ["QA-SYNC-07","SG-02 lifecycle canonical path = EMPTY → DRAFT → REVIEW → APPROVED / ACTIVE → SUPERSEDED。"],
        ["QA-SYNC-08","SG-02 Current Control Denominator = 12，包含 Submit Review；舊 11-control 描述不得作為 Current truth。"],
        ["QA-SYNC-09","只有 Word authority 已定義不等於 Runtime 已完成；Implementation / Production evidence 必須獨立標示。"],
    ],6.1)

def add_01(doc):
    heading(doc,'Nine-System QA / Criteria Synchronization Closure / 九系統品質治理同步')
    add_common_truth(doc)
    doc.add_heading('Global Ownership / 全域所有權',level=2)
    table(doc,["Concern","Canonical Owner","Consumers","Rule"],[
        ["Criteria Catalog / Profile","SG-02 Quality Governance","CORE/ASSET/VIDEO/EDIT/QA + any mapped consumer","Only SG-02 governs profile/version lifecycle; consumer systems never fork private criteria copies."],
        ["Concrete Review Items","Owning profile in SG-02","QA-01 / departmental evaluation surfaces","UI must render concrete rows, not a JSON blob/count."],
        ["Scorecard / Finding / Recheck","QA review runtime/domain","QA-01 + upstream correction owners","Finding must bind exact output/version + criteria version + evidence refs."],
        ["Business Content Completeness","Global System Logic","All 18 pages / all 9 systems","Empty static catalog blocks completion declaration."],
        ["Production Acceptance","Actual runtime/deployment owner","System Engineer / QA evidence","Word completion never substitutes runtime evidence."],
    ],6.2)
    doc.add_heading('Completion Declaration Gate',level=2)
    doc.add_paragraph('任何系統若需要固定 Catalog / Criteria / Dimension / Checklist / Profile / Preset / Stage Definition，只有 Schema、Control、API、空 Table 或空 Library 均不得宣告 complete。完成至少需要：canonical rows + rendering rule + lifecycle/state rule + owner + evidence/validation rule。')

def add_02(doc):
    heading(doc,'Conversation / Memory / Multi-AI Quality Context Synchronization')
    add_common_truth(doc)
    doc.add_heading('Conversation QA Context Contract',level=2)
    table(doc,["Field / Ref","Required meaning","Mutation rule"],[
        ["qa_profile_uid","Exact QCP-* profile selected for the discussed artifact","Conversation may display/read; cannot redefine profile."],
        ["criteria_version_id","Exact approved CriteriaVersion used for evaluation","Immutable reference in message/review lineage."],
        ["quality_gate_policy_version","Exact policy version paired with criteria","No latest/default fallback."],
        ["scorecard_id","Exact evaluation result being discussed","Chat cannot overwrite system score."],
        ["finding_ids","Exact open/closed findings referenced by discussion","Any correction decision must preserve IDs."],
        ["evidence_refs","Exact frame/time/range/checksum/provenance evidence","Conversation may explain; cannot fabricate evidence."],
        ["exact_output_ref","Exact artifact/output/version under review","No alias-only or latest pointer."],
        ["recheck_output_ref","NEW exact output/version used after correction","Must differ from superseded reviewed output when correction occurred."],
        ["manual_review_case_ref","Manual review case if required","Manual PASS cannot directly force QA PASS."],
    ],6.0)
    doc.add_heading('Conversation Behavior Invariants',level=2)
    table(doc,["Behavior","Rule"],[
        ["Explain QA","AI may summarize criteria, score, findings and evidence using exact refs."],
        ["Propose correction","AI may draft corrective instructions, but effectful correction must be handed to the owning department/runtime."],
        ["Modify criteria","FORBIDDEN from ordinary conversation. Criteria changes must enter SG-02 governed lifecycle."],
        ["Change score manually","FORBIDDEN. Human can resolve a ManualReviewCase, not type a replacement system score."],
        ["Memory/history","Conversation history must preserve criteria/profile/version/scorecard/finding/evidence refs so a later assistant can restore exact review context."],
        ["Multi-AI","All agents in a discussion share the same exact QA context refs; disagreement is discussion content, not an authority fork."],
    ],6.0)

def add_03(doc):
    heading(doc,'Execution / Compiler / Tool / AIAPI Quality Runtime Contract Synchronization')
    add_common_truth(doc)
    doc.add_heading('Quality Execution Contract',level=2)
    table(doc,["Operation class","Required input","Required output / state","Fail-closed condition"],[
        ["Resolve Criteria","artifact_type + department + exact scope","approved profile_uid + criteria_version_id + policy_version","No applicable approved profile/version"],
        ["Evaluate Output","exact_output_ref + checksum + profile/version + required_checks + evidence context","scorecard_id + dimension scores + weighted total + hard-block status","Output/version/checksum mismatch or missing criteria"],
        ["Create Finding","scorecard_id + dimension_uid/category + severity + affected scope + evidence refs","finding_id OPEN","Category/Severity/Evidence owner missing"],
        ["Create Correction Request","finding_ids + source_scorecard_ids + exact output/version + target owner + required_action","correction_request_id","No canonical target owner / no exact output"],
        ["Recheck","NEW exact output/version + prior finding/correction refs + same or explicitly superseding criteria version","fresh scorecard + finding closure eligibility","Reusing old output/version after correction"],
        ["Configure Criteria Draft","profile baseline + department/applicability + dimensions/checks/gate policy","DRAFT CriteriaVersion","Attempt to mutate APPROVED/ACTIVE in place"],
        ["Submit Criteria Review","DRAFT resource/version","REVIEW state + audit event","Not DRAFT / missing permission / invalid content hash"],
        ["Approve Criteria","REVIEW resource/version + approver evidence","APPROVED CriteriaVersion","Self/unauthorized approval, non-REVIEW state, invalid/missing evidence"],
    ],5.9)
    doc.add_heading('SG-02 Lifecycle Runtime Authority',level=2)
    table(doc,["Control / Action","Operation","Route / interface","Permission","Word authority status"],[
        ["Configure","configureGovernedResource","existing governed-resource configuration contract","quality.criteria.configure","Existing authority"],
        ["Submit Review","submitGovernedResourceReview","POST /v1/governance/resources/{id}/submit-review","quality.criteria.configure","NEW Authority after complete Word-source scan"],
        ["Approve","approveGovernedResource","existing governed-resource approval contract","governance.approve","Existing authority"],
    ],6.0)
    doc.add_paragraph('03 MUST treat the Submit Review operation as a required construction contract, but MUST NOT claim Runtime/Production implementation evidence merely because this Word contract exists.')
    doc.add_heading('Typed Payload Minimums',level=2)
    table(doc,["Payload","Minimum fields"],[
        ["EvaluateOutputRequest","exact_output_ref, output_version_ref, checksum, department, artifact_type, profile_uid, criteria_version_id, policy_version, required_check_refs, evidence_context"],
        ["CreateFindingRequest","scorecard_id, dimension_uid, category_uid, severity, affected_scope, evidence_refs, exact_output_ref"],
        ["CorrectionRequest","finding_ids, source_scorecard_ids, exact_output_ref, root_cause_classification, required_action, target_department, revalidation_requirements"],
        ["RecheckRequest","prior_scorecard_id, finding_ids, correction_request_id, NEW exact_output_ref/version, criteria_version_id, policy_version"],
        ["SubmitGovernedResourceReviewRequest","resource_id, current_draft_version, content_hash, correlation_id, idempotency_key, submission_reason/evidence refs"],
    ],5.8)

def add_04(doc):
    heading(doc,'CORE Quality Profile / Handoff Synchronization')
    add_common_truth(doc)
    doc.add_heading('CORE Profile Consumption',level=2)
    table(doc,["Rule","Exact requirement"],[
        ["Profile","CORE evaluation uses QCP-CORE-BASELINE-V1; 8 concrete dimensions; total weights=100; pass threshold=95."],
        ["Criteria version","CORE must bind exact approved CriteriaVersion; no local private copy."],
        ["Evaluation UI","CORE evaluation surface renders concrete rows: Goal/Blueprint Intent, Logic Structure, Character/DNA Plan, Scene/Object Continuity Plan, Production Requirement Completeness, Ambiguity Control, Evidence/Provenance, Rights/Policy."],
        ["Hard block","Rights/Policy failure blocks CORE approval regardless of total score."],
        ["Handoff","Approved CORE output hands downstream exact criteria/profile/version/evidence refs with Canonical Script/Blueprint lineage."],
        ["Change impact","If CORE content changes after evaluation, previous scorecard becomes stale and downstream must resolve a fresh evaluation."],
    ],6.0)

def add_05(doc):
    heading(doc,'Creative Production Quality Profile Synchronization')
    add_common_truth(doc)
    doc.add_heading('Department → Profile Mapping',level=2)
    table(doc,["Department / Artifact","Profile","Concrete review scope"],[
        ["ASSET image-compatible","QCP-ASSET-IMAGE-V1","Script/Manifest, Identity/DNA, Object, Scene, Composition, Geometry/Artifacts, Style/Color/Lighting, Technical, Rights/Policy"],
        ["ASSET audio-compatible","QCP-ASSET-AUDIO-V1","Script/Manifest, Voice/Semantic Fit, Pronunciation/Content, Timing, Noise/Clipping, Level/Dynamics, Technical Format, Rights/Policy"],
        ["VIDEO","QCP-VIDEO-BASELINE-V1","Script, Identity, Motion/Body, Camera/Gaze, Scene/Continuity, Timing/Duration, Technical, Asset/DNA Fidelity, Rights/Policy"],
        ["EDIT","QCP-EDIT-FINAL-V1","Assembly, Pacing, Visual/Frame Continuity, Voice, BGM, SFX, AV Sync, Lip Sync, Subtitle Content, Subtitle Timing/Readability, Technical, Rights/Provenance"],
        ["Final QA / Release Candidate","QCP-QA-FINAL-OUTPUT-V1","17 concrete final-output dimensions including identity, frame continuity, lip sync, audio, subtitle, composition, edit, AV sync, technical, rights/policy"],
    ],5.8)
    doc.add_heading('Creative Production QA Flow',level=2)
    table(doc,["Step","Contract"],[
        ["Generate","Department creates exact output/version; output identity/checksum/lineage immutable for that review."],
        ["Evaluate","Resolve exact approved profile/version; produce scorecard with concrete dimension rows."],
        ["Finding","Create canonical finding with category/severity/evidence and exact affected scope."],
        ["Correction","Route to original owner; do not edit output inside QA."],
        ["New Version","Owner generates NEW exact version."],
        ["Recheck","Fresh scorecard on NEW version; previous scorecard remains historical evidence."],
        ["Final Gate","Total >=95 AND all hard blocks/required checks PASS before release-candidate eligibility."],
    ],6.0)
    doc.add_paragraph('SG-02 control denominator is 12. The Submit Review control/action is part of the governing criteria lifecycle and 05 must not retain an 11-control assumption.')

def add_06(doc):
    heading(doc,'Enterprise Business Development QA Consumption Synchronization')
    add_common_truth(doc)
    doc.add_heading('Enterprise Output Quality Rule',level=2)
    table(doc,["Enterprise output case","QA requirement"],[
        ["Creative/media artifact generated for outreach/delivery","Must resolve the applicable existing QCP profile by artifact type; if no applicable profile exists, BLOCK and route a governed profile request to SG-02 rather than inventing a generic score."],
        ["Business records (lead/target/campaign/account metadata)","These are dynamic business data, not QA Criteria catalog rows; do not seed fake values. Validate with business schema/permission/runtime rules."],
        ["Externally deliverable content","Must carry exact output/version/checksum + criteria_version_id + scorecard_id + required-check state before delivery/publish handoff."],
        ["Corrected deliverable","Must use NEW exact output/version and fresh recheck evidence."],
        ["Rights/policy-sensitive delivery","Rights/Policy Required Check and provenance evidence are mandatory hard gate."],
    ],6.0)
    doc.add_heading('No Parallel QA System',level=2)
    doc.add_paragraph('06 MUST NOT create a separate enterprise-only quality score that substitutes for governed QA. Business KPIs, lead scores or campaign scores may exist for business decisions, but they are not QA Gate scores and cannot satisfy release-quality acceptance.')

def add_07(doc):
    heading(doc,'Social Publishing QA / Release Gate Synchronization')
    add_common_truth(doc)
    doc.add_heading('Publish Eligibility Contract',level=2)
    table(doc,["Required publish binding","Rule"],[
        ["exact_content_ref / version","Publish exactly the reviewed version; alias/latest-only input forbidden."],
        ["artifact_checksum","Must match QA-reviewed output checksum / manifest."],
        ["criteria_profile_uid + criteria_version_id","Exact approved criteria used for the current QA pass."],
        ["quality_gate_policy_version","Exact gate policy paired with the scorecard."],
        ["scorecard_id","Must be current for the exact version."],
        ["qa_gate_status","Must be PASS; MANUAL_REVIEW_CLOSED alone is insufficient."],
        ["required_checks","Rights/Policy/Provenance/Checksum and other applicable hard checks must PASS."],
        ["open_blocking_findings","Must equal 0."],
        ["evidence_refs","Preserve evidence references for audit / rollback / dispute review."],
    ],5.9)
    doc.add_heading('Publishing Fail-closed Rules',level=2)
    table(doc,["Condition","Required behavior"],[
        ["No QA profile mapping","BLOCK; request SG-02 governed mapping. Do not invent default criteria."],
        ["Criteria DRAFT/REVIEW only","BLOCK publication."],
        ["Artifact changed after QA","Invalidate publish eligibility; require fresh QA."],
        ["Rights/Policy check FAIL","BLOCK regardless of total score."],
        ["Scheduled post references stale/superseded version","BLOCK and require exact current reviewed version."],
    ],6.0)

def add_08(doc):
    heading(doc,'Strategy Intelligence / Human Decision vs QA Governance Synchronization')
    add_common_truth(doc)
    doc.add_heading('Decision-domain Separation',level=2)
    table(doc,["Concept","Strategy system meaning","QA meaning","Invariant"],[
        ["Candidate score / confidence","Decision-support signal for strategy comparison","Not a QA dimension score","Never reuse strategy confidence as QA PASS evidence."],
        ["Human strategy approval","Business/strategic decision authority","Not QA Manual Review completion","A strategy-approved artifact may still require QA before external/release use."],
        ["Evidence","Supports strategy facts/candidates/decisions","Supports QA dimension findings/checks","Refs may overlap but owner/state semantics stay distinct."],
        ["Risk","Strategic/business risk","Quality/rights/policy blocking condition","Do not collapse into one generic severity."],
    ],5.8)
    doc.add_heading('When Strategy Outputs Need QA',level=2)
    table(doc,["Output","Rule"],[
        ["Internal strategy candidate / fact projection","Uses strategy governance; no automatic QA profile unless explicitly mapped."],
        ["Presentation/media/report artifact generated from strategy","Resolve applicable existing QCP profile by artifact type before release/external distribution."],
        ["New recurring strategy-specific quality checklist","Must be created as a governed SG-02 profile/version; an empty 'criteria' container is not acceptable."],
        ["Artifact changed after QA","Fresh evaluation required before publication/delivery."],
    ],6.0)

def add_09(doc):
    heading(doc,'System Engineer / Self-Inspection Quality Governance Synchronization')
    add_common_truth(doc)
    doc.add_heading('System Self-Inspection Gates',level=2)
    table(doc,["Gate","Inspection rule","Failure classification"],[
        ["STATIC_CATALOG_CONTENT","Declared static Catalog/Criteria/Dimension/Checklist/Profile must contain concrete canonical rows and render contract","CONTENT_INCOMPLETE"],
        ["CRITERIA_LIFECYCLE_REACHABILITY","SG-02 must support EMPTY→DRAFT→REVIEW→APPROVED; Submit Review control/action/operation must exist in authority","LIFECYCLE_GAP"],
        ["DENOMINATOR_ALIGNMENT","SG-02 current control denominator =12; any current 11-control claim is stale","DENOMINATOR_STALE"],
        ["CONSUMER_BINDING","Quality-consuming system must carry exact profile/version/policy/scorecard/finding/evidence refs","QUALITY_CONTEXT_UNBOUND"],
        ["STALE_SCORECARD","Exact output/version changed after QA but prior scorecard still treated current","STALE_QA_EVIDENCE"],
        ["HARD_BLOCK_BYPASS","Total score or human review used to bypass Rights/Policy/Evidence/etc.","QUALITY_GATE_BYPASS"],
        ["RUNTIME_CLAIM","Word authority exists but Runtime/Production not actually validated","IMPLEMENTATION_EVIDENCE_MISSING"],
    ],5.8)
    doc.add_heading('Nine-System Synchronization Requirement',level=2)
    table(doc,["System document","Required synchronization"],[
        ["01","Global owner matrix + Business Content Completeness + completion declaration gate"],
        ["02","Conversation/memory preservation of exact QA context refs; no criteria mutation via chat"],
        ["03","Execution payloads/operations for evaluate/finding/correction/recheck and SG-02 lifecycle"],
        ["04","CORE profile consumption/handoff"],
        ["05","ASSET/VIDEO/EDIT/QA profile mapping and correction/recheck flow"],
        ["06","Enterprise outputs consume governed QA; no parallel QA score"],
        ["07","Publish only exact QA-passed version; rights/policy hard gate"],
        ["08","Strategy decision score/human approval separated from QA score/gate"],
        ["09","Self-inspection catches empty static catalogs, stale denominators, missing QA context and false runtime-completion claims"],
    ],5.8)

FUNCS={"01":add_01,"02":add_02,"03":add_03,"04":add_04,"05":add_05,"06":add_06,"07":add_07,"08":add_08,"09":add_09}
updated=[]
for k,fn in FILES.items():
    p=Path(fn)
    d=Document(p)
    if has_mark(d):
        print("ALREADY_SYNCED",fn)
        continue
    FUNCS[k](d)
    d.save(p)
    Document(p)
    updated.append(fn)

Path("__temp").mkdir(exist_ok=True)
Path("__temp/nine_system_sync_report.json").write_text(json.dumps({"marker":MARK,"updated":updated,"profiles":PROFILES},ensure_ascii=False,indent=2),encoding="utf-8")
print("UPDATED",len(updated),updated)
