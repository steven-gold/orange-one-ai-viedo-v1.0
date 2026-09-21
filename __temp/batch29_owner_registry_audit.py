from pathlib import Path
from docx import Document
import json,re,collections,subprocess

BASE_HEAD="b3d60bbe78a43aaf25171887ee5d4b93b98601e6"
SYSTEM_DOCS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
S02=SYSTEM_DOCS[1];S03=SYSTEM_DOCS[2];S04=SYSTEM_DOCS[3];S06=SYSTEM_DOCS[5];S08=SYSTEM_DOCS[7];S09=SYSTEM_DOCS[8]
PAIR_PRECEDENCE={
 tuple(sorted((S03,S09))):S03,
 tuple(sorted((S02,S04))):S02,
 tuple(sorted((S02,S06))):S06,
 tuple(sorted((S02,S08))):S02,
 tuple(sorted((S02,S09))):S02,
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]
SENTINELS={"SOURCE_NOT_DEFINED","NO_PUBLIC_API_BY_AUTHORITY","LOCAL LOCAL"}

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out
def parse_operation_registries(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        oi=hfind(h,"operation")
        if oi is None:continue
        idx={
          "operation":oi,
          "method_path":hfind(h,"method / path","method","path","route","endpoint"),
          "payload_schema":hfind(h,"payload / schema","payload","schema","request"),
          "persistence_owner":hfind(h,"persistence owner","persistence","storage owner"),
          "runtime_status":hfind(h,"runtime status","status"),
          "runtime_owner":hfind(h,"runtime owner"),
          "control":hfind(h,"control uid"),
          "gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
        }
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if oi>=len(vals):continue
            op=vals[oi]
            if not op or op in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri,"headers":h})
            out.append(rec)
    return out

assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

page_maps={p:compose(parse_controls(fn)) for p,fn in PAGES.items()}
sys_maps={sf:compose(parse_controls(sf)) for sf in SYSTEM_DOCS}
registry={sf:parse_operation_registries(sf) for sf in SYSTEM_DOCS}

targets=[]
for page,m in page_maps.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if st in {"EFFECTFUL_EXACT","READ_EXACT"} and not missing(op) and missing(r.get("method_path","")):
            targets.append(("method_path",page,uid,r))
        if st=="EFFECTFUL_EXACT" and not missing(op) and missing(r.get("payload_schema","")):
            targets.append(("payload_schema",page,uid,r))
        if st=="EFFECTFUL_EXACT" and not missing(op) and missing(r.get("persistence_owner","")):
            targets.append(("persistence_owner",page,uid,r))
assert len(targets)==324

rows=[]
for field,page,uid,r in targets:
    sources=sorted(sf for sf,sm in sys_maps.items() if uid in sm)
    if len(sources)==1: owner=sources[0]
    else:
        pair=tuple(sources);assert pair in PAIR_PRECEDENCE,(page,uid,field,sources);owner=PAIR_PRECEDENCE[pair]
    op=r["operation"];status=r["runtime_status"]
    exact=[];unscoped=[];mismatch=[]
    for rr in registry[owner]:
        if rr["operation"]!=op:continue
        v=rr.get(field,"")
        if missing(v) or v in SENTINELS:continue
        rs=rr.get("runtime_status","")
        item={"value":v,"table":rr["table"],"row":rr["row"],"runtime_status":rs,"runtime_owner":rr.get("runtime_owner",""),
              "control":rr.get("control",""),"gate":rr.get("gate",""),"permission":rr.get("permission","")}
        if rs==status:exact.append(item)
        elif missing(rs):unscoped.append(item)
        else:mismatch.append(item)
    def vals(items):
        d=collections.defaultdict(list)
        for x in items:d[x["value"]].append(x)
        return dict(d)
    ev=vals(exact);uv=vals(unscoped);mv=vals(mismatch)
    if len(ev)==1:
        cls="REGISTRY_EXACT_UNIQUE";candidate=next(iter(ev))
    elif len(ev)>1:
        cls="REGISTRY_EXACT_CONFLICT";candidate=None
    elif len(uv)==1:
        cls="REGISTRY_UNSCOPED_UNIQUE";candidate=next(iter(uv))
    elif len(uv)>1:
        cls="REGISTRY_UNSCOPED_CONFLICT";candidate=None
    elif mv:
        cls="REGISTRY_STATUS_MISMATCH_ONLY";candidate=None
    else:
        cls="NO_REGISTRY_RELATION";candidate=None
    rows.append({"field":field,"page":page,"uid":uid,"operation":op,"runtime_status":status,"canonical_owner":owner,
      "classification":cls,"candidate":candidate,"exact_values":ev,"unscoped_values":uv,"status_mismatch_values":mv})

classes=["REGISTRY_EXACT_UNIQUE","REGISTRY_EXACT_CONFLICT","REGISTRY_UNSCOPED_UNIQUE","REGISTRY_UNSCOPED_CONFLICT","REGISTRY_STATUS_MISMATCH_ONLY","NO_REGISTRY_RELATION"]
summary={
"target_cells":324,
**{c.lower():sum(1 for x in rows if x["classification"]==c) for c in classes},
"by_field_class":{f:dict(collections.Counter(x["classification"] for x in rows if x["field"]==f)) for f in ["method_path","payload_schema","persistence_owner"]},
"exact_unique_controls":len(set((x["page"],x["uid"]) for x in rows if x["classification"]=="REGISTRY_EXACT_UNIQUE")),
}
assert sum(summary[c.lower()] for c in classes)==324
report={"marker":"ACPOS-20260922-BATCH-29-OWNER-REGISTRY-AUDIT-V1","base_head":BASE_HEAD,"summary":summary,"rows":rows}
Path("__batch29_owner_registry_audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
Path("__batch29_summary.txt").write_text(json.dumps(summary,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
print("BATCH29_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
