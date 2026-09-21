from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-36-SYSTEM02-CANONICAL-CONTRACT-REMEDIATION-V1"
BASE_HEAD="d98109ef26fa8ac375b98c42a16e2fbc40fbce94"
S02="02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
CORE="ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
STR="ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
SYS="ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S02:"aa6df50d6505cc5c4848001c3c30b970c97982b8",
CORE:"07f9911a5647255f825ec3e92a3a2d79ab4869b2",
STR:"e6a5152c1abe905f749b80be9d3773a8418bfdde",
SYS:"b4d91493c8d9ec40c7d5f5cf05d9b2638eb84ad7",
LOGIC:"99030bd5790cf78815b20d0ea7547618bc74f601",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":CORE,
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":STR,
"SYS-01":SYS,
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]

OP_CONTRACTS={
"createConversationThread":{
 "method_path":"POST /v1/conversations/{CONVERSATION_ID}/threads",
 "payload_schema":"CreateConversationThreadRequest",
 "persistence_owner":"Shared Conversation Core / ConversationThread",
 "basis":"ConversationThread is the canonical ordered container with immutable thread_id; creation is owner-local and must preserve Conversation membership."
},
"analyzeConversationMessage":{
 "method_path":"POST /v1/conversations/{CONVERSATION_ID}/messages/{MESSAGE_ID}/analysis",
 "payload_schema":"AnalyzeConversationMessageRequest",
 "persistence_owner":"Shared Conversation Core / ConversationMessage + MessageRelationEdge",
 "basis":"Analysis is derived from an exact source message and must remain reconstructable through message identity and relation lineage."
},
"createConversationBranch":{
 "method_path":"POST /v1/conversations/{CONVERSATION_ID}/threads/{THREAD_ID}/branches",
 "payload_schema":"CreateConversationBranchRequest",
 "persistence_owner":"Shared Conversation Core / ConversationBranch + ConversationThread",
 "basis":"CREATE_BRANCH writes Branch + Thread from exact source message/context snapshot and fails closed on ambiguous source."
},
"createAssistantSummary":{
 "method_path":"POST /v1/conversations/{CONVERSATION_ID}/assistant-records/summaries",
 "payload_schema":"CreateAssistantSummaryRequest",
 "persistence_owner":"Shared Conversation Core / AssistantMeetingRecord + ConversationMemoryProjection",
 "basis":"Assistant summary is a derived/versioned artifact with exact source refs; silent rewrite is forbidden."
},
"createAssistantStructuredDecision":{
 "method_path":"POST /v1/conversations/{CONVERSATION_ID}/assistant-records/structured-decisions",
 "payload_schema":"CreateAssistantStructuredDecisionRequest",
 "persistence_owner":"Shared Conversation Core / AssistantMeetingRecord + DecisionMemoryLink",
 "basis":"Structured decision record is an assistant record plus a reference link to upstream decision authority; Assistant cannot auto-confirm Decision."
},
"sendConversationMessage":{
 "method_path":"POST /v1/conversations/{CONVERSATION_ID}/messages",
 "payload_schema":"SendConversationMessageRequest",
 "persistence_owner":"Shared Conversation Core / ConversationMessage + ConversationContextBinding",
 "basis":"SEND_MESSAGE writes Message; every governed AI request/response chain binds exact ContextSnapshot/Fingerprint."
},
"stopConversationGeneration":{
 "method_path":"POST /v1/conversations/{CONVERSATION_ID}/generation/stop",
 "payload_schema":"StopConversationGenerationRequest",
 "persistence_owner":"ACPOS-AI-03 Runtime Job + ACPOS-AI-02 ConversationAuditEnvelope",
 "basis":"STOP_RESPONSE requests cancellation; actual job semantics are owned by AI-03 while Conversation Core preserves audit/correlation identity."
},
"attachConversationContext":{
 "method_path":"POST /v1/conversations/{CONVERSATION_ID}/context/attachments",
 "payload_schema":"AttachConversationContextRequest",
 "persistence_owner":"Shared Conversation Core / ConversationAttachmentRef + ConversationContextBinding",
 "basis":"Attachment inclusion requires exact typed reference resolution, permission/classification, and context binding; raw object remains owned upstream."
},
"getStrategicConversationProjection":{
 "method_path":"GET /v1/conversations/{CONVERSATION_ID}/projections/strategy",
 "payload_schema":None,
 "persistence_owner":None,
 "basis":"READ_EXACT projection over persisted conversation identities; read does not create a persistence owner."
},
"getConversationAttachmentContext":{
 "method_path":"GET /v1/conversations/{CONVERSATION_ID}/attachments/{ATTACHMENT_ID}/context",
 "payload_schema":None,
 "persistence_owner":None,
 "basis":"READ_EXACT attachment/context resolution; ambiguous or unauthorized attachment blocks and no mutation occurs."
},
}

TARGETS={
"CORE-01-BTN-NEW-THREAD":{"page":"CORE-01","operation":"createConversationThread","fields":["method_path","payload_schema","persistence_owner"]},
"CORE-01-MENU-ANALYZE":{"page":"CORE-01","operation":"analyzeConversationMessage","fields":["method_path","payload_schema","persistence_owner"]},
"CORE-01-MENU-BRANCH":{"page":"CORE-01","operation":"createConversationBranch","fields":["method_path","payload_schema","persistence_owner"]},
"CORE-01-FLD-ASSISTANT-SUMMARY":{"page":"CORE-01","operation":"createAssistantSummary","fields":["method_path","payload_schema","persistence_owner"]},
"CORE-01-FLD-STRUCTURED-DECISION":{"page":"CORE-01","operation":"createAssistantStructuredDecision","fields":["method_path","payload_schema","persistence_owner"]},
"CORE-01-FLD-MESSAGE":{"page":"CORE-01","operation":"sendConversationMessage","fields":["payload_schema","persistence_owner"]},
"CORE-01-BTN-SEND":{"page":"CORE-01","operation":"sendConversationMessage","fields":["payload_schema","persistence_owner"]},
"STR-01-VIEW-CONVERSATION":{"page":"STR-01","operation":"getStrategicConversationProjection","fields":["method_path"]},
"STR-01-BTN-ATTACH":{"page":"STR-01","operation":"getConversationAttachmentContext","fields":["method_path"]},
"STR-01-INP-MESSAGE":{"page":"STR-01","operation":"sendConversationMessage","fields":["payload_schema","persistence_owner"]},
"STR-01-BTN-SEND":{"page":"STR-01","operation":"sendConversationMessage","fields":["payload_schema","persistence_owner"]},
"STR-01-BTN-STOP":{"page":"STR-01","operation":"stopConversationGeneration","fields":["payload_schema","persistence_owner"]},
"SYS-01-BTN-ATTACH":{"page":"SYS-01","operation":"attachConversationContext","fields":["method_path","payload_schema","persistence_owner"]},
"SYS-01-BTN-SEND":{"page":"SYS-01","operation":"sendConversationMessage","fields":["payload_schema","persistence_owner"]},
"SYS-01-BTN-STOP":{"page":"SYS-01","operation":"stopConversationGeneration","fields":["payload_schema","persistence_owner"]},
}

def norm(x): return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v): return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        fv={}
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            fv[f]=vals
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        base["_field_values"]=fv;out[uid]=base
    return out
def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():
    got=blob(f);assert got==s,(f,got,s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

page_files={"CORE-01":CORE,"STR-01":STR,"SYS-01":SYS}
pre_page_maps={p:compose(parse_controls(fn)) for p,fn in page_files.items()}
pre_owner=compose(parse_controls(S02))
expected_cells=sum(len(x["fields"]) for x in TARGETS.values())
assert expected_cells==34,expected_cells

# Preconditions.
for uid,t in TARGETS.items():
    op=t["operation"];page=t["page"];pr=pre_page_maps[page][uid]
    assert pr["operation"]==op,(uid,pr["operation"],op)
    assert pr["runtime_status"] in {"EFFECTFUL_EXACT","READ_EXACT"},(uid,pr["runtime_status"])
    assert uid in pre_owner,(uid,"OWNER_UID_MISSING")
    orow=pre_owner[uid]
    assert orow["operation"]==op,(uid,orow["operation"],op)
    assert orow["runtime_status"]==pr["runtime_status"],(uid,orow["runtime_status"],pr["runtime_status"])
    for f in t["fields"]:
        assert missing(pr.get(f,"")),(uid,"PAGE_NOT_MISSING",f,pr.get(f))
        # Existing owner may already carry send/stop route; target list intentionally excludes those method cells.
        if f!="method_path" or uid not in {"SYS-01-BTN-SEND","SYS-01-BTN-STOP"}:
            assert missing(orow.get(f,"")) or orow.get(f,"")==OP_CONTRACTS[op].get(f),(uid,"OWNER_CONFLICT",f,orow.get(f))

# Patch exact target fields in a Word control table; simplified/reference rows without target column are skipped.
def patch_fields(path,target_uids):
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact_index(h,"Control UID")
        if ci is None:continue
        idx={
          "method_path":hfind(h,"method / path","method","path"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "persistence_owner":hfind(h,"persistence owner"),
        }
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals):continue
            uid=vals[ci]
            if uid not in target_uids:continue
            spec=TARGETS[uid];contract=OP_CONTRACTS[spec["operation"]]
            for field in spec["fields"]:
                fi=idx[field]
                if fi is None or fi>=len(row.cells):continue
                new=contract[field]
                assert new is not None,(uid,field,"NULL_CONTRACT_VALUE")
                old=norm(row.cells[fi].text)
                if old==new:
                    hits[(uid,field)]+=1;continue
                assert missing(old),(path,uid,field,old,new)
                row.cells[fi].text=new;hits[(uid,field)]+=1
    expected={(uid,f) for uid in target_uids for f in TARGETS[uid]["fields"]}
    assert set(hits)==expected,(path,"PATCH_MISMATCH",sorted(expected-set(hits)),dict(hits))
    d.save(path);Document(path)
    return {f"{u}:{f}":n for (u,f),n in hits.items()}

patches={}
for page,path in page_files.items():
    uids=[u for u,t in TARGETS.items() if t["page"]==page]
    patches[f"PAGE_{page}"]=patch_fields(path,uids)
patches["SYSTEM02_OWNER"]=patch_fields(S02,list(TARGETS))

def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.33);sec.bottom_margin=Inches(.33);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=4.2):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for rr in rows:
        cells=t.add_row().cells
        for i,v in enumerate(rr):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
    return t

# Canonical owner registry annex.
d=Document(S02);add_landscape(d);d.add_heading("Batch 36 · Shared Conversation Core Canonical Contract Registry",level=1)
p=d.add_paragraph();p.add_run(f"[{MARK}::SYSTEM02] ").bold=True
p.add_run("This registry closes the System 02 Method/Path, Payload/Schema and Persistence Owner gaps using existing Shared Conversation Core entities, transitions and fail-closed invariants. Routes stay inside the existing /v1/conversations namespace. Read operations remain non-mutating. Runtime execution is not asserted.")
rows=[]
for op,c in OP_CONTRACTS.items():
    status="READ_EXACT" if op in {"getStrategicConversationProjection","getConversationAttachmentContext"} else "EFFECTFUL_EXACT"
    rows.append([op,status,c["method_path"],c["payload_schema"] or "N/A_READ_NO_BODY",c["persistence_owner"] or "N/A_READ_ONLY",c["basis"]])
add_table(d,["Operation","Runtime Status","Method / Path","Payload / Schema","Persistence Owner","Authority / Semantics"],rows,3.55)
d.save(S02);Document(S02)

# Page ledgers.
for page,path in page_files.items():
    d=Document(path);add_landscape(d);d.add_heading("Batch 36 · Shared Conversation Core Contract Binding Ledger",level=1)
    p=d.add_paragraph();p.add_run(f"[{MARK}::PAGE::{page}] ").bold=True
    p.add_run("Only the fields listed in the Current 324-cell denominator are synchronized. Existing Action/Gate/Permission/Operation/Runtime Owner/Runtime Status values remain unchanged.")
    rr=[]
    for uid,tg in TARGETS.items():
        if tg["page"]!=page:continue
        c=OP_CONTRACTS[tg["operation"]]
        rr.append([uid,tg["operation"],", ".join(tg["fields"]),c["method_path"],c["payload_schema"] or "N/A",c["persistence_owner"] or "N/A"])
    add_table(d,["Control UID","Operation","Fields Closed","Method / Path","Payload / Schema","Persistence Owner"],rr,3.85)
    d.save(path);Document(path)

# Post validation.
post_page_maps={p:compose(parse_controls(fn)) for p,fn in page_files.items()}
post_owner=compose(parse_controls(S02))
for uid,t in TARGETS.items():
    c=OP_CONTRACTS[t["operation"]]
    for m,label in [(post_page_maps[t["page"]],"PAGE"),(post_owner,"OWNER")]:
        r=m[uid]
        for f in t["fields"]:
            assert r[f]==c[f],(label,uid,f,r[f],c[f])

# Full remaining denominator across 18 page docs.
all_page_maps={p:compose(parse_controls(fn)) for p,fn in PAGES.items()}
remaining=[]
for page,m in all_page_maps.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if st in {"EFFECTFUL_EXACT","READ_EXACT"} and not missing(op) and missing(r.get("method_path","")):
            remaining.append(("method_path",page,uid))
        if st=="EFFECTFUL_EXACT" and not missing(op) and missing(r.get("payload_schema","")):
            remaining.append(("payload_schema",page,uid))
        if st=="EFFECTFUL_EXACT" and not missing(op) and missing(r.get("persistence_owner","")):
            remaining.append(("persistence_owner",page,uid))
rc=collections.Counter(f for f,_,_ in remaining)
assert len(remaining)==290,(len(remaining),rc)
assert rc=={"method_path":63,"payload_schema":119,"persistence_owner":108},rc

# Same-page real route collision check after new bindings.
collisions=[]
for page,m in all_page_maps.items():
    by=collections.defaultdict(list)
    for uid,r in m.items():
        mp=r.get("method_path","");op=r.get("operation","")
        if missing(mp) or mp in {"LOCAL LOCAL","NO_PUBLIC_API_BY_AUTHORITY"}:continue
        by[mp].append((uid,op))
    for mp,rr in by.items():
        ops=sorted(set(op for _,op in rr if not missing(op)))
        if len(ops)>1:collisions.append({"page":page,"method_path":mp,"operations":ops,"controls":rr})
assert not collisions,collisions[:10]

logic=Document(LOGIC);add_landscape(logic);logic.add_heading("Batch 36 · System 02 Canonical Contract Closure",level=1)
p=logic.add_paragraph();p.add_run(f"[{MARK}] ").bold=True
p.add_run("System 02 Shared Conversation Core closes its 34 contract cells from owner-local evidence. Existing send/stop transport is preserved; new operations use the same /v1/conversations namespace. Persistence owners reference canonical conversation entities or, for stop-generation runtime cancellation, the AI-03 runtime job plus System 02 audit envelope. No Production/runtime completion is claimed.")
add_table(logic,["Metric","Count / Value","Result"],[
["Pre contract denominator",324,"Method 71 + Payload 132 + Persistence 121"],
["System 02 cells closed",34,"8 Method + 13 Payload + 13 Persistence"],
["Post contract denominator",290,"Method 63 + Payload 119 + Persistence 108"],
["Same-page real route collisions",0,"PASS"],
["Runtime execution claimed","False","Design-contract closure only"],
],4.1)
machine={
"marker":MARK,"base_head":BASE_HEAD,"pre_contract_cells":324,"system02_cells_closed":34,
"method_path_closed":8,"payload_schema_closed":13,"persistence_owner_closed":13,
"post_contract_cells":290,"remaining_method_path":63,"remaining_payload_schema":119,"remaining_persistence_owner":108,
"same_page_route_collisions":0,"runtime_execution_claimed":False,
}
logic.add_paragraph("BATCH36_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

changed=[S02,CORE,STR,SYS,LOGIC]
report={"machine":machine,"contracts":OP_CONTRACTS,"targets":TARGETS,"patches":patches,
"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed}}
Path("__batch36_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH36="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
