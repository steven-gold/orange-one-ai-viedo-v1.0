from pathlib import Path
import json,re
from docx import Document

targets=[
"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
]
kw=re.compile(r"(control|action|operation|port|gate|permission|runtime|api|method|path|binding|registry|denominator|uid|控制|按鈕|功能|操作|權限|閘|接口|埠)",re.I)
out={}
for name in targets:
    d=Document(name)
    tables=[]
    for ti,t in enumerate(d.tables,1):
        rr=[[" ".join(c.text.split()) for c in r.cells] for r in t.rows]
        if not rr: continue
        hdr=rr[0]
        sample=rr[1:6]
        blob=" | ".join(" | ".join(r) for r in rr[:8])
        if kw.search(blob):
            tables.append({
              "index":ti,
              "rows":len(rr),
              "cols":max(len(r) for r in rr),
              "header":hdr,
              "sample":sample,
              "all_rows":rr if len(rr)<=220 else rr[:220],
            })
    out[name]={"paragraphs":len(d.paragraphs),"tables":len(d.tables),"candidate_tables":tables}
Path("__temp_reports__").mkdir(exist_ok=True)
Path("__temp_reports__/authority_table_shapes.json").write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding="utf-8")
