from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib

MARK='ACPOS-20260921-BATCH-16-OPERATION-OWNER-CONTRACT-CLOSURE-V1'
BASE_HEAD='dc72b52f0e147cc934ad3aaf87120ae44f2d3018'
LOGIC='ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx'
S02='02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx'
S04='04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx'
S05='05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx'
S06='06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx'
S09='09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx'
PAGES={
'AIAPI-01':'ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx',
'ASSET-01':'ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx',
'CORE-01':'ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx',
'DB-01':'ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx',
'DEV-01':'ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx',
'EDIT-01':'ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx',
'ERP-01':'ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx',
'IAM-01':'ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx',
'INFO-01':'ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx',
'KB-01':'ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx',
'QA-01':'ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx',
'SG-02':'ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx',
'SOC-01':'ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx',
'STR-01':'ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx',
'SYS-01':'ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx',
'VIDEO-01':'ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx',
'WB-01':'ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx',
'ADMIN-STR-01':'ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx',
}
OWNER_DOCS={S02:'System 02 / Shared Conversation Core',S04:'System 04 / CORE Blueprint',S05:'System 05 / Creative Production',S06:'System 06 / Enterprise Business',S09:'System 09 / System Engineer'}
FIELDS=['control','type','label','action','gate','permission','payload_schema','operation','method_path','runtime_owner','persistence_owner','runtime_status']

EXPECTED_SHA={
S02:'cdcbed04e42c206b1c99b2743d0c60cd18c42bdf',
S04:'20591ee8233ff74aeb3d82d6bb048ecc5a60d65c',
S05:'71ca2680d5bb971c87368097565f9b7f4deaa5cc',
S06:'526eed38a8be88a32be7510549b00f4e43c59390',
S09:'94295ac8f43d97b105933276ba23a279c6310885',
PAGES['ASSET-01']:'56a72348c8fbb93dafdb91bf7a5634cb682293d8',
PAGES['CORE-01']:'a140edcba1d76d5d38b088ec856f6c2bd4692978',
PAGES['DEV-01']:'220208f480b71a415f53f8413f83a3edb9b45bec',
PAGES['EDIT-01']:'6b5f465ab1eb6c445dfb4f19162c3939a961a203',
PAGES['QA-01']:'dd5533f7ab0ffec491f2d6ab16c86c84cead8a7c',
PAGES['STR-01']:'2273aa79858e3a0b4bb7c9ff99cd192d54220bc4',
PAGES['SYS-01']:'10f4eec724a08e5807633a3de009209834b9831c',
PAGES['VIDEO-01']:'5c524d869b8e255b1706fbc3faa05918d0ab9524',
LOGIC:'43cd0a065adf7a9ffb61c7c6f824107201d04258',
}

def C(page,uid,owner_doc,fields,operation=None,runtime_owner=None):
    return {'page':page,'uid':uid,'owner_doc':owner_doc,'fields':fields,'operation':operation,'runtime_owner':runtime_owner}

contracts=[]
for uid,op in [
('CORE-01-BTN-NEW-THREAD','createConversationThread'),('CORE-01-MENU-ANALYZE','analyzeConversationMessage'),
('CORE-01-MENU-BRANCH','createConversationBranch'),('CORE-01-FLD-ASSISTANT-SUMMARY','createAssistantSummary'),
('CORE-01-FLD-STRUCTURED-DECISION','createAssistantStructuredDecision'),('CORE-01-FLD-MESSAGE','sendConversationMessage'),
('CORE-01-BTN-SEND','sendConversationMessage')]: contracts.append(C('CORE-01',uid,S02,['operation'],operation=op))
contracts += [
C('STR-01','STR-01-VIEW-CONVERSATION',S02,['operation'],operation='getStrategicConversationProjection'),
C('STR-01','STR-01-BTN-ATTACH',S02,['operation'],operation='getConversationAttachmentContext'),
C('SYS-01','SYS-01-BTN-ATTACH',S02,['runtime_owner'],runtime_owner='Shared Conversation Core'),
]
for uid,op in [
('CORE-01-BTN-PROJECT-CREATE','createProjectDraft'),('CORE-01-BTN-TOPIC-CREATE','createTopicDraft'),
('CORE-01-FLD-EVALUATION','evaluateCoreCandidate'),('CORE-01-FLD-HUMAN-DECISION','recordCoreHumanDecision'),
('CORE-01-BTN-CANDIDATE-CREATE','createCoreCandidate'),('CORE-01-BTN-CANDIDATE-CONFIRM','acceptCoreCandidate'),
('CORE-01-BTN-RETURN-MODIFY','requestCoreCandidateRevision'),('CORE-01-BTN-PROJECT-VALIDATE','validateProjectDraft'),
('CORE-01-BTN-PROJECT-CONFIRM','confirmProjectDraft'),('CORE-01-BTN-STORY-CANDIDATE','createStoryCandidateSet'),
('CORE-01-BTN-DNA-LOCK','requestDNALock'),('CORE-01-BTN-CORE-REVIEW','submitCoreReview'),
('CORE-01-BTN-PROJECT-LOCK','requestMotherLock'),('CORE-01-BTN-BLUEPRINT-CREATE','createProjectBlueprintCandidate'),
('CORE-01-BTN-BLUEPRINT-VALIDATE','validateProjectBlueprint'),('CORE-01-BTN-BLUEPRINT-APPROVE','approveProjectBlueprint'),
('CORE-01-BTN-CHILD-LOCK','requestChildLock'),('CORE-01-BTN-CANONICAL-SCRIPT','getCanonicalScript'),
('CORE-01-BTN-CANDIDATE-COMPARE','compareCoreCandidates')]: contracts.append(C('CORE-01',uid,S04,['operation'],operation=op))
asset_ops=[
('ASSET-01-BTN-EXECUTE','startAssetProductionFlow'),('ASSET-01-BTN-CORRECTION-EXECUTE','executeAssetCorrection'),
('ASSET-01-BTN-LAYER-DOC-CREATE','createAssetLayerDocument'),('ASSET-01-BTN-LAYER-DOC-UPDATE','updateAssetLayerDocument'),
('ASSET-01-BTN-LAYER-ADD','addAssetLayer'),('ASSET-01-BTN-LAYER-DELETE','deleteAssetLayer'),
('ASSET-01-BTN-LAYER-DUPLICATE','duplicateAssetLayer'),('ASSET-01-BTN-LAYER-REORDER','reorderAssetLayers'),
('ASSET-01-CTL-LAYER-PROPERTIES','updateAssetLayerProperties'),('ASSET-01-CTL-LAYER-MASK','updateAssetLayerMask'),
('ASSET-01-BTN-PATCH-CREATE','createAssetPatch'),('ASSET-01-BTN-PATCH-PREVIEW','previewAssetPatch'),
('ASSET-01-BTN-PATCH-ACCEPT','acceptAssetPatch'),('ASSET-01-BTN-PATCH-REJECT','rejectAssetPatch'),
('ASSET-01-BTN-RETRY','retryAssetTask'),('ASSET-01-BTN-EVALUATE','evaluateAssetCandidate'),
('ASSET-01-BTN-CONFIRM','confirmAssetCandidate'),('ASSET-01-BTN-HANDOFF','handoffAssetToVideo')]
for uid,op in asset_ops:
    fields=['operation'] if uid=='ASSET-01-BTN-CORRECTION-EXECUTE' else ['operation','runtime_owner']
    contracts.append(C('ASSET-01',uid,S05,fields,operation=op,runtime_owner='SHARED_PRODUCTION_OPERATION_RUNTIME'))
edit_ops=[
('EDIT-01-LBL-API-JOB','getEditCorrectionJobStatus'),('EDIT-01-LBL-BINDING-FINGERPRINT','getEditCorrectionBindingFingerprint'),
('EDIT-01-LBL-CURRENT-SCRIPT-SECTION','getEditCorrectionScriptSection'),('EDIT-01-LBL-LIPSYNC-SYNC-BINDING','getEditLipSyncBinding'),
('EDIT-01-LBL-SUB-SYNC-BINDING','getEditSubtitleSyncBinding'),('EDIT-01-LST-IMPORT-QUEUE','getEditImportQueue'),
('EDIT-01-LST-VOICE-TAKES','getEditVoiceTakes'),('EDIT-01-PNL-API-CANDIDATE','getEditCorrectionCandidatePreview'),
('EDIT-01-PNL-FINAL-PREVIEW','getEditFinalPreview')]
for uid,op in edit_ops:
    fields=['operation'] if uid=='EDIT-01-PNL-FINAL-PREVIEW' else ['operation','runtime_owner']
    contracts.append(C('EDIT-01',uid,S05,fields,operation=op,runtime_owner='CLIENT/ORCHESTRATION'))
for uid,op in [('QA-01-BTN-FINDING-POINT','setQAFindingPoint'),('QA-01-BTN-EVIDENCE-IN','setQAEvidenceRangeIn'),('QA-01-BTN-EVIDENCE-OUT','setQAEvidenceRangeOut')]:
    contracts.append(C('QA-01',uid,S05,['operation'],operation=op))
for uid,op in [
('VIDEO-01-BTN-EXECUTE','startVideoGenerationFlow'),('VIDEO-01-BTN-CONFIRM','confirmVideoCandidate'),
('VIDEO-01-BTN-HANDOFF','handoffVideoToEdit'),('VIDEO-01-BTN-EXEC-CORRECTION','executeVideoCorrection'),
('VIDEO-01-BTN-RETRY','retryVideoTask'),('VIDEO-01-BTN-EVALUATE','evaluateVideoCandidate'),
('VIDEO-01-BTN-FINDING','createVideoFinding')]: contracts.append(C('VIDEO-01',uid,S05,['operation'],operation=op))
for uid,owner in [
('DEV-01-BTN-CAMPAIGN-APPROVE','Existing outreach owner'),('DEV-01-BTN-EMAIL-DISPATCH','Existing outreach owner'),
('DEV-01-BTN-CANDIDATE-CREATE','Existing enterprise development owner'),('DEV-01-BTN-DISCOVERY-PAUSE','Shared Acquisition / Knowledge owner'),
('DEV-01-BTN-DISCOVERY-RESUME','Shared Acquisition / Knowledge owner'),('DEV-01-BTN-DISCOVERY-START','Shared Acquisition / Knowledge owner'),
('DEV-01-BTN-DISCOVERY-STOP','Shared Acquisition / Knowledge owner'),('DEV-01-BTN-MERGE','Existing enterprise development owner'),
('DEV-01-BTN-MERGE-PREVIEW','Existing enterprise development owner')]: contracts.append(C('DEV-01',uid,S06,['runtime_owner'],runtime_owner=owner))
contracts += [
C('SYS-01','SYS-01-BTN-CANDIDATE-CREATE',S09,['runtime_owner'],runtime_owner='ChangePlanner'),
C('SYS-01','SYS-01-BTN-CR-CREATE',S09,['runtime_owner'],runtime_owner='SystemEngineer'),
]
assert len(contracts)==77,len(contracts)
assert sum(len(c['fields']) for c in contracts)==102,sum(len(c['fields']) for c in contracts)

def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest()
for f,s in EXPECTED_SHA.items():
    got=blob(f);assert got==s,(f,got,s)

def norm(x):return re.sub(r'\s+',' ',(x or '').replace('\n',' | ').strip())
def missing(v):return v in {'','—','-','SOURCE_NOT_DEFINED'}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def display_only_type(t):
    u=(t or '').upper()
    return u in {'READONLY','LABEL','TEXT','BADGE','STATUS','METRIC','KPI','DIVIDER','ICON','PANEL','CARD','VIEW','LIST','TABLE','CHIP','TAG','DISPLAY'} or any(x in u for x in ['READONLY','LABEL','DISPLAY','METRIC','KPI','STATUS'])
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,'control uid')
        if ci is None:continue
        idx={'control':ci,'type':hfind(h,'type'),'label':hfind(h,'label','顯示名稱'),'action':hfind(h,'action uid'),'gate':hfind(h,'gate uid','gate'),'permission':hfind(h,'permission','auth resource'),'payload_schema':hfind(h,'payload / schema','payload','schema'),'operation':hfind(h,'operation'),'method_path':hfind(h,'method / path','method','path'),'runtime_owner':hfind(h,'runtime owner'),'persistence_owner':hfind(h,'persistence owner'),'runtime_status':hfind(h,'runtime status')}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ''
            if not uid or uid in {'—','-'}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else '') for k,i in idx.items()}
            rec.update({'table':ti+1,'row':ri,'headers':h});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r['control']].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,'')) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,'') for r in rs if not missing(r.get(f,''))))
            if missing(base.get(f,'')) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out

def is_wb_na(r): return missing(r.get('action','')) and r.get('type')=='SECTION_OPEN' and r.get('gate')=='PAGE_READ' and r.get('permission')=='workspace.dashboard.view' and r.get('payload_schema')=='N/A_READ_PROJECTION' and r.get('operation')=='getDashboardReadModel' and r.get('method_path')=='GET /v1/dashboard/read-model' and r.get('runtime_owner')=='DASHBOARD_READ_MODEL' and r.get('runtime_status')=='READ_EXACT'
def is_aiapi_na(r): return missing(r.get('action','')) and r.get('type')=='BUTTON_OR_ROW_ACTION' and r.get('payload_schema')=='AIAPI Page Operation-specific Form / Provider Profile Field Contract' and not missing(r.get('permission','')) and not missing(r.get('operation','')) and not missing(r.get('method_path','')) and not missing(r.get('runtime_owner','')) and r.get('runtime_status') in {'EFFECTFUL_EXACT','READ_EXACT'}
def is_edit_na(r): return missing(r.get('operation','')) and r.get('runtime_status')=='LOCAL_WORKING_DRAFT_EXACT' and r.get('method_path')=='NO_PUBLIC_API_BY_AUTHORITY' and not missing(r.get('action','')) and not missing(r.get('gate','')) and not missing(r.get('permission','')) and not missing(r.get('runtime_owner',''))
def is_iam_op_na(r): return missing(r.get('operation','')) and r.get('runtime_status')=='ORCHESTRATES_EXISTING_EXACT_OPERATIONS' and not missing(r.get('action','')) and not missing(r.get('gate','')) and not missing(r.get('permission',''))
def is_iam_owner_na(r): return missing(r.get('runtime_owner','')) and r.get('runtime_status')=='ORCHESTRATES_EXISTING_EXACT_OPERATIONS' and not missing(r.get('action','')) and not missing(r.get('gate','')) and not missing(r.get('permission',''))
DB_OWNER='Database Read Model / DB-01';ERP_CONNECTOR_OWNER='ERPConnectorService';ERP_FINANCE_OWNER='FinanceGuardrailService'
ERP_GROUPS={
'ERP-01-ACT-READ-CONNECTOR':('ERP-01-GATE-CONNECTOR-READ','ERP-01-PERM-CONNECTOR',ERP_CONNECTOR_OWNER),
'ERP-01-ACT-READ-FINANCE':('ERP-01-GATE-FINANCE','ERP-01-PERM-FINANCE',ERP_FINANCE_OWNER),
'ERP-01-ACT-READ-SYNC':('ERP-01-GATE-SYNC-READ','ERP-01-PERM-SYNC',ERP_CONNECTOR_OWNER)}
def is_db_read_model_operation_na(r): return missing(r.get('operation','')) and missing(r.get('method_path','')) and r.get('runtime_status')=='UI_LOCAL_OR_READ_SOURCE' and r.get('runtime_owner')==DB_OWNER and not missing(r.get('action','')) and not missing(r.get('gate','')) and not missing(r.get('permission',''))
def is_erp_read_field_operation_na(r):
    action=r.get('action','')
    if action not in ERP_GROUPS:return False
    gate,perm,owner=ERP_GROUPS[action]
    return r.get('type')=='READONLY' and missing(r.get('operation','')) and missing(r.get('method_path','')) and r.get('runtime_status')=='UI_LOCAL_OR_READ_SOURCE' and r.get('gate')==gate and r.get('permission')==perm and r.get('runtime_owner')==owner
def classify_current(page,field,r):
    v=r.get(field,'');st=r.get('runtime_status','')
    if not missing(v):return 'BOUND'
    if 'UI_LOCAL_EXACT' in st:
        if field in {'operation','runtime_owner'}:return 'LEGITIMATE_NA_UI_LOCAL'
        if field=='action' and display_only_type(r.get('type','')):return 'LEGITIMATE_NA_UI_LOCAL'
    if 'OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE' in st and field=='operation':return 'CLOSED_BY_OWNER_LOCAL_COMMAND'
    if 'READ_EXACT' in st and display_only_type(r.get('type','')):
        if field=='action':return 'LEGITIMATE_NA_READ_PRESENTATION'
        if field=='operation' and not missing(r.get('runtime_owner','')):return 'READ_OWNER_BOUND_OPERATION_UNSPECIFIED'
    if field=='runtime_owner' and any(x in st for x in ['SPEC_EXACT_RUNTIME_BLOCKED','SPEC_EXACT_RUNTIME_NOT_EXECUTED','RUNTIME_IMPLEMENTATION_BLOCKED']):return 'KNOWN_RUNTIME_BLOCKER'
    if v=='SOURCE_NOT_DEFINED':return 'SOURCE_RESOLUTION_REQUIRED'
    if page=='WB-01' and field=='action' and is_wb_na(r):return 'LEGITIMATE_NA_READ_PROJECTION_ACTION'
    if page=='AIAPI-01' and field=='action' and is_aiapi_na(r):return 'LEGITIMATE_NA_DIRECT_OPERATION_ACTION'
    if page=='EDIT-01' and field=='operation' and is_edit_na(r):return 'LEGITIMATE_NA_LOCAL_WORKING_DRAFT_OPERATION'
    if page=='IAM-01' and field=='operation' and is_iam_op_na(r):return 'LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_OPERATION'
    if page=='IAM-01' and field=='runtime_owner' and is_iam_owner_na(r):return 'LEGITIMATE_NA_ORCHESTRATION_AGGREGATE_RUNTIME_OWNER'
    if page=='DB-01' and field=='operation' and is_db_read_model_operation_na(r):return 'LEGITIMATE_NA_DB_READ_MODEL_OPERATION'
    if page=='ERP-01' and field=='operation' and is_erp_read_field_operation_na(r):return 'LEGITIMATE_NA_ERP_READ_PROJECTION_FIELD_OPERATION'
    return 'DEFINITION_BINDING_GAP'
def inventory():
    rows=[];counts=collections.Counter();maps={}
    for page,fn in PAGES.items():
        bm=compose(parse_controls(fn));maps[page]=bm
        for uid,r in bm.items():
            for field in ['action','gate','permission','operation','runtime_owner']:
                cls=classify_current(page,field,r);counts[cls]+=1
                if cls=='DEFINITION_BINDING_GAP':rows.append((page,uid,field))
    return rows,counts,maps

pre_rows,pre_counts,pre_maps=inventory()
assert len(pre_rows)==122,len(pre_rows)
target_fields={(c['page'],c['uid'],f) for c in contracts for f in c['fields']}
assert len(target_fields)==102,len(target_fields)
assert target_fields.issubset(set(pre_rows)),sorted(target_fields-set(pre_rows))[:20]
st=collections.Counter()
for page,uid,field in sorted(target_fields):
    r=pre_maps[page][uid];status=r.get('runtime_status','')
    assert status in {'EFFECTFUL_EXACT','READ_EXACT'},(page,uid,field,status)
    st[status]+=1
assert st=={'EFFECTFUL_EXACT':81,'READ_EXACT':21},st
assert collections.Counter(f for _,_,f in target_fields)=={'operation':65,'runtime_owner':37}

owner_text={}
for path in OWNER_DOCS:
    d=Document(path)
    owner_text[path]='\n'.join([p.text for p in d.paragraphs]+[cell.text for t in d.tables for row in t.rows for cell in row.cells])
for c in contracts:
    assert c['uid'] in owner_text[c['owner_doc']],(c['owner_doc'],c['uid'],'UID_NOT_FOUND_IN_CANONICAL_OWNER')

by_page=collections.defaultdict(list)
for c in contracts:by_page[c['page']].append(c)
def patch_page(path,cs):
    wanted={(c['uid'],f):c[f] for c in cs for f in c['fields']}
    hits=collections.Counter();d=Document(path)
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=exact_index(h,'Control UID');usi=exact_index(h,'UID Status')
        oi=exact_index(h,'Operation');ri=exact_index(h,'Runtime Owner')
        if ci is None or usi is None:continue
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals) or usi>=len(vals):continue
            uid=vals[ci]
            if not any(k[0]==uid for k in wanted):continue
            assert vals[usi]=='EXISTING_SOURCE_UID',(path,uid,vals[usi])
            for field,idx in [('operation',oi),('runtime_owner',ri)]:
                key=(uid,field)
                if key not in wanted:continue
                assert idx is not None,(path,uid,field,'COLUMN_MISSING')
                old=norm(row.cells[idx].text);new=wanted[key]
                if missing(old):row.cells[idx].text=new
                else:assert old==new,(path,uid,field,old,new)
                hits[key]+=1
    assert set(hits)==set(wanted),(path,'PATCH_MISMATCH',sorted(set(wanted)-set(hits)),dict(hits))
    d.save(path);Document(path)
    return {f'{u}:{f}':n for (u,f),n in hits.items()}
page_patch={}
for page,cs in by_page.items():page_patch[page]=patch_page(PAGES[page],cs)

post_patch_maps={page:compose(parse_controls(PAGES[page])) for page in by_page}
def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.33);sec.bottom_margin=Inches(.33);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement('w:tblHeader');e.set(qn('w:val'),'true');trPr.append(e)
def shade(cell,fill='EDE9FE'):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement('w:shd');e.set(qn('w:fill'),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=4.4):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for rr in rows:
        cells=t.add_row().cells
        for i,v in enumerate(rr):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
    return t
def owner_basis(path):
    return {
    S02:'Shared Conversation Core semantics / Thread-Branch-Message-Assistant-Resume boundary',
    S04:'CORE Entity × Operation contract / Project-Story-DNA-Blueprint-Lock-Handoff boundary',
    S05:'Creative Production runtime / ASSET-VIDEO-EDIT-QA owner boundary',
    S06:'Existing enterprise development / outreach / shared acquisition owner boundary',
    S09:'System Engineer component/event owner boundary'}[path]
by_owner=collections.defaultdict(list)
for c in contracts:by_owner[c['owner_doc']].append(c)
for path,cs in by_owner.items():
    d=Document(path)
    alltxt='\n'.join([p.text for p in d.paragraphs]+[cell.text for t in d.tables for row in t.rows for cell in row.cells])
    assert MARK not in alltxt,(path,'BATCH16_ALREADY_APPLIED')
    add_landscape(d)
    d.add_heading('Batch 16 · Canonical Operation / Runtime Owner Contract Closure',level=1)
    p=d.add_paragraph();p.add_run('['+MARK+'::CANONICAL_OWNER] ').bold=True
    p.add_run('This ledger materializes the missing Operation and/or Runtime Owner definitions inside the already-resolved Canonical Owner. It does not create a parallel owner, and it does not synthesize public API routes. Existing Method/Path values are preserved; where Method/Path is absent, transport binding remains implementation-required.')
    rows=[]
    for c in sorted(cs,key=lambda x:(x['page'],x['uid'])):
        r=post_patch_maps[c['page']][c['uid']]
        method=r.get('method_path','')
        transport=method if not missing(method) else 'NOT_DEFINED_BY_CURRENT_WORD / implementation binding required'
        rows.append([c['page'],c['uid'],','.join(c['fields']),r.get('operation','—'),r.get('runtime_owner','—'),r.get('runtime_status',''),transport,owner_basis(path)])
    add_table(d,['Page','Target UID','Closed Field(s)','Operation','Runtime Owner','Runtime Status','Transport Handling','Authority Basis'],rows,3.9)
    d.save(path);Document(path)

for page,cs in by_page.items():
    path=PAGES[page];d=Document(path)
    alltxt='\n'.join([p.text for p in d.paragraphs]+[cell.text for t in d.tables for row in t.rows for cell in row.cells])
    assert MARK not in alltxt,(path,'BATCH16_ALREADY_APPLIED')
    add_landscape(d)
    d.add_heading('Batch 16 · Operation / Runtime Owner Binding Ledger',level=1)
    p=d.add_paragraph();p.add_run('['+MARK+'::PAGE::'+page+'] ').bold=True
    p.add_run('The existing product-control registry cells above are updated only for the exact fields authorized by the matching Canonical Owner contract. No Method/Path is invented and no Runtime/Production execution is claimed.')
    rows=[]
    final=compose(parse_controls(path))
    for c in sorted(cs,key=lambda x:x['uid']):
        r=final[c['uid']]
        rows.append([c['uid'],','.join(c['fields']),r.get('operation','—'),r.get('runtime_owner','—'),r.get('runtime_status',''),r.get('method_path','') if not missing(r.get('method_path','')) else '—','BOUND_BY_BATCH_16'])
    add_table(d,['Target UID','Closed Field(s)','Operation','Runtime Owner','Runtime Status','Existing Method / Path','Result'],rows,4.1)
    d.save(path);Document(path)

post_rows,post_counts,post_maps=inventory()
removed=set(pre_rows)-set(post_rows);added=set(post_rows)-set(pre_rows)
assert removed==target_fields,{'missing':sorted(target_fields-removed)[:20],'unexpected':sorted(removed-target_fields)[:20]}
assert not added,sorted(added)[:20]
assert len(post_rows)==20,len(post_rows)
remaining=collections.Counter(f for _,_,f in post_rows)
assert remaining=={'action':7,'gate':8,'permission':5},remaining
assert ('SYS-01','SYS-01-BTN-NAV-OPEN','gate') in set(post_rows),post_rows
for c in contracts:
    r=post_maps[c['page']][c['uid']]
    for f in c['fields']: assert r.get(f)==c[f],(c['page'],c['uid'],f,r.get(f),c[f])

logic=Document(LOGIC)
lt='\n'.join([p.text for p in logic.paragraphs]+[cell.text for t in logic.tables for row in t.rows for cell in row.cells])
assert MARK not in lt,'LOGIC_BATCH16_ALREADY_APPLIED'
add_landscape(logic)
logic.add_heading('Batch 16 · Canonical Operation / Runtime Owner Contract Closure',level=1)
p=logic.add_paragraph();p.add_run('['+MARK+'] ').bold=True
p.add_run('Batch 16 closes the 102 remaining Operation/Runtime Owner definition gaps after Batch 15. The closure is split by Current runtime semantics: 81 EFFECTFUL_EXACT fields and 21 READ_EXACT fields. Each value is materialized through its already-resolved Canonical Owner. No public Method/Path is synthesized. This is Word authority closure only, not implementation or Production execution evidence.')
add_table(logic,['Class','Count','Result'],[
['Pre-batch Definition Binding Gap',122,'Batch 15 exact denominator'],
['EFFECTFUL_EXACT Operation/Owner fields',81,'Closed through Canonical Owner contracts'],
['READ_EXACT Operation/Owner fields',21,'Closed through Canonical Owner contracts'],
['Operation fields closed',65,'Formal Operation identities materialized'],
['Runtime Owner fields closed',37,'Existing owner boundary names materialized'],
['Public Method/Path invented',0,'None; missing transports remain implementation-required'],
['Post-batch Definition Binding Gap',20,'Action 7 + Gate 8 + Permission 5'],
['Preserved SYS Gate conflict',1,'SYS-01-BTN-NAV-OPEN remains fail-closed inside Gate denominator']],4.3)
add_table(logic,['Canonical Owner','Field closures'],[[OWNER_DOCS[p],sum(len(c['fields']) for c in cs)] for p,cs in sorted(by_owner.items())],4.6)
machine={'marker':MARK,'base_head':BASE_HEAD,'pre_definition_gaps':122,'unique_controls_touched':77,'effectful_fields_closed':81,'read_fields_closed':21,'operation_fields_closed':65,'runtime_owner_fields_closed':37,'gaps_closed':102,'post_definition_gaps':20,'remaining_action':7,'remaining_gate_including_conflict':8,'remaining_permission':5,'preserved_sys_gate_conflict':1,'public_routes_invented':0,'runtime_execution_claimed':False}
logic.add_paragraph('BATCH16_MACHINE_JSON='+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(',',':')))
logic.save(LOGIC);Document(LOGIC)

changed_docs=sorted(set([LOGIC]+list(by_owner)+[PAGES[p] for p in by_page]))
assert len(changed_docs)==14,(len(changed_docs),changed_docs)
report={'machine':machine,'target_field_count':len(target_fields),'target_control_count':len(contracts),'runtime_status_counts':dict(st),'pre_gap_field_counts':dict(collections.Counter(f for _,_,f in pre_rows)),'post_gap_field_counts':dict(remaining),'contracts':contracts,'page_patch_counts':page_patch,'remaining_rows':[{'page':p,'uid':u,'field':f} for p,u,f in sorted(post_rows)],'changed_docs':changed_docs,'output_blob_sha':{f:blob(f) for f in changed_docs}}
Path('__batch16_remediation_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print('BATCH16_REMEDIATION='+json.dumps(machine,ensure_ascii=False,sort_keys=True))
