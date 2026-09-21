from pathlib import Path
from docx import Document
import json,re,collections

SYSTEMS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}

def norm(x): return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def get(vals,i): return "" if i is None or i>=len(vals) else norm(vals[i])
def missing(v): return v in {"","—","-","SOURCE_NOT_DEFINED"}

DISPLAY_ONLY_TYPES={"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"}
def display_only(r):
    typ=r.get("type","").upper()
    return typ in DISPLAY_ONLY_TYPES or any(x in typ for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]

def parse_control_rows(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows: continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(h,"control uid")
        if ci is None: continue
        idx={
          "control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),"operation":hfind(h,"operation"),
          "runtime_owner":hfind(h,"runtime owner"),"runtime_status":hfind(h,"runtime status")
        }
        for ri,r in enumerate(t.rows[1:],2):
            vals=[c.text for c in r.cells]
            uid=get(vals,ci)
            if not uid or uid in {"—","-"} or not re.search(r"[A-Za-z0-9]",uid): continue
            rec={k:get(vals,idx[k]) for k in FIELDS}
            rec.update({"source":str(path),"table":ti+1,"row":ri,"headers":h,"indices":idx})
            out.append(rec)
    return out

def compose(rows):
    groups=collections.defaultdict(list)
    for r in rows: groups[r["control"]].append(r)
    out={}
    for uid,rs in groups.items():
        ranked=sorted(rs,key=lambda r:sum(1 for k in FIELDS[1:] if not missing(r.get(k,""))),reverse=True)
        base=dict(ranked[0])
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1: base[f]=vals[0]
        base["same_uid_rows"]=rs
        out[uid]=base
    return out

def classify(field,row):
    val=row.get(field,"");status=row.get("runtime_status","")
    if not missing(val): return "BOUND"
    if "UI_LOCAL_EXACT" in status:
        if field in {"operation","runtime_owner"}: return "LEGITIMATE_NA_UI_LOCAL"
        if field=="action" and display_only(row): return "LEGITIMATE_NA_UI_LOCAL"
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in status and field=="operation": return "CLOSED_BY_OWNER_LOCAL_COMMAND"
    if "READ_EXACT" in status and display_only(row):
        if field=="action": return "LEGITIMATE_NA_READ_PRESENTATION"
        if field=="operation" and not missing(row.get("runtime_owner","")): return "READ_OWNER_BOUND_OPERATION_UNSPECIFIED"
    if field=="runtime_owner" and any(x in status for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]): return "KNOWN_RUNTIME_BLOCKER"
    if val=="SOURCE_NOT_DEFINED": return "SOURCE_RESOLUTION_REQUIRED"
    return "DEFINITION_BINDING_GAP"

# Build current 310 unresolved queue from page state.
queue=[]
for page,fn in PAGES.items():
    rows=parse_control_rows(fn);bm=compose(rows)
    for uid,t in sorted(bm.items()):
        for field in ["action","gate","permission","operation","runtime_owner"]:
            if classify(field,t)!="DEFINITION_BINDING_GAP": continue
            # Keep known SYS conflict out of unresolved queue.
            if page=="SYS-01" and uid=="SYS-01-BTN-NAV-OPEN" and field=="gate": continue
            queue.append({
              "page":page,"file":fn,"uid":uid,"field":field,
              "type":t.get("type",""),"label":t.get("label",""),
              "action":t.get("action",""),"gate":t.get("gate",""),
              "permission":t.get("permission",""),"operation":t.get("operation",""),
              "runtime_owner":t.get("runtime_owner",""),"runtime_status":t.get("runtime_status",""),
              "field_column_present":any(r["indices"].get(field) is not None for r in t["same_uid_rows"])
            })
assert len(queue)==310,len(queue)

# Read every system table once, retaining raw header/value context.
system_tables=[]
header_sets=collections.Counter()
for sf in SYSTEMS:
    d=Document(sf)
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        if not any(headers):continue
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if not any(vals):continue
            joined=" | ".join(vals)
            system_tables.append({"source":sf,"table":ti+1,"row":ri,"headers":headers,"values":vals,"joined":joined})

contexts=[]
matched_header_sets=collections.Counter()
for q in queue:
    keys=[]
    for k in ["uid","action","gate","operation"]:
        v=q.get(k,"")
        if not missing(v): keys.append((k,v))
    rows=[]
    for tr in system_tables:
        matches=[{"key":k,"value":v} for k,v in keys if v and v in tr["joined"]]
        if not matches:continue
        rows.append({
          "source":tr["source"],"table":tr["table"],"row":tr["row"],
          "matches":matches,"headers":tr["headers"],"values":tr["values"]
        })
        matched_header_sets[tuple(tr["headers"])]+=1
    contexts.append(q|{"context_rows":rows,"context_row_count":len(rows)})

# Safe alias extraction: only explicit semantic header aliases; no generic "owner" or "function".
ALIASES={
 "action":[r"^action uid$",r"^action id$",r"^action_uid$",r"^action$"],
 "gate":[r"^gate uid$",r"^gate id$",r"^gate_uid$",r"^gate$",r"^gate / policy$",r"^gate/policy$"],
 "permission":[r"^permission$",r"^permission resource$",r"^auth resource$",r"^authorization resource$",r"^required permission$"],
 "operation":[r"^operation$",r"^operation id$",r"^operation uid$",r"^operation_id$",r"^api operation$",r"^service operation$"],
 "runtime_owner":[r"^runtime owner$",r"^runtime_owner$",r"^registered runtime owner$",r"^runtime service owner$",r"^runtime / owner$",r"^runtime/owner$"],
}
def alias_indices(headers,field):
    out=[]
    for i,h in enumerate(headers):
        hl=norm(h).lower()
        if any(re.match(p,hl) for p in ALIASES[field]):out.append(i)
    return out

resolution=[]
for c in contexts:
    field=c["field"];vals=collections.defaultdict(list)
    for tr in c["context_rows"]:
        for i in alias_indices(tr["headers"],field):
            if i>=len(tr["values"]):continue
            v=tr["values"][i]
            if missing(v):continue
            vals[v].append({
              "source":tr["source"],"table":tr["table"],"row":tr["row"],
              "header":tr["headers"][i],"matches":tr["matches"]
            })
    if len(vals)==1:
        v=next(iter(vals));status="UNIQUE_EXPLICIT_ALIAS_BINDING"
        resolution.append(c|{"resolution_status":status,"candidate_value":v,"candidate_sources":vals[v]})
    elif len(vals)>1:
        resolution.append(c|{"resolution_status":"EXPLICIT_ALIAS_CONFLICT","candidate_values":dict(vals)})
    else:
        resolution.append(c|{"resolution_status":"NO_EXPLICIT_ALIAS_BINDING"})

counts=collections.Counter(r["resolution_status"] for r in resolution)
byfield=collections.defaultdict(collections.Counter)
bypage=collections.defaultdict(collections.Counter)
for r in resolution:
    byfield[r["field"]][r["resolution_status"]]+=1
    bypage[r["page"]][r["resolution_status"]]+=1

payload={
 "queue_count":len(queue),
 "resolution_counts":dict(counts),
 "by_field":{k:dict(v) for k,v in byfield.items()},
 "by_page":{k:dict(v) for k,v in bypage.items()},
 "matched_header_sets":[{"headers":list(k),"count":v} for k,v in matched_header_sets.most_common()],
 "rows":resolution
}
Path("__batch07_context_audit.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH07_CONTEXT_AUDIT="+json.dumps({
 "queue_count":len(queue),"resolution_counts":dict(counts),
 "by_field":{k:dict(v) for k,v in byfield.items()},
 "by_page":{k:dict(v) for k,v in bypage.items()}
},ensure_ascii=False,sort_keys=True))
