from pathlib import Path
from docx import Document
import json,re,collections,subprocess

BASE_HEAD="a1a54df49b7c6803d35ef3f60058f4ebfb25b4ee"
SYSTEM_DOCS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
OPS=[
"approveCorrectionScriptCandidate","configureGovernedResource","createCandidate","createChangeRequest",
"decideCandidate","exportProjection","generateCorrectionScriptCandidate","refreshProjection","runSandboxTest",
"saveDraft","searchProjection","sendConversationMessage","setKillSwitch",
]
ACTION_REVIEW=["CORE-01-ACT-CANDIDATE-CREATE","STR-01-ACT-NOOP"]
IAM_UID="IAM-01-BTN-COMPLETE"

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out

assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

page_rows=[]
for page,fn in PAGES.items():
    for r in parse_controls(fn):
        r["page"]=page;page_rows.append(r)
system_rows=[]
for sf in SYSTEM_DOCS:
    for r in parse_controls(sf):
        r["system_file"]=sf;system_rows.append(r)

# Pull explicit authority language around Operation identity / namespace / uniqueness.
authority_language=[]
keywords=re.compile(r"(operation uid|operation identity|namespace|globally unique|global unique|same operation|canonical operation|route identity|operation-specific|exact operation|domain[- ]scoped)",re.I)
for fn in SYSTEM_DOCS+[LOGIC]:
    d=Document(fn)
    paras=[norm(p.text) for p in d.paragraphs if norm(p.text)]
    for i,p in enumerate(paras):
        if keywords.search(p):
            authority_language.append({"file":fn,"paragraph_index":i,"text":p[:1500]})

op_reports=[]
for op in OPS:
    rows=[r for r in page_rows if r.get("operation")==op]
    methods=sorted(set(r["method_path"] for r in rows if not missing(r.get("method_path",""))))
    owners=sorted(set(r["runtime_owner"] for r in rows if not missing(r.get("runtime_owner",""))))
    pages=sorted(set(r["page"] for r in rows))
    sys_occ=[]
    for r in system_rows:
        if r.get("operation")==op:
            sys_occ.append({k:r.get(k,"") for k in ["system_file","control","action","gate","permission","operation","method_path","runtime_owner","runtime_status"]})
    if len(methods)>1:
        classification="OPERATION_IDENTITY_OVERLOAD_EXACT_ROUTE_DIVERGENCE"
    elif len(owners)>1:
        classification="OPERATION_OWNER_DRIFT_OR_DOMAIN_ALIAS"
    else:
        classification="CONSISTENT"
    op_reports.append({
      "operation":op,"classification":classification,"methods":methods,"owners":owners,"pages":pages,
      "controls":[{k:r.get(k,"") for k in ["page","control","action","gate","permission","operation","method_path","runtime_owner","persistence_owner","runtime_status"]} for r in rows],
      "system_occurrences":sys_occ,
    })

# Same-UID page-internal permission conflict details.
iam_rows=[r for r in page_rows if r["page"]=="IAM-01" and r["control"]==IAM_UID]
iam_perm=sorted(set(r["permission"] for r in iam_rows if not missing(r["permission"])))
iam_system=[{k:r.get(k,"") for k in ["system_file","control","action","gate","permission","operation","method_path","runtime_owner","runtime_status"]} for r in system_rows if r["control"]==IAM_UID]
iam={
 "uid":IAM_UID,"classification":"TRUE_SAME_UID_SAME_PAGE_FIELD_CONFLICT" if len(iam_perm)>1 else "CONSISTENT",
 "permission_values":iam_perm,"page_occurrences":iam_rows,"system_occurrences":iam_system,
}

# LOCAL LOCAL is a transport sentinel, not a unique route.
local_rows=[r for r in page_rows if r.get("method_path")=="LOCAL LOCAL"]
local_ops=sorted(set(r.get("operation","") for r in local_rows if not missing(r.get("operation",""))))
local_class="SCANNER_FALSE_POSITIVE_NON_UNIQUE_LOCAL_SENTINEL" if len(local_ops)>1 else "CONSISTENT"

# Action review.
action_reports=[]
for action in ACTION_REVIEW:
    rows=[r for r in page_rows if r.get("action")==action and not missing(r.get("operation",""))]
    ops=sorted(set(r["operation"] for r in rows))
    action_reports.append({
      "action":action,
      "classification":"ACTION_IDENTITY_OVERLOAD_REVIEW" if len(ops)>1 and action!="STR-01-ACT-NOOP" else ("LEGITIMATE_NOOP_SHARED_READ_IDENTITY" if action=="STR-01-ACT-NOOP" else "CONSISTENT"),
      "operations":ops,
      "controls":[{k:r.get(k,"") for k in ["page","control","gate","permission","operation","method_path","runtime_owner","runtime_status"]} for r in rows],
    })

summary={
 "operation_items":len(op_reports),
 "route_divergence":sum(1 for x in op_reports if x["classification"]=="OPERATION_IDENTITY_OVERLOAD_EXACT_ROUTE_DIVERGENCE"),
 "owner_drift_or_alias":sum(1 for x in op_reports if x["classification"]=="OPERATION_OWNER_DRIFT_OR_DOMAIN_ALIAS"),
 "operation_consistent":sum(1 for x in op_reports if x["classification"]=="CONSISTENT"),
 "iam_same_uid_conflict":1 if iam["classification"].startswith("TRUE_") else 0,
 "local_local_false_positive":1 if local_class.startswith("SCANNER_FALSE") else 0,
 "action_identity_overload_review":sum(1 for x in action_reports if x["classification"]=="ACTION_IDENTITY_OVERLOAD_REVIEW"),
 "legitimate_noop_shared_read_identity":sum(1 for x in action_reports if x["classification"]=="LEGITIMATE_NOOP_SHARED_READ_IDENTITY"),
 "authority_language_hits":len(authority_language),
}
report={
 "marker":"ACPOS-20260922-BATCH-19-OPERATION-IDENTITY-NAMESPACE-AUDIT-V1",
 "base_head":BASE_HEAD,"summary":summary,"operation_reports":op_reports,"iam_permission_conflict":iam,
 "local_local":{"classification":local_class,"operations":local_ops,"controls":[{k:r.get(k,"") for k in ["page","control","operation","method_path","runtime_owner","runtime_status"]} for r in local_rows]},
 "action_reports":action_reports,"authority_language":authority_language,
}
Path("__batch19_operation_identity_audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
Path("__batch19_summary.txt").write_text("\n".join(f"{k}={v}" for k,v in summary.items())+"\n",encoding="utf-8")
print("BATCH19_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
