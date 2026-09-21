from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,collections,hashlib,subprocess

MARK="ACPOS-20260922-BATCH-43-SYSTEM08-STRATEGY-CONTRACT-REMEDIATION-V1"
BASE_HEAD="eee89b611d5b58697dd5107a9463d98b956b14f4"
S08="08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
STR="ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
EXPECTED_SHA={
S08:"e9d4463de00e0045105a4bd10337b5249378bbc8",
STR:"8e3fa12c108a0a55f3e7cd18fa01c15b63666670",
LOGIC:"7cd4a1e30b34f0a787556e2d34942b89cec0cc29",
}
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":STR,
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","payload_schema","operation","method_path","runtime_owner","persistence_owner","runtime_status"]
TARGETS={
"STR-01-BTN-SUBMIT-REVIEW":{
 "operation":"submitStrategyReview",
 "fields":{
   "method_path":"POST /v1/state-commands/strategycandidate/submitstrategyreview",
   "payload_schema":"SubmitStrategyReviewRequest",
   "persistence_owner":"System 08 / StrategyCandidate + Review State/Audit",
 },
 "basis":"CANDIDATE_READY→REVIEW_REQUIRED requires exact candidate, expected_version, review permission and evidence; submission creates/updates governed review-request state only and must not create an approval decision."
},
"STR-01-BTN-ADOPT":{
 "operation":"adoptAsContextCandidate",
 "fields":{
   "persistence_owner":"AI-01 Context Owner / Governed Strategy Context Adoption Link + System 08 Audit",
 },
 "basis":"Adopt-as-context occurs only after human review/decision evidence and does not execute an owner module. Context authority remains AI-01; System 08 preserves candidate/decision/audit linkage without creating a second Context Engine."
},
}

def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def exact_index(headers,name):
    n=name.lower()
    for i,h in enumerate(headers):
        if norm(h).lower()==n:return i
    return None
def parse_controls(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),
          "payload_schema":hfind(h,"payload / schema","payload","schema"),
          "operation":hfind(h,"operation"),"method_path":hfind(h,"method / path","method","path"),
          "runtime_owner":hfind(h,"runtime owner"),"persistence_owner":hfind(h,"persistence owner"),
          "runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells];uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"file":path,"table":ti+1,"row":ri});out.append(rec)
    return out
def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        out[uid]=base
    return out
def blob(path):
    b=Path(path).read_bytes()
    return hashlib.sha1(b"blob "+str(len(b)).encode()+b"\0"+b).hexdigest()

for f,s in EXPECTED_SHA.items():
    got=blob(f);assert got==s,(f,got,s)
assert len(list(Path(".").glob("*.docx")))==31
assert subprocess.check_output(["git","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

pre_page=compose(parse_controls(STR));pre_owner=compose(parse_controls(S08))
for uid,s in TARGETS.items():
    for m,label in [(pre_page,"PAGE"),(pre_owner,"OWNER")]:
        assert uid in m,(label,uid,"MISSING")
        r=m[uid]
        assert r["operation"]==s["operation"],(label,uid,r["operation"],s["operation"])
        assert r["runtime_status"]=="EFFECTFUL_EXACT",(label,uid,r["runtime_status"])
        for field,new in s["fields"].items():
            assert missing(r[field]),(label,uid,field,r[field],new)
    if uid=="STR-01-BTN-ADOPT":
        assert pre_page[uid]["method_path"]=="POST /v1/state-commands/strategycandidate/adoptascontextcandidate"
        assert pre_page[uid]["payload_schema"]=="AdoptAsContextCandidateRequest"

def patch(path):
    d=Document(path);hits=collections.Counter()
    for t in d.tables:
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=exact_index(h,"Control UID")
        if ci is None:continue
        idx={"method_path":hfind(h,"method / path","method","path"),"payload_schema":hfind(h,"payload / schema","payload","schema"),"persistence_owner":hfind(h,"persistence owner")}
        for row in t.rows[1:]:
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals):continue
            uid=vals[ci]
            if uid not in TARGETS:continue
            for field,new in TARGETS[uid]["fields"].items():
                fi=idx[field]
                if fi is None or fi>=len(row.cells):continue
                old=norm(row.cells[fi].text)
                if old==new:hits[(uid,field)]+=1;continue
                assert missing(old),(path,uid,field,old,new)
                row.cells[fi].text=new;hits[(uid,field)]+=1
    expected={(u,f) for u,s in TARGETS.items() for f in s["fields"]}
    assert set(hits)==expected,(path,dict(hits),expected-set(hits))
    d.save(path);Document(path)
    return {f"{u}:{f}":n for (u,f),n in hits.items()}

patches={"SYSTEM08_OWNER":patch(S08),"STR_PAGE":patch(STR)}

def add_landscape(doc):
    sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=sec.page_height,sec.page_width
    sec.top_margin=Inches(.33);sec.bottom_margin=Inches(.33);sec.left_margin=Inches(.28);sec.right_margin=Inches(.28)
def repeat_header(row):
    trPr=row._tr.get_or_add_trPr();e=OxmlElement("w:tblHeader");e.set(qn("w:val"),"true");trPr.append(e)
def shade(cell,fill="EDE9FE"):
    tcPr=cell._tc.get_or_add_tcPr();e=OxmlElement("w:shd");e.set(qn("w:fill"),fill);tcPr.append(e)
def add_table(doc,headers,rows,fs=4.0):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";repeat_header(t.rows[0])
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h);shade(t.rows[0].cells[i])
    for rr in rows:
        cells=t.add_row().cells
        for i,v in enumerate(rr):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                for run in p.runs:run.font.size=Pt(fs)
    return t

for path,title,prefix in [
(S08,"Batch 43 · Strategy Review / Context Adoption Canonical Contracts","SYSTEM08"),
(STR,"Batch 43 · STR-01 Strategy Contract Binding Ledger","PAGE::STR-01"),
]:
    d=Document(path);add_landscape(d);d.add_heading(title,level=1)
    p=d.add_paragraph();p.add_run(f"[{MARK}::{prefix}] ").bold=True
    p.add_run("Review submission is a governed request to enter REVIEW_REQUIRED, never an approval. Context adoption requires prior human review and persists through AI-01 Context authority plus System 08 audit linkage; Strategy does not create a second Context Engine.")
    rr=[]
    for uid,s in TARGETS.items():
        rr.append([uid,s["operation"],s["fields"].get("method_path","UNCHANGED"),s["fields"].get("payload_schema","UNCHANGED"),s["fields"].get("persistence_owner","UNCHANGED"),s["basis"]])
    add_table(d,["Control UID","Operation","Method / Path","Payload / Schema","Persistence Owner","Semantics"],rr,3.5)
    d.save(path);Document(path)

post_page=compose(parse_controls(STR));post_owner=compose(parse_controls(S08))
for uid,s in TARGETS.items():
    for m,label in [(post_page,"PAGE"),(post_owner,"OWNER")]:
        for f,v in s["fields"].items():
            assert m[uid][f]==v,(label,uid,f,m[uid][f],v)

all_maps={p:compose(parse_controls(fn)) for p,fn in PAGES.items()}
remaining=[]
for page,m in all_maps.items():
    for uid,r in m.items():
        st=r.get("runtime_status","");op=r.get("operation","")
        if missing(op) or st not in {"EFFECTFUL_EXACT","READ_EXACT"}:continue
        if missing(r.get("method_path","")):remaining.append(("method_path",page,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("payload_schema","")):remaining.append(("payload_schema",page,uid))
        if st=="EFFECTFUL_EXACT" and missing(r.get("persistence_owner","")):remaining.append(("persistence_owner",page,uid))
rc=collections.Counter(f for f,_,_ in remaining)
assert len(remaining)==279,(len(remaining),rc)
assert rc=={"method_path":62,"payload_schema":118,"persistence_owner":99},rc

logic=Document(LOGIC);add_landscape(logic);logic.add_heading("Batch 43 · System 08 Strategy Contract Closure",level=1)
p=logic.add_paragraph();p.add_run(f"[{MARK}] ").bold=True
p.add_run("System 08 closes four STR-01 cells: submit-review transport/payload/persistence and adopt-as-context persistence. Human review remains mandatory, AI self-approval remains forbidden, and adopted context is owned by AI-01 Context authority with System 08 audit linkage.")
add_table(logic,["Metric","Value","Result"],[
["Pre contract denominator",283,"Method 63 + Payload 119 + Persistence 101"],
["System 08 cells closed",4,"Method 1 + Payload 1 + Persistence 2"],
["Post contract denominator",279,"Method 62 + Payload 118 + Persistence 99"],
["AI self-approval introduced",0,"Forbidden"],
["Second Context Engine created",0,"None"],
["Runtime execution claimed","False","Design-contract closure only"],
],4.1)
machine={"marker":MARK,"base_head":BASE_HEAD,"pre_contract_cells":283,"system08_cells_closed":4,
"method_path_closed":1,"payload_schema_closed":1,"persistence_owner_closed":2,
"post_contract_cells":279,"remaining_method_path":62,"remaining_payload_schema":118,"remaining_persistence_owner":99,
"ai_self_approval_introduced":0,"second_context_engine_created":0,"runtime_execution_claimed":False}
logic.add_paragraph("BATCH43_MACHINE_JSON="+json.dumps(machine,ensure_ascii=False,sort_keys=True,separators=(",",":")))
logic.save(LOGIC);Document(LOGIC)

changed=[S08,STR,LOGIC]
report={"machine":machine,"targets":TARGETS,"patches":patches,"changed_docs":changed,"output_blob_sha":{f:blob(f) for f in changed}}
Path("__batch43_remediation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH43="+json.dumps(machine,ensure_ascii=False,sort_keys=True))
