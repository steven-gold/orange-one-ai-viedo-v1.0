from pathlib import Path
from docx import Document
import json,re,collections

SYSTEMS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}

def norm(x): return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def get(vals,i): return "" if i is None or i>=len(vals) else norm(vals[i])
def missing(v): return v in {"","—","-","SOURCE_NOT_DEFINED"}

DISPLAY_ONLY_TYPES={"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"}
def display_only_type(t):
    typ=(t or "").upper()
    return typ in DISPLAY_ONLY_TYPES or any(x in typ for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]
ALIASES={
 "control":[r"^control uid$",r"^control id$",r"^control_uid$"],
 "action":[r"^action uid$",r"^action id$",r"^action_uid$",r"^action$"],
 "gate":[r"^gate uid$",r"^gate id$",r"^gate_uid$",r"^gate$",r"^gate / policy$",r"^gate/policy$"],
 "permission":[r"^permission$",r"^permission resource$",r"^auth resource$",r"^authorization resource$",r"^required permission$"],
 "operation":[r"^operation$",r"^operation id$",r"^operation uid$",r"^operation_id$",r"^api operation$",r"^service operation$"],
 "runtime_owner":[r"^runtime owner$",r"^runtime_owner$",r"^registered runtime owner$",r"^runtime service owner$",r"^runtime / owner$",r"^runtime/owner$"],
}
def alias_indices(headers,field):
    out=[]
    for i,h in enumerate(headers):
        hl=norm(h).lower()
        if any(re.match(p,hl) for p in ALIASES[field]): out.append(i)
    return out

def parse_page_rows(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(h,"control uid")
        if ci is None:continue
        idx={
          "control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),
          "action":hfind(h,"action uid"),"gate":hfind(h,"gate uid","gate"),
          "permission":hfind(h,"permission","auth resource"),"operation":hfind(h,"operation"),
          "runtime_owner":hfind(h,"runtime owner"),"runtime_status":hfind(h,"runtime status")
        }
        for ri,r in enumerate(t.rows[1:],2):
            vals=[c.text for c in r.cells]
            uid=get(vals,ci)
            if not uid or uid in {"—","-"} or not re.search(r"[A-Za-z0-9]",uid):continue
            rec={k:get(vals,idx[k]) for k in FIELDS}
            rec.update({"source":str(path),"table":ti+1,"row":ri,"headers":h,"indices":idx})
            out.append(rec)
    return out

def compose(rows):
    groups=collections.defaultdict(list)
    for r in rows:groups[r["control"]].append(r)
    out={}
    for uid,rs in groups.items():
        ranked=sorted(rs,key=lambda r:sum(1 for k in FIELDS[1:] if not missing(r.get(k,""))),reverse=True)
        base=dict(ranked[0])
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        base["same_uid_rows"]=rs
        out[uid]=base
    return out

def classify_gap(field,row):
    val=row.get(field,"");status=row.get("runtime_status","")
    if not missing(val):return "BOUND"
    if "UI_LOCAL_EXACT" in status:
        if field in {"operation","runtime_owner"}:return "LEGITIMATE_NA_UI_LOCAL"
        if field=="action" and display_only_type(row.get("type","")):return "LEGITIMATE_NA_UI_LOCAL"
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in status and field=="operation":return "CLOSED_BY_OWNER_LOCAL_COMMAND"
    if "READ_EXACT" in status and display_only_type(row.get("type","")):
        if field=="action":return "LEGITIMATE_NA_READ_PRESENTATION"
        if field=="operation" and not missing(row.get("runtime_owner","")):return "READ_OWNER_BOUND_OPERATION_UNSPECIFIED"
    if field=="runtime_owner" and any(x in status for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return "KNOWN_RUNTIME_BLOCKER"
    if val=="SOURCE_NOT_DEFINED":return "SOURCE_RESOLUTION_REQUIRED"
    return "DEFINITION_BINDING_GAP"

# Generic exact rows for every System doc.
generic=collections.defaultdict(list)
owner_index=collections.defaultdict(list)
for sf in SYSTEMS:
    d=Document(sf)
    for ti,t in enumerate(d.tables):
        if not t.rows: continue
        headers=[norm(c.text) for c in t.rows[0].cells]
        vals_control=alias_indices(headers,"control")
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if not any(vals):continue
            rec={"source":sf,"table":ti+1,"row":ri,"headers":headers,"values":vals}
            generic[sf].append(rec)
            for ci in vals_control:
                if ci<len(vals):
                    uid=vals[ci]
                    if uid and uid not in {"—","-"}:owner_index[uid].append(rec)

# Reconstruct 310 then keep 226 unique-owner rows only.
allq=[]
for page,fn in PAGES.items():
    rows=parse_page_rows(fn);bm=compose(rows)
    for uid,t in sorted(bm.items()):
        for field in ["action","gate","permission","operation","runtime_owner"]:
            if classify_gap(field,t)!="DEFINITION_BINDING_GAP":continue
            if page=="SYS-01" and uid=="SYS-01-BTN-NAV-OPEN" and field=="gate":continue
            sources=sorted(set(r["source"] for r in owner_index.get(uid,[])))
            allq.append({
              "page":page,"file":fn,"uid":uid,"field":field,
              "type":t.get("type",""),"label":t.get("label",""),
              "action":t.get("action",""),"gate":t.get("gate",""),
              "permission":t.get("permission",""),"operation":t.get("operation",""),
              "runtime_owner":t.get("runtime_owner",""),"runtime_status":t.get("runtime_status",""),
              "owner_sources":sources,
              "field_column_present":any(r["indices"].get(field) is not None for r in t["same_uid_rows"])
            })
assert len(allq)==310,len(allq)
unique=[q for q in allq if len(q["owner_sources"])==1]
multi=[q for q in allq if len(q["owner_sources"])>1]
absent=[q for q in allq if len(q["owner_sources"])==0]
assert (len(unique),len(multi),len(absent))==(226,56,28),(len(unique),len(multi),len(absent))

# Relation policy: only exact equality within the same unique owner document.
# No reverse inference from Permission, Runtime Owner, label, or free text.
MATCH_POLICY={
 "action":["control"],
 "gate":["control","action"],
 "permission":["control","action","gate"],
 "operation":["control","action"],
 "runtime_owner":["control","action","operation"],
}
def qkey(q,key):
    return q["uid"] if key=="control" else q.get(key,"")

resolution=[]
counts=collections.Counter()
by_owner=collections.defaultdict(collections.Counter)
by_field=collections.defaultdict(collections.Counter)
for q in unique:
    sf=q["owner_sources"][0];field=q["field"]
    candidates=collections.defaultdict(list)
    for key in MATCH_POLICY[field]:
        kval=qkey(q,key)
        if missing(kval):continue
        for r in generic[sf]:
            # Match key only via an explicit alias column and exact cell equality.
            matched=False
            for ki in alias_indices(r["headers"],key):
                if ki<len(r["values"]) and r["values"][ki]==kval:
                    matched=True;break
            if not matched:continue
            for fi in alias_indices(r["headers"],field):
                if fi>=len(r["values"]):continue
                v=r["values"][fi]
                if missing(v):continue
                candidates[v].append({
                  "source":sf,"table":r["table"],"row":r["row"],
                  "matched_by":key,"matched_value":kval,"target_header":r["headers"][fi]
                })
    if len(candidates)==1:
        val=next(iter(candidates))
        status="UNIQUE_SAME_OWNER_EXACT_RELATION"
        rec=q|{"resolution_status":status,"candidate_value":val,"candidate_sources":candidates[val]}
    elif len(candidates)>1:
        status="SAME_OWNER_RELATION_CONFLICT"
        rec=q|{"resolution_status":status,"candidate_values":dict(candidates)}
    else:
        status="NO_SAME_OWNER_EXACT_RELATION"
        rec=q|{"resolution_status":status}
    resolution.append(rec);counts[status]+=1;by_owner[sf][status]+=1;by_field[field][status]+=1

payload={
 "denominator":226,
 "counts":dict(counts),
 "by_owner":{k:dict(v) for k,v in by_owner.items()},
 "by_field":{k:dict(v) for k,v in by_field.items()},
 "rows":resolution
}
Path("__batch09_same_owner_relation_audit.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH09_AUDIT="+json.dumps({
 "denominator":226,
 "counts":dict(counts),
 "by_owner":{k:dict(v) for k,v in by_owner.items()},
 "by_field":{k:dict(v) for k,v in by_field.items()}
},ensure_ascii=False,sort_keys=True))
print("BATCH09_RESOLVABLE_BEGIN")
for r in resolution:
    if r["resolution_status"]=="UNIQUE_SAME_OWNER_EXACT_RELATION":
        print(json.dumps(r,ensure_ascii=False,sort_keys=True))
print("BATCH09_RESOLVABLE_END")
print("BATCH09_CONFLICT_BEGIN")
for r in resolution:
    if r["resolution_status"]=="SAME_OWNER_RELATION_CONFLICT":
        print(json.dumps(r,ensure_ascii=False,sort_keys=True))
print("BATCH09_CONFLICT_END")
