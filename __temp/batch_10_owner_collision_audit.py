from pathlib import Path
from docx import Document
import json,re,collections

SYSTEMS=[
"01_ACPOS_Global_Intelligent_Brain_Information_Lifecycle_Mother_Basic_Logic_Design_Normative_Contract_v4.docx",
"02_ACPOS_AI_Conversation_Memory_MultiAI_Assistant_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"03_ACPOS_AI_Execution_Script_Compiler_Tool_AIAPI_Runtime_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"04_ACPOS_CORE_AI_Blueprint_Canonical_Script_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"07_ACPOS_Social_Publishing_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"08_ACPOS_Strategy_Intelligence_MultiAI_Decision_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
"09_ACPOS_AI_System_Engineer_Self_Inspection_Evolution_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx",
]
PAGES={
"AIAPI-01":"ACPOS_AIAPI-01_AI_API管理_Mother_Basic_Design_OPTIMIZED.docx",
"ASSET-01":"ACPOS_ASSET-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"CORE-01":"ACPOS_CORE-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DB-01":"ACPOS_DB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"DEV-01":"ACPOS_DEV-01_企業自動開發系統_Mother_Basic_Design_OPTIMIZED.docx",
"EDIT-01":"ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ERP-01":"ACPOS_ERP-01_ERP與財務_Mother_Basic_Design_OPTIMIZED.docx",
"IAM-01":"ACPOS_IAM-01_帳戶與權限_Mother_Basic_Design_OPTIMIZED.docx",
"INFO-01":"ACPOS_INFO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"KB-01":"ACPOS_KB-01_知識庫_Mother_Basic_Design_OPTIMIZED.docx",
"QA-01":"ACPOS_QA-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SG-02":"ACPOS_SG-02_QA審查項目名單_Mother_Basic_Design_OPTIMIZED.docx",
"SOC-01":"ACPOS_SOC-01_社群發布_Mother_Basic_Design_OPTIMIZED.docx",
"STR-01":"ACPOS_STR-01_WORKSPACE_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"SYS-01":"ACPOS_SYS-01_系統維護與生命週期工作區_Mother_Basic_Design_OPTIMIZED.docx",
"VIDEO-01":"ACPOS_VIDEO-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"WB-01":"ACPOS_WB-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx",
"ADMIN-STR-01":"ACPOS_admin_STR-01_戰略中心後台_Mother_Basic_Design_OPTIMIZED.docx",
}
FIELDS=["control","type","label","action","gate","permission","operation","runtime_owner","runtime_status"]

def norm(x): return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v): return v in {"","—","-","SOURCE_NOT_DEFINED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def display_only_type(t):
    u=(t or "").upper()
    return u in {"READONLY","LABEL","TEXT","BADGE","STATUS","METRIC","KPI","DIVIDER","ICON","PANEL","CARD","VIEW","LIST","TABLE","CHIP","TAG","DISPLAY"} or any(x in u for x in ["READONLY","LABEL","DISPLAY","METRIC","KPI","STATUS"])

def parse_control_rows(path):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        ci=hfind(h,"control uid")
        if ci is None:continue
        idx={"control":ci,"type":hfind(h,"type"),"label":hfind(h,"label","顯示名稱"),"action":hfind(h,"action uid"),"gate":hfind(h,"gate"),"permission":hfind(h,"permission","auth resource"),"operation":hfind(h,"operation"),"runtime_owner":hfind(h,"runtime owner"),"runtime_status":hfind(h,"runtime status")}
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            uid=vals[ci] if ci<len(vals) else ""
            if not uid or uid in {"—","-"}:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in idx.items()}
            rec.update({"source":path,"table":ti+1,"row":ri,"headers":h,"values":vals})
            out.append(rec)
    return out

def compose(rows):
    g=collections.defaultdict(list)
    for r in rows:g[r["control"]].append(r)
    out={}
    for uid,rs in g.items():
        base=max(rs,key=lambda r:sum(not missing(r.get(f,"")) for f in FIELDS[1:])).copy()
        for f in FIELDS[1:]:
            vals=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
            if missing(base.get(f,"")) and len(vals)==1:base[f]=vals[0]
        base["same_uid_rows"]=rs;out[uid]=base
    return out

def gap(field,r):
    v=r.get(field,"");st=r.get("runtime_status","")
    if not missing(v):return False
    if "UI_LOCAL_EXACT" in st and (field in {"operation","runtime_owner"} or (field=="action" and display_only_type(r.get("type","")))):return False
    if "OWNER_ORCHESTRATED_RUNTIME_EXACT_NO_PUBLIC_ROUTE" in st and field=="operation":return False
    if "READ_EXACT" in st and display_only_type(r.get("type","")):
        if field=="action":return False
        if field=="operation" and not missing(r.get("runtime_owner","")):return False
    if field=="runtime_owner" and any(x in st for x in ["SPEC_EXACT_RUNTIME_BLOCKED","SPEC_EXACT_RUNTIME_NOT_EXECUTED","RUNTIME_IMPLEMENTATION_BLOCKED"]):return False
    if v=="SOURCE_NOT_DEFINED":return False
    return True

# System exact owner rows and document text contexts.
owner_index=collections.defaultdict(list)
doc_context={}
for sf in SYSTEMS:
    rows=parse_control_rows(sf)
    for r in rows:owner_index[r["control"]].append(r)
    d=Document(sf)
    paras=[norm(p.text) for p in d.paragraphs if norm(p.text)]
    # include tables as row text to detect explicit authority declarations outside standard control registry
    generic_rows=[]
    for ti,t in enumerate(d.tables):
        for ri,row in enumerate(t.rows,1):
            vals=[norm(c.text) for c in row.cells]
            if any(vals):generic_rows.append({"table":ti+1,"row":ri,"text":" | ".join(vals)})
    doc_context[sf]={"paras":paras,"rows":generic_rows}

# Current collision denominator.
collisions=[]
for page,fn in PAGES.items():
    bm=compose(parse_control_rows(fn))
    for uid,r in bm.items():
        for field in ["action","gate","permission","operation","runtime_owner"]:
            if not gap(field,r):continue
            if page=="SYS-01" and uid=="SYS-01-BTN-NAV-OPEN" and field=="gate":continue
            sources=sorted(set(x["source"] for x in owner_index.get(uid,[])))
            if len(sources)>1:
                collisions.append({
                  "page":page,"file":fn,"uid":uid,"field":field,
                  "type":r.get("type",""),"label":r.get("label",""),"action":r.get("action",""),"gate":r.get("gate",""),
                  "permission":r.get("permission",""),"operation":r.get("operation",""),"runtime_owner":r.get("runtime_owner",""),
                  "runtime_status":r.get("runtime_status",""),"owner_sources":sources
                })
assert len(collisions)==56,len(collisions)

AUTH_PATTERNS=[
 r"canonical[^.。;；|]{0,100}(owner|authority|source)",
 r"(authoritative|authority)[^.。;；|]{0,100}(owner|source|contract)",
 r"source[- ]of[- ]truth",
 r"single source of truth",
 r"(primary|master)[^.。;；|]{0,80}(owner|authority|contract)",
 r"(mirror|mirrored|secondary|derived|projection|reference only|non-authoritative)[^.。;；|]{0,100}(owner|authority|contract|source)?",
]
def authority_contexts(sf,uid):
    out=[]
    for i,p in enumerate(doc_context[sf]["paras"],1):
        low=p.lower()
        if uid in p or any(re.search(pt,low,re.I) for pt in AUTH_PATTERNS):
            # retain only authority-looking lines or exact UID lines
            if uid in p or any(re.search(pt,low,re.I) for pt in AUTH_PATTERNS):
                out.append({"kind":"paragraph","index":i,"text":p})
    for rr in doc_context[sf]["rows"]:
        text=rr["text"];low=text.lower()
        if uid in text or any(re.search(pt,low,re.I) for pt in AUTH_PATTERNS):
            out.append({"kind":"table","table":rr["table"],"row":rr["row"],"text":text})
    # keep bounded exact UID hits and authority declaration hits
    uid_hits=[x for x in out if uid in x["text"]][:20]
    auth_hits=[x for x in out if uid not in x["text"] and any(re.search(pt,x["text"],re.I) for pt in AUTH_PATTERNS)][:30]
    return uid_hits+auth_hits

def summarize_owner_rows(uid,sf):
    rs=[r for r in owner_index.get(uid,[]) if r["source"]==sf]
    vals={}
    for f in FIELDS[1:]:
        vv=sorted(set(r.get(f,"") for r in rs if not missing(r.get(f,""))))
        vals[f]=vv
    return {"row_count":len(rs),"rows":[{"table":r["table"],"row":r["row"],"values":{f:r.get(f,"") for f in FIELDS}} for r in rs],"field_values":vals}

def compare(a,b):
    conflicts={};only_a={};only_b={};same={}
    for f in FIELDS[1:]:
        av=set(a["field_values"][f]);bv=set(b["field_values"][f])
        if av and bv:
            if av==bv:same[f]=sorted(av)
            else:conflicts[f]={"a":sorted(av),"b":sorted(bv)}
        elif av:only_a[f]=sorted(av)
        elif bv:only_b[f]=sorted(bv)
    if conflicts:cls="DIVERGENT_NONEMPTY_VALUES"
    elif only_a and only_b:cls="COMPLEMENTARY_NO_CONFLICT"
    elif only_a:cls="OWNER_A_STRICTER_COVERAGE"
    elif only_b:cls="OWNER_B_STRICTER_COVERAGE"
    else:cls="EQUIVALENT_OR_BOTH_EMPTY"
    return cls,conflicts,only_a,only_b,same

rows=[];pair_counts=collections.Counter();comparison_counts=collections.Counter();field_counts=collections.Counter()
for c in collisions:
    a,b=c["owner_sources"]
    sa=summarize_owner_rows(c["uid"],a);sb=summarize_owner_rows(c["uid"],b)
    cls,conflicts,only_a,only_b,same=compare(sa,sb)
    rec=c|{
      "owner_a":a,"owner_b":b,
      "owner_a_summary":sa,"owner_b_summary":sb,
      "comparison_class":cls,
      "conflicts":conflicts,"only_a":only_a,"only_b":only_b,"same":same,
      "owner_a_authority_contexts":authority_contexts(a,c["uid"]),
      "owner_b_authority_contexts":authority_contexts(b,c["uid"])
    }
    rows.append(rec);pair_counts[(a,b)]+=1;comparison_counts[cls]+=1;field_counts[c["field"]]+=1

payload={
 "denominator":56,
 "pair_counts":{" + ".join(k):v for k,v in pair_counts.items()},
 "comparison_counts":dict(comparison_counts),
 "missing_field_counts":dict(field_counts),
 "rows":rows
}
Path("__batch10_owner_collision_audit.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH10_COLLISION_AUDIT="+json.dumps({
 "denominator":56,
 "pair_counts":{" + ".join(k):v for k,v in pair_counts.items()},
 "comparison_counts":dict(comparison_counts),
 "missing_field_counts":dict(field_counts)
},ensure_ascii=False,sort_keys=True))
print("BATCH10_ROWS_BEGIN")
for r in rows:
    print(json.dumps({
      "page":r["page"],"uid":r["uid"],"field":r["field"],
      "owner_a":r["owner_a"],"owner_b":r["owner_b"],"comparison_class":r["comparison_class"],
      "conflicts":r["conflicts"],"only_a":r["only_a"],"only_b":r["only_b"],
      "a_contexts":r["owner_a_authority_contexts"],"b_contexts":r["owner_b_authority_contexts"]
    },ensure_ascii=False,sort_keys=True))
print("BATCH10_ROWS_END")
