from docx import Document
import json,re
FN="06_ACPOS_Enterprise_Business_Development_AI_System_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
OPS=["getERPSyncStatus","getERPFailure","getERPFinanceFactPack","getERPForecast","getERPCapacityGuardrails","exportProjection",
"createERPConnector","updateERPConnector","validateERPConnector","validateERPMapping"]
TERMS=["ERPConnectorService","FinanceGuardrailService","Existing Projection/Export owner"]
def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
d=Document(FN);hits=[]
for i,p in enumerate(d.paragraphs,1):
    t=norm(p.text)
    if t and any(x in t for x in OPS+TERMS):
        hits.append({"kind":"paragraph","index":i,"text":t})
for ti,tb in enumerate(d.tables,1):
    if not tb.rows:continue
    headers=[norm(c.text) for c in tb.rows[0].cells]
    for ri,row in enumerate(tb.rows[1:],2):
        vals=[norm(c.text) for c in row.cells];joined=" | ".join(vals)
        if any(x in joined for x in OPS+TERMS):
            hits.append({"kind":"table","table":ti,"row":ri,"headers":headers,"values":vals})
print("BATCH15_ERP_CONTRACT="+json.dumps(hits,ensure_ascii=False,sort_keys=True))
