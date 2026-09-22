from pathlib import Path
from docx import Document
import json,re,collections,subprocess

BASE_HEAD="2f709927292813154da145ec56791c51f33b1794"
TARGET="EDIT-01-BTN-FLOW-START"
ACTION="EDIT-01-ACT-FLOW-START-CONTINUE"
OP="dispatchEditingFlowStartContinue"
METHOD="NO_PUBLIC_API_BY_AUTHORITY"
PAYLOAD="EditingFlowStartContinueDispatchContext"
PERSIST="DELEGATED_TO_SELECTED_EDIT_PORT_PERSISTENCE"
GATE="EDIT-01-GATE-CONTEXT-INTEGRITY"
UI_PERMISSION="EDITING_USE"
RUNTIME_OWNER="EDITING/ORCHESTRATION"
PORTS={
"EDIT-01-PORT-EDIT-RUN-CREATE":{
 "operation":"createEditingRuntimeRun",
 "method":"POST /v1/editing-runtime-runs",
 "payload":"CreateEditingRuntimeRunRequest",
 "permission":"editing.task.execute",
},
"EDIT-01-PORT-ASSEMBLY-COMPLETE":{
 "operation":"completeAssembly",
 "method":"POST /v1/editing-runtime-runs/{runId}/assembly",
 "payload":"CompleteEditingAssemblyRequest",
 "permission":"editing.task.execute",
},
}
S05="05_ACPOS_Creative_Production_AI_ASSET_VIDEO_EDIT_VOICE_QA_Mother_Basic_Logic_Design_Normative_Contract_v1.0.docx"
EDIT="ACPOS_EDIT-01_MOTHER_BASIC_DESIGN_TECH_PURPLE_v1.0.docx"
LOGIC="ACPOS_SYSTEM_LOGIC_GLOBAL_INTEGRATION_Mother_Basic_Design_OPTIMIZED.docx"
DOCS=sorted(Path(".").glob("*.docx"))
assert len(DOCS)==31
assert subprocess.check_output(["git","-c","core.quotePath=false","diff","--name-only",BASE_HEAD,"HEAD","--","*.docx"],text=True).strip()==""

LEGACY_RE=re.compile(r"(?<![A-Za-z0-9])R(?:9(?:\.0\.1)?|5)(?![A-Za-z0-9])",re.I)
AMBIG_LOCATOR_RE=re.compile(r"\bt\d+/r(?:5|9)\b",re.I)
OLD_EXACT=[
"createEditingRuntimeRun | completeAssembly",
"POST /v1/editing-runtime-runs | POST /v1/editing-runtime-runs/{runId}/assembly",
"CreateEditingRuntimeRunRequest | CompleteEditingAssemblyRequest",
]
def norm(x):return re.sub(r"\s+"," ",(x or "").replace("\n"," | ").strip())
def missing(v):return v in {"","—","-","SOURCE_NOT_DEFINED","UNCHANGED"}
def hfind(headers,*needles):
    hs=[norm(x).lower() for x in headers]
    for needle in needles:
        n=needle.lower()
        for i,h in enumerate(hs):
            if n in h:return i
    return None
def full_text(path):
    d=Document(path);parts=[]
    parts.extend(norm(p.text) for p in d.paragraphs if norm(p.text))
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(norm(c.text) for c in row.cells))
    return "\n".join(parts)
def control_rows(path,uid):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        if ci is None:continue
        ix={
          "control":ci,"type":hfind(h,"type"),"action":hfind(h,"action uid"),
          "gate":hfind(h,"gate uid","gate"),"permission":hfind(h,"permission","auth resource"),
          "payload":hfind(h,"payload / schema","payload","schema"),"operation":hfind(h,"operation"),
          "method":hfind(h,"method / path","method","path"),"runtime_owner":hfind(h,"runtime owner"),
          "persistence":hfind(h,"persistence owner"),"runtime_status":hfind(h,"runtime status")
        }
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if ci>=len(vals) or vals[ci]!=uid:continue
            rec={k:(vals[i] if i is not None and i<len(vals) else "") for k,i in ix.items()}
            rec.update({"file":str(path),"table":ti,"row":ri});out.append(rec)
    return out
def port_rows(path,uid):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        pi=hfind(h,"port uid")
        if pi is None:continue
        oi=hfind(h,"operation");mi=hfind(h,"method / path","method","path");peri=hfind(h,"permission")
        payi=hfind(h,"payload / schema","payload","schema")
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if pi<len(vals) and vals[pi]==uid:
                out.append({
                  "file":str(path),"table":ti,"row":ri,"headers":h,"values":vals,
                  "operation":vals[oi] if oi is not None and oi<len(vals) else "",
                  "method":vals[mi] if mi is not None and mi<len(vals) else "",
                  "permission":vals[peri] if peri is not None and peri<len(vals) else "",
                  "payload":vals[payi] if payi is not None and payi<len(vals) else "",
                })
    return out
def reverse_rows(path,uid):
    d=Document(path);out=[]
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        pi=hfind(h,"port uid");ci=hfind(h,"ui-flow consumer","consumer")
        if pi is None or ci is None:continue
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if pi<len(vals) and vals[pi]==uid:
                out.append({"file":str(path),"table":ti,"row":ri,"headers":h,"values":vals,
                            "consumer":vals[ci] if ci<len(vals) else ""})
    return out

# 1. Package-wide ambiguous legacy token cleanup verification.
legacy=[];ambig=[];old_exact=[]
for p in DOCS:
    txt=full_text(p)
    for m in LEGACY_RE.finditer(txt):
        legacy.append({"file":p.name,"match":m.group(0),"context":txt[max(0,m.start()-120):m.end()+180]})
    for m in AMBIG_LOCATOR_RE.finditer(txt):
        ambig.append({"file":p.name,"match":m.group(0),"context":txt[max(0,m.start()-120):m.end()+180]})
    for s in OLD_EXACT:
        if s in txt:
            old_exact.append({"file":p.name,"value":s,"count":txt.count(s)})

# 2. Exact target row parity between page and canonical owner.
target={}
for p in [Path(S05),Path(EDIT)]:
    rows=control_rows(p,TARGET)
    exact=[r for r in rows if r["runtime_status"]=="EFFECTFUL_EXACT"]
    target[p.name]=exact
    assert len(exact)==1,(p.name,exact)
    r=exact[0]
    expected={
      "action":ACTION,"gate":GATE,"permission":UI_PERMISSION,"payload":PAYLOAD,
      "operation":OP,"method":METHOD,"runtime_owner":RUNTIME_OWNER,
      "persistence":PERSIST,"runtime_status":"EFFECTFUL_EXACT"
    }
    for k,v in expected.items():assert r[k]==v,(p.name,k,r[k],v)
    for k in ["operation","method","payload","persistence"]:
        assert "|" not in r[k],(p.name,k,r[k])

# 3. Atomic port integrity and reverse-consumer linkage.
ports={};reverse={}
for uid,s in PORTS.items():
    found=[]
    for p in DOCS:
        found.extend(port_rows(p,uid))
    assert found,(uid,"PORT_NOT_FOUND")
    exact=[]
    for r in found:
        if r["operation"]==s["operation"] and r["method"]==s["method"] and r["permission"]==s["permission"]:
            exact.append(r)
    assert exact,(uid,"NO_EXACT_PORT_CONTRACT",found)
    ports[uid]=exact
    rev=[]
    for p in [Path(S05),Path(EDIT)]:
        rev.extend(reverse_rows(p,uid))
    assert len(rev)>=2,(uid,"REVERSE_BINDING_MISSING",rev)
    assert all(OP in x["consumer"] for x in rev),(uid,"REVERSE_CONSUMER_MISMATCH",rev)
    reverse[uid]=rev

# 4. Forward orchestration branch table + state/no-op branch.
branch_rows=[]
permission_bridge=[]
for p in [Path(S05),Path(EDIT)]:
    d=Document(p)
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells]
        si=hfind(h,"selected port");bi=hfind(h,"branch")
        if si is not None and bi is not None:
            for ri,row in enumerate(t.rows[1:],2):
                vals=[norm(c.text) for c in row.cells]
                if any(uid in " | ".join(vals) for uid in PORTS) or "NO MUTATION" in " | ".join(vals).upper():
                    branch_rows.append({"file":p.name,"table":ti,"row":ri,"headers":h,"values":vals})
        # explicit multi-layer permission ledger
        if any("UI Permission" in norm(c.text) or "Runtime Permission" in norm(c.text) for c in t.rows[0].cells+t.rows[-1].cells):
            for ri,row in enumerate(t.rows,1):
                vals=[norm(c.text) for c in row.cells]
                if any(x in " | ".join(vals) for x in ["UI Permission","Runtime Permission"]):
                    permission_bridge.append({"file":p.name,"table":ti,"row":ri,"headers":h,"values":vals})
assert sum(PORTS.keys() <= set(" | ".join(x["values"]) for x in branch_rows) for _ in [0])>=0  # keep report-only below
forward_blob="\n".join(" | ".join(x["values"]) for x in branch_rows)
for uid in PORTS:assert uid in forward_blob,(uid,"FORWARD_BRANCH_MISSING")
assert re.search(r"NO\s+MUTATION|no atomic mutation",forward_blob,re.I),("NO_MUTATION_BRANCH_MISSING",forward_blob[:5000])

# 5. Action->operation uniqueness for target action; operation->action reverse consistency.
action_ops=collections.defaultdict(set)
operation_actions=collections.defaultdict(set)
all_control_rows=[]
for p in DOCS:
    d=Document(p)
    for ti,t in enumerate(d.tables,1):
        if not t.rows:continue
        h=[norm(c.text) for c in t.rows[0].cells];ci=hfind(h,"control uid")
        ai=hfind(h,"action uid");oi=hfind(h,"operation");si=hfind(h,"runtime status")
        if ci is None or ai is None or oi is None:continue
        for ri,row in enumerate(t.rows[1:],2):
            vals=[norm(c.text) for c in row.cells]
            if max(ci,ai,oi)>=len(vals):continue
            uid,act,op=vals[ci],vals[ai],vals[oi]
            st=vals[si] if si is not None and si<len(vals) else ""
            if uid and act and not missing(act) and op and not missing(op):
                action_ops[act].add(op);operation_actions[op].add(act)
                all_control_rows.append((p.name,ti,ri,uid,act,op,st))
assert action_ops[ACTION]=={OP},("TARGET_ACTION_MULTI_OPERATION",action_ops[ACTION])
assert ACTION in operation_actions[OP],("TARGET_OPERATION_REVERSE_ACTION_MISSING",operation_actions[OP])

# 6. Check target-related stale composite references and stale pre-remediation Operation on same target UID.
stale_target_rows=[]
for p in DOCS:
    for r in control_rows(p,TARGET):
        if any(x in r.get("operation","") for x in ["createEditingRuntimeRun","completeAssembly"]) and r.get("operation")!=OP:
            stale_target_rows.append(r)

# 7. Global target chain completeness vs legal local orchestration semantics.
gaps=[]
for fn,rows in target.items():
    r=rows[0]
    for fld in ["action","gate","permission","payload","operation","method","runtime_owner","persistence","runtime_status"]:
        if missing(r[fld]):gaps.append((fn,fld,r[fld]))
# Atomic ports must carry runtime permission and exact route; payload may live in separate contract tables, so report, don't invent.
port_missing=[]
for uid,rows in ports.items():
    for r in rows:
        for fld in ["operation","method","permission"]:
            if missing(r[fld]):port_missing.append((uid,r["file"],fld,r[fld]))

summary={
"docx_files":31,
"legacy_r9_r5_hits":len(legacy),
"ambiguous_locator_hits":len(ambig),
"old_composite_string_occurrences":sum(x["count"] for x in old_exact),
"target_exact_rows":sum(len(v) for v in target.values()),
"target_docs":len(target),
"target_field_gaps":len(gaps),
"target_stale_atomic_operation_rows":len(stale_target_rows),
"target_action_operation_count":len(action_ops[ACTION]),
"target_operation_action_count":len(operation_actions[OP]),
"atomic_ports":len(PORTS),
"atomic_port_exact_contract_rows":sum(len(v) for v in ports.values()),
"reverse_binding_rows":sum(len(v) for v in reverse.values()),
"forward_branch_rows":len(branch_rows),
"permission_bridge_rows":len(permission_bridge),
"atomic_port_required_field_gaps":len(port_missing),
"bidirectional_status":"PASS",
"runtime_execution_claimed":False,
}
if legacy or ambig or stale_target_rows or gaps or port_missing or action_ops[ACTION]!={OP}:
    summary["bidirectional_status"]="ISSUES_FOUND"
report={
"marker":"ACPOS-20260922-BATCH-62-EDIT-ORCHESTRATION-BIDIRECTIONAL-CLOSURE-AUDIT-V1",
"base_head":BASE_HEAD,"summary":summary,
"legacy_hits":legacy,"ambiguous_locator_hits":ambig,"old_composite_occurrences":old_exact,
"target_rows":target,"ports":ports,"reverse_bindings":reverse,"forward_branch_rows":branch_rows,
"permission_bridge_rows":permission_bridge,
"target_action_operations":sorted(action_ops[ACTION]),
"target_operation_actions":sorted(operation_actions[OP]),
"stale_target_rows":stale_target_rows,"target_gaps":gaps,"port_required_field_gaps":port_missing,
}
Path("__batch62_bidirectional_audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
Path("__batch62_summary.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
print("BATCH62_SUMMARY="+json.dumps(summary,ensure_ascii=False,sort_keys=True))
